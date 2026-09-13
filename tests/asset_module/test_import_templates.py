"""每一種匯入都有範例檔，而且範例檔匯回來不會出事（2026-09-10）。

使用者：「每一個都要範例檔案，不然 USER 不會知道要匯入什麼資料」、「匯錯要提醒」。

要守的三件事：
1. 範例檔的表頭，匯入程式要認得（不然範例本身就是錯的示範）
2. 範例檔原封不動匯回來，範例列要被略過，不能長出假系統／假資產
3. 放錯列（例：類別表放到業務系統那列）要擋下來並講清楚，不是默默匯一半
"""
import io
import sqlite3
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import business_system as bs  # noqa: E402
import db  # noqa: E402
import excel_import  # noqa: E402
import import_templates as tpl  # noqa: E402
import rvtools_import  # noqa: E402


@pytest.fixture()
def conn(tmp_path):
    path = tmp_path / "t.db"
    db.init_db(path)
    c = sqlite3.connect(str(path))
    c.row_factory = sqlite3.Row
    c.execute("INSERT INTO business_system (api_id, name, ap_department, ap_owner, sys_class) "
              "VALUES ('N-001','甲系統','某部門','某甲','第一類')")
    c.commit()
    yield c
    c.close()


def _save(tmp_path, source):
    fname, data = tpl.build(source)
    p = tmp_path / fname
    p.write_bytes(data)
    return p


def _xlsx(tmp_path, header, rows, name="f.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    p = tmp_path / name
    wb.save(p)
    return p


@pytest.mark.parametrize("source", sorted(tpl.TEMPLATES))
def test_每一種都產得出範例檔_而且有表頭和範例列(source):
    fname, data = tpl.build(source)
    assert fname.endswith(".xlsx")
    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) >= 2, f"{source}/{ws.title} 只有表頭，使用者看不出要填什麼"
        assert any(v not in (None, "") for v in rows[1])


def test_系統類別範例_原封不動匯回來_被擋下且原本類別保留(conn, tmp_path):
    with pytest.raises(ValueError, match="範例"):
        bs.import_class_file(_save(tmp_path, "sys_class"), conn)
    assert conn.execute("SELECT sys_class FROM business_system WHERE api_id='N-001'").fetchone()[0] == "第一類"


def test_系統類別範例_加一列真資料_範例列略過真資料生效(conn, tmp_path):
    p = _save(tmp_path, "sys_class")
    wb = load_workbook(p)
    wb.active.append([4, "第二類", "N-001", "甲系統"])
    wb.save(p)
    r = bs.import_class_file(p, conn)
    assert r["updated"] == 1 and r["example_rows"] == 3
    assert r["not_in_table"] == [], "範例代碼不該被當成「對照表沒有」"


def test_業務系統範例_匯回來不會長出假系統(conn, tmp_path):
    r = bs.import_file(_save(tmp_path, "business_system"), conn)
    assert r["imported"] == 0 and r["example_rows"] == 2
    assert conn.execute("SELECT COUNT(*) FROM business_system WHERE api_id LIKE '範例%'").fetchone()[0] == 0


def test_RVTools範例_表頭認得_範例列略過(tmp_path):
    p = _save(tmp_path, "rvtools")
    assert rvtools_import.parse_rvtools(p) == []
    # 同一份表頭填真資料要解析得出來——證明範例的表頭是對的
    wb = load_workbook(p)
    ws = wb["vInfo"]
    hdr = [c.value for c in ws[1]]
    row = ["" for _ in hdr]
    row[hdr.index("VM")] = "app-vm01"
    row[hdr.index("VM UUID")] = "42000000-0000-0000-0000-000000000009"
    ws.append(row)
    wb.save(p)
    recs = rvtools_import.parse_rvtools(p)
    assert [r["vm_name"] for r in recs] == ["app-vm01"]


def test_CIA範例_分頁與表頭認得_範例列不寫入(conn, tmp_path):
    p = _save(tmp_path, "cia_excel")
    summary = excel_import.import_excel(p, conn)
    assert summary["example_rows"] >= 1
    assert conn.execute("SELECT COUNT(*) FROM hardware WHERE asset_serial LIKE '範例%'").fetchone()[0] == 0
    assert not any("找不到分頁" in e for e in summary["errors"])


# ---- 匯錯要提醒 ----

def test_類別表放到業務系統那列_擋下並指出該用哪一列(conn, tmp_path):
    p = _xlsx(tmp_path, ("項次", "系統類別", "APID"), [(1, "第一類", "N-001")])
    with pytest.raises(ValueError, match="系統類別對照表"):
        bs.import_file(p, conn)
    assert conn.execute("SELECT ap_owner FROM business_system WHERE api_id='N-001'").fetchone()[0] == "某甲"


def test_業務系統表放到類別那列_擋下並指出該用哪一列(conn, tmp_path):
    p = _xlsx(tmp_path, ("system_id", "system", "ap_department", "ap_owner"),
              [("N-001", "甲系統", "某部門", "某甲")])
    with pytest.raises(ValueError, match="業務系統對照表"):
        bs.import_class_file(p, conn)


def test_類別表代碼全對不上_不套用_不能把現有分級清光(conn, tmp_path):
    p = _xlsx(tmp_path, ("系統類別", "APID"), [("第二類", "N-999"), ("第三類", "N-998")])
    with pytest.raises(ValueError, match="對不上"):
        bs.import_class_file(p, conn)
    assert conn.execute("SELECT sys_class FROM business_system WHERE api_id='N-001'").fetchone()[0] == "第一類"


def test_類別表多數對不上_匯入但回警告(conn, tmp_path):
    p = _xlsx(tmp_path, ("系統類別", "APID"),
              [("第二類", "N-001"), ("第三類", "N-998"), ("第三類", "N-997")])
    r = bs.import_class_file(p, conn)
    assert r["updated"] == 1 and r["warning"] and "2／3" in r["warning"]


def test_不是CIA清冊_擋下並列出實際分頁(conn, tmp_path):
    p = _xlsx(tmp_path, ("VM", "Host"), [("x", "y")])
    with pytest.raises(ValueError, match="不是 CIA"):
        excel_import.import_excel(p, conn)


def test_範例檔端點_未知來源回404_已知來源回xlsx():
    from fastapi.testclient import TestClient
    import api

    api.app.dependency_overrides[api.require_auth] = lambda: {"username": "t"}
    try:
        c = TestClient(api.app)
        assert c.get("/api/import/nope/export-template").status_code == 404
        r = c.get("/api/import/sys_class/export-template")
        assert r.status_code == 200 and r.content[:2] == b"PK"
    finally:
        api.app.dependency_overrides.pop(api.require_auth, None)
