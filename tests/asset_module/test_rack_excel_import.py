"""機櫃圖 Excel 對照與回寫（2.007.20）：解析、比對狀態、硬性擋下、套用前快照、重複套用、預覽後資料變了。"""
import io
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import auth  # noqa: E402
import db  # noqa: E402
import api  # noqa: E402
import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

_STUB = "test-password-123"
ROOM_DB = "02_內湖機房"
ROOM = "內湖機房"


def make_xlsx(racks: dict[str, list[tuple[int, int, str]]], height=20) -> bytes:
    """每個櫃三欄：左 U、中設備、右 U；第 1 列放櫃名。設備佔多 U 就合併儲存格。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "4F"
    for i, (name, devs) in enumerate(racks.items()):
        c0 = 1 + i * 4
        ws.cell(1, c0 + 1, name)
        for k in range(height):
            u = height - k
            ws.cell(2 + k, c0, u)
            ws.cell(2 + k, c0 + 2, u)
        for lo, hi, label in devs:
            r_top, r_bot = 2 + (height - hi), 2 + (height - lo)
            ws.cell(r_top, c0 + 1, label)
            if r_bot > r_top:
                ws.merge_cells(start_row=r_top, start_column=c0 + 1, end_row=r_bot, end_column=c0 + 1)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _hw(conn, serial, host, ip, rack_no, model="Dell R760"):
    db.insert_hardware(conn, asset_serial=serial, hostname=host, ip=ip, physical_location=ROOM_DB,
                       rack_no=rack_no, asset_status="使用中", device_model=model)


def _client(tmp, monkeypatch):
    monkeypatch.setenv("ASSET_BACKUP_DIR", str(Path(tmp) / "bk"))
    db_path = Path(tmp) / "test.db"
    db.init_db(db_path)

    def _override():
        conn = db.get_connection(db_path)
        try:
            yield conn
        finally:
            conn.close()

    api.app.dependency_overrides[api.get_db] = _override
    c = TestClient(api.app)
    conn = db.get_connection(db_path)
    try:
        db.create_user(conn, "admin", auth.hash_password(_STUB))
        _hw(conn, "S1", "nopos1", "10.99.7.1", "")                      # 沒填櫃位
        _hw(conn, "S2", "moved2", "10.99.7.2", "B8/5U")                 # Excel 說在 7-8U
        _hw(conn, "S3", "same3", "10.99.7.3", "B8/10-11U")              # 一致
        _hw(conn, "S4", "h4-idrac", "10.99.7.4", "", model="iDRAC9")    # 管理卡
        _hw(conn, "S5", "vm5", "10.99.7.5", "", model="(VM)")           # VM
        _hw(conn, "S6", "twin6", "10.99.7.6", "")
        _hw(conn, "S7", "twin7", "10.99.7.7", "")
        _hw(conn, "S8", "occupant8", "10.99.7.8", "B8/15U")             # 佔著 15U
        _hw(conn, "S9", "other9", "10.99.7.9", "B8/18U")
        conn.commit()
    finally:
        conn.close()
    assert c.post("/api/auth/login", json={"username": "admin", "password": _STUB}).status_code == 200
    return c, db_path


SHEET = {"B8": [(2, 2, "Dell R760 10.99.7.1"),            # S1 NOPOS → 預設勾
                (7, 8, "Dell R760 10.99.7.2"),            # S2 MOVED → 不勾
                (10, 11, "Dell R760 10.99.7.3"),          # S3 OK
                (3, 3, "iDRAC 10.99.7.4"),                # 管理卡
                (4, 4, "VM 10.99.7.5"),                   # VM
                (5, 5, "Dell 10.99.7.6 / 10.99.7.7"),     # 一格兩台
                (15, 15, "Dell R760 10.99.7.1"),          # S1 又被指到 → dup
                (18, 18, "Cisco switch"),                 # S9 在 18U，名稱對不上 → POSONLY
                (13, 13, "HPE DL380 10.99.9.9")]}         # 系統沒有 → MISSING


def _by_label(rows):
    return {r["label"]: r for r in rows}


def test_解析櫃名與合併U(monkeypatch):
    import rack_excel_import as rx
    p = rx.parse_workbook(make_xlsx({"B08(VMware)": [(7, 8, "A"), (3, 3, "B")], "M1": [(1, 1, "C")]}))
    assert [r["rack"] for r in p["racks"]] == ["B8", "M1"]
    d = {x["label"]: (x["u_start"], x["u_end"]) for x in p["racks"][0]["devices"]}
    assert d == {"A": (7, 8), "B": (3, 3)}


def test_略過關機與調整分頁():
    import rack_excel_import as rx
    wb = openpyxl.load_workbook(io.BytesIO(make_xlsx({"B8": [(1, 1, "x")]})))
    wb.create_sheet("M07-08 調整")
    bio = io.BytesIO()
    wb.save(bio)
    assert rx.parse_workbook(bio.getvalue())["skipped_sheets"] == ["M07-08 調整"]


def test_比對狀態與硬性擋下(monkeypatch):
    with tempfile.TemporaryDirectory() as tmp:
        c, _ = _client(tmp, monkeypatch)
        r = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                   files={"file": ("x.xlsx", make_xlsx(SHEET), "application/octet-stream")})
        assert r.status_code == 200, r.text
        rows = r.json()["rows"]
        by = {}
        for x in rows:
            by.setdefault(x["label"], []).append(x)
        assert by["Dell R760 10.99.7.3"][0]["status"] == "OK"
        s2 = by["Dell R760 10.99.7.2"][0]
        assert s2["status"] == "MOVED" and not s2["auto"] and "moved" in s2["blockers"]
        assert by["iDRAC 10.99.7.4"][0]["status"] == "BLOCKED" and "mgmt" in by["iDRAC 10.99.7.4"][0]["blockers"]
        assert "vm" in by["VM 10.99.7.5"][0]["blockers"]
        assert "ambig" in by["Dell 10.99.7.6 / 10.99.7.7"][0]["blockers"] or by["Dell 10.99.7.6 / 10.99.7.7"][0]["status"] == "BLOCKED"
        assert by["Cisco switch"][0]["status"] == "POSONLY"
        assert by["HPE DL380 10.99.9.9"][0]["status"] == "MISSING"
        # S1 被兩格指到 → 兩格都降成 dup，不能自動寫
        dups = by["Dell R760 10.99.7.1"]
        assert len(dups) == 2 and all("dup" in d["blockers"] and d["action"] is None for d in dups)
        assert r.json()["summary"]["strict_ok_pct"] == round(1 / len(rows) * 100)


def test_套用_快照_稽核_重複套用(monkeypatch):
    with tempfile.TemporaryDirectory() as tmp:
        c, db_path = _client(tmp, monkeypatch)
        sheet = {"B8": [(2, 2, "Dell R760 10.99.7.1"), (7, 8, "Dell R760 10.99.7.2"), (3, 3, "iDRAC 10.99.7.4")]}
        pv = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                    files={"file": ("x.xlsx", make_xlsx(sheet), "application/octet-stream")}).json()
        rows = _by_label(pv["rows"])
        s1, s2, idrac = rows["Dell R760 10.99.7.1"], rows["Dell R760 10.99.7.2"], rows["iDRAC 10.99.7.4"]
        assert s1["auto"] and not s2["auto"] and not idrac["auto"]
        ap = c.post("/api/racks/excel-import/apply", json={"run_id": pv["run_id"], "row_ids": [s1["id"], idrac["id"]]})
        assert ap.status_code == 200, ap.text
        out = ap.json()
        assert out["applied"] == 1 and out["backup"] and Path(out["backup"]).exists()
        assert any(s["id"] == idrac["id"] for s in out["skipped"])          # 管理卡伺服器端也擋
        det = c.get("/api/racks/detail", params={"room": ROOM, "rack": "B8"}).json()
        s1dev = next(d for d in det["devices"] if d["asset_serial"] == "S1")
        assert (s1dev["u_start"], s1dev["u_end"]) == (2, 2)
        # 稽核紀錄留了
        conn = db.get_connection(db_path)
        try:
            n = conn.execute("SELECT COUNT(*) FROM rack_position_audit WHERE action='set' AND note LIKE 'Excel 機櫃圖匯入%'").fetchone()[0]
        finally:
            conn.close()
        assert n == 1
        # 同一格再套用一次 → 跳過
        again = c.post("/api/racks/excel-import/apply", json={"run_id": pv["run_id"], "row_ids": [s1["id"]]}).json()
        assert again["applied"] == 0 and "已經套用" in again["skipped"][0]["why"]
        # MOVED 那格使用者明確勾了才寫
        ap2 = c.post("/api/racks/excel-import/apply", json={"run_id": pv["run_id"], "row_ids": [s2["id"]]}).json()
        assert ap2["applied"] == 1


def test_預覽後資料變了_重算後擋下(monkeypatch):
    with tempfile.TemporaryDirectory() as tmp:
        c, db_path = _client(tmp, monkeypatch)
        pv = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                    files={"file": ("x.xlsx", make_xlsx({"B8": [(2, 2, "Dell R760 10.99.7.1")]}), "application/octet-stream")}).json()
        row = pv["rows"][0]
        assert row["auto"]
        # 預覽之後有人把 S1 的名稱改成管理卡寫法 → 現在重算要擋
        conn = db.get_connection(db_path)
        try:
            db.update_hardware(conn, "S1", {"hostname": "nopos1-idrac"})
        finally:
            conn.close()
        out = c.post("/api/racks/excel-import/apply", json={"run_id": pv["run_id"], "row_ids": [row["id"]]}).json()
        assert out["applied"] == 0 and out["skipped"]


def test_重疊的格預設不勾_硬擋之外可手動(monkeypatch):
    with tempfile.TemporaryDirectory() as tmp:
        c, _ = _client(tmp, monkeypatch)
        # S1 要放 15U，但 S8 已在那裡
        pv = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                    files={"file": ("x.xlsx", make_xlsx({"B8": [(15, 15, "Dell R760 10.99.7.1")]}), "application/octet-stream")}).json()
        row = pv["rows"][0]
        assert row["status"] == "NOPOS" and "overlap" in row["blockers"] and not row["auto"]


def test_不是機櫃圖的檔案給清楚訊息(monkeypatch):
    with tempfile.TemporaryDirectory() as tmp:
        c, _ = _client(tmp, monkeypatch)
        wb = openpyxl.Workbook()
        wb.active["A1"] = "hello"
        bio = io.BytesIO()
        wb.save(bio)
        r = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                   files={"file": ("x.xlsx", bio.getvalue(), "application/octet-stream")})
        assert r.status_code == 400 and "找不到機櫃圖" in r.json()["detail"]


def test_別名櫃_寫進系統原本的櫃_不新增櫃(monkeypatch):
    """內湖「1-1櫃(A1)」＝Excel 的 A1：寫入要用系統的 1-1，不能憑空多出一個 A1；重疊也要看 1-1（審查 H1）。"""
    with tempfile.TemporaryDirectory() as tmp:
        c, db_path = _client(tmp, monkeypatch)
        conn = db.get_connection(db_path)
        try:
            _hw(conn, "A-1", "alias5", "10.99.8.1", "1-1櫃(A1)/5U")
            _hw(conn, "A-2", "occ8", "10.99.8.2", "1-1/8U")
            _hw(conn, "A-3", "mover", "10.99.8.4", "1-1/30U")
            conn.commit()
        finally:
            conn.close()
        pv = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                    files={"file": ("x.xlsx", make_xlsx({"A1": [(8, 8, "Dell 10.99.8.4")]}), "application/octet-stream")}).json()
        row = pv["rows"][0]
        assert row["status"] == "MOVED" and row["to_rack"] == "1-1"
        assert "overlap" in row["blockers"] and "new_rack" not in row["blockers"]   # 1-1 的 8U 已被 occ8 佔著
        # 目標改空位：7U
        pv = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                    files={"file": ("x.xlsx", make_xlsx({"A1": [(7, 7, "Dell 10.99.8.4")]}), "application/octet-stream")}).json()
        row = pv["rows"][0]
        assert "overlap" not in row["blockers"]
        out = c.post("/api/racks/excel-import/apply", json={"run_id": pv["run_id"], "row_ids": [row["id"]]}).json()
        assert out["applied"] == 1 and "/1-1/" in out["done"][0]["to"]
        det = c.get("/api/racks/detail", params={"room": ROOM, "rack": "1-1"}).json()
        assert any(d["asset_serial"] == "A-3" and d["u_start"] == 7 for d in det["devices"])


def test_同一台已有一格OK_另一格不能搬(monkeypatch):
    with tempfile.TemporaryDirectory() as tmp:
        c, _ = _client(tmp, monkeypatch)
        pv = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                    files={"file": ("x.xlsx", make_xlsx({"B8": [(10, 11, "Dell 10.99.7.3")], "B9": [(3, 3, "Dell 10.99.7.3")]}),
                           "application/octet-stream")}).json()
        by = {r["rack"]: r for r in pv["rows"]}
        assert by["B8"]["status"] == "OK"
        assert "dup" in by["B9"]["blockers"] and by["B9"]["action"] is None


def test_U超出範圍的格不會讓整批失敗(monkeypatch):
    with tempfile.TemporaryDirectory() as tmp:
        c, _ = _client(tmp, monkeypatch)
        sheet = {"B8": [(2, 2, "Dell R760 10.99.7.1"), (65, 65, "Dell R760 10.99.7.2")]}
        pv = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                    files={"file": ("x.xlsx", make_xlsx(sheet, height=70), "application/octet-stream")}).json()
        by = {r["label"]: r for r in pv["rows"]}
        assert "urange" in by["Dell R760 10.99.7.2"]["blockers"] and by["Dell R760 10.99.7.2"]["action"] is None
        ids = [by["Dell R760 10.99.7.1"]["id"], by["Dell R760 10.99.7.2"]["id"]]
        out = c.post("/api/racks/excel-import/apply", json={"run_id": pv["run_id"], "row_ids": ids})
        assert out.status_code == 200 and out.json()["applied"] == 1


def test_多張工作表的同一個櫃合併(monkeypatch):
    import rack_excel_import as rx
    wb = openpyxl.load_workbook(io.BytesIO(make_xlsx({"B4": [(7, 8, "網路設備")]})))
    ws2 = openpyxl.load_workbook(io.BytesIO(make_xlsx({"B4": [(2, 3, "主機"), (7, 8, "重疊的")]}))).active
    ws2.title = "大型主機"
    wb2 = openpyxl.Workbook()
    wb2.remove(wb2.active)
    for src in (wb.active, ws2):
        t = wb2.create_sheet(src.title)
        for row in src.iter_rows():
            for cell in row:
                if cell.value is not None:
                    t[cell.coordinate] = cell.value
        for m in src.merged_cells.ranges:
            t.merge_cells(str(m))
    bio = io.BytesIO()
    wb2.save(bio)
    p = rx.parse_workbook(bio.getvalue())
    assert len(p["racks"]) == 1 and p["racks"][0]["rack"] == "B4"
    labels = {d["label"] for d in p["racks"][0]["devices"]}
    assert labels == {"網路設備", "主機"} and len(p["conflicts"]) == 1


def test_櫃名優先取設備欄_不被U數字樣誤判():
    import rack_excel_import as rx
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B1"] = "Colo-4Y12"
    ws["A2"] = "U 數"
    for k in range(12):
        ws.cell(3 + k, 1, 12 - k)
        ws.cell(3 + k, 3, 12 - k)
    bio = io.BytesIO()
    wb.save(bio)
    assert rx.parse_workbook(bio.getvalue())["racks"][0]["rack"] == "4Y12"


def test_機櫃正面帶出Excel層_系統沒畫到的格(monkeypatch):
    """2026-10-28：以 Excel 為主。系統面板沒畫到的格（沒登位置／系統沒有／待確認）都要在 rack detail 出現；
    OK 與 POSONLY（同一格系統面板已經在）不重複畫；不新增資產、台數不變。"""
    with tempfile.TemporaryDirectory() as tmp:
        c, db_path = _client(tmp, monkeypatch)
        before = c.get("/api/racks/detail", params={"room": ROOM, "rack": "B8"}).json()
        assert before["excel"] == {"run": None, "cells": [], "total": 0}      # 還沒上傳過 Excel
        r = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                   files={"file": ("x.xlsx", make_xlsx(SHEET), "application/octet-stream")})
        assert r.status_code == 200, r.text
        d = c.get("/api/racks/detail", params={"room": ROOM, "rack": "B8"}).json()
        ex = d["excel"]
        assert ex["run"]["filename"] == "x.xlsx" and ex["total"] == len(SHEET["B8"])
        st = {x["label"]: x["status"] for x in ex["cells"]}
        assert st["HPE DL380 10.99.9.9"] == "MISSING"
        assert st["Dell R760 10.99.7.1"] in ("NOPOS", "BLOCKED")
        assert "Dell R760 10.99.7.3" not in st                 # OK：系統面板已經畫了
        assert "Cisco switch" not in st                        # POSONLY：同一格系統面板已在
        assert all(x["suggest"] for x in ex["cells"])
        assert len(d["devices"]) == len(before["devices"])     # 沒有新增任何設備
        # 面板燈號／「確認」欄用：每台系統設備跟 Excel 對的結果（S3 一致、S9 名稱不同；S2 位置不同）
        by_key = ex["by_key"]
        assert sorted(by_key.values()).count("OK") == 1 and "POSONLY" in by_key.values() and "MOVED" in by_key.values()


def test_Excel獨有的櫃_平面圖有磚_點得開_只有Excel層_不算台數(monkeypatch):
    """2026-10-29：Excel 有 B02(VMware)，系統完全沒登記這個櫃 → 機房平面圖要有這個櫃、點開看得到 Excel 的格。"""
    import rack_layout as rl
    with tempfile.TemporaryDirectory() as tmp:
        c, db_path = _client(tmp, monkeypatch)
        sheet = {**SHEET, "B02(VMware)": [(10, 11, "VMware ESXi 10.99.8.1"), (12, 12, "Dell R760 10.99.8.2")]}
        r = c.post("/api/racks/excel-import/preview", params={"room": ROOM},
                   files={"file": ("x.xlsx", make_xlsx(sheet), "application/octet-stream")})
        assert r.status_code == 200, r.text
        dc = c.get("/api/racks/datacenter").json()
        room = next(x for x in dc["rooms"] if x["room"] == ROOM)
        t = next(t for t in room["racks"] if t["rack"] == "B2")
        assert t["excel_only"] is True and t["excel_cells"] == 2 and t["device_count"] == 0       # 不算台數
        assert not next(t for t in room["racks"] if t["rack"] == "B8").get("excel_only")
        assert room["racked"] == sum(x["device_count"] for x in room["racks"])
        assert room["power"]["free_u"] == sum(x["power"]["free_u"] for x in room["racks"] if not x.get("excel_only"))   # 合計不灌水
        det = c.get("/api/racks/detail", params={"room": ROOM, "rack": "B2"})
        assert det.status_code == 200, det.text
        d = det.json()
        assert d["excel_only"] is True and d["devices"] == [] and len(d["excel"]["cells"]) == 2
        assert all(x["status"] == "MISSING" for x in d["excel"]["cells"])
        assert "B2" in d["room_racks"]
        # 沒有 Excel 層、系統也沒有的櫃，仍然查無
        assert c.get("/api/racks/detail", params={"room": ROOM, "rack": "Z99"}).status_code in (400, 404)
        conn = db.get_connection(db_path)
        try:
            assert conn.execute("SELECT COUNT(*) FROM hardware").fetchone()[0] == 9           # 沒有新增資產
        finally:
            conn.close()
