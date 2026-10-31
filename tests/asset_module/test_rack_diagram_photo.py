"""機櫃照片對帳（2026-10-24）：清單格式＋來源照片／拍攝日期／把握度、配對順序（序號→主機名→IP→機型＋位置）、
未辨識不硬配、系統有照片沒有、逐筆確認（只增不改）、只有確認過的才進 CIA 待更正。全部假資料。"""
import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import rack_diagram as rd  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    return db.get_connection(p)


def _seed(conn):
    def hw(serial, host, ip, rack, model, hw_serial=None):
        db.insert_hardware(conn, asset_serial=serial, hostname=host, ip=ip, physical_location="02_內湖機房",
                           rack_no=rack, device_model=model, hw_serial=hw_serial)
    hw("H-1", "APPSVR003-225", None, "2-2/11-15U", "IBM S922", hw_serial="7800AAA")   # 照片在 C1 7-10U
    hw("H-2", "APPSVRB21", None, "2-2/7-10U", "IBM S922")                              # 照片在 B3 13-16U（同型號兩台）
    hw("H-3", "APPSVRB23", None, "2-2/7-10U", "IBM S922")
    hw("H-4", "sec1t", None, "2-2/12-15U", "IBM E950")                                 # 照片同位置 → 一致
    hw("H-5", "lonelyb2", None, "2-2/30U", "DELL R750")                                # 照片沒拍到（同櫃）
    hw("H-6", "only-in-b9", None, "2-9/3U", "DELL R750")                               # 整櫃沒拍到
    hw("H-7", "dellbox01", None, "2-3/20U", "DELL R660")                                 # 照片寫 R750 → 機型不同
    hw("H-8", "solo740", None, "2-4/5U", "DELL R740")                                  # 照片只寫機型，同櫃唯一 → 機型＋位置
    conn.commit()


def _flat(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "照片判讀"
    ws.append(["機房", "樓層", "機櫃", "U起", "U迄", "內容", "IP", "序號", "來源", "拍攝日期", "把握度"])
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


ROWS = [
    ["內湖", "3F", "C1", 7, 10, "IBM S922", None, "7800AAA", "IMG_0101.jpg", "2026-10-20", "高"],      # 序號配
    ["內湖", "3F", "B3", 13, 16, "IBM S922 APPSVRB21", None, None, "IMG_0102.jpg", "2026-10-20", "中"],  # 主機名配
    ["內湖", "3F", "B4", 13, 16, "IBM S922 APPSVRB23", None, None, "IMG_0103.jpg", "2026-10-20", "中"],
    ["內湖", "3F", "B2", 12, 15, "IBM E950 sec1t", None, None, "IMG_0104.jpg", "2026-10-20", "高"],
    ["內湖", "3F", "B2", 7, 10, "IBM S922", None, None, "IMG_0104.jpg", "2026-10-20", "低"],            # 只有機型、同櫃兩台同型號 → 不配
    ["內湖", "3F", "B3", 20, 20, "DELL R750 dellbox01", None, None, "IMG_0105.jpg", "2026-10-20", "高"],   # 系統在 B3 20U 是 R660
    ["內湖", "3F", "B4", 5, 5, "DELL R740", None, None, "IMG_0106.jpg", "2026-10-20", "中"],             # 機型＋位置
    ["內湖", "3F", "B4", 30, 31, "看不清楚的機器", None, None, "IMG_0106.jpg", "2026-10-20", "未辨識"],
    ["內湖", "3F", "B4", 40, 40, "Cisco SW ghost99", "10.99.9.9", None, "IMG_0106.jpg", "2026-10-20", "高"],  # 系統沒有
]


def test_清單格式_來源與把握度():
    rows, racks = rd.parse_flat([list(r) for r in [["機房", "樓層", "機櫃", "U起", "U迄", "內容", "IP", "序號", "來源", "拍攝日期", "把握度"], *ROWS]])
    r0 = rows[0]
    assert (r0["rack_norm"], r0["u_start"], r0["u_end"], r0["serials"], r0["source"], r0["shot_date"], r0["confidence"]) == \
        ("C-1", 7, 10, ["7800AAA"], "IMG_0101.jpg", "2026-10-20", "高")
    assert rows[7]["confidence"] == "未辨識"
    assert rows[8]["main_ip"] == "10.99.9.9"
    assert {r["rack_norm"] for r in racks} == {"C-1", "B-3", "B-4", "B-2"}
    assert rd.norm_confidence("high") == "高" and rd.norm_confidence("") is None


def test_配對順序與差異分類(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        imp = rd.import_diagram(conn, "內湖", "photos.xlsx", _flat(ROWS), "tester")
        rc = rd.reconcile(conn, imp["import_id"])
        by = {(x["rack"], x["u"]): x for x in rc["rows"]}
        assert by[("C1", "7-10")]["sys_name"] == "APPSVR003-225" and by[("C1", "7-10")]["match_by"].startswith("序號")
        assert by[("C1", "7-10")]["result"] == "機櫃不同"
        assert by[("B3", "13-16")]["sys_name"] == "APPSVRB21" and by[("B3", "13-16")]["match_by"].startswith("主機名")
        assert by[("B4", "13-16")]["result"] == "機櫃不同"
        assert by[("B2", "12-15")]["result"] == "一致"
        assert by[("B2", "7-10")]["sys_key"] is None                                 # 同型號兩台：不靠機型配
        assert by[("B3", "20")]["result"] == "機型不同"
        assert by[("B4", "5")]["sys_name"] == "solo740" and by[("B4", "5")]["match_by"].startswith("機型＋位置")
        assert by[("B4", "30-31")]["result"] == "未辨識" and by[("B4", "30-31")]["sys_key"] is None
        assert by[("B4", "40")]["result"] == "照片有系統沒有"
        assert by[("C1", "7-10")]["source"] == "IMG_0101.jpg" and by[("C1", "7-10")]["confidence"] == "高"
        assert {x["name"] for x in rc["system_only"]} == {"lonelyb2"}                  # 拍到的櫃裡系統有、照片沒有
        assert "不一定是錯" in rc["system_only"][0]["note"]
        assert [m["rack"] for m in rc["missing_racks"]] == ["2-9"]                    # 沒拍到的櫃（不等於空櫃）
        g = rc["groups"]
        assert (g["位置不同"], g["機型不同"], g["照片有系統沒有"], g["系統有照片沒有"], g["沒拍到的櫃"]) == (3, 1, 1, 1, 1)
    finally:
        conn.close()


def test_逐筆確認_只增不改_只有確認過的進CIA待更正(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        imp = rd.import_diagram(conn, "內湖", "photos.xlsx", _flat(ROWS), "tester")
        rc = rd.reconcile(conn, imp["import_id"])
        rows = {x["sys_name"]: x for x in rc["rows"] if x.get("sys_name")}
        a, b = rows["APPSVR003-225"]["row_id"], rows["APPSVRB21"]["row_id"]
        rd.review_row(conn, a, "confirm", "照片看得很清楚", "alice")
        rd.review_row(conn, b, "confirm", None, "alice")
        rd.review_row(conn, b, "misread", "照片其實是 B21 的鄰櫃", "bob")          # 改判：最新一筆為準
        with pytest.raises(ValueError):
            rd.review_row(conn, a, "apply_to_cia", None, "x")
        hist = rd.review_history(conn, imp["import_id"])
        assert len(hist) == 3                                                        # 只增不改
        fx = rd.fix_rows(conn, imp["import_id"])
        assert [x["sys_name"] for x in fx] == ["APPSVR003-225"]
        buf, n = rd.fix_xlsx(conn, imp["import_id"])
        ws = load_workbook(buf).active
        head = [c.value for c in ws[1]]
        row = dict(zip(head, [c.value for c in ws[2]]))
        assert n == 1 and row["資產序號"] == "H-1" and row["來源"] == "IMG_0101.jpg" and row["確認人"] == "alice"
        assert row["現場機櫃"] == "C1" and row["系統機櫃"] == "2-2"
        # 系統資料完全沒被改（不寫 CIA）
        assert conn.execute("SELECT rack_no FROM hardware WHERE asset_serial = 'H-1'").fetchone()[0] == "2-2/11-15U"
    finally:
        conn.close()


def test_審查意見_確認過期_多筆候選_重複確認_亂寫把握度_沒有確認人_別的匯入(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        imp = rd.import_diagram(conn, "內湖", "photos.xlsx", _flat(ROWS), "tester")
        iid = imp["import_id"]
        rows = {x["sys_name"]: x for x in rd.reconcile(conn, iid)["rows"] if x.get("sys_name")}
        rid = rows["APPSVR003-225"]["row_id"]
        with pytest.raises(PermissionError):
            rd.review_row(conn, rid, "confirm", None, None)                       # 沒有確認人不收
        ok_row = rows["sec1t"]["row_id"]
        with pytest.raises(ValueError):
            rd.review_row(conn, ok_row, "confirm", None, "alice")                 # 「一致」不是要確認的差異
        rd.review_row(conn, rid, "confirm", None, "alice")
        assert len(rd.fix_rows(conn, iid)) == 1
        # 確認之後系統資料變了：序號改成另一台在用 → 這列改配到別台，確認不算數
        conn.execute("UPDATE hardware SET hw_serial = NULL WHERE asset_serial = 'H-1'")
        conn.execute("UPDATE hardware SET hw_serial = '7800AAA' WHERE asset_serial = 'H-5'")
        conn.commit()
        assert rd.fix_rows(conn, iid) == []
        # 別的匯入的判定不會漏進來
        imp2 = rd.import_diagram(conn, "內湖", "photos2.xlsx", _flat(ROWS), "tester")
        assert rd.fix_rows(conn, imp2["import_id"]) == [] and rd.review_history(conn, imp2["import_id"]) == []
        with pytest.raises(LookupError):
            rd.review_history(conn, 9999)
        # 亂寫的把握度＝未辨識，不硬配
        assert rd.norm_confidence("很高") == "未辨識"
    finally:
        conn.close()


def test_同名兩台實體機_標多筆候選_不配(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        db.insert_hardware(conn, asset_serial="H-9", hostname="APPSVRB21.corp", ip=None, physical_location="02_內湖機房",
                           rack_no="2-5/1U", device_model="X")
        db.insert_hardware(conn, asset_serial="H-10", hostname="APPSVRB21", ip="10.99.4.4", physical_location="02_內湖機房",
                           rack_no="2-6/1U", device_model="X")
        conn.commit()
        imp = rd.import_diagram(conn, "內湖", "photos.xlsx", _flat(ROWS), "tester")
        x = next(x for x in rd.reconcile(conn, imp["import_id"])["rows"] if "APPSVRB21" in x["text"])
        assert x["result"] == "多筆候選（只能人工）" and x["sys_key"] is None
        with pytest.raises(ValueError):
            rd.review_row(conn, x["row_id"], "confirm", None, "alice")
    finally:
        conn.close()


def test_同一台被兩列確認_都不匯出(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        rows = ROWS + [["內湖", "3F", "C2", 1, 4, "IBM S922", None, "7800AAA", "IMG_0199.jpg", "2026-10-21", "中"]]
        imp = rd.import_diagram(conn, "內湖", "photos.xlsx", _flat(rows), "tester")
        both = [x for x in rd.reconcile(conn, imp["import_id"])["rows"] if x.get("sys_name") == "APPSVR003-225"]
        assert len(both) == 2
        for x in both:
            rd.review_row(conn, x["row_id"], "confirm", None, "alice")
        assert rd.fix_rows(conn, imp["import_id"]) == []
    finally:
        conn.close()


def test_截圖清單格式_樓層推機房_U欄_註記_預設來源(tmp_path):
    """10-23 截圖轉出來的清單：項次,樓層,分頁,機櫃,U起,U迄,內容,IP,註記（沒有機房／來源／把握度欄）；
    另一種是 U 寫在單一欄「12-15」。全部假資料。"""
    csv_text = "\n".join([
        "項次,樓層,分頁,機櫃,U起,U迄,內容,IP,註記",
        "1,3F,3F RACK B,C1,7,10,IBM S922 (7800AAA),,",          # 序號對 H-1
        "2,3F,3F RACK B,B3,13,16,IBM S922 APPSVRB21,,",
        "3,4F,4F RACK B,B2,12,15,IBM E950 sec1t,,",
        "4,4F,4F RACK B,B2,20,20,新DB Node1,,暫定",
        "5,4F,4F RACK B,B2,25,25,DELL R750 lonelyb2,,已關機",
        "6,21F,21F,1-1,10,10,DELL R750 dunnan01,10.99.21.1,",     # 敦南：選內湖時不收
        "7,板橋5F,板橋 M,M16,13,15,IBM AS400 bqhost01,,",         # 板橋：不收
    ]).encode("utf-8")
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        imp = rd.import_diagram(conn, "內湖", "截圖清單.csv", csv_text, "tester",
                                {"source": "機櫃圖截圖 2026-10-23", "shot_date": "2026-10-23", "confidence": "中"})
        assert imp["rows"] == 5 and "敦南 1 列" in imp["skipped"] and "板橋 1 列" in imp["skipped"]
        rc = rd.reconcile(conn, imp["import_id"])
        by = {(x["rack"], x["u"]): x for x in rc["rows"]}
        x = by[("C1", "7-10")]
        assert x["sys_name"] == "APPSVR003-225" and x["match_by"].startswith("序號") and x["result"] == "機櫃不同"
        assert (x["source"], x["shot_date"], x["confidence"]) == ("機櫃圖截圖 2026-10-23", "2026-10-23", "中")
        assert by[("B2", "12-15")]["result"] == "一致"
        assert by[("B2", "20")]["result"] == "規劃中"
        assert by[("B2", "25")]["state"] == "已關機" and "已關機" in by[("B2", "25")]["text"]
    finally:
        conn.close()
    # 單一 U 欄、機房欄
    rows, racks = rd.parse_flat([["項次", "機房", "樓層", "機櫃", "U", "內容", "IP", "註記", "系統U"],
                                 [1, "內湖", "4F", "C1", "12-15", "x host9 10.99.1.9", None, None, "3-4"],
                                 [2, "敦南", "21F", "1-1", "7", "y host8", None, None, None]], "全部清單與對帳", "內湖")
    assert [(r["rack"], r["u_start"], r["u_end"]) for r in rows] == [("C1", 12, 15)]
    assert rd.row_site(None, "3F") == "內湖" and rd.row_site(None, "板橋5F") == "板橋" and rd.row_site("敦南機房", None) == "敦南"


def test_可比對的編號_描述字不算_主機名貼紙算(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        db.insert_hardware(conn, asset_serial="H-20", hostname="APPSVR163-097", ip=None, physical_location="02_內湖機房",
                           rack_no="1-1/30U", device_model="X")
        conn.commit()
        sysd = rd._Sys(conn)
        assert rd.label_keys("Cisco switch KVM DELL Server", sysd) == []
        assert rd.label_keys("Lenovo Server-M3650 M5", sysd) == []                    # 機型字
        assert rd.label_keys("會館DB3_Node1", sysd) == []                             # 緊貼中文＝描述
        assert rd.label_keys("NH-R3-03-ACI-L207", sysd) == ["NH-R3-03-ACI-L207"]      # 有分段的貼紙
        assert rd.label_keys("IBM P720 APPSVRW21", sysd) == ["APPSVRW21"]             # 開頭像系統主機名
        rows, _ = rd.parse_flat([["機櫃", "U", "內容", "IP"],
                                 ["A1", "30", "APPSVR163-097(MOB-016)", None],
                                 ["A1", "31", "Cisco switch", None],
                                 ["A1", "32", "NH-R3-03-ACI-L207", None],
                                 ["A1", "33", "DELL Server", "10.99.77.7"]])
        imp = rd.import_diagram(conn, "內湖", "k.csv", ("機櫃,U,內容,IP\n" + "\n".join(
            ",".join(str(c or "") for c in r) for r in [["A1", 30, "APPSVR163-097(MOB-016)", ""], ["A1", 31, "Cisco switch", ""],
                                                          ["A1", 32, "NH-R3-03-ACI-L207", ""], ["A1", 33, "DELL Server", "10.99.77.7"]])).encode(), "t")
        rc = rd.reconcile(conn, imp["import_id"])
        res = {x["u"]: x for x in rc["rows"]}
        assert res["30"]["sys_name"] == "APPSVR163-097" and res["30"]["result"] == "一致"   # 括號裡的不擋主機名比對
        assert res["31"]["result"] == "沒有可比對的編號（只能人工）"
        assert res["32"]["result"] == "照片有系統沒有" and res["33"]["result"] == "照片有系統沒有"
        r = rc["rates"]
        assert r["placed"] == 1 and r["same_rack_rate"] == 100 and r["u_rate"] == 100
        assert rows and len(rows) == 4
    finally:
        conn.close()
