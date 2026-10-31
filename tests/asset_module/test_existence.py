"""存在可信度／CIA 完整率／已納管單一定義（決策 EXIST1，使用者 2026-10-18）。

使用者原話：「CIA 表是最致命的缺點，完整率可能只有六、七成；DY 可能九成；如果我真的能納管，
就代表 100% 沒問題。CIA 漏登是 CIA 表不正確，事實上就是有這幾台。」

釘住：
- 台數＝存在可信度 ≥ 門檻（跟 CIA 有沒有登記無關）；30 天沒有證據＝疑似下線；退役＝備查
- 計入 + 疑似下線 + 退役 = 全部機器；CIA 完整率＝有登記 ÷ 計入
- 已納管＝收集帳號 7 天內連得上（CIA 無關），首頁四態跟存在可信度同一個數
- 沒有主機名的掃描機器不可以被併成同一台（以前 machine_key(None, ip, None) 全落到 'sn:None'）
"""
import sys
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import account_inventory  # noqa: E402
import api  # noqa: E402
import auth  # noqa: E402
import db  # noqa: E402
import existence  # noqa: E402
import manage_state as ms  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

_STUB_CREDENTIAL = "test-password-123"


def _ago(days: float) -> str:
    return (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d %H:%M:%S")


def _seed(conn):
    hw = db.insert_hardware
    # A：CIA，收集帳號 2 天前連上 → 已納管 100
    hw(conn, asset_serial="HW-A", hostname="alpha", ip="10.99.1.1", environment="正式",
       collect_ok=1, collect_checked_at=_ago(2))
    # B：CIA，收集帳號 10 天前連上（過期）＋3 天前掃到 → 70、已納管但失聯
    hw(conn, asset_serial="HW-B", hostname="bravo", ip="10.99.1.2", environment="正式",
       collect_ok=1, collect_checked_at=_ago(10))
    # C：DY 空殼，DY 5 天前看到 → 90、CIA 待補登、只有 DY＝待確認
    c_id = hw(conn, asset_serial="DYN-C", hostname="charlie", ip="10.99.1.3")
    # D：vCenter 空殼，5 天前 → 90、CIA 待補登、不是待確認（vCenter 是 hypervisor 自己的清單）
    d_id = hw(conn, asset_serial="VC-D", hostname="delta", ip="10.99.1.4")
    # E：只有 CIA，什麼證據都沒有 → 60、疑似下線
    hw(conn, asset_serial="HW-E", hostname="echo", ip="10.99.1.5", environment="測試")
    # F：DY 空殼，但 DY 是 40 天前 → 0、疑似下線
    f_id = hw(conn, asset_serial="DYN-F", hostname="foxtrot", ip="10.99.1.6")
    # G：退役 → 備查
    hw(conn, asset_serial="HW-G", hostname="golf", ip="10.99.1.7", asset_status="報廢")
    for hid, src, days in ((c_id, "dynassets", 5), (d_id, "vcenter", 5), (f_id, "dynassets", 40)):
        conn.execute(
            "INSERT INTO source_record (source, source_key, payload, collected_at, resolved_status, "
            "resolved_hardware_id) VALUES (?, ?, ?, ?, 'new', ?)",
            (src, f"k{hid}", '{"department": "資訊部"}', _ago(days), hid))
    # B 掃到；H、I：資產庫沒有、沒有主機名的兩台掃描機器（不可以被併成一台）
    for ip, hn in (("10.99.1.2", "bravo"), ("10.99.2.8", None), ("10.99.2.9", None)):
        conn.execute("INSERT INTO scan_history (scan_time, ip, hostname, scan_ok) VALUES (?, ?, ?, 1)",
                     (_ago(3), ip, hn))
    # J：資產庫、掃描都沒有，但帳號收集 1 天前收得到 → 已納管（CIA 無關）
    conn.execute("INSERT INTO host_account (ip, username, last_seen) VALUES ('10.99.3.10', 'root', ?)",
                 (_ago(1),))
    conn.commit()


def _conn(tmp):
    p = Path(tmp) / "t.db"
    db.init_db(p)
    conn = db.get_connection(p)
    _seed(conn)
    return conn, p


def _by_host(out):
    return {(m["hostname"] or m["ip"]): m for m in out["machines"]}


def test_存在可信度等級與計入(tmp_path):
    tmp = tmp_path
    conn, _ = _conn(tmp)
    m = _by_host(existence.compute(conn))
    assert (m["alpha"]["level"], m["alpha"]["score"], m["alpha"]["onboarded"]) == ("onboarded", 100, True)
    assert m["bravo"]["level"] == "collect_stale" and m["bravo"]["score"] == 70
    assert m["bravo"]["onboarded_stale"] and not m["bravo"]["onboarded"]
    assert m["charlie"]["score"] == 90 and m["charlie"]["pending_confirm"] and not m["charlie"]["has_cia"]
    assert m["delta"]["score"] == 90 and not m["delta"]["pending_confirm"]
    assert m["echo"]["level"] == "cia_only" and m["echo"]["suspect_offline"] and not m["echo"]["counted"]
    assert m["foxtrot"]["level"] == "none" and m["foxtrot"]["suspect_offline"]
    assert m["golf"]["retired"] and not m["golf"]["counted"] and not m["golf"]["suspect_offline"]
    assert m["10.99.3.10"]["onboarded"] and m["10.99.3.10"]["counted"] and not m["10.99.3.10"]["has_cia"]


def test_沒主機名的掃描機器各算一台(tmp_path):
    tmp = tmp_path
    conn, _ = _conn(tmp)
    m = _by_host(existence.compute(conn))
    assert "10.99.2.8" in m and "10.99.2.9" in m
    assert m["10.99.2.8"]["key"] != m["10.99.2.9"]["key"]
    s = ms.summarize(conn)
    # 首頁四態也不可以把兩台併成一台（兩台都是未登記）
    assert s["counts"][ms.UNREGISTERED] == 2


def test_分割與CIA完整率(tmp_path):
    tmp = tmp_path
    conn, _ = _conn(tmp)
    s = existence.compute(conn)["summary"]
    # 計入：A B C D H I J＝7；疑似下線：E F＝2；退役：G＝1
    assert (s["counted"], s["suspect_offline"], s["retired"]) == (7, 2, 1)
    assert s["counted"] + s["suspect_offline"] + s["retired"] == s["total_machines"] == 10
    assert (s["cia_registered"], s["cia_pending"]) == (2, 5)
    assert s["cia_complete_pct"] == round(2 * 100 / 7, 1) == 28.6
    assert s["pending_confirm"] == 1
    # 每個數字＝下鑽筆數
    for b in existence.BUCKETS:
        key = {"all": "total_machines"}.get(b, b)
        assert s[key] == existence.machines_in(conn, b)["total"], b
    assert len(existence.cia_pending_rows(conn)) == s["cia_pending"]


def test_時效_收集看天數_匯入掃描看最新快照(tmp_path):
    conn, _ = _conn(tmp_path)
    # 把「現在」往回撥 4 天：B 的收集（10 天前）變成 6 天前 → 已納管
    m = _by_host(existence.compute(conn, now=datetime.now() - timedelta(days=4)))
    assert m["bravo"]["onboarded"]
    # F 的 DY 紀錄（40 天前）不在 DY 最新一次匯入（5 天前那批）裡 → 不算 DY 證據
    assert m["foxtrot"]["suspect_offline"] and "inv:dynassets" not in m["foxtrot"]["evidence"]
    # 往前撥 40 天：DY／vCenter 最新一次匯入還是那一批 → C／D 照樣 90、計入（不看日曆天數）
    out = existence.compute(conn, now=datetime.now() + timedelta(days=40))
    m = _by_host(out)
    assert m["charlie"]["score"] == 90 and m["delta"]["score"] == 90 and m["charlie"]["counted"]
    assert not out["summary"]["stale_sources"], "40 天＜過期門檻 60 天，不警示"
    # 收集帳號仍看天數：B 的收集變成 50 天前 → 不再算收集證據（只剩掃描 70）
    assert "collect_stale" not in m["bravo"]["evidence"] and m["bravo"]["level"] == "scan"
    # 天數從設定檔來，不寫死：把 onboarded_days 改成 1 → A（2 天前）就不算已納管
    rules = existence.load_rules()
    rules["onboarded_days"] = 1
    m = _by_host(existence.compute(conn, rules=rules))
    assert not m["alpha"]["onboarded"] and m["alpha"]["onboarded_stale"]


def test_來源過期只警示不降級(tmp_path):
    conn, _ = _conn(tmp_path)
    out = existence.compute(conn, now=datetime.now() + timedelta(days=70))
    s, m = out["summary"], _by_host(out)
    stale = {x["source"]: x for x in s["stale_sources"]}
    assert set(stale) == {"dynassets", "vcenter", "scan"}
    assert stale["dynassets"]["age_days"] >= 70
    # 不降級：C／D 還是 90、計入；依據標「來源過期」
    assert m["charlie"]["score"] == 90 and m["charlie"]["counted"]
    assert m["charlie"]["stale_evidence"] == ["dynassets"]
    assert "vcenter" in m["delta"]["stale_evidence"]
    # 門檻從設定檔來
    rules = existence.load_rules()
    rules["source_stale_days"] = 200
    assert not existence.compute(conn, now=datetime.now() + timedelta(days=70), rules=rules)["summary"]["stale_sources"]


def test_掃描看所屬網段最近一次掃描(tmp_path):
    """有涵蓋紀錄就逐網段看：A 段最近一次是舊的那次（新的那次只掃了 B 段），A 段機器照樣算掃到；
    但同一段較舊一次活著、最近一次沒回應的，不算。"""
    p = tmp_path / "s.db"
    db.init_db(p)
    conn = db.get_connection(p)
    import scan_scope
    scan_scope._ensure(conn)
    t_old, t_mid, t_new = _ago(20), _ago(10), _ago(1)
    db.insert_hardware(conn, asset_serial="HW-K", hostname="kilo", ip="10.99.5.10")
    db.insert_hardware(conn, asset_serial="HW-L", hostname="lima", ip="10.99.5.11")
    db.insert_hardware(conn, asset_serial="HW-M", hostname="mike", ip="10.99.1.20")
    for t, ip in ((t_old, "10.99.5.11"), (t_mid, "10.99.5.10"), (t_new, "10.99.1.20")):
        conn.execute("INSERT INTO scan_history (scan_time, ip, scan_ok) VALUES (?, ?, 1)", (t, ip))
    for t, cidr in ((t_old, "10.99.5.0/24"), (t_mid, "10.99.5.0/24"), (t_new, "10.99.1.0/24")):
        conn.execute("INSERT INTO scan_coverage (scan_time, cidr, ok, created_at) VALUES (?, ?, 1, ?)", (t, cidr, t))
    conn.commit()
    m = _by_host(existence.compute(conn))
    assert m["kilo"]["level"] == "scan" and m["kilo"]["counted"]     # 10.99.5 段最近一次（t_mid）活著
    assert m["lima"]["level"] == "cia_only"                          # 只在較舊那次活著
    assert m["mike"]["level"] == "scan"
    conn.close()


def test_已納管_CIA無關且首頁同一個數(tmp_path):
    tmp = tmp_path
    conn, _ = _conn(tmp)
    s = ms.summarize(conn)
    ex = existence.compute(conn)["summary"]
    assert s["counts"][ms.ONBOARDED] == ex["onboarded"] == 2   # A、J
    reps = {(it["hostname"] or it["ip"]) for it in s["rep_items"][ms.ONBOARDED]}
    assert reps == {"alpha", "10.99.3.10"}
    # B：連上過但超過 7 天 → 已納管但失聯（不是已納管、也不是失聯）
    assert {it["hostname"] for it in s["rep_items"][ms.ONBOARDED_STALE]} == {"bravo"}
    # 3-8 的旗標讀同一份：J 不在資產庫也算已納管
    hosts = {h["ip"]: h for h in account_inventory.inventoried_hosts(conn)}
    assert hosts["10.99.3.10"]["onboarded"] is True
    assert hosts["10.99.3.10"]["cia_registered"] is False


def test_classify_已納管不看登記也不看掃描():
    assert ms.classify(False, True, 1, collect_fresh=True) == ms.ONBOARDED
    assert ms.classify(True, False, 1, collect_fresh=True) == ms.ONBOARDED
    assert ms.classify(True, True, 1, collect_fresh=False) == ms.ONBOARDED_STALE
    assert ms.classify(True, False, 1, collect_fresh=False) == ms.ONBOARDED_STALE
    assert ms.classify(True, False, 0) == ms.LOST
    # 退役：7 天內收得到才算「退役但仍在線」
    assert ms.classify(True, False, 1, retired=True, collect_fresh=False) == ms.RETIRED
    assert ms.classify(True, False, 1, retired=True, collect_fresh=True) == ms.RETIRED_ALIVE


def _client(tmp):
    conn, p = _conn(tmp)
    conn.close()

    def _override_get_db():
        c = db.get_connection(p)
        try:
            yield c
        finally:
            c.close()

    api.app.dependency_overrides[api.get_db] = _override_get_db
    client = TestClient(api.app)
    c = db.get_connection(p)
    db.create_user(c, "tester", auth.hash_password(_STUB_CREDENTIAL))
    c.commit()
    c.close()
    assert client.post("/api/auth/login",
                       json={"username": "tester", "password": _STUB_CREDENTIAL}).status_code == 200
    return client


def test_API_加總等於下鑽_與匯出(tmp_path):
    tmp = tmp_path
    client = _client(tmp)
    try:
        s = client.get("/api/existence/summary").json()["summary"]
        for b in ("counted", "cia_pending", "suspect_offline", "onboarded"):
            r = client.get("/api/existence/machines", params={"bucket": b}).json()
            assert r["total"] == len(r["items"]) == s[b], b
        assert client.get("/api/existence/machines", params={"bucket": "亂寫"}).status_code == 400
        comp = client.get("/api/dashboard/composition").json()
        assert comp["existence"]["counted"] == s["counted"]
        x = client.get("/api/existence/cia-pending/export")
        assert x.status_code == 200 and x.content[:2] == b"PK"
        import io
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(x.content)).active
        assert ws.max_row - 1 == s["cia_pending"]
        a = client.get("/api/existence/asset/HW-A").json()
        assert a["score"] == 100 and a["onboarded"]
        assert client.get("/api/existence/asset/NOPE").status_code == 404
    finally:
        api.app.dependency_overrides.clear()


def test_月報總台數跟首頁同一份(tmp_path):
    import monthly_report

    tmp = tmp_path
    conn, _ = _conn(tmp)
    sec = next(x["data"] for x in monthly_report.collect(conn)["sections"] if x["name"] == "資產總覽")
    s = existence.compute(conn)["summary"]
    assert sec["總台數"] == s["counted"] == 7
    assert sec["CIA 完整率"] == "28.6%"
    assert sec["CIA 有登記"] + sec["CIA 待補登"] == sec["總台數"]
    assert sum(sec["依環境"].values()) == sec["總台數"]
    assert sec["依環境"]["(CIA 待補登)"] == 5   # 業務欄位只從 CIA 來


def test_2_1_已納管主機表_台數不濾掉沒IP且頂端有跳轉鈕():
    """2-1 /adopt「已納管主機」表跟首頁同一份 rep_items（state=已納管）；不可以濾掉沒 IP 的
    （濾了台數就比首頁少）。頂端頁簽「已納管主機 N」（v1.459.1 前是跳轉鈕）
    （使用者 2026-10-18 在公司機找不到全選＋勾選的一次重新納管）。"""
    src = (ROOT / "APP/asset-module/frontend/pages/adopt.vue").read_text(encoding="utf-8")
    assert "params: { state: '已納管' }" in src
    assert "managed.value = r.items || []" in src
    assert ".filter((x) => x.ip)" not in src
    # v1.459.1 起改成頁簽：「已納管主機 N」是第三個頁簽，舊錨點 #managed-hosts 視同 #onboarded
    assert 'id="managed-hosts"' in src
    assert "t === 'managed-hosts'" in src
    assert "已納管主機 {{ managed.length }}" in src and "setTab('onboarded')" in src
    # 已納管但失聯也列在表上、可勾選重新納管，但不算進「已納管」台數（managed 只放 state=已納管）
    assert "params: { state: '已納管但失聯' }" in src
    assert "managedAll.value.filter" in src
