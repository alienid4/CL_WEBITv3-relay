"""W221-009（Codex 53）：所有登記來源要有原始位置／位置分組／機櫃機房鍵／資料時間；純讀、原欄位不變、匯出含新欄位。假庫。"""
import io
import sqlite3
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import blast_radius  # noqa: E402
import db  # noqa: E402
import manage_state as ms  # noqa: E402
import v2_sections as vs  # noqa: E402

OLD = ["serial", "source", "hostname", "ip", "os", "status", "api_id", "rack"]
NEW = ["raw_location", "location_group", "rack_room", "source_updated_at", "manual_updated_at"]
CASES = [("S1", "h-ban", "01_板橋機房", "A1/10U"), ("S2", "h-branch", "分公司-台中", "B2"),
         ("S3", "h-unk", "某某不明館", None), ("S4", "h-empty", "", None)]


def _mk(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    c = db.get_connection(p)
    for i, (s, h, loc, rack) in enumerate(CASES):
        db.insert_hardware(c, asset_serial=s, hostname=h, ip=f"10.99.8.{i + 1}", physical_location=loc, rack_no=rack,
                           device_model="Dell R750", api_id="N-920", asset_name="x")
    c.execute("UPDATE hardware SET manual_updated_at = '2026-10-01 09:00:00' WHERE asset_serial = 'S1'")
    c.commit()
    c.close()
    return p


def _ro(p):
    c = sqlite3.connect(f"file:{p.as_posix()}?mode=ro", uri=True)
    c.row_factory = sqlite3.Row
    return c


def test_欄位契約_原值_正規化_分點_時間(tmp_path):
    p = _mk(tmp_path)
    c = _ro(p)
    try:
        for s, h, loc, _ in CASES:
            sec = vs.section(c, s, "sources")
            r = sec["rows"][0]
            assert r["raw_location"] == (loc or None) or r["raw_location"] == loc          # 原值不改寫
            assert r["location_group"] == ms.group_location(loc)
            assert r["rack_room"] == blast_radius.rack_room(loc)
            assert r["source_updated_at"]                                              # hardware.updated_at
        r1 = vs.section(c, "S1", "sources")["rows"][0]
        assert r1["raw_location"] == "01_板橋機房" and r1["manual_updated_at"] == "2026-10-01 09:00:00"
        assert vs.section(c, "S2", "sources")["rows"][0]["rack_room"] == "分公司-台中"      # 分點保留
        assert vs.section(c, "S3", "sources")["rows"][0]["location_group"] == "（未歸類：某某不明館）"
        assert vs.section(c, "S4", "sources")["rows"][0]["location_group"] == "未填"
        assert vs.section(c, "S2", "sources")["rows"][0]["manual_updated_at"] is None
    finally:
        c.close()


def test_原欄位與順序不變_新欄位只新增(tmp_path):
    p = _mk(tmp_path)
    c = _ro(p)
    try:
        sec = vs.section(c, "S1", "sources")
        keys = [x["k"] for x in sec["cols"]]
        assert keys[:len(OLD)] == OLD and keys[len(OLD):] == NEW
        assert set(OLD + NEW) <= set(sec["rows"][0])
        assert sec["rows"][0]["serial"] == "S1" and sec["rows"][0]["source"] == "CIA" and sec["rows"][0]["rack"] == "A1/10U"
    finally:
        c.close()


def test_純讀_唯讀連線mode_ro與query_only都成功且庫不變(tmp_path):
    p = _mk(tmp_path)
    before = p.read_bytes()
    c = _ro(p)
    try:
        c.execute("PRAGMA query_only=ON")
        sec = vs.section(c, "S1", "sources")
        assert sec["count"] == 1 and sec["unknown"] is None
        vs.all_sections(c, "S1")            # 整張病歷表載入也不報寫入錯（s_location 的 ensure_* 不在本段範圍，錯誤會被 _run 吞成 unknown）
    finally:
        c.close()
    assert p.read_bytes() == before


def test_匯出含新欄位(tmp_path):
    from openpyxl import load_workbook
    p = _mk(tmp_path)
    c = _ro(p)
    try:
        buf, _ = vs.export_xlsx(c, "S1", "sources")
    finally:
        c.close()
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    texts = [str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    assert any("原始位置" in t for t in texts) and any("01_板橋機房" in t for t in texts) and any("2026-10-01 09:00:00" in t for t in texts)
