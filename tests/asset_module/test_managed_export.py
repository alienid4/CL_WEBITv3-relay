"""2-1「已納管主機」匯出 Excel（2026-10-20）：篩選規則跟畫面一字不差、項次跟畫面同號、匯出篩完的全部列。"""
import io
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import managed_export as me  # noqa: E402

ROWS = [
    {"no": 1, "ip": "10.99.0.1", "hostname": "web-a", "asset_serial": "SN-1", "os_type": "Linux"},
    {"no": 2, "ip": "10.99.0.2", "hostname": "WIN-B", "asset_serial": "SN-2", "os_type": "Windows"},
    {"no": 3, "ip": "10.99.0.3", "hostname": None, "asset_serial": None, "os_type": None},
]
TAGS = {"SN-1": ["遷移"], "SN-2": ["遷移", "核心"]}


def test_篩選規則跟畫面一樣():
    assert [r["no"] for r in me.filter_rows(ROWS, TAGS)] == [1, 2, 3]
    assert [r["no"] for r in me.filter_rows(ROWS, TAGS, os="Windows")] == [2]
    assert [r["no"] for r in me.filter_rows(ROWS, TAGS, os="未填")] == [3], "空 OS 算「未填」"
    assert [r["no"] for r in me.filter_rows(ROWS, TAGS, tag="核心")] == [2]
    assert [r["no"] for r in me.filter_rows(ROWS, TAGS, kw="win")] == [2], "關鍵字不分大小寫"
    assert [r["no"] for r in me.filter_rows(ROWS, TAGS, kw=" sn-1 ")] == [1]
    assert [r["no"] for r in me.filter_rows(ROWS, TAGS, kw="10.99.0.3")] == [3]


def test_前端篩選也是同一套欄位():
    """畫面（adopt.vue managedShown）改了篩選欄位，這裡要跟著改——不然匯出跟畫面對不上。"""
    vue = (ROOT / "APP" / "asset-module" / "frontend" / "pages" / "adopt.vue").read_text(encoding="utf-8")
    assert "[m.ip, m.hostname, m.asset_serial].some" in vue
    assert "(m.os_type || '未填') !== os" in vue


def test_匯出檔_表頭凍結_有自動篩選_項次照畫面(monkeypatch):
    from openpyxl import load_workbook

    import db

    monkeypatch.setattr(me, "base_rows", lambda conn: [dict(r) for r in ROWS])
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "t.db"
        db.init_db(p)
        conn = db.get_connection(p)
        db.insert_hardware(conn, asset_serial="SN-2", ip="10.99.0.2", environment="正式",
                           asset_status="使用中", physical_location="機房A", os="Windows Server 2019")
        data = me.workbook_bytes(conn, os="Windows")
        conn.close()
    ws = load_workbook(io.BytesIO(data)).active
    head = [c.value for c in ws[1]]
    assert head[:8] == ["項次", "IP", "主機名稱", "OS", "資產序號", "標籤", "最後確認收得到", "上次失敗"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    r = dict(zip(head, rows[0]))
    assert r["項次"] == 2 and r["IP"] == "10.99.0.2" and r["機房"] == "機房A"
    assert r["納管帳號"] == "WinRM 憑證", "Windows 不列帳號（也絕不列密碼）"
    assert r["硬體規格（最近結果）"] == "尚未收集"
