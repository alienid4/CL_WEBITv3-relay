"""位置更正（現場盤點 U10 其實在 U11）：套用、稽核、重疊警告、批次、CIA 匯入不蓋掉、CIA 同步後自動退場、匯出。"""
import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import rack_layout as rl  # noqa: E402


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    return db.get_connection(p)


def _seed(conn):
    for serial, host, ip, rack in (("Q-1", "q1", "10.99.6.1", "K1/10U"), ("Q-2", "q2", "10.99.6.2", "K1/11-12U"),
                                   ("Q-3", "q3", "10.99.6.3", "網路櫃")):
        db.insert_hardware(conn, asset_serial=serial, hostname=host, ip=ip, physical_location="01_甲機房",
                           rack_no=rack, api_id="N-930", asset_name="位置系統")
    conn.commit()


def _m(conn, serial):
    return next(m for m in rl.machines(conn) if m["asset_serial"] == serial)


def test_更正位置_套用_重疊警告_稽核(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        k1 = _m(conn, "Q-1")["key"]
        r = rl.set_position(conn, k1, "甲機房", "K1", 11, 1, "10/23 現場盤點", "tester")
        assert [c["asset_serial"] for c in r["collisions"]] == ["Q-2"]           # 只警告、照樣存
        m = _m(conn, "Q-1")
        assert (m["rack"], m["u_start"], m["u_end"]) == ("K1", 11, 11) and m["position_corrected"]
        assert m["cia_position"]["u_start"] == 10                                  # 舊 CIA 值留著給 tooltip
        audit = rl.position_audit(conn, k1)
        assert audit[0]["action"] == "set" and audit[0]["note"] == "10/23 現場盤點"
        with pytest.raises(ValueError):
            rl.set_position(conn, k1, "甲機房", "K1", 11, 1, "  ", "tester")       # 備註必填
        with pytest.raises(ValueError):
            rl.set_position(conn, k1, "甲機房", "", 11, 1, "x", "tester")
    finally:
        conn.close()


def test_認不出的機櫃也能更正上去_CIA匯入不蓋掉_CIA同步後退場(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        k3 = _m(conn, "Q-3")["key"]
        rl.set_position(conn, k3, "甲機房", "K2", 20, 2, "現場盤點", "tester")
        m = _m(conn, "Q-3")
        assert m["kind"] == "rack" and (m["rack"], m["u_start"], m["u_end"]) == ("K2", 20, 21)
        conn.execute("UPDATE hardware SET rack_no = 'K1/30U' WHERE asset_serial = 'Q-3'")    # CIA 重匯另一個值
        conn.commit()
        assert _m(conn, "Q-3")["rack"] == "K2"                                       # 更正照樣套用
        assert [r["資產序號"] for r in rl.pending_cia_rows(conn)] == ["Q-3"]
        conn.execute("UPDATE hardware SET rack_no = 'K2/20-21U' WHERE asset_serial = 'Q-3'")  # CIA 改成跟現場一樣
        conn.commit()
        m = _m(conn, "Q-3")
        assert not m["position_corrected"] and m["rack"] == "K2"
        st = conn.execute("SELECT status FROM rack_position_override WHERE key = ?", (k3,)).fetchone()[0]
        assert st == "synced"
        assert rl.pending_cia_rows(conn) == []
        assert rl.position_audit(conn, k3)[0]["action"] == "cia_synced"
    finally:
        conn.close()


def test_批次更正_清除_匯出xlsx(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        ks = {s: _m(conn, s)["key"] for s in ("Q-1", "Q-2")}
        res = rl.set_positions_bulk(conn, [
            {"key": ks["Q-1"], "room": "甲機房", "rack": "K1", "u_start": 20, "u_height": 1},
            {"key": ks["Q-2"], "room": "甲機房", "rack": "K1", "u_start": 21, "u_height": 2},
        ], "盤點", "tester")
        assert res["count"] == 2 and all(not r["collisions"] for r in res["results"])
        assert [r["新U位"] for r in rl.pending_cia_rows(conn)] == ["20U", "21-22U"]
        buf, n = rl.pending_cia_xlsx(conn)
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(buf.read())).active
        assert n == 2 and [c.value for c in ws[1]][:3] == ["資產序號", "主機名稱", "原機房"]
        rl.clear_position(conn, ks["Q-1"], "打錯", "tester")
        assert _m(conn, "Q-1")["u_start"] == 10
        with pytest.raises(ValueError):
            rl.clear_position(conn, ks["Q-1"], None, "tester")
        d = rl.datacenter(conn)                                                      # 每台只在一個地方
        keys = [k for r in d["rooms"] for t in r["racks"] for k in t["keys"]]
        assert len(keys) == len(set(keys))
    finally:
        conn.close()
