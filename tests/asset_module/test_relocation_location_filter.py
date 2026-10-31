"""機房搬遷匯出的機房篩選（Codex 網站審查 W221-017A／017B，2026-10-31）。

A：下拉的「（未填機房）」是顯示字串，不能拿去 LIKE——要轉成 NULL／空白條件。
B：「CIA資產基準」分頁刻意維持全部資產（原始登記快照），不受機房篩選；前端說明要講明。
"""
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import relocation_export as rx  # noqa: E402


def _seed(tmp_path):
    p = tmp_path / "reloc.db"
    db.init_db(p)
    c = db.get_connection(p)
    rows = [
        ("T-N1", "n1", "10.99.5.1", None),            # 地點 NULL
        ("T-E1", "e1", "10.99.5.2", ""),               # 地點空字串
        ("T-S1", "s1", "10.99.5.3", "   "),            # 地點只有空白
        ("T-A1", "a1", "10.99.5.4", "板橋機房"),
        ("T-A2", "a2", "10.99.5.5", "板橋機房二館"),    # 相似子地點（包含比對會一起命中）
        ("T-B1", "b1", "10.99.5.6", "松山機房"),
    ]
    for serial, host, ip, loc in rows:
        c.execute("INSERT INTO hardware (asset_serial, hostname, ip, physical_location, device_model, is_vm) "
                  "VALUES (?,?,?,?,?,0)", (serial, host, ip, loc, "SRV-X"))
    c.commit()
    return c


def _serials(rs):
    return sorted(r["asset_serial"] for r in rs)


def test_未填機房選項_匹配NULL_空字串_空白(tmp_path):
    c = _seed(tmp_path)
    assert _serials(rx._all_rows(c, rx.NO_LOCATION)) == ["T-E1", "T-N1", "T-S1"]
    assert _serials(rx._rows(c, rx.NO_LOCATION)) == ["T-E1", "T-N1", "T-S1"]


def test_一般機房是包含比對_空值是全部(tmp_path):
    c = _seed(tmp_path)
    assert _serials(rx._all_rows(c, "板橋機房")) == ["T-A1", "T-A2"]       # 包含：二館也命中（已在前端說明）
    assert _serials(rx._all_rows(c, "松山機房")) == ["T-B1"]
    assert len(rx._all_rows(c, None)) == 6 and len(rx._all_rows(c, "")) == 6


def test_下拉計數與篩選同一定義(tmp_path):
    c = _seed(tmp_path)
    locs = {l["loc"]: l["n"] for l in rx.locations(c)}
    assert locs[rx.NO_LOCATION] == 3                                          # NULL／空字串／空白併成同一格
    assert locs[rx.NO_LOCATION] == len(rx._all_rows(c, rx.NO_LOCATION))       # 選了這格匯出的台數＝下拉顯示的台數


def test_基準分頁刻意全部資產_不受機房篩選(tmp_path):
    from openpyxl import load_workbook
    c = _seed(tmp_path)
    buf, n = rx.build(c, "松山機房")
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    base = wb["CIA資產基準"]
    assert n == 1                                                              # 匯出主表只有該機房 1 台
    assert base.max_row - 1 == 6                                               # 基準分頁仍是全部 6 台


def test_前端說明講明基準分頁範圍():
    s = (ROOT / "APP" / "asset-module" / "frontend" / "pages" / "relocation.vue").read_text(encoding="utf-8")
    assert "CIA資產基準" in s and "刻意維持全部資產" in s and "（未填機房）" in s
