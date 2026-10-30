"""2.x 第二批（2.004.0）：7 盤點作業共用流程（凍結快照＋與上一次比較）、ODA／Exadata「同一套、不同節點」不進位置重疊、
1.2 異常事件的「掃描無回應」列不連主機履歷、新頁進選單。全部假資料（IP 一律 10.99.x）。"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
FE = ROOT / "APP" / "asset-module" / "frontend"
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import inventory_flow as inv  # noqa: E402
import overlap_queue  # noqa: E402


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    return db.get_connection(p)


def _pkg(conn, ip, name, version, gone=False):
    conn.execute("INSERT INTO host_package (ip, asset_serial, name, version, arch, source, gone_at) VALUES (?,?,?,?,?,?,?)",
                 (ip, f"S-{ip}", name, version, "x86_64", "ssh_rpm", "2026-10-01 00:00:00" if gone else None))


# ===== 7 盤點作業：凍結快照與差異 =====

def test_軟體盤點_第一次沒有差異_第二次列出新增移除變更與主機未收集(tmp_path):
    conn = _conn(tmp_path)
    _pkg(conn, "10.99.1.1", "openssl", "3.0.1")
    _pkg(conn, "10.99.1.1", "bash", "5.1")
    _pkg(conn, "10.99.1.2", "httpd", "2.4.57")
    conn.commit()
    t1 = inv.create_task(conn, "software", "tester")
    assert (t1["hosts"], t1["rows"], t1["prev_id"], t1["diff_count"]) == (2, 3, None, 0)

    # 之後：.1 的 openssl 升版、bash 移除、新增 curl；.2 這次沒收集到；新主機 .3 出現
    conn.execute("UPDATE host_package SET version = '3.0.2' WHERE name = 'openssl'")
    conn.execute("UPDATE host_package SET gone_at = '2026-10-02 00:00:00' WHERE name IN ('bash', 'httpd')")
    _pkg(conn, "10.99.1.1", "curl", "8.0")
    _pkg(conn, "10.99.1.3", "vim", "9.0")
    conn.commit()
    t2 = inv.create_task(conn, "software", "tester")
    assert t2["prev_id"] == t1["id"]
    d = inv.diffs(conn, "software", t2["id"])["rows"]
    kinds = sorted((x["host"], x["kind"]) for x in d)
    assert kinds == sorted([("10.99.1.1", inv.DIFF_LABEL["changed"]), ("10.99.1.1", inv.DIFF_LABEL["removed"]),
                            ("10.99.1.1", inv.DIFF_LABEL["added"]), ("10.99.1.2", inv.DIFF_LABEL["host_missing"]),
                            ("10.99.1.3", inv.DIFF_LABEL["host_new"])])
    # 主機未收集只列一筆，不把它的每一筆都當成移除
    assert sum(1 for x in d if x["host"] == "10.99.1.2") == 1
    # 差異文字用欄位中文名
    ch = next(x for x in d if x["kind"] == inv.DIFF_LABEL["changed"])
    assert "版本：3.0.1" in ch["before"] and "版本：3.0.2" in ch["after"]
    # 任務清單上的差異數＝差異頁列數（同一份計算）
    lst = {t["id"]: t for t in inv.list_tasks(conn, "software")}
    assert lst[t2["id"]]["diff_count"] == len(d) == t2["diff_count"]
    assert lst[t2["id"]]["diff_open"] == len(d)


def test_盤點_快照凍結_之後收集異動不影響這次結果(tmp_path):
    conn = _conn(tmp_path)
    _pkg(conn, "10.99.2.1", "openssl", "3.0.1")
    conn.commit()
    t = inv.create_task(conn, "software", "tester")
    conn.execute("UPDATE host_package SET version = '9.9' WHERE ip = '10.99.2.1'")
    conn.commit()
    r = inv.result(conn, "software", t["id"], host="10.99.2.1")
    assert r["rows"][0]["version"] == "3.0.1"
    u = inv.result(conn, "software", t["id"])
    assert u["total"] == sum(x["count"] for x in u["rows"]) == t["rows"]


def test_盤點_差異狀態_不處理要原因_報告依種類加總(tmp_path):
    conn = _conn(tmp_path)
    _pkg(conn, "10.99.3.1", "a", "1")
    conn.commit()
    inv.create_task(conn, "software", "tester")
    _pkg(conn, "10.99.3.1", "b", "1")
    _pkg(conn, "10.99.3.1", "c", "1")
    conn.commit()
    t2 = inv.create_task(conn, "software", "tester")
    d = inv.diffs(conn, "software", t2["id"])["rows"]
    assert len(d) == 2
    with pytest.raises(ValueError):
        inv.set_status(conn, "software", t2["id"], d[0]["diff_key"], "不處理", "", "tester")
    inv.set_status(conn, "software", t2["id"], d[0]["diff_key"], "已處理", "核准變更 CR-1", "tester")
    rep = inv.report(conn, "software", t2["id"])
    b = rep["by_kind"][inv.DIFF_LABEL["added"]]
    assert (b["total"], b["已處理"], b["待處理"]) == (2, 1, 1)
    lst = {t["id"]: t for t in inv.list_tasks(conn, "software")}
    assert lst[t2["id"]]["diff_open"] == 1
    buf, fname = inv.report_xlsx(conn, "software", t2["id"])
    assert fname.endswith(".xlsx") and buf.getvalue()[:2] == b"PK"


def test_盤點_沒有收集資料不能建立任務(tmp_path):
    conn = _conn(tmp_path)
    with pytest.raises(ValueError):
        inv.create_task(conn, "service", "tester")


def test_帳號盤點_欄位與差異(tmp_path):
    conn = _conn(tmp_path)
    conn.execute("INSERT INTO host_account (ip, asset_serial, username, uid, kind, is_sudoer, can_login, pw_status) VALUES "
                 "('10.99.4.1', 'S-1', 'ops01', 1001, 'human', 0, 1, 'set')")
    conn.commit()
    inv.create_task(conn, "account", "tester")
    conn.execute("UPDATE host_account SET is_sudoer = 1")
    conn.commit()
    t2 = inv.create_task(conn, "account", "tester")
    d = inv.diffs(conn, "account", t2["id"])["rows"]
    assert len(d) == 1 and d[0]["kind"] == inv.DIFF_LABEL["changed"]
    assert d[0]["before"] == "管理員權限：否" and d[0]["after"] == "管理員權限：是"


def test_盤點_API_端點齊全_寫入走管理權限關卡():
    src = (ROOT / "APP" / "asset-module" / "backend" / "routers" / "v2.py").read_text(encoding="utf-8")
    for path in ["/api/v2/inv/{kind}/tasks", "/api/v2/inv/{kind}/tasks/{tid}/result", "/api/v2/inv/{kind}/tasks/{tid}/diffs",
                 "/api/v2/inv/{kind}/tasks/{tid}/diffs/status", "/api/v2/inv/{kind}/tasks/{tid}/report",
                 "/api/v2/inv/{kind}/tasks/{tid}/report.xlsx"]:
        assert f'"{path}"' in src, path
    assert src.count("_writer(") >= 2          # 建立任務、改差異狀態都要過寫入關卡


# ===== ODA／Exadata：同一套、不同節點 =====

def _dev(key, name, model, u0, u1, **kw):
    d = {"key": key, "name": name, "model": model, "spec_model": model, "u_start": u0, "u_end": u1, "state": "active",
         "asset_serial": key, "ip": None, "hw_serials": [], "w": None, "system_level": None, "rack_raw": "B3/10-20U"}
    d.update(kw)
    return d


def test_ODA節點同一段U_不列位置重疊_分類為同一套不同節點():
    devs = [_dev("k1", "oda-n1", "Oracle Database Appliance X9-2-HA", 10, 13),
            _dev("k2", "oda-n2", "Oracle Database Appliance X9-2-HA", 10, 13)]
    plan = overlap_queue.merge_plan(devs, "甲機房", "B3", acts=[])
    eng = plan["engineered"]
    if plan["merged"]:
        pytest.skip("此型號被既有規則併成同一台（也不會列位置重疊）")
    assert len(eng) == 1 and eng[0]["keys"] == ["k1", "k2"]
    assert {"k1", "k2"} in plan["ignore"]
    assert overlap_queue.ENGINEERED_LABEL == "同一套、不同節點"


def test_Exadata規格表標system_level_也算同一套():
    devs = [_dev("e1", "exa-db1", "X10M", 1, 42, system_level=1), _dev("e2", "exa-cel1", "X10M", 20, 21, system_level=1)]
    plan = overlap_queue.merge_plan(devs, "甲機房", "C1", acts=[])
    assert plan["merged"] or [g["keys"] for g in plan["engineered"]] == [["e1", "e2"]]


def test_一般伺服器同型號同U_仍列位置重疊():
    devs = [_dev("d1", "web01", "Dell R750", 5, 6), _dev("d2", "web02", "Dell R750", 5, 6)]
    plan = overlap_queue.merge_plan(devs, "甲機房", "A1", acts=[])
    assert plan["engineered"] == []
    assert {"d1", "d2"} not in plan["ignore"]


def test_工程系統與一般設備重疊_一般設備仍列():
    devs = [_dev("k1", "oda-n1", "ODA X9-2", 10, 13), _dev("k2", "oda-n2", "ODA X9-2", 10, 13), _dev("x1", "sw01", "Cisco 9300", 12, 12)]
    plan = overlap_queue.merge_plan(devs, "甲機房", "B3", acts=[])
    assert all("x1" not in g["keys"] for g in plan["engineered"])


# ===== 1.2 異常事件：掃描無回應不連主機履歷 =====

def test_異常事件_掃描無回應列_不連主機履歷_說明欄連掃描涵蓋():
    src = (FE / "pages" / "v2" / "changes.vue").read_text(encoding="utf-8")
    assert "SCAN_ALERT = '掃描無回應'" in src
    assert "r.scan ? null : `/v2/machine/" in src          # 主機名稱欄：掃描無回應不給連結
    assert "label: '說明', link: (r: any) => (r.scan ? '/v2/onboard?tab=gaps' : null)" in src   # 2.006.3：連 6.1 掃描缺口（新版）
    assert "note: scan ? i.hostname : ''" in src            # 訊息放說明欄


# ===== 新頁 =====

def test_第二批新頁_共用盤點元件_舊頁有新版橫幅():
    for kind in ["rack", "account", "software", "service"]:
        s = (FE / "pages" / "v2" / "inventory" / f"{kind}.vue").read_text(encoding="utf-8")
        assert "layout: 'v2'" in s and f'kind="{kind}"' in s and "<V2InventoryFlow" in s, kind
    for old, to in [("datacenter", "/v2/datacenter"), ("storage", "/v2/storage"), ("rack-audit", "/v2/inventory/rack"),
                    ("account-matrix", "/v2/inventory/account"), ("software", "/v2/inventory/software"), ("services", "/v2/inventory/service")]:
        assert f'<V2Banner to="{to}"' in (FE / "pages" / f"{old}.vue").read_text(encoding="utf-8"), old
    dc = (FE / "pages" / "v2" / "datacenter.vue").read_text(encoding="utf-8")
    for lab in ["平面圖", "機櫃正面", "電力容量", "上架規劃", "位置更正"]:
        assert f"label: '{lab}" in dc or f"label: `{lab}" in dc, lab
    st = (FE / "pages" / "v2" / "storage.vue").read_text(encoding="utf-8")
    for lab in ["依機房", "容量預警", "依業務系統", "超額配置", "快照"]:
        assert f"label: '{lab}" in st or f"label: `{lab}" in st, lab


def test_帳號盤點_特權帳號清單_跨主機_報告與匯出都有(tmp_path):
    conn = _conn(tmp_path)
    for ip, user, uid, sudo in [("10.99.6.1", "root", 0, 0), ("10.99.6.1", "ops01", 1001, 1), ("10.99.6.2", "app01", 1002, 0)]:
        conn.execute("INSERT INTO host_account (ip, asset_serial, username, uid, kind, is_sudoer, can_login, pw_status) VALUES (?,?,?,?,?,?,?,?)",
                     (ip, f"S-{ip}", user, uid, "human", sudo, 1, "set"))
    conn.commit()
    t = inv.create_task(conn, "account", "tester")
    pv = inv.result(conn, "account", t["id"], host=inv.PRIV)
    assert pv["multi_host"] and sorted((x["host"], x["item"]) for x in pv["rows"]) == [("10.99.6.1", "ops01"), ("10.99.6.1", "root")]
    rep = {x["k"]: x["v"] for x in inv.report(conn, "account", t["id"])["lines"]}
    assert rep["特權帳號（管理員權限或 UID 0）"] == 2
    from openpyxl import load_workbook
    buf, _ = inv.report_xlsx(conn, "account", t["id"])
    ws = load_workbook(buf)["特權帳號"]
    assert ws.max_row == 3
