"""匯出防公式注入（2.007.8，審查 B1）：資產資料來自外部（CIA 匯入、掃描、收集），主機名稱或備註裡的
「=HYPERLINK(...)」匯出成 xlsx／CSV 不可以變成公式。"""
import csv
import io
import sys
import tempfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import api  # noqa: E402
import db  # noqa: E402
import export_safe  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

EVIL = '=HYPERLINK("http://example.invalid/x","click")'


@pytest.fixture()
def client():
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "t.db"
        db.init_db(p)

        def _get_db():
            c = db.get_connection(p)
            try:
                yield c
            finally:
                c.close()
        api.app.dependency_overrides[api.get_db] = _get_db
        api.app.dependency_overrides[api.require_auth] = lambda: {"username": "tester", "role": "admin"}
        try:
            yield TestClient(api.app)
        finally:
            api.app.dependency_overrides.pop(api.get_db, None)
            api.app.dependency_overrides.pop(api.require_auth, None)


def test_表格匯出_公式字串回來是純文字(client):
    r = client.post("/api/export/table-xlsx", json={
        "title": "2.1 主機查詢[全部]:a/b*?", "headers": ["項次", "=SUM(A1:A9)"],
        "rows": [["H1", EVIL], ["H2", "+1+1"], ["H3", "-2+3"], ["H4", "@cmd"], ["H5", "正常"], ["H6", -5]]})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "2.1 主機查詢全部ab"[:31]                 # [ ] \ : * ? / 拿掉
    cells = {c.coordinate: c for row in ws.iter_rows() for c in row}
    assert cells["B1"].value == "'=SUM(A1:A9)" and cells["B1"].data_type == "s"
    assert cells["B2"].value == "'" + EVIL and cells["B2"].data_type == "s"
    assert [cells[f"B{i}"].value for i in (3, 4, 5, 6, 7)] == ["'+1+1", "'-2+3", "'@cmd", "正常", -5]
    assert not any(c.data_type == "f" for c in cells.values())


def test_活頁簿掃描_公式改文字():
    wb = Workbook()
    ws = wb.active
    ws.append([EVIL, "ok"])
    assert ws["A1"].data_type == "f"
    export_safe.neutralize_workbook(wb)
    assert ws["A1"].data_type == "s" and ws["A1"].value.startswith("'=")


def test_事件聯絡名單CSV_公式字元開頭加單引號(monkeypatch):
    import v2_incident

    monkeypatch.setattr(v2_incident, "notify", lambda conn, eid: {"contacts": [
        {"kind": "業務系統", "ref": "N-1", "name": EVIL, "level": None, "machines": 1, "oncall_phone": "+886-2-0000",
         "people": "@x", "machine_names": "a"}]})
    text = v2_incident.contacts_csv(None, 1)
    row = list(csv.reader(io.StringIO(text.lstrip("﻿"))))[1]
    assert row[3] == "'" + EVIL and row[6].startswith("'+") and row[7] == "'@x"
