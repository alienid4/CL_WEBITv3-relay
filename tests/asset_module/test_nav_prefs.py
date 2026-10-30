"""選單重整：套用新分類先備份（最近 5 版）、還原上一版、我的最愛每人一份、新舊編號對照匯出。"""
import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import api  # noqa: E402
import db  # noqa: E402
import nav_prefs  # noqa: E402
from fastapi import HTTPException  # noqa: E402

ADMIN = {"username": "admin"}
USER = {"username": "alice"}
LAYOUT = [{"id": "g0", "label": "總覽", "icon": "◈", "to": "/", "items": ["/issues"]}]


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    return db.get_connection(p)


def test_套用新分類_先備份_還原上一版(tmp_path):
    c = _conn(tmp_path)
    api.set_nav_layout(api.NavLayoutBody(layout=LAYOUT), session=ADMIN, conn=c)
    r = api.nav_layout_apply_default(session=ADMIN, conn=c)
    assert r["backed_up"] is True
    assert api.get_nav_layout(session=ADMIN, conn=c)["layout"] is None          # 改用內建新分類
    bk = api.nav_layout_backups(session=ADMIN, conn=c)
    assert len(bk) == 1 and bk[0]["groups"] == 1 and bk[0]["by"] == "admin"
    r = api.nav_layout_restore(session=ADMIN, conn=c)
    assert r["restored_default"] is False
    assert api.get_nav_layout(session=ADMIN, conn=c)["layout"] == LAYOUT
    api.nav_layout_restore(session=ADMIN, conn=c)                                # 再按一次＝回到還原前
    assert api.get_nav_layout(session=ADMIN, conn=c)["layout"] is None


def test_備份只留5版_非管理員不能改(tmp_path):
    c = _conn(tmp_path)
    for i in range(7):
        api.set_nav_layout(api.NavLayoutBody(layout=[{**LAYOUT[0], "label": f"版{i}"}]), session=ADMIN, conn=c)
        api.nav_layout_apply_default(session=ADMIN, conn=c)
    bk = nav_prefs.backups(c)
    assert len(bk) == 5 and bk[0]["layout"][0]["label"] == "版6"
    with pytest.raises(HTTPException):
        api.nav_layout_apply_default(session=USER, conn=c)
    with pytest.raises(HTTPException):
        api.nav_layout_restore(session=USER, conn=c)
    c2 = _conn(tmp_path / "b")
    with pytest.raises(HTTPException):
        api.nav_layout_restore(session=ADMIN, conn=c2)                            # 沒有備份


def test_我的最愛_每人一份_預設與往返(tmp_path):
    c = _conn(tmp_path)
    r = api.nav_favorites_get(session=USER, conn=c)
    assert r["items"] == ["/changes", "/oncall", "/assets", "/datacenter", "/topology"] and r["is_default"]
    api.nav_favorites_put(api.NavFavoritesBody(items=["/storage", "/assets", "/storage", "bad"],
                                               recent=["/a", "/b", "/c", "/d", "/e", "/f"]), session=USER, conn=c)
    r = api.nav_favorites_get(session=USER, conn=c)
    assert r["items"] == ["/storage", "/assets"] and len(r["recent"]) == 5 and not r["is_default"]
    assert api.nav_favorites_get(session=ADMIN, conn=c)["is_default"]            # 別人的不受影響
    assert r["storage"] == "server"


def test_新舊編號對照匯出(tmp_path):
    buf = nav_prefs.mapping_xlsx([{"no": 1, "label": "資產查詢", "to": "/assets", "old_no": "3-1", "cur_no": "3-1",
                                   "new_no": "2-1", "new_group": "查詢"}])
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(buf.read())).active
    assert [c.value for c in ws[1]] == nav_prefs.MAPPING_COLS
    assert [c.value for c in ws[2]][3:6] == ["3-1", "3-1", "2-1"]
