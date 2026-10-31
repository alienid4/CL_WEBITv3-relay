"""搬遷盤點表回填：隨檔『CIA資產基準』三方比對（基準／現場／當前 DB）。只讀 DB、只寫輸出 xlsx。"""
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import relocation_import  # noqa: E402
from test_relocation_import import _bytes, _diff_rows, _export, _set, _tmpconn, _wb  # noqa: E402


def _three(res):
    sh = res["三方差異"]
    head = [c.value for c in sh[1]]
    return [dict(zip(head, [c.value for c in r])) for r in sh.iter_rows(min_row=2)]


def _cat(res, serial, col):
    got = [r for r in _three(res) if r["資產序號"] == serial and r["欄位"] == col]
    return got[0]["分類"] if got else None


def _unsign(wb):
    """模擬舊版匯出檔：沒有簽章、也沒有簽章標記。"""
    del wb["_meta"]
    wb.properties.description = None


def _db_set(conn, serial, field, value):
    conn.execute(f"UPDATE hardware SET {field}=? WHERE asset_serial=?", (value, serial))
    conn.commit()


def test_僅現場改():
    with _tmpconn() as conn:
        wb = _export(conn)
        _set(wb["Server Master"], "RL-1", "環境別", "測試")
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        res = _wb(out)
        assert _cat(res, "RL-1", "環境別") == "僅現場改"
        tw = summ["three_way"]
        assert tw["baseline_status"] == "ok" and tw["counts"]["site_only"] == 1
        assert tw["counts"]["unchanged"] > 0 and summ["units"]["three_way"] == "格"
        row = [r for r in _three(res) if r["資產序號"] == "RL-1"][0]
        assert (row["基準值"], row["現場值"], row["當前DB值"]) == ("正式", "測試", "正式")
        assert row["說明"]


def test_僅DB後來變_現場未動():
    with _tmpconn() as conn:
        wb = _export(conn)
        _db_set(conn, "RL-1", "environment", "DR")
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        res = _wb(out)
        assert _cat(res, "RL-1", "環境別") == "僅DB後來變"
        assert summ["three_way"]["counts"]["db_only"] == 1
        # 既有二方邏輯不受影響：現場值＝基準≠DB → 仍列進 CIA待異動
        assert [r["資產序號"] for r in _diff_rows(res)] == ["RL-1"]


def test_雙方同改():
    with _tmpconn() as conn:
        wb = _export(conn)
        _set(wb["Server Master"], "RL-1", "環境別", "DR")
        _db_set(conn, "RL-1", "environment", "DR")
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        assert _cat(_wb(out), "RL-1", "環境別") == "雙方同改"
        assert summ["three_way"]["counts"]["both_same"] == 1


def test_雙方衝突():
    with _tmpconn() as conn:
        wb = _export(conn)
        _set(wb["Server Master"], "RL-1", "環境別", "測試")
        _db_set(conn, "RL-1", "environment", "DR")
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        assert _cat(_wb(out), "RL-1", "環境別") == "雙方衝突"
        assert summ["three_way"]["counts"]["conflict"] == 1


def test_缺基準分頁_不當成差異():
    with _tmpconn() as conn:
        wb = _export(conn)
        del wb["CIA資產基準"]
        _set(wb["Server Master"], "RL-1", "環境別", "測試")
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        tw = summ["three_way"]
        assert tw["baseline_status"] == "missing"
        assert sum(v for k, v in tw["counts"].items() if k != "no_baseline") == 0
        rows = _three(_wb(out))
        assert len(rows) == 1 and rows[0]["分類"] == "無基準" and "沒有" in rows[0]["說明"]
        assert summ["diff"] == 1      # 既有二方比對照常


def test_缺欄_標疑似且不當成差異():
    with _tmpconn() as conn:
        wb = _export(conn)
        base = wb["CIA資產基準"]
        head = {c.value: i for i, c in enumerate(base[1], 1)}
        base.delete_cols(head["環境別"])
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        tw = summ["three_way"]
        assert tw["baseline_status"] == "疑似被改" and any("缺欄" in x for x in tw["baseline_reasons"])
        assert _cat(_wb(out), "RL-1", "環境別") is None      # 現場＝DB，缺欄不當成差異
        assert tw["counts"]["unchanged"] == 0


def test_基準被改_格式不符標疑似且不崩():
    with _tmpconn() as conn:
        wb = _export(conn)
        wb["CIA資產基準"].cell(1, 1).value = "被改的欄名"
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        assert summ["three_way"]["baseline_status"] == "疑似被改"
        res = _wb(out)
        assert "基準疑似被改，三方比對僅供參考" in _three(res)[0]["資產序號"]
        assert summ["three_way"]["note"] == "基準疑似被改，三方比對僅供參考"
        # 基準內容被改成別的值（欄名、筆數都不變，只改一格值）：簽章不符 -> 偵測
        wb2 = _export(conn)
        b2 = wb2["CIA資產基準"]
        h2 = {c.value: i for i, c in enumerate(b2[1], 1)}
        for r in range(2, b2.max_row + 1):
            if b2.cell(r, h2["資產序號"]).value == "RL-1":
                b2.cell(r, h2["環境別"]).value = "亂改"
        _set(wb2["Server Master"], "RL-1", "環境別", "測試")
        out2, s2 = relocation_import.reconcile(conn, _bytes(wb2))
        assert s2["three_way"]["baseline_status"] == "疑似被改" and s2["three_way"]["signature"] == "mismatch"
        assert _cat(_wb(out2), "RL-1", "環境別") == "基準不可信"
        c2 = s2["three_way"]["counts"]
        assert c2["unchanged"] == 0 and c2["untrusted"] > 0     # 簽章不符時任何格都不得判「未改」


def test_基準缺某台_無基準():
    with _tmpconn() as conn:
        wb = _export(conn)
        _unsign(wb)
        base = wb["CIA資產基準"]
        h = {c.value: i for i, c in enumerate(base[1], 1)}
        for r in range(2, base.max_row + 1):
            if base.cell(r, h["資產序號"]).value == "RL-2":
                base.delete_rows(r)
                break
        _set(wb["Server Master"], "RL-2", "環境別", "測試")
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        assert _cat(_wb(out), "RL-2", "環境別") == "無基準"
        assert summ["three_way"]["baseline_status"] == "疑似被改"


def test_刪簽章_偵測():
    with _tmpconn() as conn:
        wb = _export(conn)
        del wb["_meta"]                      # 只刪簽章分頁，文件標記仍在
        _set(wb["Server Master"], "RL-1", "環境別", "測試")
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        tw = summ["three_way"]
        assert tw["baseline_status"] == "疑似被改" and tw["signature"] == "deleted"
        assert tw["counts"]["unchanged"] == 0
        assert _cat(_wb(out), "RL-1", "環境別") == "基準不可信"


def test_舊版檔無簽章_僅供參考但仍比對():
    with _tmpconn() as conn:
        wb = _export(conn)
        _unsign(wb)
        _set(wb["Server Master"], "RL-1", "環境別", "測試")
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        tw = summ["three_way"]
        assert tw["baseline_status"] == "僅供參考" and tw["signature"] == "legacy"
        assert _cat(_wb(out), "RL-1", "環境別") == "僅現場改"
        assert tw["counts"]["unchanged"] > 0
        assert "無簽章（舊版檔）" in _three(_wb(out))[0]["資產序號"]


def test_金鑰與簽章值不外洩():
    import json
    import zipfile
    import db
    with _tmpconn() as conn:
        wb = _export(conn)
        key = db.get_setting(conn, "export_hmac_key")
        assert key and len(key) >= 32
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        assert key not in json.dumps(summ, ensure_ascii=False)
        for blob in (_bytes(wb), out.getvalue()):
            with zipfile.ZipFile(io.BytesIO(blob)) as z:
                for n in z.namelist():
                    assert key.encode() not in z.read(n)
        assert _wb(out)["_meta"].sheet_state == "veryHidden"


def test_重複回填冪等():
    with _tmpconn() as conn:
        wb = _export(conn)
        _set(wb["Server Master"], "RL-1", "環境別", "測試")       # 僅現場改
        _db_set(conn, "RL-2", "environment", "DR")               # 僅 DB 變
        _set(wb["Server Master"], "RL-2", "環境別", None)        # 現場留空 -> 系統補藍字
        out1, s1 = relocation_import.reconcile(conn, _bytes(wb))
        out2, s2 = relocation_import.reconcile(conn, out1.getvalue())
        out3, s3 = relocation_import.reconcile(conn, out2.getvalue())
        assert s1["three_way"]["counts"] == s2["three_way"]["counts"] == s3["three_way"]["counts"]
        for o in (out1, out2, out3):
            res = _wb(o)
            assert _cat(res, "RL-1", "環境別") == "僅現場改"
            assert _cat(res, "RL-2", "環境別") == "僅DB後來變"
            assert res.sheetnames.count("三方差異") == 1


def test_列序打亂仍正確():
    with _tmpconn() as conn:
        wb = _export(conn)
        ws = wb["Server Master"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        ws.delete_rows(2, ws.max_row)
        for r in reversed(rows):
            ws.append(r)
        base = wb["CIA資產基準"]
        brows = [[c.value for c in r] for r in base.iter_rows(min_row=2)]
        base.delete_rows(2, base.max_row)
        for r in reversed(brows):
            base.append(r)
        _set(ws, "RL-1", "環境別", "測試")
        _db_set(conn, "RL-2", "environment", "DR")
        out, _ = relocation_import.reconcile(conn, _bytes(wb))
        res = _wb(out)
        assert _cat(res, "RL-1", "環境別") == "僅現場改"
        assert _cat(res, "RL-2", "環境別") == "僅DB後來變"


def test_處理歷史不受影響():
    with _tmpconn() as conn:
        wb = _export(conn)
        _set(wb["Server Master"], "RL-1", "環境別", "測試")
        out1, _ = relocation_import.reconcile(conn, _bytes(wb))
        w1 = _wb(out1)
        sh = w1["CIA待異動"]
        head = {c.value: i for i, c in enumerate(sh[1], 1)}
        sh.cell(2, head["申請單編號"]).value = "ES800099"
        sh.cell(2, head["處理狀態"]).value = "已送件"
        out2, _ = relocation_import.reconcile(conn, _bytes(w1))
        rows = _diff_rows(_wb(out2))
        assert len(rows) == 1 and rows[0]["申請單編號"] == "ES800099" and rows[0]["處理狀態"] == "已送件"


# ---- 上傳資源上限 ----
def test_超大檔_413類例外():
    big = io.BytesIO(b"0" * (relocation_import.MAX_UPLOAD_BYTES + 5))
    try:
        relocation_import.read_upload(big)
        raise AssertionError("應該拒絕")
    except relocation_import.RelocationTooLarge as e:
        assert "太大" in str(e)
    with _tmpconn() as conn:
        try:
            relocation_import.reconcile(conn, b"0" * (relocation_import.MAX_UPLOAD_BYTES + 5))
            raise AssertionError("應該拒絕")
        except relocation_import.RelocationTooLarge:
            pass


def test_高壓縮比zip與過多檔案_被擋():
    import zipfile
    with _tmpconn() as conn:
        b = io.BytesIO()
        with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("xl/big.bin", b"\0" * (relocation_import.MAX_ZIP_TOTAL + 1024 * 1024))
        assert len(b.getvalue()) < 2 * 1024 * 1024
        try:
            relocation_import.reconcile(conn, b.getvalue())
            raise AssertionError("應該拒絕")
        except relocation_import.RelocationTooLarge as e:
            assert "解壓後" in str(e)
        b2 = io.BytesIO()
        with zipfile.ZipFile(b2, "w") as z:
            for i in range(relocation_import.MAX_ZIP_FILES + 1):
                z.writestr(f"f{i}.txt", "x")
        try:
            relocation_import.reconcile(conn, b2.getvalue())
            raise AssertionError("應該拒絕")
        except relocation_import.RelocationTooLarge:
            pass


def test_列數過多_被擋():
    from openpyxl import Workbook
    with _tmpconn() as conn:
        wb = Workbook()
        ws = wb.active
        ws.title = "Server Master"
        ws.append(["資產序號"])
        for i in range(relocation_import.MAX_SHEET_ROWS + 5):
            ws.append([f"S{i}"])
        try:
            relocation_import.reconcile(conn, _bytes(wb))
            raise AssertionError("應該拒絕")
        except relocation_import.RelocationTooLarge as e:
            assert "列數太多" in str(e)


# ---- 公式字元 ----
def test_公式開頭的值輸出為文字():
    with _tmpconn() as conn:
        wb = _export(conn)
        _db_set(conn, "RL-2", "environment", '=HYPERLINK("http://x")')
        c = wb["Server Master"]
        r = _set(c, "RL-1", "環境別", "+cmd|calc")
        head = {str(x.value).strip(): i for i, x in enumerate(c[1], 1) if x.value}
        c.cell(r, head["環境別"]).data_type = "s"
        out, _ = relocation_import.reconcile(conn, _bytes(wb))
        res = _wb(out)
        found = 0
        for name in ("三方差異", "CIA待異動"):
            for row in res[name].iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value[:1] in "=+-@":
                        found += 1
                        assert cell.data_type == "s", (name, cell.value)
        assert found >= 2


def test_簽章_重複序號與空序號不誤報():
    """獨立審查（2026-10-31）：匯出含空序號、匯入跳過空序號會讓簽章誤報；重複序號排序打平也會。"""
    import relocation_export as rx
    hdr = ["資產序號", "主機名稱"]
    a = [["S1", "x"], ["S1", "y"], ["S2", "z"]]
    b = [["S2", "z"], ["S1", "y"], ["S1", "x"]]                     # 同樣內容、列序不同（含重複序號）
    assert rx.base_signature("k", hdr, a) == rx.base_signature("k", hdr, b)
    c = [["S1", "x"], ["S1", "Y"], ["S2", "z"]]                     # 只改一格
    assert rx.base_signature("k", hdr, a) != rx.base_signature("k", hdr, c)
