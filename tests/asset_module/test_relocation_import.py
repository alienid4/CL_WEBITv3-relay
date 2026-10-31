"""機房搬遷盤點表回填（relocation_import.reconcile）：W221-017C。

守：重複回填不丟人填的處理歷史、不疊加備註；未知序號逐筆列出；摘要單位講清楚。
"""
import io
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import relocation_export  # noqa: E402
import relocation_import  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


import contextlib  # noqa: E402


@contextlib.contextmanager
def _tmpconn():
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "t.db"
        db.init_db(p)
        c = db.get_connection(p)
        try:
            db.insert_hardware(c, asset_serial="RL-1", hostname="web01", ip="10.99.1.1",
                               os="Rocky Linux 9.7", environment="正式", asset_status="使用中")
            db.insert_hardware(c, asset_serial="RL-2", hostname="web02", ip="10.99.1.2",
                               os="Rocky Linux 9.7", environment="正式", asset_status="使用中")
            yield c
        finally:
            c.close()


def _wb(buf):
    return load_workbook(io.BytesIO(buf.getvalue() if hasattr(buf, "getvalue") else buf))


def _set(ws, serial, col_name, value):
    head = {str(c.value).strip(): i for i, c in enumerate(ws[1], 1) if c.value}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, head["資產序號"]).value == serial:
            ws.cell(r, head[col_name]).value = value
            return r
    raise AssertionError(serial)


def _note(ws, row):
    return ws.cell(row, [c.value for c in ws[1]].index("備註") + 1).value


def _bytes(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _export(conn):
    buf, _ = relocation_export.build(conn)
    return _wb(buf)


def _diff_rows(wb):
    sh = wb["CIA待異動"]
    head = [c.value for c in sh[1]]
    return [dict(zip(head, [c.value for c in r])) for r in sh.iter_rows(min_row=2)
            if r[2].value]


def test_匯出改一格回填_產生差異且備註寫入():
    with _tmpconn() as conn:
        wb = _export(conn)
        _set(wb["Server Master"], "RL-1", "環境別", "測試")
        out, summ = relocation_import.reconcile(conn, _bytes(wb))
        res = _wb(out)
        rows = _diff_rows(res)
        assert [(r["資產序號"], r["欄位名稱"], r["CIA原值"], r["現場盤點值"]) for r in rows] == \
            [("RL-1", "環境別", "正式", "測試")]
        assert summ["diff"] == 1 and summ["diff_rows"] == 1 and summ["diff_machines"] == 1
        assert summ["units"]["diff"] == "格" and summ["units"]["unknown_serial"] == "列"
        ws = res["Server Master"]
        r = _set(ws, "RL-1", "環境別", "測試")
        assert _note(ws, r).count("【系統對帳不一致】") == 1


def test_重複回填_保留處理歷史_備註不疊加():
    with _tmpconn() as conn:
        wb = _export(conn)
        _set(wb["Server Master"], "RL-1", "環境別", "測試")
        out1, _ = relocation_import.reconcile(conn, _bytes(wb))
        # 人在 CIA待異動 填處理資料
        w1 = _wb(out1)
        sh = w1["CIA待異動"]
        head = {c.value: i for i, c in enumerate(sh[1], 1)}
        sh.cell(2, head["申請單編號"]).value = "ES800099"
        sh.cell(2, head["處理狀態"]).value = "已送件"
        sh.cell(2, head["完成日期"]).value = "2026-11-01"
        out2, _ = relocation_import.reconcile(conn, _bytes(w1))
        out3, _ = relocation_import.reconcile(conn, out2.getvalue())   # 第三次
        for out in (out2, out3):
            res = _wb(out)
            rows = _diff_rows(res)
            assert len(rows) == 1
            assert rows[0]["申請單編號"] == "ES800099"
            assert rows[0]["處理狀態"] == "已送件"
            assert str(rows[0]["完成日期"]).startswith("2026-11-01")
            ws = res["Server Master"]
            r = _set(ws, "RL-1", "環境別", "測試")
            assert _note(ws, r).count("【系統對帳不一致】") == 1


def test_已處理過的列即使本次已一致也保留_沒處理過的就不留():
    with _tmpconn() as conn:
        wb = _export(conn)
        _set(wb["Server Master"], "RL-1", "環境別", "測試")
        _set(wb["Server Master"], "RL-2", "環境別", "測試")
        out1, _ = relocation_import.reconcile(conn, _bytes(wb))
        w1 = _wb(out1)
        sh = w1["CIA待異動"]
        head = {c.value: i for i, c in enumerate(sh[1], 1)}
        for r in range(2, sh.max_row + 1):
            if sh.cell(r, 1).value == "RL-1":
                sh.cell(r, head["申請單編號"]).value = "ES1"
                sh.cell(r, head["處理狀態"]).value = "已完成"
        # 現場把兩台都改回與系統一致
        _set(w1["Server Master"], "RL-1", "環境別", "正式")
        _set(w1["Server Master"], "RL-2", "環境別", "正式")
        out2, summ = relocation_import.reconcile(conn, _bytes(w1))
        rows = _diff_rows(_wb(out2))
        assert [r["資產序號"] for r in rows] == ["RL-1"]
        assert rows[0]["申請單編號"] == "ES1" and "已一致" in rows[0]["差異類型"]
        assert summ["history_kept_rows"] == 1


def test_未知序號逐筆列出且重複回填不疊加():
    with _tmpconn() as conn:
        wb = _export(conn)
        ws = wb["Server Master"]
        ws.append(["GHOST-1", "使用中", "實體Server", "", "ghost1"])
        ws.append(["GHOST-2", "使用中", "實體Server", "", "ghost2"])
        out1, summ = relocation_import.reconcile(conn, _bytes(wb))
        assert summ["unknown_serial"] == 2
        out2, _ = relocation_import.reconcile(conn, out1.getvalue())
        res = _wb(out2)
        assert "未知序號" in res.sheetnames
        rows = [r[0].value for r in res["未知序號"].iter_rows(min_row=2)]
        assert rows == ["GHOST-1", "GHOST-2"]


def test_沒差異時不報錯且有佔位列():
    with _tmpconn() as conn:
        out, summ = relocation_import.reconcile(conn, _bytes(_export(conn)))
        assert summ["diff_rows"] == 0
        res = _wb(out)
        assert res["CIA待異動"].cell(2, 1).value.startswith("（本次比對沒有不一致")
        assert res["未知序號"].cell(2, 1).value.startswith("（本次沒有未知序號")
