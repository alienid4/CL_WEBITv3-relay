"""資產主檔匯出（2026-09-18）：欄位說明要讀得到 schema 註解、NA 不算有效值、
同台筆數不跨退役。"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import hardware_master as hm  # noqa: E402


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    c = db.get_connection(p)
    db.insert_hardware(c, asset_serial="A-1", hostname="h1", ip="192.0.2.1",
                       asset_status="使用中", request_no="NA")
    db.insert_hardware(c, asset_serial="A-2", hostname="h1", ip="192.0.2.1",
                       asset_status="使用中", request_no="E600")
    db.insert_hardware(c, asset_serial="A-3", hostname="h1", ip="192.0.2.1",
                       asset_status="報廢")                    # 上一台，不是同一台
    db.insert_hardware(c, asset_serial="DYN-9", hostname="h9", ip="192.0.2.9")
    c.commit()
    return c


def test_schema註解讀得到():
    n = hm.schema_comments()
    assert len(n) > 30, f"只讀到 {len(n)} 欄，schema.sql 解析多半壞了"
    assert n.get("asset_serial") == "資產序號"
    assert n.get("vi_sdk_server"), "寫在上方的整段註解也要讀得到"


def test_NA不算有效值(tmp_path):
    c = _conn(tmp_path)
    d = {x["column"]: x for x in hm.field_dictionary(c)}
    assert d["request_no"]["filled"] == 2, "NA 算已填（跟原本那張表一致）"
    assert d["request_no"]["valid"] == 1, "但有效值要排除 NA"


def test_同台筆數不跨退役(tmp_path):
    c = _conn(tmp_path)
    header, body = hm.master_rows(c)
    i_sn, i_n = header.index("asset_serial"), header.index("同台筆數")
    i_src = header.index("來源")
    got = {r[i_sn]: r[i_n] for r in body}
    assert got["A-1"] == 2 and got["A-2"] == 2, "兩筆使用中同 key → 同台 2 筆"
    assert got["A-3"] == 1, "報廢那筆是上一台，不可以算進使用中的同台"
    assert {r[i_sn]: r[i_src] for r in body}["DYN-9"] == "DY"


def test_匯出檔有兩個分頁(tmp_path):
    from openpyxl import load_workbook
    c = _conn(tmp_path)
    wb = load_workbook(hm.export_xlsx(c))
    assert wb.sheetnames == ["資產主檔", "欄位說明"]
    assert wb["資產主檔"].max_row == 5, "表頭＋4 筆，一筆都不能少"
