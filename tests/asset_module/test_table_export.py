"""2-1 每個頁簽都有搜尋＋匯出 Excel（2026-10-22 使用者：「每個頁面都要有」）。

通用匯出 /api/export/table-xlsx：前端送畫面上的列，後端只轉格式、不查資料庫。
"""
import io
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "APP" / "asset-module" / "backend"
sys.path.insert(0, str(BACKEND))

ADOPT_VUE = (ROOT / "APP" / "asset-module" / "frontend" / "pages" / "adopt.vue").read_text(encoding="utf-8")


def _tab(name, nxt):
    return ADOPT_VUE.split(f"<div v-show=\"activeTab === '{name}'\">")[1].split(f"<div v-show=\"activeTab === '{nxt}'\">")[0]


def test_2_1四個清單頁簽都有搜尋與匯出():
    tabs = {"pending": _tab("pending", "unregistered"), "unregistered": _tab("unregistered", "failures"),
            "failures": _tab("failures", "onboarded"), "onboarded": ADOPT_VUE.split("<div v-show=\"activeTab === 'onboarded'\">")[1]}
    for name, body in tabs.items():
        assert 'type="search"' in body, f"{name} 沒有搜尋框"
        assert "匯出 Excel" in body, f"{name} 沒有匯出"


def test_結果篩選有待開WinRM_待開SSH_待佈身分():
    """2026-10-22 使用者：「（待開 WinRM 這一類匯出 Excel）沒這功能啊」。"""
    for k, label in (("needs_winrm_port", "待開 WinRM"), ("needs_ssh_port", "待開 SSH"),
                     ("needs_credential", "待佈身分")):
        assert f"{k}: '{label}'" in ADOPT_VUE
        assert f"{k}: (r) => r.status === '{k}'" in ADOPT_VUE
        assert f"'{k}'" in ADOPT_VUE.split("const D_RESULT_KEYS: DResult[] = [")[1].split("]")[0]


def test_搬到Windows的_不給一鍵納管_下一步講WinRM():
    """2026-10-22：2.38 開 22／445／5985，機器判 SSH；人工搬到 Windows 後不該再有 Linux 的一鍵納管。"""
    import re
    src = re.search(r"function canProbe\(r: DispatchRow\): boolean \{(.*?)\n\}", ADOPT_VUE, re.S).group(1)
    assert "if (r.tab === 'winrm') return false" in src
    assert "(r.route === 'winrm' || r.tab === 'winrm')) return '請設定 WinRM 憑證'" in ADOPT_VUE


def test_通用匯出_畫面送什麼就匯什麼(monkeypatch):
    from fastapi.testclient import TestClient
    import api
    from _app_shared import require_auth

    api.app.dependency_overrides[require_auth] = lambda: {"username": "tester"}
    try:
        c = TestClient(api.app)
        r = c.post("/api/export/table-xlsx", json={"title": "未登記主機", "headers": ["IP", "主機名稱"],
                                                  "rows": [["192.0.2.1", "h1"], ["192.0.2.2", None]]})
        assert r.status_code == 200, r.text
        ws = load_workbook(io.BytesIO(r.content)).active
        assert [[c.value for c in row] for row in ws.iter_rows()] == [
            ["IP", "主機名稱"], ["192.0.2.1", "h1"], ["192.0.2.2", None]]
    finally:
        api.app.dependency_overrides.clear()
