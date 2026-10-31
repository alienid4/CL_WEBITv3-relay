"""2026-10-29 機櫃面板雙擊改機型／IP／序號／備註：存自己的表、不動 CIA、留稽核、可還原、CIA 跟上就功成身退。全部假資料。"""
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import rack_field_edit as fe  # noqa: E402
import rack_layout as rl  # noqa: E402


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    return db.get_connection(p)


def _hw(conn, serial, host, ip, rack, model=None, hw_serial=None):
    db.insert_hardware(conn, asset_serial=serial, hostname=host, ip=ip, physical_location="01_甲機房",
                       rack_no=rack, device_model=model, hw_serial=hw_serial)


def _m(conn, serial):
    return next(x for x in rl.machines(conn) if x["asset_serial"] == serial)


def _seed(conn):
    _hw(conn, "F-1", "app01", "10.99.8.1", "B10/10U", "DELL R750", "ABC1234")
    conn.commit()


def test_改機型IP序號備註_不動CIA_畫面套用更正(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        k = _m(conn, "F-1")["key"]
        fe.set_edit(conn, rl.machines(conn), k, "model", "dell r760", "現場看到是 R760", "alice")
        fe.set_edit(conn, rl.machines(conn), k, "ip", "10.99.8.9", None, "alice")
        fe.set_edit(conn, rl.machines(conn), k, "serial", "xyz-9876", None, "alice")
        fe.set_edit(conn, rl.machines(conn), k, "remark", "  第二顆電源壞  ", None, "alice")
        m = _m(conn, "F-1")
        assert m["model"] == "DELL R760" and m["ip"] == "10.99.8.9" and m["hw_serials"] == ["XYZ9876"]
        assert m["remark"] == "第二顆電源壞"
        assert set(m["field_edits"]) == {"model", "ip", "serial", "remark"} and m["field_edits"]["model"]["by"] == "alice"
        assert m["field_edits"]["model"]["cia"] == "DELL R750"                    # 記得 CIA 原本是什麼
        row = conn.execute("SELECT device_model, ip, hw_serial FROM hardware WHERE asset_serial='F-1'").fetchone()
        assert tuple(row) == ("DELL R750", "10.99.8.1", "ABC1234")                 # CIA 的 hardware 完全沒被改
        acts = fe.audit(conn, k)
        assert len(acts) == 4 and all(a["updated_by"] == "alice" and a["action"] == "set" for a in acts)
        pend = fe.pending_cia(conn, rl.machines(conn))
        assert {p["label"] for p in pend} == {"機型", "IP", "序號", "備註"}
    finally:
        conn.close()


def test_驗證_錯的值被擋_沒改的也擋(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        ms = rl.machines(conn)
        k = ms[0]["key"]
        with pytest.raises(ValueError):
            fe.set_edit(conn, ms, k, "ip", "10.99.8", None, "a")                    # 不是有效 IP
        with pytest.raises(ValueError):
            fe.set_edit(conn, ms, k, "model", "HPE Server", None, "a")              # 機型要「廠牌 型號」
        with pytest.raises(ValueError):
            fe.set_edit(conn, ms, k, "ip", "  ", None, "a")                         # 不能空白（要清除請還原）
        with pytest.raises(ValueError):
            fe.set_edit(conn, ms, k, "ip", "10.99.8.1", None, "a")                  # 跟現在一樣
        with pytest.raises(ValueError):
            fe.set_edit(conn, ms, k, "remark", "x" * 201, None, "a")
        with pytest.raises(ValueError):
            fe.set_edit(conn, ms, k, "hostname", "x", None, "a")                    # 沒開放的欄位
        with pytest.raises(LookupError):
            fe.set_edit(conn, ms, "no-such-key", "ip", "10.99.8.2", None, "a")
        assert fe.audit(conn) == [] and fe.pending_cia(conn, ms) == []
    finally:
        conn.close()


def test_還原_回到CIA的值_再改也沿用第一次的CIA值(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        k = _m(conn, "F-1")["key"]
        fe.set_edit(conn, rl.machines(conn), k, "ip", "10.99.8.5", None, "alice")
        fe.set_edit(conn, rl.machines(conn), k, "ip", "10.99.8.6", None, "bob")
        assert _m(conn, "F-1")["field_edits"]["ip"]["cia"] == "10.99.8.1"            # 第二次改，CIA 值不會變成第一次的更正
        out = fe.clear_edit(conn, rl.machines(conn), k, "ip", "carol")
        assert out["after"] == "10.99.8.1" and _m(conn, "F-1")["ip"] == "10.99.8.1"
        assert "ip" not in _m(conn, "F-1")["field_edits"]
        with pytest.raises(LookupError):
            fe.clear_edit(conn, rl.machines(conn), k, "ip", "carol")                # 沒更正可還原
        acts = [(a["action"], a["updated_by"]) for a in fe.audit(conn, k)]
        assert acts == [("clear", "carol"), ("set", "bob"), ("set", "alice")]       # 誰改的、誰還原的都留著
    finally:
        conn.close()


def test_CIA改成跟現場一樣_更正功成身退(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        k = _m(conn, "F-1")["key"]
        fe.set_edit(conn, rl.machines(conn), k, "model", "DELL R760", None, "alice")
        fe.set_edit(conn, rl.machines(conn), k, "remark", "備註留著", None, "alice")
        conn.execute("UPDATE hardware SET device_model = 'DELL R760' WHERE asset_serial = 'F-1'")   # CIA 負責人照改了
        conn.commit()
        m = _m(conn, "F-1")
        assert m["model"] == "DELL R760" and "model" not in m["field_edits"]          # 不再套用
        assert "remark" in m["field_edits"]                                            # 備註沒有 CIA 對應，繼續留著
        st = conn.execute("SELECT status FROM rack_field_override WHERE field='model'").fetchone()[0]
        assert st == "synced"
        assert [p["label"] for p in fe.pending_cia(conn, rl.machines(conn))] == ["備註"]
        assert any(a["action"] == "cia_synced" for a in fe.audit(conn))
    finally:
        conn.close()


def test_CIA待更正xlsx_有項次_欄位對(tmp_path):
    import io

    from openpyxl import load_workbook
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        k = _m(conn, "F-1")["key"]
        fe.set_edit(conn, rl.machines(conn), k, "ip", "10.99.8.5", "現場標籤", "alice")
        buf, n = fe.pending_xlsx(conn, rl.machines(conn))
        ws = load_workbook(io.BytesIO(buf.getvalue())).active
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        assert n == 1 and rows[0] == fe.XLSX_COLS
        assert rows[1][0] == 1 and rows[1][4] == "app01" and rows[1][5] == "IP"
        assert rows[1][6] == "10.99.8.1" and rows[1][7] == "10.99.8.5" and rows[1][8] == "alice" and rows[1][10] == "現場標籤"
    finally:
        conn.close()


def test_API_限管理員_有過守門(tmp_path, monkeypatch):
    from fastapi import HTTPException
    from routers import overlaps as ro

    conn = _conn(tmp_path)
    try:
        _seed(conn)
        k = _m(conn, "F-1")["key"]
        req = type("R", (), {"headers": {}, "client": type("C", (), {"host": "10.99.0.9"})()})()
        body = ro.FieldEditBody(key=k, field="ip", value="10.99.8.7")
        monkeypatch.setattr(ro, "_can_edit_nav", lambda s, c: False)
        with pytest.raises(HTTPException) as e:
            ro.rack_field_edit_endpoint(body, req, None, conn)
        assert e.value.status_code == 403
        with pytest.raises(HTTPException) as e:
            ro.rack_field_edit_clear_endpoint(k, "ip", req, None, conn)
        assert e.value.status_code == 403
        monkeypatch.setattr(ro, "_can_edit_nav", lambda s, c: True)
        called = []
        monkeypatch.setattr(ro, "_guard", lambda c, rq: called.append(1))
        monkeypatch.setattr(ro, "_actor", lambda s, rq: "alice")
        out = ro.rack_field_edit_endpoint(body, req, None, conn)
        assert out["after"] == "10.99.8.7" and called == [1]
        with pytest.raises(HTTPException) as e:
            ro.rack_field_edit_endpoint(ro.FieldEditBody(key=k, field="ip", value="bad"), req, None, conn)
        assert e.value.status_code == 400
        with pytest.raises(HTTPException) as e:
            ro.rack_field_edit_endpoint(ro.FieldEditBody(key="zz", field="ip", value="10.99.8.8"), req, None, conn)
        assert e.value.status_code == 404
        lst = ro.rack_field_edits_endpoint(None, conn)
        assert lst["items"][0]["value"] == "10.99.8.7" and lst["audit"][0]["updated_by"] == "alice"
        assert ro.rack_field_edit_clear_endpoint(k, "ip", req, None, conn)["after"] == "10.99.8.1"
    finally:
        conn.close()


def test_審查補強_機型只差大小寫也退場_IP不撞_重改用最新CIA值_xlsx不當公式(tmp_path):
    import io

    from openpyxl import load_workbook
    conn = _conn(tmp_path)
    try:
        _hw(conn, "F-1", "app01", "10.99.8.1", "B10/10U", "Dell r740", "ABC1234")
        _hw(conn, "F-2", "app02", "10.99.8.2", "B10/12U", "DELL R750", "ZZZ9999")
        conn.commit()
        k = _m(conn, "F-1")["key"]
        fe.set_edit(conn, rl.machines(conn), k, "model", "Dell R740 ", None, "a")          # 更正值 → "DELL R740"（跟 CIA 只差大小寫）
        conn.execute("UPDATE hardware SET device_model = 'Dell  R740' WHERE asset_serial='F-1'")
        conn.commit()
        assert "model" not in _m(conn, "F-1")["field_edits"]                                # 大小寫／空白不同也算 CIA 跟上了
        with pytest.raises(ValueError):
            fe.set_edit(conn, rl.machines(conn), k, "ip", "10.99.8.2", None, "a")           # 撞到使用中的 app02
        fe.set_edit(conn, rl.machines(conn), k, "ip", "10.99.8.3", None, "a")
        conn.execute("UPDATE hardware SET ip='10.99.8.3' WHERE asset_serial='F-1'")
        conn.commit()
        assert "ip" not in _m(conn, "F-1")["field_edits"]                                   # 已退場
        k = _m(conn, "F-1")["key"]                                                          # 機器的 key 可能隨 CIA 的 IP 改變，重取
        fe.set_edit(conn, rl.machines(conn), k, "ip", "10.99.8.4", None, "b")               # 退場後再改：CIA 值是 .3，不是更早的 .1
        assert _m(conn, "F-1")["field_edits"]["ip"]["cia"] == "10.99.8.3"
        fe.set_edit(conn, rl.machines(conn), k, "remark", "=HYPERLINK(\"x\")", "+cmd", "a")
        buf, n = fe.pending_xlsx(conn, rl.machines(conn))
        ws = load_workbook(io.BytesIO(buf.getvalue())).active
        vals = [c.value for r in ws.iter_rows(min_row=2) for c in r if isinstance(c.value, str)]
        assert any(v.startswith("'=") for v in vals) and not any(v.startswith("=") for v in vals)
    finally:
        conn.close()
