"""機櫃圖對帳：解析現場格式 Excel（合併格、一格多 IP、已拔電／暫定、略過理線槽、前／後）、
櫃名別名（內湖 1-2 ＝ A2）、對帳分類、匯出再匯入要一致。全部假資料（10.99.x、假主機名）。"""
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import rack_diagram as rd  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    return db.get_connection(p)


def _site_book() -> bytes:
    """現場格式：櫃頭（藍底、跨 U｜設備｜U 三欄）、兩側 U 數字 42→1、設備合併格。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "4F RACK A"
    blue = PatternFill("solid", fgColor="00B0F0")

    def rack(col, name, devices):
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 2)
        ws.cell(1, col, name).fill = blue
        for u in range(42, 0, -1):
            r = 2 + 42 - u
            ws.cell(r, col, u)
            ws.cell(r, col + 2, u)
        for top_u, bot_u, text in devices:
            r1, r2 = 2 + 42 - top_u, 2 + 42 - bot_u
            if r2 > r1:
                ws.merge_cells(start_row=r1, start_column=col + 1, end_row=r2, end_column=col + 1)
            ws.cell(r1, col + 1, text)

    rack(1, "A2", [(42, 42, "理線槽"), (40, 39, "DELL R750 app01 10.99.5.1"), (35, 35, "後:Cisco SW 10.99.5.2"),
                   (30, 29, "fakesrv02 10.99.5.3\nfakesrv03 10.99.5.4"), (20, 20, "Lenovo SR630 10.99.5.9(已拔電)"),
                   (10, 10, "(暫定)新DB Node1"), (5, 5, "IBM S922 (ABC1234)"), (3, 3, "電源排插x20")])
    rack(5, "A3(期自營)", [(40, 40, "app01-dup 10.99.5.1")])          # 同一個 IP 又出現：自相矛盾
    return _save(wb)


def _save(wb) -> bytes:
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_解析現場格式():
    rows = rd.parse_workbook(_site_book())
    by = {(r["rack"], r["u_start"]): r for r in rows}
    assert all(r["floor"] == "4F" for r in rows)
    assert by[("A2", 39)]["u_end"] == 40 and by[("A2", 39)]["main_ip"] == "10.99.5.1"     # 合併格跨 2U
    assert by[("A2", 35)]["side"] == "後" and by[("A2", 35)]["main_ip"] == "10.99.5.2"
    assert {r["main_ip"] for r in rows if r["u_start"] == 29} == {"10.99.5.3", "10.99.5.4"}   # 一格兩台拆開
    assert by[("A2", 20)]["state"] == "已拔電"
    assert by[("A2", 10)]["planned"] is True
    assert by[("A2", 5)]["serials"] == ["ABC1234"]
    texts = " ".join(r["text"] for r in rows)
    assert "理線槽" not in texts and "電源排插" not in texts
    assert by[("A3", 40)]["rack_norm"] == "A-3"                                            # 「A3(期自營)」只取櫃名


def test_別名_內湖排櫃數字():
    assert rd.alias("內湖", "1-2") == "A-2" and rd.alias("內湖", "3-10") == "C-10"
    assert rd.alias("敦南", "1-2") == "1-2" and rd.norm_rack("M01") == "M-1"
    assert rd.alias("內湖", "2-1", {"內湖": {"row_letters": "XYZ"}}) == "Y-1"


def _seed(conn):
    def hw(serial, host, ip, rack, model=None, hw_serial=None, vm=False):
        db.insert_hardware(conn, asset_serial=serial, hostname=host, ip=ip, physical_location="02_內湖機房",
                           rack_no=rack, device_model=model, hw_serial=hw_serial)
    hw("H-1", "app01", "10.99.5.1", "1-2/39-40U", "DELL R750")         # 一致（別名 1-2＝A2）
    hw("H-2", "sw02", "10.99.5.2", "1-2/30U", "Cisco SW")              # 同櫃但 U 位不同
    hw("H-3", "fakesrv02", "10.99.5.3", "1-4/29U", "X")                # 機櫃不同
    hw("H-5", "s922", None, "1-2/5U", "IBM S922", hw_serial="ABC1234")  # 用序號對到
    hw("H-6", "lonely", "10.99.5.88", "1-7/10U", "X")                  # 系統有、現場圖沒有的櫃
    conn.commit()


def test_對帳分類(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        imp = rd.import_diagram(conn, "內湖", "fake.xlsx", _site_book(), "tester")
        rc = rd.reconcile(conn, imp["import_id"])
        res = {(x["rack"], x["u_start"], x.get("ip")): x["result"] for x in rc["rows"]}
        assert res[("A2", 39, "10.99.5.1")] == "現場圖自相矛盾"           # 同 IP 也在 A3
        assert res[("A2", 35, "10.99.5.2")] == "同櫃但U位不同"
        assert res[("A2", 29, "10.99.5.3")] == "機櫃不同"
        assert res[("A2", 29, "10.99.5.4")] == "照片有系統沒有"
        assert res[("A2", 10, None)] == "規劃中"
        assert res[("A2", 5, None)] == "一致"                              # 序號對到
        assert [m["rack"] for m in rc["missing_racks"]] == ["1-4", "1-7"]
        assert rc["counts"]["一致"] == 1 and any(r["rack"] == "A2" for r in rc["rate"])
        buf, n = rd.fix_xlsx(conn, imp["import_id"])
        assert n == 0                                                    # 沒人按「確認（照片對）」→ 不進待更正
    finally:
        conn.close()


def test_匯出現場格式_再匯入一致_並列需比對(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        buf, _ = rd.export_xlsx(conn)
        wb = load_workbook(buf)
        assert wb.sheetnames[0] == "需比對清單" and "內湖機房_A排" in wb.sheetnames          # 一個機房＋排一張
        assert "還沒匯入過現場圖" in str(wb["需比對清單"].cell(2, 3).value)
        data = buf.getvalue()
        imp = rd.import_diagram(conn, "內湖", "roundtrip.xlsx", data, "tester")
        rc = rd.reconcile(conn, imp["import_id"])
        assert rc["counts"]["一致"] == len(rc["rows"]) == 5 and rc["missing_racks"] == []
        # 再匯入一份少一台的現場圖：匯出時那台標「現場圖找不到」，那一櫃整櫃沒了就標「現場圖沒有這一櫃」
        imp2 = rd.import_diagram(conn, "內湖", "fake.xlsx", _site_book(), "tester")
        buf2, n2 = rd.export_xlsx(conn)
        todo = [[c.value for c in r] for r in load_workbook(buf2)["需比對清單"].iter_rows(min_row=2)]
        why = {t[5]: (t[1], t[8]) for t in todo}
        assert why["lonely"][0] == "先看" and "現場圖沒有這一櫃" in why["lonely"][1]
        assert "U位不同" in why["sw02"][1]
        d = rd.diff_previous(conn, imp2["import_id"])
        assert d["previous"]["id"] == imp["import_id"] and d["moved"]
    finally:
        conn.close()


def test_櫃下方沒U位的設備與數字開頭櫃名():
    wb = Workbook()
    ws = wb.active
    ws.title = "3F"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    ws.cell(1, 1, "2J06")
    for u in range(42, 0, -1):
        ws.cell(2 + 42 - u, 1, u)
        ws.cell(2 + 42 - u, 3, u)
    ws.cell(46, 2, "位置不明（有機櫃、沒 U 位）")
    ws.cell(47, 2, "⚠ X1 nouhost 10.99.6.1")
    ws.cell(48, 2, "IBM X346 (ZZZ9999)(已關機可下架)")
    ws.cell(52, 2, "網路相關設備")                                     # 圖例：沒有 IP／序號，不算設備
    racks: list = []
    rows = rd.parse_workbook(_save(wb), racks)
    assert [r["rack_norm"] for r in racks] == ["2-J-6"]
    assert {(r["main_ip"], r["u_start"]) for r in rows} == {("10.99.6.1", None), (None, None)}
    assert next(r for r in rows if not r["main_ip"])["state"] == "已關機"


def test_0000不算IP():
    assert rd.parse_device("Cisco SW 0.0.0.0")["ips"] == []


def test_需比對清單分頁不當機櫃圖讀():
    wb = Workbook()
    ws = wb.active
    ws.title = "需比對清單"
    ws.append(["項次", "優先", "機房", "機櫃", "U", "主機", "IP", "機型", "為什麼要比對"])
    for i in range(5):
        ws.append([i + 1, "先看", "內湖", "C4", "37", f"fakeoob{i}", "", "", "現場圖找不到"])
    racks: list = []
    assert rd.parse_workbook(_save(wb), racks) == [] and racks == []
