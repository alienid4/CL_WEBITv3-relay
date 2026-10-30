"""機櫃實體盤點（2026-10-27）：開始盤點凍結基準、離線 xlsx 來回、驗檔錯誤、部分上傳與重交、差異清單、
一鍵處理、基準不跟著後來的異動變。全部假資料。"""
import io
import sys
from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import rack_inventory as ri  # noqa: E402
import rack_layout as rl  # noqa: E402


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    return db.get_connection(p)


def _hw(conn, serial, host, ip, rack, model=None, os=None, hw_serial=None):
    db.insert_hardware(conn, asset_serial=serial, hostname=host, ip=ip, physical_location="01_甲機房",
                       rack_no=rack, device_model=model, os=os, hw_serial=hw_serial)


def _seed(conn):
    _hw(conn, "A-1", "app01", "10.99.6.1", "B1/30-31U", "DELL R750", hw_serial="SN0001")
    _hw(conn, "A-2", "app01-idrac", "10.99.7.1", "B1/30U", "DELL R750")                 # 管理口 → 併進 app01
    _hw(conn, "A-3", "app02", "10.99.6.2", "B1/20U", "HPE DL360")                       # 沒有原廠序號
    _hw(conn, "F-1", "FakeStore-FS5300", "10.99.5.1", "B1/10U", "IBM FS5300", hw_serial="78F0000")
    _hw(conn, "F-2", "FakeStore-FS5300", "10.99.5.2", "B1/10U", "IBM FS5300")           # 同一台、多管理 IP
    _hw(conn, "C-1", "app03", "10.99.6.3", "B2/5U", "DELL R660", hw_serial="SN0003")
    _hw(conn, "C-2", "app04", "10.99.6.4", "B2/8U", "DELL R660", hw_serial="SN0004")
    conn.commit()


def _key(conn, serial):
    return next(x["key"] for x in rl.machines(conn) if x["asset_serial"] == serial)


def _wb(buf):
    return load_workbook(io.BytesIO(buf.getvalue()))


def _save(wb) -> bytes:
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _row_of(ws, asset_serial):
    for r in range(ri.HEAD_ROW + 1, ws.max_row + 1):
        if ws.cell(r, 4).value == asset_serial:
            return r
    raise AssertionError(asset_serial)


def test_開始盤點_基準一箱一列_下載檔格式(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        a = ri.start_audit(conn, "甲機房", None, "alice")
        assert a["racks"] == 2 and a["boxes"] == 5                              # B1：app01、app02、FS5300；B2：2 台
        det = ri.audit_detail(conn, a["id"])
        b1 = next(r for r in det["racks"] if r["rack"] == "B1")
        assert b1["book_count"] == 3 and not b1["done"]
        sh = ri.rack_sheet(conn, a["id"], "甲機房", "B1")
        assert [r["name"] for r in sh["rows"]] == ["app01", "app02", "FakeStore-FS5300"]   # 從上到下
        assert "管理 IP：10.99.7.1（iDRAC）" in sh["rows"][0]["name_text"]
        assert "多管理 IP" in sh["rows"][2]["name_text"]
        buf, fname = ri.build_xlsx(conn, a["id"])
        wb = _wb(buf)
        assert wb.sheetnames[:2] == ["B1", "B2"] and wb[ri.META_SHEET].sheet_state == "veryHidden"
        ws = wb["B1"]
        assert [c.value for c in ws[ri.HEAD_ROW]] == ri.COLS
        assert ws.cell(5, 2).value == "30-31U" and ws.cell(5, 5).value == "SN0001" and ws.cell(5, 7).value == "在"
        assert ws.protection.sheet is True
        assert ws.cell(5, 4).protection.locked is True and ws.cell(5, 7).protection.locked is False
        assert ws.cell(5 + 3 + 1, 1).value == "新增1"                            # 3 列之後一列標題、再 5 列空白
        assert any("在,不在,不符" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)
    finally:
        conn.close()


def test_xlsx來回_部分上傳_重交取代_差異清單(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        aid = ri.start_audit(conn, "甲機房", None, "alice")["id"]
        wb = _wb(ri.build_xlsx(conn, aid)[0])
        ws = wb["B1"]
        ws["B2"], ws["E2"] = "王小明", "2026-10-27"
        ws.cell(_row_of(ws, "A-3"), 7).value = "不在"
        r = _row_of(ws, "A-1")
        ws.cell(r, 7).value, ws.cell(r, 8).value = "不符", "32-33"
        r = _row_of(ws, "F-1")
        ws.cell(r, 10).value = "78F9999"                                         # 結果還是「在」→ 自動視為不符
        e = 5 + 3 + 1                                                           # 第一列「新增」
        ws.cell(e, 2).value, ws.cell(e, 3).value, ws.cell(e, 5).value = "40", "Cisco C9300", "FOC123"
        res = ri.parse_upload(conn, aid, _save(wb), "alice", "a.xlsx")
        assert res["ok"], res
        assert [(s["rack"], s["version"]) for s in res["saved"]] == [("B1", 1)]
        assert res["skipped"][0]["rack"] == "B2"                                 # 沒填盤點人＝還沒盤
        assert any("視為「不符」" in w for w in res["warnings"])
        d = {x["kind"]: x for x in ri.diffs(conn, aid)}
        assert set(d) == {"missing", "u", "serial", "extra"}
        assert d["u"]["found"] == "32-33U" and d["missing"]["name"] == "app02" and d["extra"]["due"] == "2026-11-03"
        det = ri.audit_detail(conn, aid)
        assert [r["done"] for r in det["racks"]] == [True, False]
        # 重交 B1：只剩一個不在 → 取代，舊的留歷史
        ri.submit_rack(conn, aid, "甲機房", "B1", {_key(conn, "A-3"): {"result": "不在"}}, [], "王小明", "2026-10-28", "alice")
        assert {x["kind"] for x in ri.diffs(conn, aid)} == {"missing"}
        hist = ri.rack_sheet(conn, aid, "甲機房", "B1")["history"]
        assert [(h["version"], h["superseded"]) for h in hist] == [(2, 0), (1, 1)]
        # 網頁填 B2（全部在）→ 進度 2/2
        ri.submit_rack(conn, aid, "甲機房", "B2", {}, [], "李大華", None, "bob")
        s = ri.summary(conn, aid)
        assert (s["rack_done"], s["rack_total"], s["counted"], s["book_count"]) == (2, 2, 5, 5)
        assert s["diff_open"] == 1 and [x["name"] for x in s["serial_missing"]] == ["app02"]
        assert sorted(s["people"]) == ["李大華", "王小明"]
        # 上次盤點日期：整箱（含管理口）
        last = ri.last_audit_map(conn)
        assert last[_key(conn, "A-2")] == "2026-10-28" and last[_key(conn, "C-1")] == date.today().isoformat()
        buf, _ = ri.report_xlsx(conn, aid)
        assert set(_wb(buf).sheetnames) == {"摘要", "各櫃", "差異", "序號待補"}
    finally:
        conn.close()


def test_驗檔錯誤講到工作表與列(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        aid = ri.start_audit(conn, "甲機房", None, "alice")["id"]
        good = ri.build_xlsx(conn, aid)[0]
        wb = _wb(good)
        ws = wb["B1"]
        ws["B2"] = "王小明"
        ws.cell(5, 7).value = "壞了"
        ws.cell(6, 7).value = "不符"                                            # 不符但沒填正確值
        ws.cell(7, 8).value = "abc"
        res = ri.parse_upload(conn, aid, _save(wb), "alice")
        assert not res["ok"] and not res["saved"]
        assert "工作表 B1，第 5 列：結果欄只能是 在/不在/不符" in res["errors"]
        assert any(e.startswith("工作表 B1，第 6 列：") and "不符" in e for e in res["errors"])
        assert any(e.startswith("工作表 B1，第 7 列：正確 U") for e in res["errors"])
        assert not ri._active_submits(conn, aid)                                 # 有錯整份不存
        # 列被搬動（資產編號對不上）
        wb = _wb(good)
        ws = wb["B1"]
        ws["B2"] = "王小明"
        ws.cell(5, 4).value = "X-9"
        assert "被改過" in ri.parse_upload(conn, aid, _save(wb), "alice")["errors"][0]
        # 別次盤點的檔
        aid2 = ri.start_audit(conn, None, [{"room": "甲機房", "rack": "B2"}], "alice")["id"]
        with pytest.raises(ValueError, match="不是這次"):
            ri.parse_upload(conn, aid2, good.getvalue(), "alice")
        with pytest.raises(ValueError, match="不是 xlsx|讀不了"):
            ri.parse_upload(conn, aid, b"not a workbook", "alice")
        with pytest.raises(ValueError, match="沒有這一櫃"):
            ri.start_audit(conn, None, [{"room": "甲機房", "rack": "Z9"}], "alice")
    finally:
        conn.close()


def test_一鍵處理_差異狀態_CIA待更正(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        aid = ri.start_audit(conn, "甲機房", None, "alice")["id"]
        k1, k3, f1 = _key(conn, "A-1"), _key(conn, "A-3"), _key(conn, "F-1")
        ri.submit_rack(conn, aid, "甲機房", "B1", {
            k3: {"result": "不在"}, k1: {"result": "不符", "correct_u": "34-35"},
            f1: {"result": "不符", "correct_model": "IBM FS5200"}}, [{"model": "Cisco C9300", "serial": "FOC1", "u": "40"}],
            "王小明", "2026-10-27", "alice")
        d = {x["kind"]: x for x in ri.diffs(conn, aid)}
        out = ri.act(conn, aid, d["missing"]["diff_key"], "mark_removed", "alice")
        assert out["keys"] == [k3]
        assert next(m for m in rl.machines(conn) if m["key"] == k3)["kind"] == "removed"
        ri.act(conn, aid, d["u"]["diff_key"], "fix_position", "alice")
        m = next(m for m in rl.machines(conn) if m["key"] == k1)
        assert (m["u_start"], m["u_end"]) == (34, 35)
        m2 = next(m for m in rl.machines(conn) if m["key"] == _key(conn, "A-2"))
        assert m2["u_start"] == 34                                               # 管理口跟著整箱走
        d = {x["kind"]: x for x in ri.diffs(conn, aid)}
        assert d["missing"]["status"] == d["u"]["status"] == "已處理"
        assert not d["u"]["during_change"]                                       # 自己的一鍵處理不算盤點期間異動
        with pytest.raises(ValueError):
            ri.act(conn, aid, d["u"]["diff_key"], "fix_position", "alice")
        with pytest.raises(ValueError, match="原因"):
            ri.set_diff_status(conn, aid, d["model"]["diff_key"], "不處理", "", "alice")
        cia = [r for r in rl.pending_cia_rows(conn) if r["更正類型"].startswith("盤點")]
        assert {r["更正類型"] for r in cia} == {"盤點：型號不符", "盤點：櫃上有但沒登記"}
        assert next(r for r in cia if r["更正類型"] == "盤點：型號不符")["正確型號"] == "IBM FS5200"
        ex = next(r for r in cia if r["更正類型"] == "盤點：櫃上有但沒登記")
        assert (ex["正確型號"], ex["正確序號"], ex["新U位"]) == ("Cisco C9300", "FOC1", "40U")
        ri.set_diff_status(conn, aid, d["model"]["diff_key"], "不處理", "CIA 已經是對的", "alice")
        assert all(r["更正類型"] != "盤點：型號不符" for r in rl.pending_cia_rows(conn))
        assert all(set(rl.PENDING_CIA_COLS) <= set(r) for r in rl.pending_cia_rows(conn))
        buf, _ = rl.pending_cia_xlsx(conn)
        assert buf.getvalue()
    finally:
        conn.close()


def test_基準凍結_盤點期間異動_逾期(tmp_path, monkeypatch):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        aid = ri.start_audit(conn, "甲機房", None, "alice")["id"]
        k = _key(conn, "C-1")
        rl.move_box(conn, k, "甲機房", "B2", 12, 1, "有人搬了", "bob")              # 基準之後在系統裡搬
        _hw(conn, "N-1", "new01", "10.99.6.9", "B2/30U", "DELL R660")            # 基準之後新登記
        conn.commit()
        sh = ri.rack_sheet(conn, aid, "甲機房", "B2")
        assert [(r["name"], r["u"]) for r in sh["rows"]] == [("app04", "8U"), ("app03", "5U")]   # 基準沒變
        ri.submit_rack(conn, aid, "甲機房", "B2", {}, [{"model": "X", "serial": "S1"}], "王小明",
                       (date.today() - timedelta(days=10)).isoformat(), "alice")
        d = {x["kind"]: x for x in ri.diffs(conn, aid)}
        assert d["changed"]["key"] == k and "位置更正" in d["changed"]["during_change"][0]
        assert d["extra"]["overdue"] is True and d["extra"]["aging_days"] == 10
        dc = rl.datacenter(conn)
        room = next(r for r in dc["rooms"] if r["room"] == "甲機房")
        # B1 的 3 台、B2 新登記的 1 台沒盤過；B2 的兩台剛盤 → 不算
        assert len(room["audit_stale_keys"]) == 4 and room["audit_never_count"] == 4
        det = rl.rack_detail(conn, "甲機房", "B2")
        assert det["last_audit"][k] == (date.today() - timedelta(days=10)).isoformat()
        ri.close_audit(conn, aid, "alice")
        with pytest.raises(ValueError, match="結案"):
            ri.submit_rack(conn, aid, "甲機房", "B2", {}, [], "王小明", None, "alice")
    finally:
        conn.close()


def test_API_權限與守門(tmp_path, monkeypatch):
    from fastapi import HTTPException
    from routers import rackaudit as ra

    conn = _conn(tmp_path)
    try:
        _seed(conn)
        req = type("R", (), {"headers": {}, "client": type("C", (), {"host": "10.99.0.9"})()})()
        monkeypatch.setattr(ra, "_can_edit_nav", lambda s, c: False)
        with pytest.raises(HTTPException) as e:
            ra.rack_audit_start_endpoint(ra.AuditStartBody(room="甲機房"), req, None, conn)
        assert e.value.status_code == 403
        monkeypatch.setattr(ra, "_can_edit_nav", lambda s, c: True)
        called = []
        monkeypatch.setattr(ra, "_guard", lambda c, rq: called.append(1))
        a = ra.rack_audit_start_endpoint(ra.AuditStartBody(room="甲機房"), req, None, conn)
        assert called == [1] and a["boxes"] == 5
        with pytest.raises(HTTPException) as e:
            ra.rack_audit_detail_endpoint(999, None, conn)
        assert e.value.status_code == 404
        with pytest.raises(HTTPException) as e:
            ra.rack_audit_rack_save_endpoint(a["id"], ra.AuditRackBody(room="甲機房", rack="B1", counted_by=""), req, None, conn)
        assert e.value.status_code == 400
    finally:
        conn.close()


def test_verify_numbers_盤點不變式(tmp_path):
    import verify_numbers as v

    conn = _conn(tmp_path)
    try:
        _seed(conn)
        aid = ri.start_audit(conn, "甲機房", None, "alice")["id"]
        ri.submit_rack(conn, aid, "甲機房", "B1", {_key(conn, "A-3"): {"result": "不在"}}, [{"model": "X"}], "王小明", None, "alice")
        v.RESULTS.clear()
        v.check_rack_audit(conn)
        got = [r for r in v.RESULTS if r["page"] == "機櫃盤點"]
        assert len(got) == 5 and all(r["ok"] for r in got), got
    finally:
        conn.close()


def test_盤點結果點一櫃_列出帳上每一箱_未盤標未盤(tmp_path):
    """2.007.13 使用者：「盤點出來我看不到裡面的台數和狀況」——7.1 盤點結果點機櫃要列出那一櫃的設備。"""
    import inventory_flow as inv
    conn = _conn(tmp_path)
    _seed(conn)
    aid = ri.start_audit(conn, "甲機房", None, "alice")["id"]
    top = inv.result(conn, "rack", aid)
    b1 = next(r for r in top["rows"] if r["name"].endswith("B1"))
    d = inv.result(conn, "rack", aid, host=b1["key"])
    assert d["level"] == "item" and d["done"] is False
    assert d["total"] == b1["count"] and {r["result"] for r in d["rows"]} == {"未盤"}
    assert any("app01" in r["name"] and r["serial"] == "A-1" for r in d["rows"])
