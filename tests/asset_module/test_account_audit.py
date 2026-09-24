"""帳號盤點：匯入上次盤點／匯出本次（2026-09-18）。全部用假資料（192.0.2.x）。

使用者：「帳號盤點格式是用這個為準……匯出可選標準 A-R，跟全部兩種」。
欄位與規則一律用 account_export（09-03 已定的 18 欄與 type 規則）——這裡驗的是「只有一套」
加上今天新的三件事：沿用上次人工 type、保留不認得的欄、跟上次比較。
"""
import io
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import account_audit as aa  # noqa: E402
import account_export  # noqa: E402
import db  # noqa: E402

STD = list(account_export.STANDARD_COLUMNS)
# 使用者的檔案：type_id 欄被 Excel 截斷顯示成 type_（別名要認得），外加一欄系統不認得的 h_col
FILE_HEADER = [("type_" if c == "type_id" else c) for c in STD] + ["h_col"]
NORM_HEADER = STD + ["h_col"]


def _xlsx(rows, title_row=True) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    if title_row:
        ws.append(["上次帳號盤點（標題列，系統要能自己跳過）"])
    ws.append(FILE_HEADER)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _base_row(user, uid, shell, t_code, t_info, host="web1", ip="192.0.2.10", h="X"):
    d = {"system_id": "N-001", "system": "範例系統", "ap_department": "資訊部", "ap_owner": "王小明",
         "hostname": host, "ip_addr": ip, "username": user, "password": "x", "uid": uid, "gid": uid,
         "gecos": user, "home": "/home/" + user, "shell": shell, "type_": t_code, "type_info": t_info,
         "department": "資訊處架構部", "owner": "李小華",
         "login_status": "無法登入" if "nologin" in shell else "可登入", "h_col": h}
    return [d.get(c, "") for c in FILE_HEADER]


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    c = db.get_connection(p)
    db.insert_hardware(c, asset_serial="A-1", hostname="web1", ip="192.0.2.10", asset_status="使用中",
                       api_id="N-001", usage_unit="資訊部", user_name="王小明", custodian="李小華",
                       inventory_division="資訊處", inventory_department="架構部")
    acc = [("root", 0, "/bin/bash", "default"), ("alice", 1001, "/bin/zsh", "human"),   # alice 的 shell 變了
           ("bob", 1002, "/bin/bash", "human"), ("webit3scan", 1500, "/bin/bash", "service"),
           ("oracle", 1600, "/bin/bash", "service"), ("weird", 1700, "/bin/bash", None)]
    for u, uid, sh, kind in acc:
        c.execute("INSERT INTO host_account (ip, asset_serial, username, uid, gid, gecos, home, shell, "
                  "can_login, kind, is_sudoer, never_logged_in, source, first_seen, last_seen) "
                  "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                  ("192.0.2.10", "A-1", u, uid, uid, u, "/home/" + u, sh, 1, kind,
                   1 if u == "root" else 0, 0, "ssh", "2026-09-01", "2026-09-18 10:00:00"))
    c.commit()
    return c


def _import_baseline(c):
    rows = [
        _base_row("root", 0, "/bin/bash", "1", "最高權限帳號", h="H-root"),
        _base_row("alice", 1001, "/bin/bash", "3", "程式運行帳號", h="H-alice"),   # 人工判成 3
        _base_row("oracle", 1600, "/bin/bash", "4", "資料庫運行帳號"),            # 人工判成 4
        _base_row("carol", 1003, "/bin/bash", "7", "人員使用帳號"),               # 這次不見了
        _base_row("dave", 1004, "/bin/bash", "7", "人員使用帳號", host="db9", ip="192.0.2.99"),  # 整台沒收
    ]
    rows.append([""] * len(FILE_HEADER))                                            # 空白列
    blank = dict(zip(FILE_HEADER, [""] * len(FILE_HEADER)), hostname="web1", ip_addr="192.0.2.10")
    rows.append([blank[c] for c in FILE_HEADER])                                    # 沒帳號
    return aa.import_file(c, _xlsx(rows), "上次盤點.xlsx", "tester")


def test_欄位就是account_export那18欄_只有一套():
    assert aa.STANDARD_DEFAULT == STD
    assert STD[7] == "password" and "type_id" in STD, "H 欄＝password；type_ 是被截斷的 type_id"


def test_匯入_自己找表頭_認得截斷的type_保留不認得的欄(tmp_path):
    c = _conn(tmp_path)
    r = _import_baseline(c)
    assert r["rows"] == 5 and r["hosts"] == 2
    assert r["header"] == NORM_HEADER, "type_ 要正規化成 type_id；其餘順序原樣"
    assert r["unknown_columns"] == ["h_col"]
    assert r["skipped_count"] == 1 and "username" in r["skipped"][0]["reason"]


def test_標準匯出_順序照匯入檔_值來自account_export_未知欄沿用上次(tmp_path):
    c = _conn(tmp_path)
    _import_baseline(c)
    res = aa.build_export(c, "standard")
    assert res["columns"] == NORM_HEADER
    rows = {r[NORM_HEADER.index("username")]: r for r in res["rows"]}
    root = rows["root"]
    assert root[NORM_HEADER.index("h_col")] == "H-root"
    assert root[NORM_HEADER.index("password")] == "x"
    assert root[NORM_HEADER.index("department")] == "資訊處架構部"
    assert root[NORM_HEADER.index("owner")] == "李小華"
    assert rows["bob"][NORM_HEADER.index("h_col")] == "", "上次沒有這個帳號就沒得沿用"


def test_type_只自動填1257_上次人工判的沿用_判不出留空(tmp_path):
    """使用者 09-03：3／4／6「目前沒有邏輯，都要人工判斷」——系統不猜。"""
    c = _conn(tmp_path)
    _import_baseline(c)
    rows = {r[NORM_HEADER.index("username")]: r for r in aa.build_export(c, "standard")["rows"]}
    t = lambda u: (rows[u][NORM_HEADER.index("type_id")], rows[u][NORM_HEADER.index("type_info")])
    assert t("root") == ("1", "最高權限帳號")
    assert t("bob") == ("7", "人員使用帳號")
    assert t("webit3scan") == ("5", "自動化盤點帳號")
    assert t("alice") == ("3", "程式運行帳號"), "人工判過的比系統規則優先"
    assert t("oracle") == ("4", "資料庫運行帳號"), "系統判不出、上次人工填過 → 沿用"
    assert t("weird") == ("", ""), "判不出、上次也沒填 → 留空，不猜"


def test_比較_分清楚帳號不見與整台沒收(tmp_path):
    c = _conn(tmp_path)
    _import_baseline(c)
    res = aa.build_export(c, "standard")
    cmp = {rec.get("username"): (k, ch) for rec, k, ch in res["compare"]}
    assert cmp["root"][0] == "不變"
    assert cmp["alice"] == ("有變更", "shell")
    assert cmp["bob"][0] == "新增"
    assert cmp["carol"][0] == "上次有這次沒有", "同一台有收集、帳號不見了"
    assert cmp["dave"][0] == "本次未盤點到這台", "整台沒收集，不可以說帳號被刪"


def test_全部匯出用account_export全欄位加比較欄_標準版另開比較分頁(tmp_path):
    c = _conn(tmp_path)
    _import_baseline(c)
    std = aa.build_export(c, "standard")["columns"]
    allc = aa.build_export(c, "all")["columns"]
    assert "與上次盤點比較" not in std
    assert "與上次盤點比較" in allc and "is_sudoer" in allc, "全部＝系統知道的全部欄位"
    buf, _ = aa.export_xlsx(c, "standard")
    from openpyxl import load_workbook
    wb = load_workbook(buf)
    assert wb.sheetnames[:2] == ["帳號盤點", "與上次比較"], "比較結果另開分頁，不弄髒要交出去的那張"


def test_沒匯入過_用標準欄位且講明無從比較(tmp_path):
    c = _conn(tmp_path)
    res = aa.build_export(c, "standard")
    assert res["columns"] == STD
    assert all("無從比較" in k for _, k, _ in res["compare"])


def test_範例檔_欄位與代碼表():
    from openpyxl import load_workbook
    wb = load_workbook(aa.template_xlsx())
    hdr = [c.value for c in wb["帳號盤點"][1]]
    assert hdr == STD
    codes = [r[0] for r in wb["type 代碼"].iter_rows(min_row=2, values_only=True)]
    assert codes == [1, 2, 3, 4, 5, 6, 7, 8]


def test_dump匯出再匯回_內容一樣(tmp_path):
    """dump＝整包搬到另一台再匯回。帳號盤點的匯入要吃得回自己的 dump。"""
    import import_export
    c = _conn(tmp_path)
    _import_baseline(c)
    dump = import_export.build_dump(c, "account_audit")
    header, rows = aa.baseline_rows(c)
    assert header == NORM_HEADER and len(rows) == 5
    p2 = tmp_path / "other"
    p2.mkdir()
    db.init_db(p2 / "t.db")
    c2 = db.get_connection(p2 / "t.db")
    r = aa.import_file(c2, dump, "account_audit_20260918.webit3dump", "tester")
    assert r["rows"] == 5 and r["header"] == NORM_HEADER
    assert aa.baseline_rows(c2) == (header, rows), "搬過去要一模一樣"
    c2.close()


def test_別的來源的dump不收(tmp_path):
    import gzip
    import pytest
    c = _conn(tmp_path)
    fake = gzip.compress(json.dumps({"magic": "WEBIT3-IMPORT-DUMP", "version": 1, "source": "dynassets",
                                     "headers": ["ip"], "rows": [["192.0.2.1"]]}).encode())
    with pytest.raises(ValueError, match="不是帳號盤點"):
        aa.import_file(c, fake, "x.webit3dump", "t")
