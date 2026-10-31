"""A2 完整病歷表：舊資產頁每一個頁簽在 A2 都有一段（釘住清單，拿掉就紅）；每一段的筆數＝列數；
查不到要寫原因；收集過的機器軟體／服務／帳號／重建資料都看得到；匯出完整病歷。全部假資料。"""
import io
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "APP" / "asset-module" / "backend"))

import db  # noqa: E402
import v2_sections as vs  # noqa: E402

OLD_PAGE = ROOT / "APP" / "asset-module" / "frontend" / "pages" / "assets" / "[serial].vue"


def _conn(tmp_path):
    p = tmp_path / "t.db"
    db.init_db(p)
    return db.get_connection(p)


def _acc(name, **kw):
    a = {"username": name, "uid": 1000, "kind": "human", "can_login": True, "pw_status": "set", "is_sudoer": False,
         "sudo_nopasswd": False, "never_logged_in": False, "authorized_keys": 0}
    a.update(kw)
    return a


def _seed(conn):
    import account_inventory
    import service_inventory
    import software_inventory

    conn.execute("INSERT OR REPLACE INTO business_system (api_id, name) VALUES ('N-099', '假交易系統')")
    db.insert_hardware(conn, asset_serial="A-1", hostname="aix01", ip="10.99.1.18", physical_location="01_甲機房",
                       rack_no="B2/12-15U", device_model="IBM S1024", os="AIX 7.2", api_id="N-099")
    db.insert_hardware(conn, asset_serial="L-1", hostname="lnx01", ip="10.99.2.14", physical_location="01_甲機房",
                       rack_no="VM", device_model="(VM)", os="Red Hat Enterprise Linux 9", api_id="N-099", is_vm="Y")
    conn.commit()
    software_inventory.upsert_packages(conn, "10.99.2.14", "L-1", [
        {"name": "bash", "version": "5.1", "arch": "x86_64"}, {"name": "openssl", "version": "3.0", "arch": "x86_64"}], "ssh_rpm")
    service_inventory.upsert_services(conn, "10.99.2.14", "L-1", [{"port": 22, "proto": "tcp", "bind": "0.0.0.0"},
                                                                  {"port": 443, "proto": "tcp", "bind": "0.0.0.0"}], "ssh_ss")
    account_inventory.upsert_accounts(conn, "10.99.2.14", "L-1", [_acc("root", uid=0, is_sudoer=True), _acc("alice")],
                                      "ssh", os_info={"family": "rhel"})
    data = {"groups": [{"group": "儲存／檔案系統", "items": [
        {"id": "storage.fstab", "group": "儲存／檔案系統", "title": "fstab", "necessity": "A", "state": "collected", "label": "已收集",
         "files": [{"path": "/etc/fstab", "content": "/dev/mapper/rhel-root / xfs defaults 0 0"}]}]},
        {"group": "排程", "items": [{"id": "cron.user", "group": "排程", "title": "使用者 crontab", "necessity": "A",
                                    "state": "denied", "label": "權限不足", "files": []}]}],
        "summary": {"total_a": 2, "a_collected": 1}}
    conn.execute("INSERT INTO host_rebuild (asset_serial, ip, platform, data_json, collected_at) VALUES (?,?,?,?,?)",
                 ("L-1", "10.99.2.14", "linux", json.dumps(data, ensure_ascii=False), "2026-10-20 13:30:59"))
    conn.commit()


def test_舊資產頁每個頁簽在A2都有一段():
    src = OLD_PAGE.read_text(encoding="utf-8")
    m = re.search(r"const ASSET_TABS: AssetTab\[\] = \[([^\]]*)\]", src)
    old_tabs = re.findall(r"'([a-z_]+)'", m.group(1))
    assert len(old_tabs) >= 14
    extra = {"iis", "ocp_siblings"}                               # 舊頁的 IIS 樹、OCP 同叢集區塊
    for t in old_tabs + sorted(extra):
        assert t in vs.OLD_TAB_MAP, f"舊資產頁的「{t}」在 A2 沒有對應的段"
        for k in vs.OLD_TAB_MAP[t]:
            assert k in vs.SECTION_KEYS, f"{t} → {k} 不存在"
    # 2.x 新增的也要在
    for k in ("conn", "changes", "sources", "location", "existence", "rebuild", "same_system", "same_rack", "ocp"):
        assert k in vs.SECTION_KEYS
    assert {g for g, _, _ in vs.SECTIONS} == {g for g, _ in vs.GROUPS}
    # 主機履歷頁籤順序（藍圖）：概要（前端）之後依序
    assert [lab for _, lab in vs.GROUPS] == ["影響範圍", "硬體與位置", "軟體與服務", "帳號與安全", "關聯設備", "異動紀錄", "原始資料", "系統重建資料"]


def test_每段筆數等於列數_查不到要講原因(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        for q in ("aix01", "lnx01"):
            secs = vs.all_sections(conn, q)
            assert [s["key"] for s in secs] == vs.SECTION_KEYS
            for s in secs:
                assert not (s["unknown"] or "").startswith("讀取失敗"), (q, s["key"], s["unknown"])
                if s["unknown"]:
                    assert s["count"] is None and s["rows"] == []
                else:
                    assert s["count"] == len(s["rows"]), (q, s["key"])
            summ = {x["key"]: x for x in vs.summary(conn, q)["sections"]}
            for s in secs:
                assert summ[s["key"]]["count"] == s["count"] and summ[s["key"]]["unknown"] == s["unknown"]
                assert vs.section(conn, q, s["key"])["count"] == s["count"]
    finally:
        conn.close()


def test_收集過的機器_軟體服務帳號重建都看得到(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        s = {x["key"]: x for x in vs.all_sections(conn, "lnx01")}
        assert s["software"]["count"] == 2 and {r["name"] for r in s["software"]["rows"]} == {"bash", "openssl"}
        assert s["services"]["count"] == 2 and {r["name"] for r in s["services"]["rows"]} == {"tcp/22", "tcp/443"}
        assert s["accounts"]["count"] == 2 and {r["username"] for r in s["accounts"]["rows"]} == {"root", "alice"}
        rb = s["rebuild"]
        assert rb["count"] == 2 and rb["kind"] == "rebuild" and "/etc/fstab" in rb["rows"][0]["content"]
        assert rb["rows"][1]["state"] == "權限不足"
        assert s["same_system"]["count"] == 1 and s["same_system"]["rows"][0]["hostname"] == "aix01"
        a = {x["key"]: x for x in vs.all_sections(conn, "aix01")}
        assert a["software"]["unknown"].startswith("未納管")             # 沒收集＝講原因，不留白
        assert a["same_rack"]["unknown"] is None and a["cia"]["count"] > 5
        buf, name = vs.export_xlsx(conn, "lnx01")
        wb = load_workbook(io.BytesIO(buf.getvalue()))
        assert wb.sheetnames[0] == "摘要" and len(wb.sheetnames) == len(vs.SECTIONS) + 1 and name.startswith("主機履歷_lnx01")
        buf, name = vs.export_xlsx(conn, "lnx01", only="rebuild")
        assert load_workbook(io.BytesIO(buf.getvalue())).sheetnames == ["摘要", "系統重建資料"]
    finally:
        conn.close()


def test_verify_numbers_頁簽數字等於列數(tmp_path):
    import verify_numbers as vn

    conn = _conn(tmp_path)
    try:
        _seed(conn)
        vn.RESULTS.clear()
        vn.check_v2_sections(conn)
        got = [r for r in vn.RESULTS if r["page"] == "2.x 主機履歷"]
        assert got and all(r["ok"] for r in got), got
    finally:
        conn.close()


def test_原因不會出現兩層括號_重建項目有精簡狀態(tmp_path):
    assert vs.flat_reason("沒有觀察到連線（系統連不進去、或還沒讀取過服務）") == "沒有觀察到連線：系統連不進去、或還沒讀取過服務"
    assert vs.flat_reason("不在機櫃上（VM 或 CIA 沒填機櫃）") == "不在機櫃上：VM 或 CIA 沒填機櫃"
    assert "（" not in vs.flat_reason("讀取失敗：X（a（b））")
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        for q in ("aix01", "lnx01"):
            for s in vs.all_sections(conn, q):
                assert "（" not in (s["unknown"] or ""), (q, s["key"], s["unknown"])
        rb = vs.section(conn, "lnx01", "rebuild")["rows"]
        assert [(r["status"], r["problem"]) for r in rb] == [("ok", False), ("bad", True)]   # 必要項缺＝有問題
    finally:
        conn.close()


def _esxi_world(conn):
    """一台 ESXi（有資產登記）上跑 3 台 VM，2 台放在 ds:D1、1 台放在 ds:D2。"""
    db.insert_hardware(conn, asset_serial="E-1", hostname="esx01", ip="10.99.9.1", physical_location="01_甲機房",
                       rack_no="B5/10-11U", device_model="Dell R750", os="VMware ESXi 7.0", api_id="N-099")
    conn.execute("INSERT INTO ci_node (node_id,node_type,label,asset_serial,source,created_at,updated_at) "
                 "VALUES ('esxi:ESX01','esxi','esx01','E-1','t','t','t')")
    for i, ds in ((1, "D1"), (2, "D1"), (3, "D2")):
        db.insert_hardware(conn, asset_serial=f"V-{i}", hostname=f"vm0{i}", ip=f"10.99.9.1{i}", is_vm=1,
                           os="Red Hat Enterprise Linux 8.10 (Ootpa)", device_model="(VM)", rack_no="VM")
        conn.execute("INSERT INTO ci_node (node_id,node_type,label,asset_serial,source,created_at,updated_at) "
                     "VALUES (?, 'host', ?, ?, 't','t','t')", (f"hw:V-{i}", f"vm0{i}", f"V-{i}"))
        conn.execute("INSERT INTO ci_edge (src_node_id,dst_node_id,edge_type,confidence,source,created_at) "
                     "VALUES (?, 'esxi:ESX01', 'runs_on', '證據', 't', 't')", (f"hw:V-{i}",))
        conn.execute("INSERT OR IGNORE INTO ci_node (node_id,node_type,label,source,created_at,updated_at) "
                     "VALUES (?, 'datastore', ?, 't','t','t')", (f"ds:{ds}", ds))
        conn.execute("INSERT INTO ci_edge (src_node_id,dst_node_id,edge_type,confidence,source,created_at) "
                     "VALUES (?, ?, 'stores_on', '證據', 't', 't')", (f"hw:V-{i}", f"ds:{ds}"))
    conn.commit()


def test_掛了會影響_ESXi上的VM與儲存_筆數等於下鑽(tmp_path):
    conn = _conn(tmp_path)
    try:
        _seed(conn)
        _esxi_world(conn)
        s = {x["key"]: x for x in vs.all_sections(conn, "esx01")}
        assert s["imp_on"]["count"] == 3 and {r["hostname"] for r in s["imp_on"]["rows"]} == {"vm01", "vm02", "vm03"}
        assert {r["datastore"]: r["vms"] for r in s["imp_storage"]["rows"]} == {"D1": 2, "D2": 1}
        assert s["imp_callers"]["unknown"].startswith("沒有觀察到連入的連線")
        summ = vs.summary(conn, "esx01")
        t = summ["impact_total"]
        assert t["known"] and t["hosts"] == 3 and "承載／相依主機 3 台" in t["text"] and t["link"] is None   # 2.006.5：不再連舊版 3-2
        # 受影響業務系統：本系統一定列、寫重要度（找碴 C9：原本「影響 1 套」看不出那就是自己）
        rel = {r["api_id"]: r for r in s["imp_systems"]["rows"]}
        own = {(r["api_id"] or "").strip() for r in vs.Ctx(conn, "esx01").rows if (r["api_id"] or "").strip()}
        assert all(rel[a]["relation"] == "本系統" and rel[a]["severity"] in ("重大", "一般") for a in own)
        for x in summ["sections"]:
            if x["group"] == "impact" and x["count"] is not None:
                assert x["count"] == len(s[x["key"]]["rows"])
        v = {x["key"]: x for x in vs.all_sections(conn, "vm01")}
        dep = {r["name"]: r["kind"] for r in v["imp_depends"]["rows"]}
        assert dep.get("esx01") == "ESXi" and dep.get("D1") == "Datastore"
        assert v["imp_on"]["unknown"]                                       # VM 上面沒有東西
        a = {x["key"]: x for x in vs.all_sections(conn, "aix01")}
        assert a["imp_on"]["unknown"] and "CI 圖譜沒有此主機" in a["imp_on"]["unknown"]
        assert vs.summary(conn, "aix01")["impact_total"]["known"] is False
        assert {g for g, _ in vs.GROUPS} >= {"impact", "near"}
        assert {k for g, k, _ in vs.SECTIONS if g == "near"} == {"same_system", "same_rack", "ocp", "conn"}
    finally:
        conn.close()


def test_類型logo與版本_ESXi被標VM_NA主機名(tmp_path):
    import v2_views as v

    b = v.type_badge({"os": "Red Hat Enterprise Linux 8.10 (Ootpa)", "device_model": "Dell R750"})
    assert (b["logo"], b["version_label"]) == ("rhel", "RHEL 8.10")
    assert v.type_badge({"os": "CentOS 8 (64-bit)", "device_model": "(VM)", "is_vm": 1})["logo"] == "centos"
    assert v.type_badge({"os": "Ubuntu 22.04 LTS"})["version_label"] == "Ubuntu 22.04"
    assert v.type_badge({"os": "Linux"})["logo"] == "linux"
    w = v.type_badge({"os": "Windows Server 2019"})
    assert (w["logo"], w["version_label"]) == ("windows", "Windows Server 2019")
    a = v.type_badge({"os": "AIX 7.2"})
    assert (a["logo"], a["tile"], a["version_label"]) == (None, "AIX", "AIX 7.2")
    i = v.type_badge({"os": "V7R4M0", "device_model": "IBM AS400"})
    assert (i["logo"], i["tile"], i["version_label"]) == (None, "A400", "IBM i 7.4")
    e = v.type_badge({"os": "VMWare ESXi 6.7", "device_model": "(VM)", "is_vm": 1})
    assert (e["code"], e["logo"], e["vm"], e["note"]) == ("ESX", "esxi", False, "CIA 標成 VM，請確認")
    assert v.type_badge({"os": "Red Hat Enterprise Linux CoreOS 4.14", "device_model": "(VM)"})["logo"] == "openshift"
    assert v.type_badge({"os": None, "device_model": "Cisco C9300", "group_name": "網路設備"})["logo"] is None
    conn = _conn(tmp_path)
    try:
        db.insert_hardware(conn, asset_serial="N-1", hostname="N/A", ip="10.99.8.8", asset_name="核心交換器-01")
        db.insert_hardware(conn, asset_serial="N-2", hostname="N/A", ip="10.99.8.9")
        conn.commit()
        rows = {r["asset_serial"]: r for r in conn.execute("SELECT * FROM hardware")}
        assert v.name_of(rows["N-1"]) == "核心交換器-01" and v.name_of(rows["N-2"]) == "10.99.8.9"
    finally:
        conn.close()


def test_logo檔與後端代號一致():
    src = (ROOT / "APP" / "asset-module" / "frontend" / "composables" / "v2Logos.ts").read_text(encoding="utf-8")
    import v2_views as v
    keys = {k for _, k, _ in v._DISTRO} | {"linux", "windows", "esxi", "openshift", "oracle"}
    for k in keys:
        assert f'"{k}": {{' in src, f"後端會回 logo={k}，前端 v2Logos.ts 沒有"
    assert "<script" not in src and "href" not in src
