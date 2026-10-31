"""機櫃圖現場更正：在面板上雙擊改「機型／IP／序號／備註」（2026-10-29）。

以現場實際為準，但**不直接改 CIA 的 hardware 表**：更正存自己的表（跟 rack_position_override 同一套做法），
CIA 匯入不會蓋掉；CIA 之後改成跟更正一樣就自動標「CIA 已同步」、不再套用。
每次更正／還原都寫 rack_field_audit（誰、何時、改前改後），進「CIA 待更正」清單給 CIA 負責人照改。
還原＝刪掉更正，回到 CIA 的值。
"""
from __future__ import annotations

import ipaddress
import sqlite3

FIELDS = ("model", "ip", "serial", "remark")
FIELD_LABEL = {"model": "機型", "ip": "IP", "serial": "序號", "remark": "備註"}
REMARK_MAX = 200


def ensure_tables(conn: sqlite3.Connection) -> None:
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_field_override ("
        " key TEXT NOT NULL, field TEXT NOT NULL, value TEXT NOT NULL, cia_value TEXT,"
        " status TEXT NOT NULL DEFAULT 'active', note TEXT, updated_by TEXT, updated_at TEXT, synced_at TEXT,"
        " PRIMARY KEY (key, field))")
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_field_audit ("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, key TEXT NOT NULL, field TEXT NOT NULL, action TEXT NOT NULL,"
        " before_value TEXT, after_value TEXT, note TEXT, updated_by TEXT,"
        " updated_at TEXT NOT NULL DEFAULT (datetime('now','localtime')))")


def _table_exists(conn: sqlite3.Connection) -> bool:
    return conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='rack_field_override'").fetchone() is not None


def active(conn: sqlite3.Connection) -> dict[str, dict[str, dict]]:
    """所有生效中的更正：{key: {field: row}}。只讀（不建表），表還沒建就是沒有。"""
    if not _table_exists(conn):
        return {}
    out: dict[str, dict[str, dict]] = {}
    for r in conn.execute("SELECT * FROM rack_field_override WHERE status = 'active'"):
        out.setdefault(r["key"], {})[r["field"]] = dict(r)
    return out


def mark_synced(conn: sqlite3.Connection, items: list[tuple[str, str, str]]) -> None:
    """CIA 已經改成跟現場更正一樣：更正功成身退（紀錄留著）。items = [(key, field, value)]。"""
    for k, f, v in items:
        conn.execute("UPDATE rack_field_override SET status = 'synced', synced_at = datetime('now','localtime') "
                     "WHERE key = ? AND field = ? AND status = 'active'", (k, f))
        conn.execute("INSERT INTO rack_field_audit (key, field, action, before_value, note, updated_by) "
                     "VALUES (?,?,'cia_synced',?,?,'system')", (k, f, v, "CIA 已同步：已跟現場更正一致"))
    conn.commit()


def clean_value(conn: sqlite3.Connection, field: str, raw: str | None, model_std: str | None = None) -> str:
    """驗證並整理使用者輸入。錯就丟 ValueError（訊息直接給人看）。"""
    import model_naming

    if field not in FIELDS:
        raise ValueError("這個欄位不能在這裡改")
    s = (raw or "").strip()
    if not s:
        raise ValueError(f"{FIELD_LABEL[field]}不能空白（要回到 CIA 的值請按「還原」）")
    if field == "model":
        return model_naming.normalize(conn, s)
    if field == "serial":
        return model_naming.normalize_serial(conn, model_std, s)
    if field == "ip":
        try:
            return str(ipaddress.ip_address(s))
        except ValueError:
            raise ValueError(f"不是有效的 IP：{s}")
    if len(s) > REMARK_MAX:
        raise ValueError(f"備註最多 {REMARK_MAX} 字")
    return s


def _retired_sql() -> str:
    import lifecycle
    return lifecycle.retired_sql_list()


def _same_model(a: str | None, b: str | None) -> bool:
    """機型比對不分大小寫、空白折疊（更正值已被統一成全大寫，CIA 的寫法可能是 Dell r740）。"""
    return " ".join((a or "").split()).upper() == " ".join((b or "").split()).upper()


def _current(m: dict, field: str) -> str:
    """這台目前畫面上看到的值（已套用更正）。"""
    if field == "model":
        return m.get("model") or ""
    if field == "ip":
        return m.get("ip") or ""
    if field == "serial":
        return (m.get("hw_serials") or [""])[0]
    return m.get("remark") or ""


def set_edit(conn: sqlite3.Connection, ms: list[dict], key: str, field: str, value: str, note: str | None,
             by: str) -> dict:
    ensure_tables(conn)
    m = next((x for x in ms if x["key"] == key), None)
    if m is None:
        raise LookupError("查無這台")
    new = clean_value(conn, field, value, m.get("model"))
    before = _current(m, field)
    if new == before:
        raise ValueError("跟目前的值一樣，沒有改")
    old = conn.execute("SELECT cia_value, status FROM rack_field_override WHERE key = ? AND field = ?", (key, field)).fetchone()
    # 第一次更正時，「目前的值」就是 CIA 的值；之後沿用第一次記的。已退場（synced）又重新改的，CIA 現在的值就是舊更正值
    cia = old["cia_value"] if old and old["status"] == "active" else before
    if field == "ip":
        used = conn.execute("SELECT asset_serial, hostname FROM hardware WHERE ip = ? AND COALESCE(asset_status,'') NOT IN "
                            + _retired_sql(), (new,)).fetchall()
        mine = set(m.get("serials") or [])
        other = [u for u in used if u["asset_serial"] not in mine]
        if other:
            raise ValueError(f"IP {new} 已經被使用中的資產 {other[0]['asset_serial']}（{other[0]['hostname'] or '未填主機名稱'}）使用，請先確認")
    conn.execute(
        "INSERT INTO rack_field_override (key, field, value, cia_value, status, note, updated_by, updated_at) "
        "VALUES (?,?,?,?, 'active', ?, ?, datetime('now','localtime')) "
        "ON CONFLICT(key, field) DO UPDATE SET value = excluded.value, cia_value = excluded.cia_value, "
        "status = 'active', note = excluded.note, updated_by = excluded.updated_by, "
        "updated_at = excluded.updated_at, synced_at = NULL",
        (key, field, new, cia, (note or "").strip() or None, by))
    conn.execute("INSERT INTO rack_field_audit (key, field, action, before_value, after_value, note, updated_by) "
                 "VALUES (?,?,'set',?,?,?,?)", (key, field, before, new, (note or "").strip() or None, by))
    conn.commit()
    return {"key": key, "field": field, "label": FIELD_LABEL[field], "before": before, "after": new, "cia": cia,
            "name": m["name"]}


def clear_edit(conn: sqlite3.Connection, ms: list[dict], key: str, field: str, by: str) -> dict:
    """還原：刪掉這台這欄的更正，回到 CIA 的值。"""
    ensure_tables(conn)
    if field not in FIELDS:
        raise ValueError("這個欄位不能在這裡改")
    row = conn.execute("SELECT * FROM rack_field_override WHERE key = ? AND field = ? AND status = 'active'",
                       (key, field)).fetchone()
    if row is None:
        raise LookupError("這一欄沒有現場更正，不用還原")
    m = next((x for x in ms if x["key"] == key), None)
    conn.execute("DELETE FROM rack_field_override WHERE key = ? AND field = ?", (key, field))
    conn.execute("INSERT INTO rack_field_audit (key, field, action, before_value, after_value, note, updated_by) "
                 "VALUES (?,?,'clear',?,?,'還原回 CIA 的值',?)", (key, field, row["value"], row["cia_value"], by))
    conn.commit()
    return {"key": key, "field": field, "label": FIELD_LABEL[field], "before": row["value"], "after": row["cia_value"],
            "name": m["name"] if m else key}


def pending_cia(conn: sqlite3.Connection, ms: list[dict]) -> list[dict]:
    """「CIA 待更正」清單：生效中的現場更正（CIA 還沒跟上），給 CIA 負責人照改。"""
    if not _table_exists(conn):
        return []
    by_k = {m["key"]: m for m in ms}
    out = []
    for r in conn.execute("SELECT * FROM rack_field_override WHERE status = 'active' ORDER BY updated_at DESC"):
        m = by_k.get(r["key"])
        out.append({"key": r["key"], "name": m["name"] if m else r["key"], "room": m["room"] if m else None,
                    "rack": m["rack"] if m else None, "u_start": m["u_start"] if m else None,
                    "field": r["field"], "label": FIELD_LABEL.get(r["field"], r["field"]),
                    "cia_value": r["cia_value"], "value": r["value"], "note": r["note"],
                    "by": r["updated_by"], "at": r["updated_at"]})
    return out


XLSX_COLS = ["項次", "機房", "機櫃", "起始U", "主機名稱", "欄位", "CIA原值", "現場實際值", "更正人", "更正時間", "原因"]


def pending_xlsx(conn: sqlite3.Connection, ms: list[dict]):
    """CIA 待更正（機型／IP／序號／備註）的 xlsx：給 CIA 負責人一次照改。"""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = pending_cia(conn, ms)
    wb = Workbook()
    ws = wb.active
    ws.title = "CIA待更正_欄位"
    for c, name in enumerate(XLSX_COLS, 1):
        cell = ws.cell(1, c, name)
        cell.fill = PatternFill("solid", fgColor="00806A")
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = max(10, min(28, len(name) + 8))
    for i, r in enumerate(rows, 2):
        for c, v in enumerate([i - 1, r["room"], r["rack"], r["u_start"], r["name"], r["label"], r["cia_value"],
                               r["value"], r["by"], r["at"], r["note"]], 1):
            if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
                v = "'" + v                      # 自由輸入的備註／原因不能被 Excel 當公式執行
            ws.cell(i, c, v)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, len(rows)


def audit(conn: sqlite3.Connection, key: str | None = None, limit: int = 100) -> list[dict]:
    if not conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='rack_field_audit'").fetchone():
        return []
    q, args = "SELECT * FROM rack_field_audit", []
    if key:
        q += " WHERE key = ?"
        args.append(key)
    return [dict(r) for r in conn.execute(q + " ORDER BY id DESC LIMIT ?", (*args, max(1, min(int(limit), 500))))]

