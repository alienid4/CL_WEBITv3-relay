"""機櫃：系統主機部署圖的「機櫃」頁簽、4-5 機房架構圖、機櫃反查共用的一份計算。

資料只有兩欄：hardware.physical_location（機房）與 hardware.rack_no（機櫃／U 位）。
rack_no 寫法很雜（「M10/31U」「C5/34-35U」「1-1櫃(A1)/10U」「M03-M04」「218號-4F」「(VM)」…），
所以 parse_rack() 只認看得懂的幾種，**認不出來就保留原值、不猜**（rack=None, kind='unparsed'）。

一台＝system_stats.machine_key（同一台多筆登記只算一台，退役不列）。
VM（rack_no 寫「(VM)」或掛在 ESXi 上）不在機櫃裡：機房架構圖排除並計數；
系統主機部署圖的機櫃頁簽則讓 VM 沿用它所在 ESXi 主機的機櫃（標「經 ESXi xxx」）。

樓層：資料裡沒有。留 app_settings 掛鉤：
  rack_floor_map  JSON {"機櫃代碼前綴": "樓層"}，空的就不顯示樓層
  rack_floor_layout JSON（日後機櫃在樓面的排／位置），目前只原樣回傳，不畫假座標
"""
from __future__ import annotations

import json
import re
import sqlite3
from typing import Any

import fleet_cache as _fleet_cache

DEFAULT_RACK_U = 42           # 機櫃高度沒有資料，畫立面圖用 42U（可由 app_settings rack_height_u 改）
_NA = {"", "N/A", "NA", "-", "--", "無", "NONE", "NULL"}

_RE_RACK_U = re.compile(r"^\s*(?P<rack>[^/]+?)\s*/\s*(?P<u1>\d{1,2})\s*(?:-\s*(?P<u2>\d{1,2}))?\s*U\s*(?P<rest>[、,，].*)?$", re.I)
_RE_RACK_ONLY = re.compile(
    r"^\s*(?P<rack>[A-Za-z]{1,4}-?\d{1,3}|\d[A-Za-z]\d{1,3}|\d{1,2}-{1,2}\d{1,2}\s*櫃?\s*(?:\([A-Za-z]\d{1,3}\))?)\s*櫃?\s*$")
_RE_FLOOR_RACK = re.compile(
    r"^\s*(?P<floor>\d{1,2})\s*F\s*機房\s*(?:(?P<a>\d{1,2})-{1,2}(?P<b>\d{1,2})|(?P<code>[A-Za-z]\d{1,3}))\s*$", re.I)
_RE_DASH_RACK = re.compile(r"^(\d{1,2})-{1,2}(\d{1,2})\s*櫃?\s*(?:\([A-Za-z]\d{1,3}\))?$")


def display_name(hostname: str | None, asset_name: str | None = None, fallback: str | None = None) -> str:
    """主機名是空的或「N/A」時改用資產名稱（CIA 很多機器主機名填 N/A）。"""
    h = (hostname or "").strip()
    if h.upper() in _NA:
        return (asset_name or "").strip() or h or (fallback or "")
    return h


def parse_rack(raw: str | None) -> dict[str, Any]:
    """rack_no → {kind, rack, u_start, u_end, raw}。kind：empty／vm／rack／unparsed。

    認得的寫法：「機櫃/起-迄U」（之後用「、」接的其他 U 段只保留在 raw）、純機櫃代碼（M11、A-9櫃、1-2櫃）。
    其餘（大樓／樓層、多櫃範圍「M03-M04」、「網路櫃」）一律 unparsed，畫面顯示原值。"""
    v = (raw or "").strip()
    if v.upper() in _NA:
        v = ""
    if not v:
        return {"kind": "empty", "rack": None, "u_start": None, "u_end": None, "raw": None}
    if v.upper() in ("(VM)", "VM", "（VM）"):
        return {"kind": "vm", "rack": None, "u_start": None, "u_end": None, "raw": v}
    m = _RE_RACK_U.match(v)
    if m:
        u1 = int(m.group("u1"))
        u2 = int(m.group("u2")) if m.group("u2") else u1
        lo, hi = min(u1, u2), max(u1, u2)
        extra = []
        for a, b in re.findall(r"(\d{1,2})\s*(?:-\s*(\d{1,2}))?\s*U", m.group("rest") or "", re.I):
            a1, b1 = int(a), int(b) if b else int(a)
            extra.append([min(a1, b1), max(a1, b1)])
        return {"kind": "rack", "rack": _norm_rack(m.group("rack")), "u_start": lo, "u_end": hi, "raw": v,
                "extra_ranges": extra}
    m = _RE_RACK_ONLY.match(v)
    if m:
        return {"kind": "rack", "rack": _norm_rack(m.group("rack")), "u_start": None, "u_end": None, "raw": v}
    m = _RE_FLOOR_RACK.match(v)
    if m:
        rack = _norm_rack(m.group("code")) if m.group("code") else f"{int(m.group('a'))}-{int(m.group('b'))}"
        return {"kind": "rack", "rack": rack, "u_start": None, "u_end": None, "raw": v,
                "floor_hint": f"{int(m.group('floor'))}F"}
    return {"kind": "unparsed", "rack": None, "u_start": None, "u_end": None, "raw": v}


def _norm_rack(r: str) -> str:
    """同一個櫃的不同寫法收成一種：B08＝B8、A-9＝A9、「1--4櫃(A4)」＝「1-4」。"""
    r = r.strip()
    m = _RE_DASH_RACK.match(r)
    if m:
        return f"{int(m.group(1))}-{int(m.group(2))}"
    r = r.rstrip("櫃").strip()
    m = re.fullmatch(r"([A-Za-z]+)-?(\d+)", r)
    if m:
        return f"{m.group(1).upper()}{int(m.group(2))}"
    return r.upper()


def rack_prefix(rack: str | None) -> str:
    """機櫃排（M10→M、C5→C、1-3→1）；認不出排就回「其他」。"""
    m = re.match(r"^([A-Za-z]+)\d", rack or "")
    if m:
        return m.group(1).upper()
    m = re.match(r"^(\d+)-\d+$", rack or "")
    return m.group(1) if m else "其他"


def _setting_json(conn: sqlite3.Connection, key: str) -> Any:
    try:
        r = conn.execute("SELECT value FROM app_settings WHERE key = ?", (key,)).fetchone()
        return json.loads(r[0]) if r and r[0] else None
    except (sqlite3.OperationalError, json.JSONDecodeError, TypeError):  # 沒有表／值不是 JSON：當沒設定
        return None


def floor_of(floor_map: dict | None, rack: str | None) -> str | None:
    if not floor_map or not rack:
        return None
    best = None
    for prefix, floor in floor_map.items():
        if rack.upper().startswith(str(prefix).upper()) and (best is None or len(prefix) > len(best[0])):
            best = (prefix, floor)
    return best[1] if best else None


def rack_height(conn: sqlite3.Connection) -> int:
    v = _setting_json(conn, "rack_height_u")
    try:
        return int(v) if v else DEFAULT_RACK_U
    except (TypeError, ValueError):
        return DEFAULT_RACK_U


OFF_STATUS = ("停用", "閒置")       # 有機櫃＝關機未下架
SCRAP_STATUS = ("報廢",)            # 有機櫃＝待確認（可能已搬走只是沒改資料）
STATE_LABEL = {"active": "在用", "off": "關機未下架", "pending": "待確認"}
_GARBLED = re.compile(r"\ufffd|[\u00c0-\u00ff]{2,}")


OVERRIDE_STATES = ("active", "off", "removed")
OVERRIDE_LABEL = {"active": "確認在用", "off": "標記已關機未下架", "removed": "標記已不在櫃上", None: "清除標記"}


def ensure_state_table(conn: sqlite3.Connection) -> None:
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_state_override ("
        " key TEXT PRIMARY KEY, state TEXT NOT NULL, note TEXT, updated_by TEXT, updated_at TEXT)")
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_state_audit ("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, key TEXT NOT NULL, old_state TEXT, new_state TEXT,"
        " note TEXT, updated_by TEXT, updated_at TEXT NOT NULL DEFAULT (datetime('now','localtime')))")


def set_state_override(conn: sqlite3.Connection, key: str, state: str | None, note: str | None, by: str,
                       commit: bool = True) -> dict:
    """機櫃現場盤點結果（人工標記，每次都寫稽核紀錄）：
      active＝確認在用（報廢但其實還在用）；off＝標記已關機未下架；removed＝標記已不在櫃上（從機櫃拿掉）；
      None＝清除標記，回到依資產狀態判斷。"""
    ensure_state_table(conn)
    if state is not None and state not in OVERRIDE_STATES:
        raise ValueError("state 只能是 active、off、removed 或清除")
    old = conn.execute("SELECT state FROM rack_state_override WHERE key = ?", (key,)).fetchone()
    if state is None:
        conn.execute("DELETE FROM rack_state_override WHERE key = ?", (key,))
    else:
        conn.execute(
            "INSERT INTO rack_state_override (key, state, note, updated_by, updated_at) VALUES (?,?,?,?,"
            "datetime('now','localtime')) ON CONFLICT(key) DO UPDATE SET state = excluded.state, "
            "note = excluded.note, updated_by = excluded.updated_by, updated_at = excluded.updated_at",
            (key, state, note, by))
    conn.execute("INSERT INTO rack_state_audit (key, old_state, new_state, note, updated_by) VALUES (?,?,?,?,?)",
                 (key, old["state"] if old else None, state, note, by))
    if commit:
        conn.commit()
    return {"key": key, "state": state}


def set_state_bulk(conn: sqlite3.Connection, keys: list[str], state: str | None, note: str | None, by: str) -> dict:
    """一次套用現場盤點結果到多台（同一個交易；每台各一筆稽核紀錄）。"""
    if not keys:
        raise ValueError("沒有選任何設備")
    if state is not None and state not in OVERRIDE_STATES:
        raise ValueError("state 只能是 active、off、removed 或清除")
    for k in keys:
        set_state_override(conn, k, state, note, by, commit=False)
    conn.commit()
    return {"count": len(keys), "state": state}


def state_audit(conn: sqlite3.Connection, limit: int = 200) -> list[dict]:
    ensure_state_table(conn)
    return [dict(r) for r in conn.execute("SELECT * FROM rack_state_audit ORDER BY id DESC LIMIT ?", (limit,))]


def machines(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    """全站機器（跨請求快取，2.006.2）：主機履歷一次請求會呼叫三、四次（每次約 0.5 秒）。
    資料庫沒變才命中（fleet_cache，回深拷貝）；算法在 _machines，不變。"""
    return _machines_cached(conn)


@_fleet_cache.cached   # 跨請求快取：資料庫沒變就不重算（fleet_cache）
def _machines_cached(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    return _machines(conn)


def _machines(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    """全站機器（machine_key 逐台）＋機房／機櫃解析＋所屬系統＋存在可信度＋型號規格＋機櫃狀態。

    機櫃狀態（實體在不在櫃上的口徑，2026-10-22 定案；跟全站台數用的存在可信度是兩回事）：
      - 使用中＋有機櫃＝在用（active）：算電、算 U、算重量、算機櫃台數，不管存在可信度看不看得到；
        看不到的另標「未觀察到」（unobserved）
      - 停用／閒置＋有機櫃＝關機未下架（off）；報廢＋有機櫃＝待確認（pending，可能還在櫃上）
      - 人工標記（rack_state_override，有稽核紀錄）：active／off／removed（已不在櫃上，從機櫃拿掉）
      off／pending 佔 U、算重量、用電算 0（另列「若全部開機 +X W」），不算台數。
    退役又沒有可判讀機櫃的機器不列（備查，不在任何機櫃裡）。
    型號寫「(VM)」卻填了機櫃的，不放進機櫃（kind='vm_model'，資料品質清單列出來）。"""
    import blast_radius
    import existence
    import model_specs
    import system_stats

    groups: dict[str, list] = {}
    for r in conn.execute(
            "SELECT asset_serial, hostname, ip, asset_name, api_id, physical_location, rack_no, asset_status, "
            "device_model, hw_serial, custodian, owner, pending_status, os, environment "
            "FROM hardware ORDER BY asset_serial"):
        k = system_stats.lifecycle_prefix(r) + system_stats.machine_key(r["hostname"], r["ip"], r["asset_serial"])
        groups.setdefault(k, []).append(r)
    ex = {m["key"]: m for m in existence.compute(conn)["machines"]}
    specs = model_specs.index(conn)
    ensure_state_table(conn)
    overrides = {r["key"]: dict(r) for r in conn.execute("SELECT * FROM rack_state_override")}
    ensure_position_table(conn)
    positions = {r["key"]: dict(r) for r in conn.execute(
        "SELECT * FROM rack_position_override WHERE status = 'active'")}
    newly_synced: list[tuple[str, dict]] = []
    import rack_field_edit
    fedits = rack_field_edit.active(conn)
    fsynced: list[tuple[str, str, str]] = []
    out = []
    for k, rows in groups.items():
        retired = k.startswith("ret:")
        parsed = [parse_rack(r["rack_no"]) for r in rows]
        # 代表值：第一筆看得懂的機櫃；沒有就第一筆認不出的原值；都沒有才看是不是 VM
        pick = next((p for p in parsed if p["kind"] == "rack"), None) \
            or next((p for p in parsed if p["kind"] == "unparsed"), None) \
            or next((p for p in parsed if p["kind"] == "vm"), None) \
            or parsed[0]
        loc = next((r["physical_location"] for r in rows if (r["physical_location"] or "").strip()), None)
        cia = {"room": blast_radius.rack_room(loc), "rack": pick["rack"] if pick["kind"] == "rack" else None,
               "u_start": pick["u_start"], "u_end": pick["u_end"], "raw": pick["raw"]}
        pos = positions.get(k)
        if pos and (cia["room"], cia["rack"], cia["u_start"], cia["u_end"]) == (
                pos["room"], pos["rack"], pos["u_start"], pos["u_end"]):
            newly_synced.append((k, pos))      # CIA 已經改成跟現場一樣：更正功成身退
            pos = None
        if pos:
            pick = {"kind": "rack", "rack": pos["rack"], "u_start": pos["u_start"], "u_end": pos["u_end"],
                    "raw": pick["raw"], "floor_hint": pick.get("floor_hint")}
        if retired and pick["kind"] != "rack":
            continue
        rep = rows[0]
        model = next((" ".join(r["device_model"].split()) for r in rows if (r["device_model"] or "").strip()), None)
        # 現場更正（面板雙擊改的機型／IP／序號／備註）：CIA 已經一樣就功成身退，否則蓋過 CIA 的值
        cia_hw_serials = sorted({r["hw_serial"].strip() for r in rows if (r["hw_serial"] or "").strip()})
        fe = {}
        for f, row in fedits.get(k, {}).items():
            same = ((f == "model" and rack_field_edit._same_model(model, row["value"])) or (f == "ip" and (rep["ip"] or "") == row["value"])
                    or (f == "serial" and row["value"] in {s.upper().replace("-", "").replace(" ", "") for s in cia_hw_serials}))
            if same:
                fsynced.append((k, f, row["value"]))
            else:
                fe[f] = row
        if "model" in fe:
            model = fe["model"]["value"]
        kind = pick["kind"]
        if kind == "rack" and model and model.strip().upper() in ("(VM)", "VM", "（VM）"):
            kind = "vm_model"
        statuses = sorted({(r["asset_status"] or "").strip() for r in rows})
        e = ex.get(k)
        seen = bool(e and e["counted"])
        state = None
        ov = overrides.get(k)
        if kind == "rack":
            if retired:
                state = "off" if any(st in OFF_STATUS for st in statuses) else "pending"
            else:
                state = "active"
            if ov and ov["state"] in ("active", "off"):
                state = ov["state"]
            elif ov and ov["state"] == "removed":
                kind = "removed"
                state = None
        spec = model_specs.lookup(specs, model)
        apis = sorted({(r["api_id"] or "").strip() for r in rows if (r["api_id"] or "").strip()})
        out.append({
            "key": k, "asset_serial": rep["asset_serial"], "serials": [r["asset_serial"] for r in rows],
            "name": display_name(rep["hostname"], rep["asset_name"], rep["asset_serial"]),
            "ip": fe["ip"]["value"] if "ip" in fe else rep["ip"], "room": pos["room"] if pos else cia["room"], "api_ids": apis,
            "remark": fe["remark"]["value"] if "remark" in fe else None,
            "field_edits": {f: {"value": r["value"], "cia": r["cia_value"], "by": r["updated_by"], "at": r["updated_at"],
                                "note": r["note"]} for f, r in fe.items()},
            "position_corrected": bool(pos), "cia_position": cia if pos else None,
            "position_note": pos["note"] if pos else None, "position_by": pos["updated_by"] if pos else None,
            "position_at": pos["updated_at"] if pos else None,
            "model": model, "model_garbled": bool(model and _GARBLED.search(model)),
            "spec_model": spec["model"] if spec else None,
            "spec_generic": bool(spec and spec.get("is_generic")),
            "w": model_specs.known_w(spec),
            "kg": float(spec["weight_kg"]) if spec and spec.get("weight_kg") not in (None, "") and not spec.get("is_generic") else None,
            "system_level": spec.get("system_level") if spec else None,
            "vendor_move_required": bool(spec and spec.get("vendor_move_required")),
            "hw_serials": [fe["serial"]["value"]] if "serial" in fe else cia_hw_serials,
            "custodian": next((r["custodian"].strip() for r in rows if (r["custodian"] or "").strip()), None),
            "owner": next((r["owner"].strip() for r in rows if (r["owner"] or "").strip()), None),
            "pending_status": next((r["pending_status"] for r in rows if (r["pending_status"] or "").strip()), None),
            "pending_confirm": bool(e and e.get("pending_confirm")),
            "asset_names": sorted({r["asset_name"] for r in rows if r["asset_name"]}),
            "asset_status": ",".join(x for x in statuses if x) or None, "retired": retired,
            "state": state, "state_label": STATE_LABEL.get(state or "", None),
            "override": ov["state"] if ov else None,
            "kind": kind, "rack": pick["rack"] if kind in ("rack", "removed") else None,
            "u_start": pick["u_start"] if kind == "rack" else None, "u_end": pick["u_end"] if kind == "rack" else None,
            "unobserved": kind == "rack" and state == "active" and not seen,
            "rack_raw": pick["raw"], "floor_hint": pick.get("floor_hint"),
            # 機櫃欄一格寫了好幾段 U（「M10/27U、28U」）：其他段另外畫成一塊（同一筆登記，台數一筆）
            "extra_ranges": (pick.get("extra_ranges") or []) if kind == "rack" else [],
            "exist_level": e["level"] if e else None, "exist_score": e["score"] if e else None,
            "exist_label": e["level_label"] if e else "不明", "seen": seen,
            "_oses": sorted({(r["os"] or "").strip() for r in rows if (r["os"] or "").strip()}),
            "_envs": sorted({(r["environment"] or "").strip() for r in rows if (r["environment"] or "").strip()}),
            "_hw_ips": [r["ip"] for r in rows if (r["ip"] or "").strip()],
        })
    if newly_synced:
        _mark_synced(conn, newly_synced)
    if fsynced:
        rack_field_edit.mark_synced(conn, fsynced)
    _attach_vm_systems(conn, out)
    _attach_type_level_ips(conn, out)
    out.sort(key=lambda m: (m["room"], m["rack"] or "", m["u_start"] or 0, m["name"].lower()))
    return out


# ===== 位置更正（現場盤點發現 U10 其實在 U11）=====
# 跟其他 *_override 一樣存自己的表，CIA 匯入不會蓋掉；CIA 之後改成跟更正一樣時自動標「CIA 已同步」、
# 不再套用（紀錄留著）。每次更正／清除都寫 rack_position_audit。

def ensure_position_table(conn: sqlite3.Connection) -> None:
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_position_override ("
        " key TEXT PRIMARY KEY, asset_serial TEXT, room TEXT NOT NULL, rack TEXT NOT NULL,"
        " u_start INTEGER, u_end INTEGER, note TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'active',"
        " cia_room TEXT, cia_rack TEXT, cia_u_start INTEGER, cia_u_end INTEGER, cia_raw TEXT,"
        " updated_by TEXT, updated_at TEXT, synced_at TEXT)")
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_position_audit ("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, key TEXT NOT NULL, action TEXT NOT NULL,"
        " before_json TEXT, after_json TEXT, note TEXT, updated_by TEXT,"
        " updated_at TEXT NOT NULL DEFAULT (datetime('now','localtime')))")


def _mark_synced(conn: sqlite3.Connection, items: list[tuple[str, dict]]) -> None:
    for k, pos in items:
        conn.execute("UPDATE rack_position_override SET status = 'synced', synced_at = datetime('now','localtime') "
                     "WHERE key = ? AND status = 'active'", (k,))
        conn.execute("INSERT INTO rack_position_audit (key, action, before_json, note, updated_by) "
                     "VALUES (?, 'cia_synced', ?, 'CIA 已同步：機櫃欄已跟現場更正一致', 'system')",
                     (k, json.dumps(pos, ensure_ascii=False)))
    conn.commit()


def _overlap(a1, a2, b1, b2) -> bool:
    return bool(a1 and b1) and a1 <= b2 and b1 <= a2


def collisions(ms: list[dict], key: str, room: str, rack: str, u_start: int | None, u_end: int | None) -> list[dict]:
    """同櫃 U 位重疊的其他設備（只警告、不擋）。"""
    if not u_start:
        return []
    return [{k: m[k] for k in ("key", "asset_serial", "name", "u_start", "u_end")}
            for m in ms if m["key"] != key and m["kind"] == "rack" and m["room"] == room and m["rack"] == rack
            and _overlap(u_start, u_end, m["u_start"], m["u_end"])]


def _validate_position(room: str, rack: str, u_start, u_height, note: str) -> tuple[str, str, int | None, int | None]:
    room, rack, note = (room or "").strip(), (rack or "").strip(), (note or "").strip()
    if not room or not rack:
        raise ValueError("機房和機櫃都要填")
    if not note:
        raise ValueError("要寫備註（例：10/23 現場盤點）")
    if u_start in (None, ""):
        return room, rack, None, None
    try:
        u1 = int(u_start)
        h = int(u_height or 1)
    except (TypeError, ValueError) as exc:
        raise ValueError("U 位和 U 數要是整數") from exc
    if u1 < 1 or h < 1 or u1 + h - 1 > 60:
        raise ValueError("U 位超出範圍")
    return room, rack, u1, u1 + h - 1


def set_position(conn: sqlite3.Connection, key: str, room: str, rack: str, u_start, u_height, note: str, by: str,
                 ms: list[dict] | None = None, commit: bool = True) -> dict:
    """更正一台的位置（機房／機櫃／起始 U／U 數）。回這台更正後與同櫃 U 位重疊的設備（只警告）。"""
    ensure_position_table(conn)
    room, rack, u1, u2 = _validate_position(room, rack, u_start, u_height, note)
    rack = _norm_rack(rack)
    ms = ms if ms is not None else machines(conn)
    m = next((x for x in ms if x["key"] == key), None)
    if m is None:
        raise ValueError(f"查無此設備：{key}")
    cia = m["cia_position"] or {"room": m["room"], "rack": m["rack"] if m["kind"] == "rack" else None,
                                "u_start": m["u_start"], "u_end": m["u_end"], "raw": m["rack_raw"]}
    old = conn.execute("SELECT * FROM rack_position_override WHERE key = ?", (key,)).fetchone()
    conn.execute(
        "INSERT INTO rack_position_override (key, asset_serial, room, rack, u_start, u_end, note, status, cia_room, "
        "cia_rack, cia_u_start, cia_u_end, cia_raw, updated_by, updated_at, synced_at) "
        "VALUES (?,?,?,?,?,?,?,'active',?,?,?,?,?,?,datetime('now','localtime'),NULL) "
        "ON CONFLICT(key) DO UPDATE SET room=excluded.room, rack=excluded.rack, u_start=excluded.u_start, "
        "u_end=excluded.u_end, note=excluded.note, status='active', cia_room=excluded.cia_room, "
        "cia_rack=excluded.cia_rack, cia_u_start=excluded.cia_u_start, cia_u_end=excluded.cia_u_end, "
        "cia_raw=excluded.cia_raw, updated_by=excluded.updated_by, updated_at=excluded.updated_at, synced_at=NULL",
        (key, m["asset_serial"], room, rack, u1, u2, note, cia["room"], cia["rack"], cia["u_start"], cia["u_end"],
         cia["raw"], by))
    after = {"room": room, "rack": rack, "u_start": u1, "u_end": u2}
    conn.execute("INSERT INTO rack_position_audit (key, action, before_json, after_json, note, updated_by) "
                 "VALUES (?, 'set', ?, ?, ?, ?)",
                 (key, json.dumps(dict(old) if old else cia, ensure_ascii=False),
                  json.dumps(after, ensure_ascii=False), note, by))
    if commit:
        conn.commit()
    return {"key": key, **after, "collisions": collisions(ms, key, room, rack, u1, u2)}


class Occupied(ValueError):
    """目標 U 位已經有別台。"""


def _move_targets(conn: sqlite3.Connection, ms: list[dict], key: str, room: str | None, rack: str, u_start,
                  u_height) -> dict:
    """搬移的試算（不寫入）：哪些筆一起搬、各搬到哪、撞到誰。結構性錯誤（查無、櫃不存在、超出高度…）丟 ValueError，
    錯誤種類放在 exc.kind（unknown_rack／out_of_range／no_change／bad_input），讓確認框當場講。"""
    import overlap_queue

    def err(kind, msg):
        e = ValueError(msg)
        e.kind = kind
        return e
    m = next((x for x in ms if x["key"] == key), None)
    if m is None or m["kind"] != "rack" or not m.get("rack"):
        raise err("bad_input", "這台沒有登記在機櫃裡，不能拖拉搬移（請用「更正位置」）")
    unplaced = not m.get("u_start")        # 2026-10-28：「位置不明（只知道在這櫃）」的設備可以從清單拖進空的 U 位
    room = room or m["room"]
    if room != m["room"]:
        raise err("bad_input", "拖拉只能在同一個機房裡搬（跨機房請用「更正位置」）")
    new_rack = _norm_rack(str(rack or "").strip())
    if new_rack not in {x["rack"] for x in ms if x["room"] == room and x.get("rack")}:
        raise err("unknown_rack", f"{room} 沒有「{rack}」這一櫃")
    try:
        u1, h = int(u_start), int(u_height)
    except (TypeError, ValueError):
        raise err("bad_input", "起始 U 和 U 數要是整數") from None
    if h < 1 or h > 20 or u1 < 1:
        raise err("out_of_range", "U 數要 1–20、起始 U 至少 1")
    top = max(rack_height(conn), 1)
    devs = [x for x in ms if x["kind"] == "rack" and x["room"] == m["room"] and x["rack"] == m["rack"]]
    plan = overlap_queue.merge_plan(devs, m["room"], m["rack"], conn=conn)
    grp = next((g for g in plan["merged"] if key in g["keys"]), None)
    keys = set(grp["keys"]) if grp else {key}
    ports = {o for o, v in mgmt_pairs(devs).items() if v["host"] in keys}
    keys |= ports
    if unplaced:          # 位置不明的拖進來只搬它自己，不連帶搬走同櫃已有位置的同群組設備
        keys, ports = {key}, set()
    by_k = {x["key"]: x for x in ms}
    resize = (not unplaced) and new_rack == m["rack"] and h != (m["u_end"] - m["u_start"] + 1)
    shift = 0 if unplaced else u1 - m["u_start"]
    targets = {}
    for k in keys:
        x = by_k[k]
        if not x.get("u_start") and k != key:
            continue
        same_span = (x["u_start"], x["u_end"]) == (m["u_start"], m["u_end"])
        nh = h if (same_span or k == key) else (x["u_end"] - x["u_start"] + 1)
        ns = u1 if (same_span or k == key) else x["u_start"] + shift
        targets[k] = (ns, ns + nh - 1)
    for k in ports & set(targets):                    # 管理口不佔 U：跟著搬，超出櫃頂／櫃底就貼著本體
        a1, a2 = targets[k]
        if a1 < 1 or a2 > top:
            targets[k] = (u1, u1 + (a2 - a1))
    body = {k: v for k, v in targets.items() if k not in ports}
    if any(a1 < 1 or a2 > top for a1, a2 in body.values()):
        raise err("out_of_range", f"超出機櫃高度（1–{top}U）")
    if new_rack == m["rack"] and all(targets[k] == (by_k[k]["u_start"], by_k[k]["u_end"]) for k in targets):
        raise err("no_change", "位置沒有改變")
    tdevs = [x for x in ms if x["kind"] == "rack" and x["room"] == room and x.get("rack") == new_rack]
    tports = set(mgmt_pairs(tdevs))
    hit = []
    for x in tdevs:
        if x["key"] not in keys and x["key"] not in tports and x.get("u_start"):
            if any(_overlap(a1, a2, x["u_start"], x["u_end"]) for a1, a2 in body.values()):
                hit.append({"key": x["key"], "name": x["name"], "u": _u_text(x), "what": "device"})
    res_all = reservations(conn)
    for rv in [r for r in res_all if r["room"] == room and r["rack"] == new_rack] + legacy_reservations(res_all, room, new_rack):
        if not rv.get("expired") and \
                any(_overlap(a1, a2, rv["u_start"], rv["u_end"]) for a1, a2 in body.values()):
            hit.append({"key": None, "name": f"預留「{rv['label']}」", "u": f"{rv['u_start']}-{rv['u_end']}U",
                        "what": "reservation"})
    return {"m": m, "room": room, "rack": new_rack, "keys": keys, "ports": ports, "targets": targets, "hit": hit,
            "resize": resize, "top": top, "by_k": by_k}


def move_check(conn: sqlite3.Connection, key: str, room: str | None, rack: str, u_start, u_height) -> dict:
    """拖拉確認框邊打邊驗（2026-10-26）：不寫入，回 ok／錯誤種類／撞到誰。"""
    ms = machines(conn)
    try:
        t = _move_targets(conn, ms, key, room, rack, u_start, u_height)
    except ValueError as exc:
        return {"ok": False, "kind": getattr(exc, "kind", "bad_input"), "error": str(exc), "collisions": []}
    hd = t["targets"][key]
    if t["hit"]:
        return {"ok": False, "kind": "collision", "error": "目標位置已經有：" + "、".join(f"{h['name']}（{h['u']}）" for h in t["hit"][:6]),
                "collisions": t["hit"], "rack": t["rack"], "u_start": hd[0], "u_end": hd[1]}
    return {"ok": True, "kind": None, "error": None, "collisions": [], "rack": t["rack"], "u_start": hd[0], "u_end": hd[1],
            "count": len(t["targets"]), "names": sorted(t["by_k"][k]["name"] for k in t["targets"]), "top": t["top"]}


def ensure_move_log(conn: sqlite3.Connection) -> None:
    """拖拉搬移紀錄（2026-10-26 「回到上一步」）：一次拖拉一列，記下每筆搬之前的位置更正（沒有＝照 CIA）。"""
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_move_log ("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, room TEXT NOT NULL, head_key TEXT NOT NULL, head_name TEXT,"
        " from_json TEXT NOT NULL, to_json TEXT NOT NULL, before_json TEXT NOT NULL, after_json TEXT NOT NULL,"
        " note TEXT, moved_by TEXT, moved_at TEXT NOT NULL DEFAULT (datetime('now','localtime')),"
        " undone_at TEXT, undone_by TEXT, undo_note TEXT)")


def move_box(conn: sqlite3.Connection, key: str, room: str, rack: str, u_start, u_height, note: str, by: str) -> dict:
    """拖拉搬移／調高度（2026-10-25）：搬「一個箱子」——這台＋畫成同一台的其他筆＋它的管理卡。
    搬移：每一筆保持原本相對這台的位移與自己的高度；調高度：只改這台（以及跟它同一段 U 的同箱登記）的高度，
    管理卡這類自己 U 段不同的照舊。走既有的位置更正（逐筆稽核、進 CIA 待更正、可以清除還原），一個交易。
    目標 U 有別台或預留就拒絕（講出是誰）；目標櫃必須是這個機房已有的櫃；超出機櫃高度、位置沒變都拒絕。
    2026-10-26：每次搬移記進 rack_move_log（含每筆搬之前的位置更正），可以「復原」。"""
    ensure_position_table(conn)
    ensure_move_log(conn)
    ms = machines(conn)
    t = _move_targets(conn, ms, key, room, rack, u_start, u_height)
    if t["hit"]:
        raise Occupied("目標位置已經有：" + "、".join(f"{h['name']}（{h['u']}）" for h in t["hit"][:6]))
    m, by_k, targets = t["m"], t["by_k"], t["targets"]
    names = []
    before = {}
    for k in targets:
        r = conn.execute("SELECT * FROM rack_position_override WHERE key = ?", (k,)).fetchone()
        before[k] = dict(r) if r else None
    try:
        for k, (a1, a2) in sorted(targets.items()):
            set_position(conn, k, t["room"], t["rack"], a1, a2 - a1 + 1, note, by, ms=ms, commit=False)
            names.append(by_k[k]["name"])
        hd = targets[key]
        frm = {"rack": m["rack"], "u_start": m["u_start"], "u_end": m["u_end"]}
        to = {"rack": t["rack"], "u_start": hd[0], "u_end": hd[1]}
        after = {k: {"room": t["room"], "rack": t["rack"], "u_start": a1, "u_end": a2} for k, (a1, a2) in targets.items()}
        cur = conn.execute(
            "INSERT INTO rack_move_log (room, head_key, head_name, from_json, to_json, before_json, after_json, note, moved_by) "
            "VALUES (?,?,?,?,?,?,?,?,?)",
            (t["room"], key, m["name"], json.dumps(frm, ensure_ascii=False), json.dumps(to, ensure_ascii=False),
             json.dumps(before, ensure_ascii=False), json.dumps(after, ensure_ascii=False), note, by))
        conn.commit()
    except Exception:
        conn.rollback()          # 一筆失敗就整批退回，不留半套
        raise
    return {"moved": names, "room": t["room"], "rack": t["rack"], "u_start": hd[0], "u_end": hd[1], "resize": t["resize"],
            "from": frm, "move_id": cur.lastrowid, "name": m["name"]}


class MoveChanged(LookupError):
    """要復原的搬移之後，這些設備又被改過：要使用者確認才復原。"""

    def __init__(self, msg: str, changes: list[dict]):
        super().__init__(msg)
        self.changes = changes


def recent_moves(conn: sqlite3.Connection, room: str | None = None, rack: str | None = None, limit: int = 30) -> list[dict]:
    """最近的拖拉搬移（新到舊），每列可復原。rack 有給就只列搬進或搬出這櫃的。"""
    ensure_move_log(conn)
    q, args = "SELECT * FROM rack_move_log", []
    if room:
        q += " WHERE room = ?"
        args.append(room)
    rows = [dict(r) for r in conn.execute(q + " ORDER BY id DESC LIMIT ?", (*args, max(1, min(int(limit), 200)) * 4))]
    out = []
    for r in rows:
        frm, to = json.loads(r["from_json"]), json.loads(r["to_json"])
        if rack and _norm_rack(rack) not in (frm.get("rack"), to.get("rack")):
            continue
        out.append({"id": r["id"], "room": r["room"], "key": r["head_key"], "name": r["head_name"],
                    "from": frm, "to": to, "from_text": f"{frm['rack']}／{_u_txt(frm['u_start'], frm['u_end'])}",
                    "to_text": f"{to['rack']}／{_u_txt(to['u_start'], to['u_end'])}",
                    "count": len(json.loads(r["after_json"])), "note": r["note"], "by": r["moved_by"], "at": r["moved_at"],
                    "undone": bool(r["undone_at"]), "undone_at": r["undone_at"], "undone_by": r["undone_by"]})
        if len(out) >= limit:
            break
    return out


def undo_move(conn: sqlite3.Connection, move_id: int, by: str, force: bool = False, note: str | None = None) -> dict:
    """復原一次拖拉搬移：每一筆回到搬之前的樣子（之前有位置更正就還原那一筆；之前沒有＝刪掉更正、回 CIA 位置）。
    之後這台又被改過（又搬、清除、或原位置現在有別台）→ 丟 MoveChanged，force=True（使用者確認）才照做。逐筆進稽核。"""
    ensure_position_table(conn)
    ensure_move_log(conn)
    # 2026-10-26 審查 L4：「有沒有人改過」的檢查跟寫入要在同一個交易裡（BEGIN IMMEDIATE 先拿寫鎖），
    # 不然兩個人同時按復原、或復原時有人剛好又搬了，會照舊資料判斷然後蓋掉別人的改動。
    # machines() 可能自己 commit（CIA 已同步的位置更正），所以先算好再拿鎖；拿鎖之後的判斷全部重讀資料表。
    ms = machines(conn)
    if conn.in_transaction:
        conn.commit()
    conn.execute("BEGIN IMMEDIATE")
    try:
        return _undo_move_locked(conn, move_id, by, force, note, ms)
    except BaseException:
        if conn.in_transaction:
            conn.rollback()
        raise


def _undo_move_locked(conn: sqlite3.Connection, move_id: int, by: str, force: bool, note: str | None,
                      ms: list[dict]) -> dict:
    r = conn.execute("SELECT * FROM rack_move_log WHERE id = ?", (move_id,)).fetchone()
    if not r:
        raise LookupError("查無這筆搬移")
    if r["undone_at"]:
        raise ValueError("這筆已經復原過")
    before, after = json.loads(r["before_json"]), json.loads(r["after_json"])
    changes = []
    for k, a in after.items():
        cur = conn.execute("SELECT room, rack, u_start, u_end, status FROM rack_position_override WHERE key = ?", (k,)).fetchone()
        if not cur or cur["status"] not in ("active", "synced") or (cur["room"], cur["rack"], cur["u_start"], cur["u_end"]) != \
                (a["room"], a["rack"], a["u_start"], a["u_end"]):
            changes.append({"key": k, "what": "搬移之後這台的位置又被改過"})
    later = conn.execute("SELECT id, head_name, after_json FROM rack_move_log WHERE id > ? AND undone_at IS NULL", (move_id,)).fetchall()
    for lr in later:
        if set(json.loads(lr["after_json"])) & set(after):
            changes.append({"key": None, "what": f"之後又搬過（{lr['head_name']}，第 {lr['id']} 筆）"})
    # 原位置現在有沒有別台（搬走之後才有人搬進來：位置更正時間晚於這次搬移；原本就重疊的不算）
    keys = set(after)
    moved_in = {x["key"] for x in conn.execute("SELECT key FROM rack_position_override WHERE updated_at >= ?",
                                                 (r["moved_at"],))}
    for k, b in before.items():
        pos = b if b else None
        m = next((x for x in ms if x["key"] == k), None)
        if pos is None and m:
            cia = m.get("cia_position") or {}
            pos = {"room": cia.get("room"), "rack": cia.get("rack"), "u_start": cia.get("u_start"), "u_end": cia.get("u_end")}
        if not pos or not pos.get("u_start") or not pos.get("rack"):
            continue
        for x in ms:
            if x["key"] not in keys and x["key"] in moved_in and x["kind"] == "rack" and x["room"] == pos["room"] and x.get("rack") == pos["rack"] \
                    and _overlap(pos["u_start"], pos["u_end"], x.get("u_start"), x.get("u_end")) \
                    and not is_mgmt_record(x):
                changes.append({"key": x["key"], "what": f"原位置現在有 {x['name']}（{_u_text(x)}）"})
    if changes and not force:
        raise MoveChanged("這次搬移之後有變動，確認後才能復原", changes)
    # 寫入：跟上面的檢查同一個交易（呼叫端 undo_move 負責 BEGIN IMMEDIATE／失敗 rollback）
    for k, b in before.items():
        now = conn.execute("SELECT * FROM rack_position_override WHERE key = ?", (k,)).fetchone()
        if b is None:
            conn.execute("DELETE FROM rack_position_override WHERE key = ?", (k,))
        else:
            cols = list(b)
            conn.execute(f"INSERT OR REPLACE INTO rack_position_override ({', '.join(cols)}) VALUES "
                         f"({', '.join('?' for _ in cols)})", [b[c] for c in cols])
        conn.execute("INSERT INTO rack_position_audit (key, action, before_json, after_json, note, updated_by) "
                     "VALUES (?, 'undo_move', ?, ?, ?, ?)",
                     (k, json.dumps(dict(now) if now else None, ensure_ascii=False),
                      json.dumps(b, ensure_ascii=False), f"復原第 {move_id} 筆搬移" + (f"：{note}" if note else ""), by))
    conn.execute("UPDATE rack_move_log SET undone_at = datetime('now','localtime'), undone_by = ?, undo_note = ? WHERE id = ?",
                 (by, note, move_id))
    conn.commit()
    frm, to = json.loads(r["from_json"]), json.loads(r["to_json"])
    return {"id": move_id, "undone": True, "name": r["head_name"], "count": len(before), "forced": bool(changes),
            "back_to": frm, "from": to}


def clear_position(conn: sqlite3.Connection, key: str, note: str | None, by: str, commit: bool = True) -> dict:
    ensure_position_table(conn)
    old = conn.execute("SELECT * FROM rack_position_override WHERE key = ?", (key,)).fetchone()
    if not old:
        raise ValueError("這台沒有位置更正")
    conn.execute("DELETE FROM rack_position_override WHERE key = ?", (key,))
    conn.execute("INSERT INTO rack_position_audit (key, action, before_json, note, updated_by) VALUES (?, 'clear', ?, ?, ?)",
                 (key, json.dumps(dict(old), ensure_ascii=False), note, by))
    if commit:
        conn.commit()
    return {"key": key, "cleared": True}


def set_positions_bulk(conn: sqlite3.Connection, items: list[dict], note: str, by: str) -> dict:
    """依盤點結果一次更正多台（同一份備註；全部寫完才算重疊，回每台的重疊清單）。"""
    if not items:
        raise ValueError("沒有要更正的設備")
    ms = machines(conn)
    for it in items:
        set_position(conn, it["key"], it["room"], it["rack"], it.get("u_start"), it.get("u_height"), note, by,
                     ms=ms, commit=False)
    conn.commit()
    after = machines(conn)
    by_key = {m["key"]: m for m in after}
    out = []
    for it in items:
        m = by_key.get(it["key"])
        out.append({"key": it["key"], "collisions": collisions(after, it["key"], m["room"], m["rack"], m["u_start"],
                                                                m["u_end"]) if m else []})
    return {"count": len(items), "results": out}


def position_audit(conn: sqlite3.Connection, key: str | None = None, limit: int = 200) -> list[dict]:
    ensure_position_table(conn)
    if key:
        rows = conn.execute("SELECT * FROM rack_position_audit WHERE key = ? ORDER BY id DESC LIMIT ?", (key, limit))
    else:
        rows = conn.execute("SELECT * FROM rack_position_audit ORDER BY id DESC LIMIT ?", (limit,))
    return [dict(r) for r in rows]


def _u_txt(a, b) -> str:
    return "" if not a else (f"{a}U" if a == b else f"{a}-{b}U")


def pending_cia_rows(conn: sqlite3.Connection) -> list[dict]:
    """「CIA 待更正」：還在套用（CIA 還沒改成一樣）的位置更正。欄名沿用機房搬遷盤點表。"""
    ms = {m["key"]: m for m in machines(conn)}      # 先跑一次，CIA 已同步的會被標掉
    out = []
    for r in conn.execute("SELECT * FROM rack_position_override WHERE status = 'active' ORDER BY updated_at"):
        m = ms.get(r["key"])
        same_place = (r["room"] == r["cia_room"] and r["rack"] == r["cia_rack"])
        h_new = (r["u_end"] - r["u_start"]) if r["u_start"] else None
        h_old = (r["cia_u_end"] - r["cia_u_start"]) if r["cia_u_start"] else None
        kind = "高度更正" if same_place and h_new is not None and h_old is not None and h_new != h_old else "位置更正"
        out.append({"更正類型": kind, "資產序號": r["asset_serial"], "主機名稱": m["name"] if m else "",
                    "原機房": r["cia_room"] or "", "原Rack": r["cia_rack"] or (r["cia_raw"] or ""),
                    "原U位": _u_txt(r["cia_u_start"], r["cia_u_end"]),
                    "新機房": r["room"], "新Rack": r["rack"], "新U位": _u_txt(r["u_start"], r["u_end"]),
                    "設備U數": (r["u_end"] - r["u_start"] + 1) if r["u_start"] else "",
                    "盤點人": r["updated_by"] or "", "現場盤點日期": (r["updated_at"] or "")[:10],
                    "備註": r["note"] or "", "CIA原值（機櫃欄）": r["cia_raw"] or ""})
    # 現場標「已不在櫃上」的（例如位置重疊處理成「這筆已不在櫃上」）：CIA 要確認是下架、搬走還是重複登記
    ensure_state_table(conn)
    for r in conn.execute("SELECT * FROM rack_state_override WHERE state = 'removed' ORDER BY updated_at"):
        m = ms.get(r["key"])
        if not m:
            continue
        cr = m.get("cia_position") or {}
        out.append({"更正類型": "已不在櫃上", "資產序號": m["asset_serial"], "主機名稱": m["name"],
                    "原機房": m["room"], "原Rack": m["rack"] or (m["rack_raw"] or ""),
                    "原U位": _u_txt(cr.get("u_start"), cr.get("u_end")) if cr else "",
                    "新機房": "", "新Rack": "", "新U位": "", "設備U數": "",
                    "盤點人": r["updated_by"] or "", "現場盤點日期": (r["updated_at"] or "")[:10],
                    "備註": f"現場標已不在櫃上，請 CIA 確認下架／搬走／重複登記：{r['note'] or ''}",
                    "CIA原值（機櫃欄）": m.get("rack_raw") or ""})
    # 畫成同一台的（系統判斷或人工合併）：CIA 裡是好幾筆資產，其實是同一台 → 建議 CIA 合併
    import overlap_queue
    acts = overlap_queue.active_decisions(conn)
    by_rack: dict[tuple, list] = {}
    for m in ms.values():
        if m["kind"] == "rack" and m.get("rack"):
            by_rack.setdefault((m["room"], m["rack"]), []).append(m)
    for (rm, rk), devs in sorted(by_rack.items()):
        for mg in overlap_queue.merge_plan(devs, rm, rk, acts)["merged"]:
            mem = [ms[k] for k in mg["keys"] if k in ms]
            head = ms.get(mg["head"])
            out.append({"更正類型": f"建議合併（{len(mem)} 筆 CIA 資產其實是同一台）",
                        "資產序號": "、".join(x["asset_serial"] for x in mem), "主機名稱": "、".join(x["name"] for x in mem),
                        "原機房": rm, "原Rack": rk, "原U位": _u_txt(head["u_start"], head["u_end"]) if head else "",
                        "新機房": "", "新Rack": "", "新U位": "", "設備U數": "",
                        "盤點人": "系統判斷" if mg["auto"] else "人工合併", "現場盤點日期": "",
                        "備註": "依據：" + "、".join(mg["reasons"]) + (f"；代表那筆 {head['asset_serial']}" if head else ""),
                        "CIA原值（機櫃欄）": "、".join(sorted({x.get("rack_raw") or "" for x in mem}))})
    # 機櫃實體盤點（2026-10-27）：型號不符、序號不符／補登、櫃上有但沒登記 → 也要 CIA 改
    import rack_inventory
    out += rack_inventory.cia_rows(conn)
    for r in out:
        r.setdefault("正確型號", "")
        r.setdefault("正確序號", "")
    return out


PENDING_CIA_COLS = ["資產序號", "主機名稱", "原機房", "原Rack", "原U位", "新機房", "新Rack", "新U位", "設備U數",
                    "盤點人", "現場盤點日期", "備註", "CIA原值（機櫃欄）", "更正類型", "正確型號", "正確序號"]


def pending_cia_xlsx(conn: sqlite3.Connection):
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = pending_cia_rows(conn)
    wb = Workbook()
    ws = wb.active
    ws.title = "CIA待更正"
    for c, name in enumerate(PENDING_CIA_COLS, 1):
        cell = ws.cell(1, c, name)
        cell.fill = PatternFill("solid", fgColor="00806A")
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = max(10, min(24, len(name) + 4))
    for i, r in enumerate(rows, 2):
        for c, name in enumerate(PENDING_CIA_COLS, 1):
            ws.cell(i, c, r[name])
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, len(rows)


def room_volts(conn: sqlite3.Connection, room: str) -> float:
    """機房電壓（app_settings room_volts JSON {機房: V}；沒設定 220V）。"""
    v = (_setting_json(conn, "room_volts") or {}).get(room)
    try:
        return float(v) if v else 220.0
    except (TypeError, ValueError):
        return 220.0


def system_level_min(conn: sqlite3.Connection) -> int:
    try:
        return int(_setting_json(conn, "system_level_min_records") or 3)
    except (TypeError, ValueError):
        return 3


def rack_power(devs: list[dict], height: int, volts: float, sl_min: int = 3,
               same_box: dict[str, str] | None = None, no_u: set[str] | None = None) -> dict[str, Any]:
    """一櫃的用電／散熱／重量／U。只有規格有 PSU 瓦數的在用設備算進用電（每台一顆 PSU）；
    其他在用設備是「未知規格」。關機未下架／待確認：用電 0，另列「若全部開機 +X W」；佔 U、算重量。
    工程系統（Exadata／ODA 這類 system_level）：同櫃同型號 ≥ sl_min 筆時，整套數字只算一次並標明。
    known_count + unknown_count = active_count。"""
    active = [d for d in devs if d["state"] == "active"]
    dormant = [d for d in devs if d["state"] in ("off", "pending")]
    # 畫成同一台的（同一台的多筆登記）：用電、重量只算代表那一筆；台數照舊（known＋unknown 仍＝active）
    twin = {k for k, h in (same_box or {}).items() if k != h}
    sl_groups: dict[str, list] = {}
    for d in devs:
        if d.get("system_level") and d.get("spec_model"):
            sl_groups.setdefault(d["spec_model"], []).append(d)
    sl_used = {m: g for m, g in sl_groups.items() if len(g) >= sl_min}
    sl_keys = {d["key"] for g in sl_used.values() for d in g}
    w = 0.0
    known, unknown = [], []
    for d in active:
        if d["key"] in sl_keys:
            continue
        if d["key"] in twin:
            head = next((x for x in devs if x["key"] == same_box[d["key"]]), None)
            (known if head is not None and head["w"] is not None else unknown).append(d["key"])
            continue
        if d["w"] is not None:
            w += d["w"]
            known.append(d["key"])
        else:
            unknown.append(d["key"])
    sl_out = []
    for m, g in sl_used.items():
        sl = g[0]["system_level"] or {}
        sw = sl.get("psu_max_w")
        act = [d for d in g if d["state"] == "active"]
        if sw not in (None, "") and act:
            w += float(sw)
            known.extend(d["key"] for d in act)
        else:
            unknown.extend(d["key"] for d in act)
        sl_out.append({"model": m, "records": len(g), "w": sw, "kg": sl.get("weight_kg_max"), "u": sl.get("rack_u"),
                       # Codex 21 號 #3：這是「推定」，不是確認。規則與門檻明講，數字不變
                       "inferred": True, "sl_min": sl_min,
                       "rule": f"同櫃同型號≥{sl_min}筆推定為同一套"})
    if_on = sum(d["w"] for d in dormant if d["w"] is not None and d["key"] not in sl_keys and d["key"] not in twin)
    kg = sum(d["kg"] for d in devs if d["kg"] is not None and d["key"] not in sl_keys and d["key"] not in twin)
    kg += sum(float(x["kg"]) for x in sl_out if x["kg"] not in (None, ""))
    occupied: set[int] = set()
    for d in devs:
        if d["key"] in (no_u or set()):
            continue                       # 併進主機的管理口：不佔 U
        if d["u_start"]:
            occupied.update(range(d["u_start"], d["u_end"] + 1))
        for a, b in d.get("extra_ranges") or []:
            occupied.update(range(a, b + 1))
    best = run = 0
    for u in range(1, height + 1):
        run = 0 if u in occupied else run + 1
        best = max(best, run)
    used = len(occupied & set(range(1, height + 1)))
    # W221-001：在用、沒有 U 位、也不是併進主機的管理口 → 「未定位」。空 U 只是上限（這些設備可能佔了未知的 U）
    unplaced = [d["key"] for d in active if not d["u_start"] and d["key"] not in (no_u or set())]
    return {
        "unplaced_count": len(unplaced), "unplaced_keys": unplaced,
        "active_count": len(active), "off_count": sum(1 for d in dormant if d["state"] == "off"),
        "pending_count": sum(1 for d in dormant if d["state"] == "pending"),
        "w": round(w, 1), "amps": round(w / volts, 2) if volts else None, "btu_h": round(w * 3.412),
        "volts": volts, "kg": round(kg, 1), "if_all_on_w": round(if_on, 1),
        "used_u": used, "free_u": max(height - used, 0), "largest_free_u": best, "height_u": height,
        "known_count": len(known), "unknown_count": len(unknown), "unknown_keys": unknown,
        "unobserved_count": sum(1 for d in active if d.get("unobserved")),
        "unobserved_keys": [d["key"] for d in active if d.get("unobserved")],
        "active_keys": [d["key"] for d in active],
        "off_keys": [d["key"] for d in dormant if d["state"] == "off"],
        "pending_keys": [d["key"] for d in dormant if d["state"] == "pending"],
        "system_level": sl_out,
        "vendor_move": sorted({d["spec_model"] for d in devs if d.get("vendor_move_required") and d.get("spec_model")}),
    }


# ===== 4-1 平面圖的「顏色依」與網段（2026-10-23）=====
# 機器類型：只看資料（hardware.os、device_model、型號規格的 system_level、CI 圖的 ESXi 節點、san_switch），
# 對不到就是「不明」，不猜。順序有意義：一體機 > ESXi > 作業系統 > SAN > 儲存 > 網路 > Power 型號。
MACHINE_TYPES = ("esxi", "linux", "windows", "power", "oracle", "network", "san", "storage", "unknown")
MACHINE_TYPE_LABEL = {"esxi": "VM 主機（ESXi）", "linux": "實體 Linux", "windows": "實體 Windows",
                      "power": "AIX／Power", "oracle": "Oracle 一體機", "network": "網路設備",
                      "san": "SAN 交換器", "storage": "儲存", "unknown": "不明"}
_RE_ORACLE = re.compile(r"exadata|\bODA\b|oracle\s*database\s*appliance", re.I)
_RE_ESXI_OS = re.compile(r"esxi|vmware\s*esx", re.I)
_RE_AIX_OS = re.compile(r"\bAIX\b|VIOS|\bIBM\s*i\b|OS/?400", re.I)
_RE_WIN_OS = re.compile(r"windows", re.I)
_RE_LINUX_OS = re.compile(r"linux|red\s*hat|redhat|rhel|centos|rocky|debian|ubuntu|suse|sles|coreos|photon|alma", re.I)
# SAN 交換器：型號寫 SAN Switch／Brocade／Cisco MDS／Brocade 機型（DS-300B、G620…）；san_switch 表有這個 IP 也算
_RE_SAN = re.compile(r"san\s*switch|brocade|\bMDS\s*9\d{3}|\bDS-?\d{3}B\b|\bG[67]\d0\b", re.I)
_RE_STORAGE = re.compile(r"storage|flash\s*system|\bFS\d{4}\b|unity|powerstore|netapp|storwize|3par|"
                         r"\bDS8\d{3}\b|\bV[57]000\b|data\s*domain|avamar|pure\s*storage|isilon|powerscale|"
                         r"\bVNX\b|\bVMAX\b|tape|\bTS\d{4}\b|nimble", re.I)
_RE_NET = re.compile(r"cisco|aruba|fortinet|fortigate|\bFG-?\d|juniper|palo\s*alto|paloalto|\bPA-\d|\bF5\b|big-?ip|"
                     r"nexus|\bN9K|catalyst|\bWS-C|\bC9[0-9]{3}|\bISR\d|router|firewall|switch|"
                     r"\bAP\s*\d{3}\b|radware|check\s*point", re.I)
_RE_POWER_MODEL = re.compile(r"\bS9\d{2}\b|\bS10\d{2}\b|\bE9\d{2}\b|\bE10\d{2}\b|\bpower\s*\d|\bAS/?400\b", re.I)


def classify_machine_type(model: str | None, oses: list[str] | None = None, system_level: Any = None,
                          is_esxi: bool = False, is_san: bool = False) -> str:
    """一台實體機的類型（MACHINE_TYPES 其中一個）。認不出來回 'unknown'，不猜。"""
    md = model or ""
    os_txt = " ".join(oses or [])
    if system_level or _RE_ORACLE.search(md):
        return "oracle"
    if is_esxi or _RE_ESXI_OS.search(os_txt):
        return "esxi"
    if _RE_AIX_OS.search(os_txt):
        return "power"
    if _RE_WIN_OS.search(os_txt):
        return "windows"
    if _RE_LINUX_OS.search(os_txt):
        return "linux"
    if is_san or _RE_SAN.search(md):
        return "san"
    if _RE_STORAGE.search(md):
        return "storage"
    if "網路設備" in os_txt or _RE_NET.search(md):
        return "network"
    if _RE_POWER_MODEL.search(md):
        return "power"
    return "unknown"


SYS_LEVELS = ("core", "noncore", "unclassified")
SYS_LEVEL_LABEL = {"core": "核心", "noncore": "非核心", "unclassified": "未分級"}


def sys_class_level(sys_class: str | None) -> str:
    """business_system 的分類字串 → core／noncore／unclassified。
    221 實際值：「S1-第一類系統軟體（核心）」「S1-第一類系統軟體」「S2-…」「S3-…」、空白。
    S1／第一類／含「核心」＝核心；其他有寫的＝非核心；空白＝未分級。"""
    v = (sys_class or "").strip()
    if not v:
        return "unclassified"
    if "核心" in v or "第一類" in v or v.upper().startswith("S1"):
        return "core"
    return "noncore"


def machine_sys_level(api_ids, levels: dict[str, str]) -> str:
    """一台掛多個系統：取最高（核心 > 非核心 > 未分級）。"""
    got = {levels.get(a, "unclassified") for a in api_ids}
    for lv in SYS_LEVELS:
        if lv in got:
            return lv
    return "unclassified"


def _ipv4_int(ip: str | None) -> int | None:
    import ipaddress
    try:
        a = ipaddress.ip_address(str(ip or "").strip())
    except ValueError:
        return None          # 不是合法 IP（亂填、空白）＝沒有網段可對，不是程式錯誤
    return int(a) if a.version == 4 else None


class _SegIndex:
    """IP → 最小的那一段（network_segment，只認 IPv4）。一次讀進來、用 /16 分桶，避免逐 IP 查 SQL。"""

    def __init__(self, conn: sqlite3.Connection):
        self.buckets: dict[int, list[dict]] = {}
        self.wide: list[dict] = []
        rows = conn.execute("SELECT cidr, raw_cidr, net_start, net_end, purpose_desc, usage, vlan, environment, "
                            "environment_raw, location FROM network_segment WHERE net_start IS NOT NULL").fetchall()
        for r in rows:
            d = {"cidr": r["cidr"] or r["raw_cidr"], "start": r["net_start"], "end": r["net_end"],
                 "purpose": (r["purpose_desc"] or r["usage"] or "").strip() or None,
                 "vlan": (str(r["vlan"]).strip() if r["vlan"] not in (None, "") else None),
                 "env": r["environment"], "env_raw": r["environment_raw"], "location": r["location"]}
            b1, b2 = d["start"] >> 16, d["end"] >> 16
            if b2 - b1 > 64:
                self.wide.append(d)
            else:
                for b in range(b1, b2 + 1):
                    self.buckets.setdefault(b, []).append(d)

    def find(self, ip: str) -> dict | None:
        n = _ipv4_int(ip)
        if n is None:
            return None
        cand = [d for d in self.buckets.get(n >> 16, []) + self.wide if d["start"] <= n <= d["end"]]
        return min(cand, key=lambda d: d["end"] - d["start"]) if cand else None


def _split_ips(v: str | None) -> list[str]:
    return [x for x in re.split(r"[\s,;，、]+", v or "") if x and x.count(".") == 3]


def _attach_type_level_ips(conn: sqlite3.Connection, ms: list[dict]) -> None:
    """每台加：machine_type（＋label）、sys_level（＋label）、ips（全部 IP 與所屬網段）、envs（環境別）、os。
    IP＝hardware.ip（同一台的每筆登記）＋asset_nic.ip／ips（用這台全部資產序號去對）。"""
    # 這幾張表 db.init_db 都會建；查詢出錯就讓它報錯，不吞成「不明」（壞掉不能看起來像沒資料）
    def q(sql: str) -> list:
        return conn.execute(sql).fetchall()
    esxi_serials = {r[0] for r in q("SELECT asset_serial FROM ci_node WHERE node_type = 'esxi' AND gone_at IS NULL") if r[0]}
    san_ips = {(r[0] or "").strip() for r in q("SELECT ip FROM san_switch")}
    levels = {r["api_id"]: sys_class_level(r["sys_class_override"] or r["sys_class"])
              for r in q("SELECT api_id, sys_class, sys_class_override FROM business_system")}
    nic: dict[str, list[str]] = {}
    for r in q("SELECT asset_serial, ip, ips FROM asset_nic"):
        nic.setdefault(r["asset_serial"], []).extend([*_split_ips(r["ip"]), *_split_ips(r["ips"])])
    seg = _SegIndex(conn)
    for m in ms:
        ips: list[str] = []
        for ip in [*m.pop("_hw_ips", []), *(x for s in m["serials"] for x in nic.get(s, []))]:
            ip = (ip or "").strip()
            if ip and ip not in ips and not ip.startswith(("127.", "169.254.")):
                ips.append(ip)
        m["ips"] = []
        envs = set(m.pop("_envs", []))
        for ip in ips:
            d = seg.find(ip)
            m["ips"].append({"ip": ip, "cidr": d["cidr"] if d else None, "purpose": d["purpose"] if d else None,
                             "vlan": d["vlan"] if d else None, "env": d["env"] if d else None})
            if d:
                envs.update(x for x in (d["env"], d["env_raw"]) if x)
        m["envs"] = sorted(envs)
        oses = m.pop("_oses", [])
        m["os"] = ", ".join(oses) or None
        m["machine_type"] = classify_machine_type(
            m["model"], oses, m.get("system_level"),
            is_esxi=bool(esxi_serials & set(m["serials"])) or bool(m.get("vm_api_ids")),
            is_san=any(i in san_ips for i in ips))
        m["machine_type_label"] = MACHINE_TYPE_LABEL[m["machine_type"]]
        m["sys_level"] = machine_sys_level([*m["api_ids"], *m.get("vm_api_ids", [])], levels)
        m["sys_level_label"] = SYS_LEVEL_LABEL[m["sys_level"]]


# ===== 管理卡（iDRAC／iLO／OOB）併到主機面板（2026-10-24，只改顯示）=====
# CIA 常把一台機器的遠端管理卡另外登記成一筆（主機名含 OOB、OS 寫 iDRAC、資產名稱寫遠端管理），
# 跟主機同櫃同 U。畫面上那是同一個實體箱子：併進主機面板顯示「管理卡 <ip>」，不算第二台、不算位置重疊。
# 台數／去重／CIA 一律不動（那是危險區域 D3），這裡只回「哪一筆是哪一台的管理卡」。
_RE_MGMT_OS = re.compile(r"idrac|\bilom?\b|\bimm\b|\bxcc\b|cimc|\bbmc\b", re.I)
# 2026-10-26：主機名尾巴是 -ilom／-ilo／-idrac／-imm／-xcc／-bmc（也接受 _）＝那台的管理口
_RE_MGMT_SUFFIX = re.compile(r"[-_](ilom|ilo|idrac|imm|xcc|bmc)$", re.I)
_MGMT_LABEL = {"ilom": "ILOM", "ilo": "iLO", "idrac": "iDRAC", "imm": "IMM", "xcc": "XCC", "bmc": "BMC", "cimc": "CIMC"}


def mgmt_suffix_base(name: str | None) -> str | None:
    """「nh-oda-t1-ilom」→「nh-oda-t1」；不是管理口命名回 None。"""
    n = (name or "").strip()
    m = _RE_MGMT_SUFFIX.search(n)
    return n[:m.start()] if m else None


def mgmt_kind(m: dict) -> str:
    """管理口種類（顯示用）：主機名尾巴優先，再看 OS，都沒有就是 OOB。"""
    hit = _RE_MGMT_SUFFIX.search((m.get("name") or "").strip())
    if hit:
        return _MGMT_LABEL[hit.group(1).lower()]
    hit = _RE_MGMT_OS.search(m.get("os") or "")
    return _MGMT_LABEL.get(hit.group(0).lower(), hit.group(0).upper()) if hit else "OOB"


# 網路／資安／機房設備本身（OOB 交換器、OOB 防火牆、PDU／KVM）名字也常帶 OOB 或掛「遠端管理」，但它們是實體設備，不是某台的管理口
_RE_NOT_MGMT_MODEL = re.compile(r"cisco|fortinet|fortigate|\bFG-?\d|\basa\b|raritan|juniper|arista|brocade|palo ?alto|\bpdu\b|\bkvm\b|"
                                r"\bc9[0-9]{3}|\bn9k|nexus|catalyst", re.I)


# 2026-10-26 審查 M1／再審：管理口訊號分強弱。
#   強訊號：主機名尾巴 -ilom／-ilo／-idrac／-imm／-xcc／-bmc，或 OS／型號本身寫 iDRAC／iLO／ILOM／BMC／IMM／XCC
#           → 一定是管理口（ILOM 常被登記成本體的 OS「Oracle Linux」，不能因為 OS 是伺服器 OS 就不認）
#   弱訊號：名稱含「OOB」、資產名稱寫「遠端管理」→ SECOOB／SVROOB 開頭的有不少是 VM 或裝了完整伺服器 OS 的主機
#           （放在 OOB 網段的管理用伺服器，不是某台的管理卡），所以型號寫 (VM) 或 OS 是完整伺服器 OS 時不算。
_RE_SERVER_OS = re.compile(r"windows\s*server|red\s*hat|\brhel\b|centos|oracle\s*linux|\bsuse\b|\bsles\b|ubuntu|"
                           r"rocky|alma|debian|\baix\b|esxi|vmware", re.I)


def _mgmt_os_text(m: dict) -> str:
    return " ".join(x for x in [m.get("os") or ""] + list(m.get("_oses") or []) if x)


def mgmt_signal(m: dict) -> str | None:
    """管理口訊號：'strong'／'weak'／None（不是管理口）。網路／資安／機房設備本身一律 None。"""
    if _RE_NOT_MGMT_MODEL.search(m.get("model") or ""):
        return None
    os_txt = _mgmt_os_text(m)
    # OS 寫 iDRAC…：一台機器（machine_key）的登記裡同時有完整伺服器 OS 的，是「主機＋它的 iDRAC 用同一個主機名登記」，
    # 那一台是主機（SECSVR194-021：Windows Server 2022 一筆＋Idrac 一筆）→ 不當強訊號
    if mgmt_suffix_base(m.get("name")) or _RE_MGMT_OS.search(m.get("model") or "") \
            or (_RE_MGMT_OS.search(os_txt) and not _RE_SERVER_OS.search(os_txt)):
        return "strong"
    if "OOB" in (m.get("name") or "").upper() or any("遠端管理" in (a or "") for a in m.get("asset_names") or []):
        return "weak"
    return None


def is_mgmt_guarded_out(m: dict) -> bool:
    """弱訊號要排除的：型號寫 (VM)、是 VM、或 OS 是完整伺服器 OS。"""
    return ((m.get("model") or "").strip().upper() in ("(VM)", "VM", "（VM）") or m.get("kind") in ("vm", "vm_model")
            or bool(_RE_SERVER_OS.search(_mgmt_os_text(m))))


def is_mgmt_record(m: dict) -> bool:
    s = mgmt_signal(m)
    return s == "strong" or (s == "weak" and not is_mgmt_guarded_out(m))


#: 廠商（型號／管理口種類看得出來的才回；看不出來回 None）。2026-10-26 審查 M3：Dell R660 的 iDRAC
#: 被「同 U 只有一台」併進 Lenovo x3550——管理口跟本體廠商不同就不是它的管理口。
_VENDORS = (
    ("dell", re.compile(r"\bdell\b|poweredge|\bidrac\b|\br[67]\d0\b", re.I)),
    ("hpe", re.compile(r"\bhpe?\b|proliant|\bdl\d{3}\b|\bml\d{3}\b|\bilo\b", re.I)),
    ("lenovo", re.compile(r"lenovo|thinksystem|\bx3\d{3}\b|system\s*x|\bxcc\b|\bimm\d?\b", re.I)),
    ("oracle", re.compile(r"oracle|exadata|\boda\b|\bsun\b|\bilom\b", re.I)),
    ("cisco", re.compile(r"cisco|\bucs\b|\bcimc\b", re.I)),
    ("ibm", re.compile(r"\bibm\b|\bpower\s*\d|\bs9\d{2}\b|\be9\d{2}\b|as/?400", re.I)),
    ("supermicro", re.compile(r"supermicro", re.I)),
    ("fujitsu", re.compile(r"fujitsu|primergy|\birmc\b", re.I)),
    ("huawei", re.compile(r"huawei|\bibmc\b", re.I)),
)


def vendor_of(m: dict) -> str | None:
    """一台（或管理口）的廠商。先看型號，型號看不出來才看管理口種類（iDRAC＝Dell、iLO＝HPE…）。"""
    for txt in (m.get("model") or "", f"{m.get('name') or ''} {_mgmt_os_text(m)}"):
        for v, rx in _VENDORS:
            if rx.search(txt):
                return v
    return None


def vendor_compatible(port: dict, host: dict, strict: bool = False) -> bool:
    """管理口跟本體的廠商對不對得上。兩邊都看得出來而且不同＝不對；strict＝一定要兩邊都看得出來而且相同。"""
    a, b = vendor_of(port), vendor_of(host)
    if a and b:
        return a == b
    return not strict


def mgmt_pairs(devs: list[dict]) -> dict[str, dict]:
    """管理卡那一筆 key → {"host": 主機 key, "confirmed": 服務標籤對得上}。
    條件：同櫃、U 範圍完全相同的非管理卡設備；有好幾台時要靠服務標籤（管理卡名稱最後一段＝主機硬體序號）挑，
    挑不出來就不併（照舊分開顯示）。"""
    import blast_radius

    out: dict[str, dict] = {}
    hosts_by_name: dict[str, list] = {}
    for h in devs:
        if not is_mgmt_record(h):
            hosts_by_name.setdefault(blast_radius.host_match_key(h.get("name")) or "", []).append(h)
    for o in devs:
        if not is_mgmt_record(o):
            continue
        # 規則一（2026-10-26）：主機名去掉 -ilom／-ilo／-idrac／-imm／-xcc／-bmc 後＝同櫃某台的主機名（U 不必相同）；
        # 管理口直接用主機的名字登記（OS 寫 iDRAC、資產名稱「硬體遠端管理」）也算
        base = mgmt_suffix_base(o.get("name")) or o.get("name")
        hit = [h for h in hosts_by_name.get(blast_radius.host_match_key(base) or "\0", []) if vendor_compatible(o, h)]
        if len(hit) == 1:
            out[o["key"]] = {"host": hit[0]["key"], "confirmed": True, "rule": "hostname", "kind": mgmt_kind(o)}
            continue
        # 規則二：服務標籤（管理口名稱最後一段＝同櫃某台的硬體序號），U 不必相同
        tag = re.split(r"[-_\s]+", (o.get("name") or "").strip())[-1].upper()
        tagged_any = [h for h in devs if h is not o and not is_mgmt_record(h) and tag
                      and tag in {x.upper() for x in h.get("hw_serials") or []} and vendor_compatible(o, h)]
        if len(tagged_any) == 1:
            out[o["key"]] = {"host": tagged_any[0]["key"], "confirmed": True, "rule": "tag", "kind": mgmt_kind(o)}
            continue
        # 規則三（v1.501.16）：同櫃、U 範圍完全相同而且只有一台；2026-10-26 審查 M3：沒有名字／序號佐證，
        # 一定要廠商看得出來而且相同才併，否則留成「找不到本體」（不能把一台真的箱子藏進別台）
        if not o.get("u_start"):
            continue
        cands = [h for h in devs if h is not o and not is_mgmt_record(h)
                 and (h.get("u_start"), h.get("u_end")) == (o.get("u_start"), o.get("u_end"))]
        if len(cands) == 1 and vendor_compatible(o, cands[0], strict=True):
            out[o["key"]] = {"host": cands[0]["key"], "confirmed": False, "rule": "same_u", "kind": mgmt_kind(o)}
    return out


COUNTED = ("box", "orphan", "suspect")      # 算台數的類別（在用的）
ROLE_LABEL = {"head": "本體", "lpar": "分割區 LPAR", "component": "元件", "mgmt": "管理口", "multi_ip": "多管理 IP"}


def rack_boxes(devs: list[dict], plan: dict | None = None, room: str | None = None, rack: str | None = None,
               conn: sqlite3.Connection | None = None) -> dict[str, Any]:
    """機櫃畫面上「一個實體箱子」的歸類（2026-10-26）。每一筆登記剛好落在一個地方：
      box     箱子的本體（代表那筆；單獨一筆也是一個箱子）
      member  同一個箱子的其他登記（分割區 LPAR／元件／多管理 IP），U、用電、台數只算一次
      port    管理口（ILOM／iDRAC…）併進主機：不佔一列、不佔 U、不算電、不算台數
      orphan  管理口但找不到本體：照樣列出、標「管理介面・找不到本體」；照舊算一台、算電、佔 U
              （它有自己的 U 位，代表那裡有一台實體機，只是本體沒登記或登記在別處——不能讓它從台數消失）
      suspect 疑似登記錯誤（同一段 1–2U 擠了 4 筆以上）：不合併，各算一台、列出來請人看
    回 {category, head_of, ports_of, same_box, roles, boxes, suspect, orphans}。台數＝box＋orphan＋suspect（只算在用）。"""
    import blast_radius
    import overlap_queue

    if plan is None:
        plan = overlap_queue.merge_plan(devs, room or "", rack or "", conn=conn)
    mp = mgmt_pairs(devs)
    by = {d["key"]: d for d in devs}
    cat: dict[str, str] = {}
    head_of: dict[str, str] = {}
    roles: dict[str, str] = {}
    suspect_keys = {k for sg in plan.get("suspect", []) for k in sg["keys"]}
    for k in mp:
        cat[k] = "port"
        roles[k] = "mgmt"
    for d in devs:
        if is_mgmt_record(d) and d["key"] not in mp:
            cat[d["key"]] = "orphan"
    for k in suspect_keys:
        if k in by and k not in cat:
            cat[k] = "suspect"
    boxes = []
    for mg in plan["merged"]:
        keys = [k for k in mg["keys"] if k in by and k not in cat]
        if len(keys) < 2:
            continue
        head = mg["head"] if mg["head"] in keys else keys[0]
        hd = by[head]
        hkey = blast_radius.host_match_key(hd.get("name"))
        for k in keys:
            if k == head:
                cat[k] = "box"
                roles[k] = "head"
                continue
            cat[k] = "member"
            head_of[k] = head
            x = by[k]
            if x.get("machine_type") == "power" or re.search(r"\bAIX\b|OS/?400|V\dR\d", x.get("os") or "", re.I):
                roles[k] = "lpar"
            elif hkey and blast_radius.host_match_key(x.get("name")) == hkey:
                roles[k] = "multi_ip"
            else:
                roles[k] = "component"
        boxes.append({"head": head, "keys": [head] + [k for k in keys if k != head], "reasons": mg["reasons"],
                      "auto": mg["auto"], "decision_id": mg.get("decision_id")})
    for d in devs:
        if d["key"] not in cat:
            cat[d["key"]] = "box"
            roles[d["key"]] = "head"
    for k, v in mp.items():
        head_of[k] = head_of.get(v["host"], v["host"])         # 管理口歸到箱子的本體
    ports_of: dict[str, list] = {}
    for k, v in mp.items():
        ports_of.setdefault(head_of[k], []).append(k)
    same_box = {k: h for k, h in head_of.items()}
    return {"category": cat, "head_of": head_of, "ports_of": ports_of, "same_box": same_box, "roles": roles,
            "boxes": boxes, "suspect": plan.get("suspect", []), "orphans": [k for k, c in cat.items() if c == "orphan"],
            "mgmt": mp}


def mgmt_host_index(ms: list[dict]) -> dict[str, str]:
    """全站：管理口 key → 本體 key（搜尋 ILOM 的 IP 要找得到主機）。"""
    racks: dict[tuple, list] = {}
    for m in ms:
        if m["kind"] == "rack":
            racks.setdefault((m["room"], m["rack"]), []).append(m)
    out = {}
    for devs in racks.values():
        for k, v in mgmt_pairs(devs).items():
            out[k] = v["host"]
    return out


def overlaps(devs: list[dict], ignore: list[set] | None = None) -> dict[str, list[str]]:
    """同櫃 U 位重疊：key → 跟它重疊的其他設備 key（只看有 U 位的）。
    ignore：已經在「位置重疊待處理」按過「合併為同一台／確認是兩台」的組，組內互相不算重疊。"""
    out: dict[str, list[str]] = {}
    placed = [d for d in devs if d.get("u_start")]
    for i, a in enumerate(placed):
        for b in placed[i + 1:]:
            if ignore and any(a["key"] in g and b["key"] in g for g in ignore):
                continue
            if _overlap(a["u_start"], a["u_end"], b["u_start"], b["u_end"]):
                out.setdefault(a["key"], []).append(b["key"])
                out.setdefault(b["key"], []).append(a["key"])
    return out


def _attach_vm_systems(conn: sqlite3.Connection, ms: list[dict]) -> None:
    """ESXi 實體主機上跑的 VM 屬於哪些系統（vm_api_ids）：機房架構圖「標示系統」要連 VM 所在的櫃一起亮。
    VM→ESXi 走 CI 圖的 runs_on 邊，ESXi 節點的 asset_serial 對回實體機。"""
    by_serial = {s: m for m in ms for s in m["serials"]}
    esxi_serial = {r["node_id"]: r["asset_serial"] for r in conn.execute(
        "SELECT node_id, asset_serial FROM ci_node WHERE node_type = 'esxi' AND gone_at IS NULL")}
    for m in ms:
        m["vm_api_ids"] = []
    acc: dict[str, set] = {}
    for r in conn.execute("SELECT src_node_id, dst_node_id FROM ci_edge WHERE edge_type = 'runs_on' "
                          "AND gone_at IS NULL AND src_node_id LIKE 'hw:%'"):
        vm = by_serial.get(r["src_node_id"][3:])
        host = by_serial.get(esxi_serial.get(r["dst_node_id"]) or "")
        if vm and host and vm["api_ids"]:
            acc.setdefault(host["key"], set()).update(vm["api_ids"])
    for m in ms:
        m["vm_api_ids"] = sorted(acc.get(m["key"], set()) - set(m["api_ids"]))


def _system_guess(conn: sqlite3.Connection, apis) -> set[str]:
    """系統名稱是「從設備名稱猜的」（business_system.needs_review）的那些代碼：這種名稱常常是某個機型
    （例：I-010「Cisco Catalyst C2960X-48TD-L」），畫在面板的「系統」欄會被當成第二個機型，所以那欄只寫代碼。"""
    apis = sorted({a for a in apis if a})
    if not apis:
        return set()
    q = ",".join("?" * len(apis))
    return {r[0] for r in conn.execute(f"SELECT api_id FROM business_system WHERE needs_review = 1 AND api_id IN ({q})", apis)}


def _system_names(conn: sqlite3.Connection, apis) -> dict[str, str]:
    import blast_radius

    return {a: blast_radius._system_name(conn, a) for a in apis}


def datacenter(conn: sqlite3.Connection, ms: list[dict] | None = None) -> dict[str, Any]:
    """4-5 機房架構圖：每個機房的機櫃磚（在用台數、系統數、已用 U、用電）、無機櫃資料、排除的 VM 數、
    資料品質清單。room 台數＝在用機櫃台數＋無機櫃資料台數（VM、關機未下架、待確認不算台數，另計）。"""
    ms = ms if ms is not None else machines(conn)
    height = rack_height(conn)
    floor_map = _setting_json(conn, "rack_floor_map") or {}
    sl_min = system_level_min(conn)
    rooms: dict[str, dict] = {}
    dq = {"vm_model": [], "garbled": [], "generic": [], "unknown_model": []}
    for m in ms:
        r = rooms.setdefault(m["room"], {"room": m["room"], "racks": {}, "no_rack": [], "vm_excluded": 0})
        if m["kind"] == "removed":
            t = r["racks"].setdefault(m["rack"], {"rack": m["rack"], "prefix": rack_prefix(m["rack"]),
                                                  "floor": floor_of(floor_map, m["rack"]), "devs": [],
                                                  "systems": set(), "vm_systems": set(), "removed": []})
            t["removed"].append(m["key"])
            continue
        if m["kind"] == "vm":
            r["vm_excluded"] += 1
            continue
        if m["kind"] == "vm_model":
            dq["vm_model"].append(_brief(m))
            continue
        if m["kind"] != "rack":
            if not m["retired"]:
                # 2026-10-26 跟 7 對帳橋同一個說法：機櫃欄填了但認不出＝有機櫃位置、只是寫法看不懂（對帳橋照算實體設備）
                r["no_rack"].append({**_brief(m), "no_rack_reason":
                                     "機櫃欄認不出" if m["kind"] == "unparsed" else "缺機櫃資料"})
            continue
        if m["model_garbled"]:
            dq["garbled"].append(_brief(m))
        elif m["spec_generic"]:
            dq["generic"].append(_brief(m))
        elif not m["spec_model"]:
            dq["unknown_model"].append(_brief(m))
        t = r["racks"].setdefault(m["rack"], {"rack": m["rack"], "prefix": rack_prefix(m["rack"]),
                                              "floor": floor_of(floor_map, m["rack"]), "devs": [],
                                              "systems": set(), "vm_systems": set(), "removed": []})
        t["devs"].append(m)
        if not t["floor"] and m.get("floor_hint"):
            t["floor"] = m["floor_hint"]          # 沒設樓層對照時，用機櫃欄自己寫的樓層（例「21F 機房 2-1」）
        if m["state"] == "active":
            t["systems"].update(m["api_ids"])
            t["vm_systems"].update(m["vm_api_ids"])
    import overlap_queue
    acts = overlap_queue.active_decisions(conn)
    import rack_inventory
    last_audit = rack_inventory.last_audit_map(conn)
    out = []
    for r in rooms.values():
        volts = room_volts(conn, r["room"])
        tiles = []
        for t in sorted(r["racks"].values(), key=lambda t: (t["prefix"], _natural(t["rack"]))):
            plan = overlap_queue.merge_plan(t["devs"], r["room"], t["rack"], acts)
            bx = rack_boxes(t["devs"], plan)
            pw = rack_power(t["devs"], height, volts, sl_min, bx["same_box"], set(bx["mgmt"]))
            pw["removed_count"] = len(t["removed"])
            pw["removed_keys"] = t["removed"]
            cat = bx["category"]
            act = [d["key"] for d in t["devs"] if d["state"] == "active"]
            # 台數＝箱子（本體、找不到本體的管理口、疑似登記錯誤各算一台）；同箱其他登記、併進主機的管理口不算（2026-10-26）
            keys = [k for k in act if cat[k] in COUNTED]
            breakdown = {c: sum(1 for k in act if cat[k] == c) for c in ("box", "suspect", "member", "port", "orphan")}
            mp = bx["mgmt"]
            ov = overlaps([d for d in t["devs"] if d["key"] not in mp], plan["ignore"])   # 管理卡、同一台、已處理的組不算重疊
            tiles.append({"overlap_count": len(ov), "overlap_keys": sorted(ov), "all_keys": [d["key"] for d in t["devs"]],
                          "record_count": len(act), "breakdown": breakdown,
                          "category": {k: c for k, c in cat.items() if c != "box"}, "head_of": bx["head_of"],
                          "suspect": bx["suspect"], "orphan_keys": bx["orphans"],
                          "rack": t["rack"], "prefix": t["prefix"], "floor": t["floor"],
                          "device_count": len(keys), "system_count": len(t["systems"]),
                          "used_u": pw["used_u"], "height_u": height, "keys": keys,
                          "api_ids": sorted(t["systems"]), "vm_api_ids": sorted(t["vm_systems"] - t["systems"]),
                          "power": pw})
        if not tiles and not r["no_rack"] and not r["vm_excluded"]:
            continue
        # Excel 有、系統完全沒登記的櫃（例：B02）：也畫一個磚，點得開（只有 Excel 層；不算台數、用電）
        import rack_excel_import
        for eo in rack_excel_import.excel_only_racks(conn, r["room"], ms):
            pw0 = rack_power([], height, volts, sl_min, {}, set())
            pw0["removed_count"], pw0["removed_keys"] = 0, []
            tiles.append({"overlap_count": 0, "overlap_keys": [], "all_keys": [], "record_count": 0,
                          "breakdown": {c: 0 for c in ("box", "suspect", "member", "port", "orphan")},
                          "category": {}, "head_of": {}, "suspect": [], "orphan_keys": [],
                          "rack": eo["rack"], "prefix": rack_prefix(eo["rack"]), "floor": floor_of(floor_map, eo["rack"]),
                          "device_count": 0, "system_count": 0, "used_u": 0, "height_u": height, "keys": [],
                          "api_ids": [], "vm_api_ids": [], "power": pw0,
                          "excel_only": True, "excel_cells": eo["cells"]})
        tiles.sort(key=lambda t: (t["prefix"], _natural(t["rack"])))
        racked = sum(t["device_count"] for t in tiles)
        real = [t for t in tiles if not t.get("excel_only")]     # 合計（用電、可用 U…）只算系統真的有的櫃，Excel 獨有的櫃不灌水
        tot = {k: round(sum(t["power"][k] for t in real), 1)
               for k in ("w", "kg", "if_all_on_w", "used_u", "free_u", "known_count", "unknown_count",
                         "off_count", "pending_count", "active_count", "unobserved_count", "removed_count",
                         "unplaced_count")}
        tot["unplaced_keys"] = [k for t in real for k in t["power"]["unplaced_keys"]]
        tot["amps"] = round(tot["w"] / volts, 2) if volts else None
        tot["btu_h"] = round(tot["w"] * 3.412)
        tot["volts"] = volts
        tot["unknown_keys"] = [k for t in real for k in t["power"]["unknown_keys"]]
        tot["off_keys"] = [k for t in real for k in t["power"]["off_keys"]]
        tot["pending_keys"] = [k for t in real for k in t["power"]["pending_keys"]]
        tot["unobserved_keys"] = [k for t in real for k in t["power"]["unobserved_keys"]]
        tot["removed_keys"] = [k for t in real for k in t["power"]["removed_keys"]]
        # 機櫃實體盤點（2026-10-27）：超過一年未盤（含從沒盤過）的在用箱子
        room_keys = [k for t in real for k in t["keys"]]
        stale = rack_inventory.stale_keys(conn, room_keys, last_audit)
        out.append({"room": r["room"], "device_count": racked + len(r["no_rack"]), "racked": racked,
                    "racks": tiles, "no_rack": r["no_rack"], "vm_excluded": r["vm_excluded"], "power": tot,
                    "orphan_count": sum(t["breakdown"]["orphan"] for t in real),
                    "orphan_keys": [k for t in real for k in t["orphan_keys"]],
                    "audit_stale_keys": stale, "audit_never_count": sum(1 for k in stale if k not in last_audit)})
    out.sort(key=lambda r: (r["room"] == "未填", -r["device_count"], r["room"]))
    return {"rooms": out, "height_u": height, "floor_map_set": bool(floor_map),
            "floor_layout": _setting_json(conn, "rack_floor_layout"),
            "total_racked": sum(r["racked"] for r in out),
            "total_orphan": sum(r["orphan_count"] for r in out),
            "total_devices": sum(r["device_count"] for r in out),
            "data_quality": {k: v for k, v in dq.items()},
            "devices": {m["key"]: _brief(m) for m in ms if m["kind"] in ("rack", "removed")},
            "system_level_min": sl_min}


def _brief(m: dict) -> dict:
    return {k: m.get(k) for k in ("key", "asset_serial", "name", "ip", "api_ids", "rack_raw", "model", "spec_model",
                                  "room", "rack", "u_start", "u_end", "state", "state_label", "asset_status", "w", "unobserved", "kind",
                                  "position_corrected", "cia_position",
                                  "exist_level", "exist_score", "exist_label", "seen", "override",
                                  "machine_type", "machine_type_label", "sys_level", "sys_level_label", "ips", "envs")}


def _natural(s: str) -> list:
    return [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", s or "")]


def rack_detail(conn: sqlite3.Connection, room: str, rack: str, highlight_api: str | None = None,
                ms: list[dict] | None = None) -> dict[str, Any]:
    """一個機櫃的立面：每台設備的 U 位、名稱、所屬系統、存在可信度；highlight_api 的設備標出來，
    其餘算「同櫃其他設備」。也就是機櫃反查（這個櫃裡有哪些系統）。"""
    ms = ms if ms is not None else machines(conn)
    devs = [m for m in ms if m["kind"] == "rack" and m["room"] == room and m["rack"] == rack]
    import rack_excel_import
    excel_only = {e["rack"] for e in rack_excel_import.excel_only_racks(conn, room, ms)}
    if not devs and rack not in excel_only and not any(r["room"] == room and r["rack"] == rack for r in reservations(conn)) \
            and not legacy_reservations(reservations(conn), room, rack):
        raise ValueError(f"查無此機櫃：{room} {rack}")
    apis = sorted({a for m in devs for a in m["api_ids"]})
    names = _system_names(conn, apis)
    guess = _system_guess(conn, apis)
    height = rack_height(conn)
    items = []
    for m in devs:
        items.append({**{k: m[k] for k in ("key", "asset_serial", "name", "ip", "u_start", "u_end", "rack_raw", "model",
                                           "exist_level", "exist_score", "exist_label", "seen", "state",
                                           "state_label", "asset_status", "w", "kg", "spec_model", "override",
                                           "unobserved", "position_corrected", "cia_position", "position_note",
                                           "position_by", "position_at", "machine_type", "machine_type_label",
                                           "sys_level", "sys_level_label", "ips", "envs", "os", "hw_serials",
                                           "remark", "field_edits")},
                      "overlap_with": [],
                      "systems": [{"api_id": a, "name": names[a], "guess": a in guess} for a in m["api_ids"]],
                      "mine": bool(highlight_api and highlight_api in m["api_ids"])})
    items.sort(key=lambda d: (-(d["u_end"] or 0), d["name"].lower()))
    import overlap_queue
    plan = overlap_queue.merge_plan(devs, room, rack, conn=conn)
    bx = rack_boxes(devs, plan)
    mp = bx["mgmt"]
    for d in items:
        d["mgmt_of"] = mp[d["key"]]["host"] if d["key"] in mp else None
        d["mgmt"] = [{"key": d2["key"], "name": d2["name"], "ip": d2["ip"], "confirmed": mp[d2["key"]]["confirmed"],
                      "kind": mp[d2["key"]]["kind"], "rule": mp[d2["key"]]["rule"]}
                     for d2 in items if mp.get(d2["key"], {}).get("host") == d["key"]]
        d["box_role"] = bx["category"][d["key"]]
        d["role_label"] = ROLE_LABEL.get(bx["roles"].get(d["key"], ""), "")
        d["orphan_port"] = bx["category"][d["key"]] == "orphan"
        d["mgmt_kind"] = mgmt_kind(d) if is_mgmt_record(d) else None
        d["suspect"] = bx["category"][d["key"]] == "suspect"
        d["extra_ranges"] = next((m.get("extra_ranges") or [] for m in devs if m["key"] == d["key"]), [])
    ov = overlaps([d for d in items if d["key"] not in mp], plan["ignore"])
    for d in items:
        d["overlap_with"] = ov.get(d["key"], [])
    # 同一台（人工合併＋系統判斷）：畫面上併成一塊面板（代表＝有硬體序號的那筆），其他列在面板裡；台數、CIA 不動
    here = {d["key"]: d for d in items}
    for mg in plan["merged"]:
        members = [here[k] for k in mg["keys"] if k in here]
        if len(members) < 2:
            continue
        head = here[mg["head"]]
        head["merged"] = [{"key": d["key"], "name": d["name"], "ip": d["ip"],
                           "role": ROLE_LABEL.get(bx["roles"].get(d["key"], ""), "")} for d in members if d is not head]
        head["merge_reason"] = "、".join(mg["reasons"])
        head["merge_auto"] = mg["auto"]
        head["merge_keys"] = mg["keys"]
        head["merge_decision_id"] = mg["decision_id"]
        for d in members:
            if d is not head:
                d["merged_into"] = head["key"]
    systems = [{"api_id": a, "name": names[a], "device_count": sum(1 for d in items if a in
                                                                    [s["api_id"] for s in d["systems"]])}
               for a in apis]
    floor_map = _setting_json(conn, "rack_floor_map") or {}
    floor = floor_of(floor_map, rack) or next((m["floor_hint"] for m in devs if m.get("floor_hint")), None)
    pw = rack_power(devs, height, room_volts(conn, room), system_level_min(conn), bx["same_box"], set(mp))
    res_all = reservations(conn)
    rsv = [r for r in res_all if r["room"] == room and r["rack"] == rack]
    legacy_rsv = legacy_reservations(res_all, room, rack)
    legacy_aud = legacy_audit_snapshots(conn, room, rack)
    key_warnings = []
    if legacy_rsv:
        key_warnings.append(f"有 {len(legacy_rsv)} 筆舊鍵預留（機房「分公司」，分點不明，待確認）；容量保守視為佔用，不歸給任何分點")
    if legacy_aud:
        key_warnings.append(f"有 {len(legacy_aud)} 份舊鍵盤點快照（機房「分公司」，分點不明，待確認）；未改寫、未歸給任何分點")
    if room == LEGACY_BRANCH_ROOM:
        key_warnings.append("分點不明（登記的機房只寫「分公司」沒有分點名），待確認")
    import rack_inventory
    last_audit = rack_inventory.last_audit_map(conn)      # 每一箱的上次盤點日期（2026-10-27）
    item_keys = {d["key"] for d in items}
    import rack_excel_import
    return {"room": room, "rack": rack, "floor": floor, "height_u": height, "power": pw, "reservations": rsv,
            "legacy_reservations": legacy_rsv, "legacy_audits": legacy_aud, "key_warnings": key_warnings,
            "excel": rack_excel_import.rack_cells(conn, room, rack, ms),
            "devices": items, "systems": systems,
            "last_audit": {k: v for k, v in last_audit.items() if k in item_keys},
            "box_count": sum(1 for d in devs if d["state"] == "active" and bx["category"][d["key"]] in COUNTED),
            "suspect": [{"keys": sg["keys"], "reason": sg["reason"], "names": [x["name"] for x in items if x["key"] in sg["keys"]]}
                        for sg in bx["suspect"]],
            "orphan_ports": bx["orphans"],
            "excel_only": rack in excel_only,
            "room_racks": sorted({m["rack"] for m in ms if m["kind"] == "rack" and m["room"] == room and m.get("rack")} | excel_only,
                                 key=_natural),
            "mine_count": sum(1 for d in items if d["mine"]),
            "others_count": sum(1 for d in items if not d["mine"]),
            "no_u_count": sum(1 for d in items if not d["u_start"])}


def system_racks(conn: sqlite3.Connection, smap: dict[str, Any], ms: list[dict] | None = None) -> dict[str, Any]:
    """系統主機部署圖「機櫃」頁簽：本系統每台主機落在哪個機房／機櫃／U。
    實體機用自己的機櫃；VM 沿用它所在 ESXi 主機（ESXi 的資產登記）的機櫃，標 via_esxi；
    都沒有＝「無機櫃資料」。每台只歸一組，組加總＝系統台數。"""
    ms = ms if ms is not None else machines(conn)
    by_serial = {}
    for m in ms:
        for s in m["serials"]:
            by_serial[s] = m
    esxi_serial = {e["node_id"]: e["asset_serial"] for e in smap["esxi"]}
    esxi_label = {e["node_id"]: e["label"] for e in smap["esxi"]}
    rooms: dict[str, dict] = {}
    no_rack: list[dict] = []
    for h in smap["hosts"]:
        own = next((by_serial[s] for s in h["serials"] if s in by_serial), None)
        src, via = None, None
        if own and own["kind"] == "rack":
            src = own
        else:
            for eid in h["esxi"]:
                em = by_serial.get(esxi_serial.get(eid) or "")
                if em and em["kind"] == "rack":
                    src, via = em, esxi_label.get(eid, eid)
                    break
        if not src:
            no_rack.append({"key": h["key"], "raw": own["rack_raw"] if own else None,
                            "via_esxi": [esxi_label.get(e, e) for e in h["esxi"]]})
            continue
        room = rooms.setdefault(src["room"], {"room": src["room"], "racks": {}})
        rk = room["racks"].setdefault(src["rack"], {"rack": src["rack"], "rows": []})
        rk["rows"].append({"key": h["key"], "u_start": src["u_start"], "u_end": src["u_end"],
                           "via_esxi": via, "rack_raw": src["rack_raw"]})
    floor_map = _setting_json(conn, "rack_floor_map") or {}
    height = rack_height(conn)
    sl_min = system_level_min(conn)
    out = []
    for r in rooms.values():
        racks = []
        for rk in sorted(r["racks"].values(), key=lambda x: _natural(x["rack"])):
            rk["rows"].sort(key=lambda x: (-(x["u_end"] or 0), x["key"]))
            devs = [m for m in ms if m["kind"] == "rack" and m["room"] == r["room"] and m["rack"] == rk["rack"]]
            racks.append({"rack": rk["rack"], "floor": floor_of(floor_map, rk["rack"]),
                          "host_count": len(rk["rows"]), "rows": rk["rows"],
                          "power": rack_power(devs, height, room_volts(conn, r["room"]), sl_min)})
        out.append({"room": r["room"], "host_count": sum(x["host_count"] for x in racks), "racks": racks})
    out.sort(key=lambda r: (-r["host_count"], r["room"]))
    return {"rooms": out, "no_rack": no_rack, "no_rack_count": len(no_rack),
            "racked_count": sum(r["host_count"] for r in out)}



def _count_racks(lst: list[dict]) -> dict[tuple, int]:
    out: dict[tuple, int] = {}
    for m in lst:
        if m["kind"] == "rack":
            out[(m["room"], m["rack"])] = out.get((m["room"], m["rack"]), 0) + 1
    return out


ANGLES = ("hardware", "host", "system")
ANGLE_LABEL = {"hardware": "硬體", "host": "主機", "system": "系統"}


def _pack(lst: list[dict], limit: int, systems=()) -> dict[str, Any]:
    racks: dict[tuple, int] = {}
    for m in lst:
        if m["kind"] == "rack":
            racks[(m["room"], m["rack"])] = racks.get((m["room"], m["rack"]), 0) + 1
    return {"total": len(lst),
            "matches": [{k: m[k] for k in ("key", "asset_serial", "name", "ip", "model", "room", "kind", "rack",
                                           "u_start", "u_end", "rack_raw", "api_ids", "exist_level",
                                           "exist_score", "exist_label", "seen")} for m in lst[:limit]],
            "racks": [{"room": r, "rack": k, "count": n} for (r, k), n in sorted(racks.items())],
            "systems": list(systems)}


def search(conn: sqlite3.Connection, q: str, ms: list[dict] | None = None, limit: int = 300) -> dict[str, Any]:
    """4-1 全域搜尋，一個框三個角度（不分大小寫、忽略空白連字號）：
      硬體：機型、硬體序號、資產序號
      主機：主機名、資產名稱、IP
      系統：APID 或系統名稱 → 該系統的全部實體設備
    VM 不列（不在機櫃裡）。每個角度各自回設備清單與「含符合設備的機櫃」，畫面分組顯示筆數。

    多個詞（空白／逗號分開，2026-10-23）：整串當一個詞**先找**，找得到就照舊（單詞行為不變，
    「IBM S922」這種帶空白的型號照樣整串比）；整串找不到、又有兩個以上的詞，才改成組合搜尋（search_multi）。"""
    import blast_radius

    all_ms = ms if ms is not None else machines(conn)
    single = _search_single(conn, q, all_ms, limit)
    terms = [t for t in re.split(r"[\s,，、;；]+", (q or "").strip()) if t]
    if len(terms) >= 2 and not any(single["counts"].values()) \
            and not single["angles"]["system"].get("racks_via_vm") and len(blast_radius.search_norm(q)) >= 2:
        return search_multi(conn, terms, all_ms, limit, q)
    return single


MULTI_FIELDS = ("system", "env", "model", "serial", "ip", "name")
MULTI_FIELD_LABEL = {"system": "系統", "env": "環境", "model": "機型", "serial": "序號", "ip": "IP", "name": "主機名"}


def search_multi(conn: sqlite3.Connection, terms: list[str], ms: list[dict], limit: int = 300,
                 q: str | None = None) -> dict[str, Any]:
    """組合搜尋：每個詞先看它打中哪個欄位（系統 > 環境 > 機型 > 序號 > IP > 主機名，取第一個當它的欄位；
    詞本身的結果＝它打中的所有欄位的聯集，不丟東西）。同欄位的詞 OR、不同欄位的詞 AND。
    例：「S922 S1022」＝機型 S922 或 S1022；「uat S922」＝UAT 環境的 S922。
    什麼都沒打中的詞列在 unknown（畫面顯示「找不到：xxx」），**不參與** AND，不會讓整個結果變空。"""
    import blast_radius

    norm = blast_radius.search_norm
    ms = [m for m in ms if m["kind"] not in ("vm", "removed", "vm_model")]
    apis = sorted({a for m in ms for a in [*m["api_ids"], *m.get("vm_api_ids", [])]})
    names = _system_names(conn, apis)
    by_key = {m["key"]: m for m in ms}
    info = []
    unknown = []
    for t in terms:
        tn = norm(t)
        if len(tn) < 2:
            unknown.append(t)
            continue

        def hit(texts):
            return any(tn in norm(x) for x in texts if x)

        sys_hit = {a for a in apis if tn in norm(a) or tn in norm(names[a])}
        hits = {
            "system": {m["key"] for m in ms if sys_hit & set(m["api_ids"])},
            "env": {m["key"] for m in ms if hit(m.get("envs", []))},
            "model": {m["key"] for m in ms if hit([m["model"]])},
            "serial": {m["key"] for m in ms if hit([*m["hw_serials"], *m["serials"]])},
            "ip": {m["key"] for m in ms if any(blast_radius.ip_match(x["ip"], t) for x in m.get("ips", []))
                   or blast_radius.ip_match(m["ip"], t)},
            "name": {m["key"] for m in ms if hit([m["name"], *m["asset_names"]])},
        }
        via = {m["key"] for m in ms if sys_hit & set(m.get("vm_api_ids", [])) and not (sys_hit & set(m["api_ids"]))}
        fields = [f for f in MULTI_FIELDS if hits[f]]
        if not fields and not via:
            unknown.append(t)
            continue
        field = fields[0] if fields else "system"
        info.append({"term": t, "field": field, "fields": fields, "keys": set().union(*hits.values()),
                     "via": via if field == "system" else set(), "systems": sorted(sys_hit)})
    groups: dict[str, set] = {}
    via_groups: dict[str, set] = {}
    for x in info:
        groups.setdefault(x["field"], set()).update(x["keys"])
        if x["field"] == "system":
            via_groups.setdefault("system", set()).update(x["via"])
    combined: set = set.intersection(*groups.values()) if groups else set()
    via_ok: set = set()
    if "system" in via_groups:
        others = [g for f, g in groups.items() if f != "system"]
        via_ok = via_groups["system"].intersection(*others) if others else set(via_groups["system"])
        via_ok -= combined
    order = {m["key"]: i for i, m in enumerate(ms)}
    lst = sorted((by_key[k] for k in combined), key=lambda m: order[m["key"]])
    pk = _pack(lst, limit, [{"api_id": a, "name": names[a]} for x in info for a in x["systems"]])
    pk["racks_via_vm"] = [{"room": r, "rack": k, "count": n} for (r, k), n in sorted(
        _count_racks([by_key[k] for k in via_ok]).items())]
    angles = {"hardware": {"label": "組合", **pk},
              "host": {"label": ANGLE_LABEL["host"], **_pack([], limit)},
              "system": {"label": ANGLE_LABEL["system"], **_pack([], limit), "racks_via_vm": []}}
    return {"q": q if q is not None else " ".join(terms), "mode": "multi", "angles": angles,
            "counts": {"hardware": len(lst), "host": 0, "system": 0},
            "multi": {"terms": [{"term": x["term"], "field": x["field"], "field_label": MULTI_FIELD_LABEL[x["field"]],
                                 "fields": [MULTI_FIELD_LABEL[f] for f in x["fields"]],
                                 "total": len(x["keys"]),
                                 "count": len(x["keys"] & combined),
                                 "keys": sorted(x["keys"] & combined, key=lambda k: order[k])} for x in info],
                      "unknown": unknown,
                      "rule": "同欄位 OR、不同欄位 AND"}}


def _search_single(conn: sqlite3.Connection, q: str, ms: list[dict], limit: int = 300) -> dict[str, Any]:
    import blast_radius

    norm = blast_radius.search_norm
    qn = norm(q)
    qi = (q or "").strip()
    empty = {a: {"label": ANGLE_LABEL[a], "total": 0, "matches": [], "racks": [], "systems": []} for a in ANGLES}
    if len(qn) < 2:
        return {"q": q, "angles": empty, "counts": {a: 0 for a in ANGLES}}
    ms = [m for m in ms if m["kind"] not in ("vm", "removed", "vm_model")]

    def hit_any(texts):
        return any(qn in norm(t) for t in texts if t)

    hw = [m for m in ms if hit_any([m["model"], *m["hw_serials"], *m["serials"]])]
    host = [m for m in ms if hit_any([m["name"], *m["asset_names"]]) or blast_radius.ip_match(m["ip"], qi)]
    # 管理口（ILOM／iDRAC…）的 IP 或名稱命中：把它的本體也帶出來（2026-10-26）
    mgi = mgmt_host_index(ms)
    by_k = {m["key"]: m for m in ms}
    for m in list(host):
        hk = mgi.get(m["key"])
        if hk and hk in by_k and by_k[hk] not in host:
            host.append(by_k[hk])
    apis = sorted({a for m in ms for a in m["api_ids"]})
    names = _system_names(conn, apis)
    sys_hit = [a for a in apis if qn in norm(a) or qn in norm(names[a])]
    sys_set = set(sys_hit)
    sysm = [m for m in ms if sys_set & set(m["api_ids"])]
    via = [m for m in ms if sys_set & set(m["vm_api_ids"]) and not (sys_set & set(m["api_ids"]))]

    angles = {
        "hardware": {"label": ANGLE_LABEL["hardware"], **_pack(hw, limit)},
        "host": {"label": ANGLE_LABEL["host"], **_pack(host, limit)},
        "system": {"label": ANGLE_LABEL["system"],
                   **_pack(sysm, limit, [{"api_id": a, "name": names[a]} for a in sys_hit]),
                   # 系統的 VM 所在的 ESXi 主機那幾櫃（設備本身不屬於這個系統，所以不算進 total）
                   "racks_via_vm": [{"room": r, "rack": k, "count": n} for (r, k), n in sorted(
                       _count_racks(via).items())]},
    }
    return {"q": q, "angles": angles, "counts": {a: angles[a]["total"] for a in ANGLES}}



# ===== 列印：A4 盤點表（一櫃一頁）與標籤貼紙 =====

def _u_text(m: dict) -> str:
    if not m["u_start"]:
        return ""
    return f"{m['u_start']}U" if m["u_start"] == m["u_end"] else f"{m['u_start']}-{m['u_end']}U"


def _marks(m: dict) -> list[str]:
    """盤點表上的標記：人工標的待 CIA 異動（例：關機未下架）、只有登記（看不到＝待確認）、只有 DY 看得到。"""
    out = []
    if m.get("state") == "off":
        out.append("關機未下架")
    elif m.get("state") == "pending":
        out.append("待確認")
    if m.get("pending_status"):
        out.append(str(m["pending_status"]))
    if m.get("unobserved"):
        out.append("未觀察到（只有登記）")
    elif m.get("pending_confirm"):
        out.append("待確認（只有 DY）")
    return out


def print_sheets(conn: sqlite3.Connection, room: str | None = None, rack: str | None = None,
                 api_id: str | None = None, ms: list[dict] | None = None) -> dict[str, Any]:
    """A4 盤點表資料。room+rack＝那一櫃；只給 room＝那個機房每一櫃；給 api_id＝這個系統碰到的每一櫃
    （含 VM 經 ESXi 的那櫃）。每一櫃列**櫃裡全部設備**（盤點看的是櫃子），U 由上往下，沒寫 U 的排最後。
    列數＝那一櫃的設備數（跟 rack_detail 同一份 machines）。"""
    import blast_radius

    ms = ms if ms is not None else machines(conn)
    if api_id:
        sm = blast_radius.system_map(conn, api_id)
        targets = [(r["room"], k["rack"]) for r in sm["racks"]["rooms"] for k in r["racks"]]
    elif room and rack:
        targets = [(room, rack)]
    elif room:
        targets = sorted({(m["room"], m["rack"]) for m in ms if m["kind"] == "rack" and m["room"] == room},
                         key=lambda t: (rack_prefix(t[1]), _natural(t[1])))
    else:
        raise ValueError("要指定機櫃、機房或系統")
    floor_map = _setting_json(conn, "rack_floor_map") or {}
    import overlap_queue
    acts = overlap_queue.active_decisions(conn)
    sheets = []
    for rm, rk in targets:
        devs = [m for m in ms if m["kind"] == "rack" and m["room"] == rm and m["rack"] == rk]
        devs.sort(key=lambda m: (m["u_start"] is None, -(m["u_end"] or 0), m["name"].lower()))
        bx = rack_boxes(devs, overlap_queue.merge_plan(devs, rm, rk, acts))
        by = {m["key"]: m for m in devs}
        rows = []
        for m in devs:
            c = bx["category"][m["key"]]
            if c in ("member", "port"):
                continue                     # 同一個箱子：寫在本體那一列
            marks = _marks(m)
            for pk in bx["ports_of"].get(m["key"], []):
                marks.append(f"管理 IP：{by[pk]['ip'] or '無 IP'}（{bx['mgmt'][pk]['kind']}）")
            mem = [by[k] for b in bx["boxes"] if b["head"] == m["key"] for k in b["keys"][1:]]
            if mem:
                marks.append(f"推定同一台，另 {len(mem)} 筆：" + "、".join(f"{x['name']}（{ROLE_LABEL[bx['roles'][x['key']]]}）" for x in mem))
            if c == "orphan":
                marks.append("管理介面・找不到本體")
            if c == "suspect":
                marks.append("疑似登記錯誤（同一段 U 太多筆）")
            for a, b in m.get("extra_ranges") or []:
                marks.append(f"另一段：{_u_txt(a, b)}")
            rows.append({"u": _u_text(m), "model": m["model"], "serial": ", ".join(m["hw_serials"]) or None,
                         "asset_serial": m["asset_serial"], "name": m["name"], "api_ids": m["api_ids"],
                         "marks": marks, "highlight": bool(api_id and api_id in m["api_ids"])})
        floor = floor_of(floor_map, rk) or next((m["floor_hint"] for m in devs if m.get("floor_hint")), None)
        sheets.append({"room": rm, "rack": rk, "floor": floor, "device_count": len(rows), "record_count": len(devs),
                       "rows": rows})
    return {"sheets": sheets, "blank_rows": 5}


#: 標籤紙規格。使用者確認的紙：A4 24 格（3×8，70×37mm）。要加別種在 app_settings label_presets 放 JSON 陣列。
LABEL_PRESETS_DEFAULT = [
    {"id": "a4_3x8", "name": "A4 3×8（70×37mm）", "cols": 3, "rows": 8, "w_mm": 70, "h_mm": 37,
     "margin_top_mm": 0.5, "margin_left_mm": 0, "gap_x_mm": 0, "gap_y_mm": 0},
]


def label_presets(conn: sqlite3.Connection) -> list[dict]:
    v = _setting_json(conn, "label_presets")
    return v if isinstance(v, list) and v else LABEL_PRESETS_DEFAULT


def label_items(conn: sqlite3.Connection, room: str | None = None, rack: str | None = None,
                api_id: str | None = None, q: str | None = None, angle: str | None = None,
                ms: list[dict] | None = None) -> dict[str, Any]:
    """標籤貼紙要印的設備（只有實體設備；VM 貼不了貼紙）。來源跟畫面一樣：一櫃、一個機房、一個系統、或搜尋結果。"""
    ms = ms if ms is not None else machines(conn)
    phys = [m for m in ms if m["kind"] not in ("vm", "removed", "vm_model")]
    if q:
        res = search(conn, q, ms=ms, limit=100000)
        angles = [angle] if angle in ANGLES else list(ANGLES)
        keys = {x["key"] for a in angles for x in res["angles"][a]["matches"]}
        sel = [m for m in phys if m["key"] in keys]
    elif api_id:
        sel = [m for m in phys if api_id in m["api_ids"]]
    elif room and rack:
        sel = [m for m in phys if m["room"] == room and m["rack"] == rack and m["kind"] == "rack"]
    elif room:
        sel = [m for m in phys if m["room"] == room]
    else:
        raise ValueError("要指定機櫃、機房、系統或搜尋字")
    # 一個箱子一張：管理口、同箱其他登記不另外印（2026-10-26）
    skip = set()
    mgmt_txt: dict[str, str] = {}
    racks = {(m["room"], m["rack"]) for m in sel if m["kind"] == "rack"}
    if racks:
        import overlap_queue
        acts = overlap_queue.active_decisions(conn)
        for rm, rk in racks:
            devs = [m for m in ms if m["kind"] == "rack" and m["room"] == rm and m["rack"] == rk]
            bx = rack_boxes(devs, overlap_queue.merge_plan(devs, rm, rk, acts))
            skip |= {k for k, c in bx["category"].items() if c in ("member", "port")}
            byk = {m["key"]: m for m in devs}
            for h, pks in bx["ports_of"].items():
                mgmt_txt[h] = "、".join(f"管理 IP：{byk[pk]['ip'] or '無 IP'}（{bx['mgmt'][pk]['kind']}）" for pk in pks)
    sel = [m for m in sel if m["key"] not in skip]
    sel.sort(key=lambda m: (m["room"], m["rack"] or "~", -(m["u_end"] or 0), m["name"].lower()))
    return {"presets": label_presets(conn), "items": [{
        "key": m["key"], "name": m["name"], "contact": m["custodian"] or m["owner"],
        "contact_is_owner": not m["custodian"] and bool(m["owner"]), "ip": m["ip"],
        "asset_serial": m["asset_serial"], "room": m["room"], "rack": m["rack"], "u": _u_text(m),
        "mgmt": mgmt_txt.get(m["key"]),
    } for m in sel]}



# ===== 上架助手：找得下的機櫃＋預留（2026-10-23）=====
# 用的是跟機櫃正面同一份資料：在用（含未觀察到）、關機未下架、待確認都佔 U；預留也佔 U。
# 櫃裡有設備沒寫 U 位＝這櫃的空 U「不確定」，排後面並講原因。

def ensure_reservation_table(conn: sqlite3.Connection) -> None:
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_reservation ("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, room TEXT NOT NULL, rack TEXT NOT NULL,"
        " u_start INTEGER NOT NULL, u_end INTEGER NOT NULL, label TEXT NOT NULL, model TEXT,"
        " w REAL, kg REAL, note TEXT, expires_at TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'active',"
        " created_by TEXT, created_at TEXT NOT NULL DEFAULT (datetime('now','localtime')),"
        " closed_by TEXT, closed_at TEXT)")


LEGACY_BRANCH_ROOM = "分公司"   # 舊鍵：W221-016 之前所有分點共用的機房名（分點不明）


def is_branch_site_room(room: str | None) -> bool:
    """「分公司-台中」這種有分點名的機房鍵（不含舊鍵「分公司」本身）。"""
    import blast_radius
    return bool(room) and bool(blast_radius.BRANCH_SITE_RE.match(room))


def legacy_reservations(res: list[dict], room: str, rack: str) -> list[dict]:
    """某分點櫃對應的舊鍵預留（room='分公司' 同櫃號）：分點不明，不歸給任何分點、不當空閒，
    保守視為佔用並標「待確認」。其他機房回空。"""
    if not is_branch_site_room(room):
        return []
    return [{**r, "legacy_key": True, "legacy_note": "舊鍵資料（分點不明，待確認）"}
            for r in res if r["room"] == LEGACY_BRANCH_ROOM and r["rack"] == rack]


def legacy_audit_snapshots(conn: sqlite3.Connection, room: str, rack: str) -> list[dict]:
    """某分點櫃對應的舊鍵盤點快照（rack_audit_rack room='分公司' 同櫃號）。只讀、不改寫、不歸給分點。"""
    if not is_branch_site_room(room):
        return []
    try:
        rows = conn.execute("SELECT audit_id, book_count FROM rack_audit_rack WHERE room = ? AND rack = ? ORDER BY audit_id",
                            (LEGACY_BRANCH_ROOM, rack)).fetchall()
    except sqlite3.OperationalError:
        return []                      # 還沒建過盤點表＝沒有舊快照
    return [{"audit_id": r["audit_id"], "book_count": r["book_count"], "legacy_key": True,
             "legacy_note": "舊鍵資料（分點不明，待確認）"} for r in rows]


def reservations(conn: sqlite3.Connection, include_closed: bool = False) -> list[dict]:
    """預留中（含已過期但還沒取消的，標 expired）。"""
    ensure_reservation_table(conn)
    q = "SELECT * FROM rack_reservation" + ("" if include_closed else " WHERE status = 'active'") + " ORDER BY room, rack, u_start"
    today = datetime_now()[:10]
    return [{**dict(r), "expired": (r["expires_at"] or "") < today} for r in conn.execute(q)]


def datetime_now() -> str:
    from datetime import datetime
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def budgets(conn: sqlite3.Connection) -> dict:
    """機櫃上限：app_settings rack_budget＝{"default": {"w": 5000, "kg": null}, "racks": {"機房/機櫃": {"w":, "kg":}}}。
    沒設定時用電上限預設 5000 W，重量上限未設定。"""
    v = _setting_json(conn, "rack_budget") or {}
    d = {"w": 5000, "kg": None, **(v.get("default") or {})}
    return {"default": d, "racks": v.get("racks") or {}}


def _budget_for(b: dict, room: str, rack: str) -> dict:
    return {**b["default"], **(b["racks"].get(f"{room}/{rack}") or {})}


def _free_runs(occupied: set[int], height: int) -> list[tuple[int, int]]:
    runs, start = [], None
    for u in range(1, height + 1):
        if u not in occupied:
            start = u if start is None else start
        elif start is not None:
            runs.append((start, u - 1))
            start = None
    if start is not None:
        runs.append((start, height))
    return runs


def placement(conn: sqlite3.Connection, u: int, qty: int = 1, w: float | None = None, kg: float | None = None,
              rooms: list[str] | None = None, ms: list[dict] | None = None) -> dict[str, Any]:
    """找放得下 qty 台（每台 u 個 U、w 瓦、kg 公斤）的機櫃。每櫃回：放得下幾台、建議起始 U（最緊的連續空間優先）、
    用電／重量放了之後會不會超過上限、空 U 是否不確定，以及原因。"""
    if u < 1 or qty < 1:
        raise ValueError("U 數與台數都要 ≥ 1")
    ms = ms if ms is not None else machines(conn)
    height = rack_height(conn)
    b = budgets(conn)
    sl_min = system_level_min(conn)
    res = reservations(conn)
    racks: dict[tuple, list] = {}
    for m in ms:
        if m["kind"] == "rack":
            racks.setdefault((m["room"], m["rack"]), []).append(m)
    for r in res:
        racks.setdefault((r["room"], r["rack"]), [])
    out = []
    for (room, rack), devs in racks.items():
        if rooms and room not in rooms:
            continue
        occupied: set[int] = set()
        for d in devs:
            if d["u_start"]:
                occupied.update(range(d["u_start"], d["u_end"] + 1))
        rsv = [r for r in res if r["room"] == room and r["rack"] == rack]
        lrsv = legacy_reservations(res, room, rack)      # 舊鍵預留：保守視為佔用（W221-016）
        rsv = rsv + lrsv
        for r in rsv:
            occupied.update(range(r["u_start"], r["u_end"] + 1))
        no_u = [d for d in devs if not d["u_start"]]
        runs = _free_runs(occupied, height)
        fits = sum((e - s0 + 1) // u for s0, e in runs)
        cand = sorted([(e - s0 + 1, s0, e) for s0, e in runs if e - s0 + 1 >= u])
        best = cand[0] if cand else None
        pw = rack_power(devs, height, room_volts(conn, room), sl_min)
        rsv_w = sum(r["w"] or 0 for r in rsv)
        rsv_kg = sum(r["kg"] or 0 for r in rsv)
        bud = _budget_for(b, room, rack)
        need_w = (w or 0) * qty
        need_kg = (kg or 0) * qty
        w_after = pw["w"] + rsv_w + need_w
        kg_after = pw["kg"] + rsv_kg + need_kg
        reasons = []
        ok = fits >= qty
        if not ok:
            reasons.append(f"連續空間只放得下 {fits} 台（要 {qty} 台、每台 {u}U）")
        if no_u:
            reasons.append(f"有 {len(no_u)} 台沒寫 U 位，空 U 不確定")
        if lrsv:
            reasons.append(f"有 {len(lrsv)} 筆舊鍵預留（機房「分公司」，分點不明，待確認），保守視為佔用")
        if room == LEGACY_BRANCH_ROOM:
            reasons.append("分點不明（機房只寫「分公司」），待確認")
        w_ok = None if bud.get("w") in (None, "") else w_after <= float(bud["w"])
        if w_ok is False:
            reasons.append(f"用電放了之後 {round(w_after)} W 超過上限 {bud['w']} W")
        if pw["unknown_count"]:
            reasons.append(f"有 {pw['unknown_count']} 台未知規格，用電可能更高")
        kg_ok = None if bud.get("kg") in (None, "") else kg_after <= float(bud["kg"])
        if kg_ok is False:
            reasons.append(f"重量放了之後 {round(kg_after)} kg 超過上限 {bud['kg']} kg")
        if w is None:
            reasons.append("型號用電未知，沒算用電")
        out.append({
            "room": room, "rack": rack, "fits": ok and w_ok is not False and kg_ok is not False,
            "fit_count": fits, "suggest_u_start": best[1] if best else None,
            "suggest_u_end": best[1] + u - 1 if best else None, "best_run": best[0] if best else 0,
            "leftover": (best[0] - u) if best else None, "free_u": height - len(occupied & set(range(1, height + 1))),
            "uncertain_u": bool(no_u), "no_u_count": len(no_u), "reserved_u": sum(r["u_end"] - r["u_start"] + 1 for r in rsv),
            "legacy_reserved_count": len(lrsv), "site_pending": room == LEGACY_BRANCH_ROOM,
            "unplaced_count": pw["unplaced_count"], "fits_state": ("unsure" if (ok and no_u) else ("yes" if ok and w_ok is not False and kg_ok is not False else "no")),
            "w_now": round(pw["w"] + rsv_w, 1), "w_after": round(w_after, 1), "w_basis": "額定", "w_budget": bud.get("w"), "w_ok": w_ok,
            "amps_after": round(w_after / pw["volts"], 2) if pw["volts"] else None,
            "kg_now": round(pw["kg"] + rsv_kg, 1), "kg_after": round(kg_after, 1), "kg_budget": bud.get("kg"), "kg_ok": kg_ok,
            "unknown_count": pw["unknown_count"], "reasons": reasons, "occupied": sorted(occupied),
        })
    out.sort(key=lambda r: (not r["fits"], r["uncertain_u"], r["leftover"] if r["leftover"] is not None else 999,
                            -(float(r["w_budget"]) - r["w_after"]) if r["w_budget"] not in (None, "") else 0,
                            r["room"], _natural(r["rack"])))
    return {"request": {"u": u, "qty": qty, "w": w, "kg": kg, "rooms": rooms or []}, "racks": out,
            "budget_default": b["default"], "height_u": height}


def reserve(conn: sqlite3.Connection, room: str, rack: str, u_start: int, u_height: int, label: str, expires_at: str,
            by: str, model: str | None = None, w: float | None = None, kg: float | None = None,
            note: str | None = None) -> dict:
    """預留一段 U（還沒登記的設備）。有到期日；跟現有設備或其他預留重疊只警告。"""
    ensure_reservation_table(conn)
    room, rack, label = (room or "").strip(), _norm_rack(rack or ""), (label or "").strip()
    if not room or not rack or not label:
        raise ValueError("機房、機櫃、名稱都要填")
    try:
        u1, h = int(u_start), int(u_height or 1)
    except (TypeError, ValueError) as exc:
        raise ValueError("U 位和 U 數要是整數") from exc
    if u1 < 1 or h < 1 or u1 + h - 1 > 60:
        raise ValueError("U 位超出範圍")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", expires_at or ""):
        raise ValueError("到期日格式要是 YYYY-MM-DD")
    u2 = u1 + h - 1
    ms = machines(conn)
    hit = collisions(ms, "", room, rack, u1, u2)
    res_all = reservations(conn)
    hit += [{"key": f"rsv:{r['id']}", "asset_serial": None, "name": f"預留：{r['label']}", "u_start": r["u_start"],
             "u_end": r["u_end"]} for r in [x for x in res_all if x["room"] == room and x["rack"] == rack]
            + legacy_reservations(res_all, room, rack)
            if _overlap(u1, u2, r["u_start"], r["u_end"])]
    cur = conn.execute(
        "INSERT INTO rack_reservation (room, rack, u_start, u_end, label, model, w, kg, note, expires_at, created_by) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?)", (room, rack, u1, u2, label, model, w, kg, note, expires_at, by))
    conn.commit()
    return {"id": cur.lastrowid, "collisions": hit}


def close_reservation(conn: sqlite3.Connection, rid: int, by: str) -> dict:
    ensure_reservation_table(conn)
    n = conn.execute("UPDATE rack_reservation SET status = 'cancelled', closed_by = ?, closed_at = datetime('now','localtime') "
                     "WHERE id = ? AND status = 'active'", (by, rid)).rowcount
    conn.commit()
    if not n:
        raise ValueError("查無這筆預留（或已取消）")
    return {"id": rid, "status": "cancelled"}


def find_assets(conn: sqlite3.Connection, q: str, limit: int = 20) -> list[dict]:
    """登記位置用：用資產序號／主機名找既有設備（實體機）。"""
    import blast_radius

    qn = blast_radius.search_norm(q)
    if len(qn) < 2:
        return []
    out = []
    for m in machines(conn):
        if m["kind"] in ("vm", "vm_model"):
            continue
        if any(qn in blast_radius.search_norm(t) for t in [m["name"], *m["serials"], *m["asset_names"]] if t):
            out.append({k: m[k] for k in ("key", "asset_serial", "name", "ip", "model", "room", "rack",
                                          "u_start", "u_end", "kind")})
        if len(out) >= limit:
            break
    return out
