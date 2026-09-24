"""機房搬遷盤點表匯入＋全欄位對帳。

把現場填回來的 Master（.xlsx）讀進來，逐台、逐欄跟系統現有值比對，回一份「校對版」：
Server Master 每格上色——
  綠＝匯入值與系統一致；紅＝不一致（現場校正/與系統不同，值以匯入的為準）；
  藍＝現場沒填、系統有值（幫你補上）；黑／留白＝兩邊都沒有或系統無此欄可比。
所有不一致同時列進「CIA待異動」分頁。**只讀不寫回資料庫**（要不要 Renew 由人另外決定）。

只比「系統也有的欄位」（見 CMP_MAP）；純人事/管理欄（系統沒有獨立來源的）不在比對範圍。
"""
from __future__ import annotations

import io

# Server Master 欄名 → hardware 欄位（只列系統也有、可比對的）
CMP_MAP = {
    "資產狀態": "asset_status", "資產名稱": "asset_name", "主機名稱": "hostname",
    "APID": "api_id", "環境別": "environment", "使用單位": "usage_unit", "擁有者": "owner",
    "保管者": "custodian", "使用者": "user_name", "所屬公司": "owning_company",
    "型號": "device_model", "Serial Number": "hw_serial", "OS": "os",
    "IP(管理/主要)": "ip", "原機房": "physical_location", "原Rack": "rack_no",
}


def _norm(v) -> str:
    return str(v).strip().lower() if v not in (None, "") else ""


def reconcile(conn, file_bytes: bytes) -> tuple[io.BytesIO, dict]:
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    RED = Font(color="C00000")
    GREEN = Font(color="2E7D32")
    BLUE = Font(color="1565C0")

    wb = load_workbook(io.BytesIO(file_bytes))
    if "Server Master" not in wb.sheetnames:
        raise ValueError("這個檔案沒有「Server Master」分頁，請確認是搬遷盤點表")
    ws = wb["Server Master"]

    # 表頭 → 欄位索引
    header = {}
    for c, cell in enumerate(ws[1], 1):
        if cell.value:
            header[str(cell.value).strip()] = c
    if "資產序號" not in header:
        raise ValueError("Server Master 找不到「資產序號」欄")
    serial_c = header["資產序號"]
    note_c = header.get("備註")   # 不一致原因寫進最後一欄「備註」

    # 系統資料：以資產序號為 key
    sys_rows = {r["asset_serial"]: r for r in conn.execute(
        "SELECT * FROM hardware WHERE asset_serial IS NOT NULL")}

    n_same = n_diff = n_fill = n_unknown = 0
    diffs: list[list] = []
    for r in range(2, ws.max_row + 1):
        serial = ws.cell(r, serial_c).value
        if serial in (None, ""):
            continue
        hw = sys_rows.get(str(serial).strip())
        host = ws.cell(r, header.get("主機名稱", serial_c)).value or ""
        if hw is None:
            n_unknown += 1
            continue
        reasons: list[str] = []
        for col_name, hw_field in CMP_MAP.items():
            c = header.get(col_name)
            if not c:
                continue
            imported = ws.cell(r, c).value
            system = hw[hw_field] if hw_field in hw.keys() else None
            iv, sv = _norm(imported), _norm(system)
            if not iv and not sv:
                continue
            if not iv and sv:                    # 現場沒填、系統有 → 幫補上（藍）
                ws.cell(r, c).value = system
                ws.cell(r, c).font = BLUE
                n_fill += 1
            elif iv and not sv:                  # 現場有、系統沒有 → 保留現場值（不變色）
                continue
            elif iv == sv:                       # 一致（綠）
                ws.cell(r, c).font = GREEN
                n_same += 1
            else:                                # 不一致（紅）→ 值以匯入的為準、記差異
                ws.cell(r, c).font = RED
                n_diff += 1
                diffs.append([serial, host, col_name, system, imported,
                              "不一致", "是", "", "待處理", "", ""])
                reasons.append(f"{col_name}：{system or '(空)'} → {imported}")
        # 這台有不一致 → 在「備註」寫清楚原因（原本→要改成），該列看了就懂
        if reasons and note_c:
            cell = ws.cell(r, note_c)
            prefix = (str(cell.value) + "；") if cell.value else ""
            cell.value = prefix + "【系統對帳不一致】" + "；".join(reasons)
            cell.font = RED

    # 寫回 CIA待異動分頁（有就清內容重寫、沒有就建）
    from relocation_export import CIA_DIFF_COLS
    if "CIA待異動" in wb.sheetnames:
        cd = wb["CIA待異動"]
        cd.delete_rows(1, cd.max_row)
    else:
        cd = wb.create_sheet("CIA待異動")
    cd.append(CIA_DIFF_COLS)
    for d in diffs:
        cd.append(d)
    if not diffs:
        cd.append(["（本次比對沒有不一致）"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    summary = {"same": n_same, "diff": n_diff, "filled": n_fill, "unknown_serial": n_unknown,
               "diff_rows": len(diffs)}
    return buf, summary
