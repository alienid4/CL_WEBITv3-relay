"""機櫃圖對帳（2026-10-23）：現場人工維護的 Excel 機櫃圖 ↔ 系統（CIA＋收集）。

背景：各機房（內湖 3F/4F、敦南 21F、板橋 5F B 排＋M01–M16/L16）都有人手畫的 Excel 機櫃圖，原檔不能出公司，
所以要在公司端匯入。兩個方向：
  1. 匯出「跟現場同格式」的 Excel：一個機房＋排一張分頁，櫃並排（U｜設備｜U｜間隔），依類別上色，
     需要現場比對的格子紅框＋⚠＋註解；第一張分頁「需比對清單」（先看／其次）。
  2. 匯入現場圖 Excel → 解析每一張分頁的櫃、U、設備文字（IP、序號、已拔電／暫定…）→ 對帳。
v1 只產報告，不寫 CIA、不寫位置更正（更正照既有「位置更正」流程人工做）。
規則來自 2026-10-23 用真實機櫃圖驗過的兩支原型（rack_sheet.py／rack_ocr_compare.py）。
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import datetime
from typing import Any

SITES = ("內湖", "敦南", "板橋")

SQL = [
    """CREATE TABLE IF NOT EXISTS rack_diagram_import (
        id INTEGER PRIMARY KEY AUTOINCREMENT, site TEXT NOT NULL, file_name TEXT, sha256 TEXT,
        sheets INTEGER, rows INTEGER, imported_by TEXT, imported_at TEXT NOT NULL, racks_json TEXT)""",
    """CREATE TABLE IF NOT EXISTS rack_diagram_row (
        id INTEGER PRIMARY KEY AUTOINCREMENT, import_id INTEGER NOT NULL,
        sheet TEXT, floor TEXT, rack TEXT, rack_norm TEXT, u_start INTEGER, u_end INTEGER,
        text TEXT, ips TEXT, main_ip TEXT, serials TEXT, state TEXT, planned INTEGER, side TEXT,
        source TEXT, shot_date TEXT, confidence TEXT)""",
    # 逐筆判定（確認（照片對）／判讀錯了）：只增不改，最新一筆為準；稽核看誰、何時、依哪張照片
    """CREATE TABLE IF NOT EXISTS rack_diagram_review (
        id INTEGER PRIMARY KEY AUTOINCREMENT, row_id INTEGER NOT NULL, import_id INTEGER NOT NULL,
        action TEXT NOT NULL, note TEXT, source TEXT, shot_date TEXT, confidence TEXT,
        reviewed_by TEXT, reviewed_at TEXT NOT NULL, sys_key TEXT, sys_serial TEXT, result TEXT,
        sys_rack TEXT, sys_u TEXT, sys_model TEXT)""",
]


def ensure_tables(conn: sqlite3.Connection) -> None:
    for s in SQL:
        conn.execute(s)
    cols = {r[1] for r in conn.execute("PRAGMA table_info(rack_diagram_import)")}
    if "racks_json" not in cols:      # 10-23 第一版沒有這欄（221 上已建過表）
        conn.execute("ALTER TABLE rack_diagram_import ADD COLUMN racks_json TEXT")
    rcols = {r[1] for r in conn.execute("PRAGMA table_info(rack_diagram_row)")}
    for c in ("source", "shot_date", "confidence"):     # 10-24 加：來源照片／拍攝日期／判讀把握度
        if c not in rcols:
            conn.execute(f"ALTER TABLE rack_diagram_row ADD COLUMN {c} TEXT")
    vcols = {r[1] for r in conn.execute("PRAGMA table_info(rack_diagram_review)")}
    for c in ("sys_key", "sys_serial", "result", "sys_rack", "sys_u", "sys_model"):   # 判定當時配到哪台、什麼結果、系統當時的值（匯出時要對得上才收）
        if c not in vcols:
            conn.execute(f"ALTER TABLE rack_diagram_review ADD COLUMN {c} TEXT")


# ===================================================================== 櫃名
def norm_rack(r: str | None) -> str:
    """C1 → C-1、M01 → M-1、D-1 櫃 → D-1、1-1 → 1-1（英文字與數字分段，數字去前導零）。"""
    toks = re.findall(r"[A-Za-z]+|\d+", r or "")
    return "-".join(str(int(t)) if t.isdigit() else t.upper() for t in toks)


DEFAULT_ALIAS = {"內湖": {"row_letters": "ABC"}}


def alias_rules(conn: sqlite3.Connection | None) -> dict:
    """CIA 櫃名 → 現場圖櫃名的別名規則（app_settings rack_diagram_alias 可覆寫）。
    預設：內湖 CIA 的「排-櫃」數字寫法，第 1 排＝A、第 2 排＝B、第 3 排＝C（1-5＝A5、2-1＝B1）。"""
    if conn is not None:
        try:
            r = conn.execute("SELECT value FROM app_settings WHERE key = 'rack_diagram_alias'").fetchone()
        except sqlite3.OperationalError:
            r = None
        if r and r[0]:
            try:
                return json.loads(r[0])
            except ValueError:
                pass      # 設定值壞掉＝用預設，畫面上的規則說明仍照預設寫
    return DEFAULT_ALIAS


def alias(site: str | None, rack: str | None, rules: dict | None = None) -> str:
    n = norm_rack(rack)
    rule = (rules or DEFAULT_ALIAS).get(site or "") or {}
    letters = rule.get("row_letters")
    if letters:
        m = re.fullmatch(r"(\d)-(\d+)", n)
        if m and 1 <= int(m.group(1)) <= len(letters):
            return f"{letters[int(m.group(1)) - 1]}-{m.group(2)}"
    return n


def label_rack(site: str | None, rack: str, rules: dict | None = None) -> str:
    a = alias(site, rack, rules)
    return a.replace("-", "") if re.fullmatch(r"[A-Z]+-\d+", a) else rack


def site_of(room: str | None) -> str | None:
    for s in SITES:
        if s in (room or ""):
            return s
    return None


# ===================================================================== 解析設備文字
_RE_IP = re.compile(r"(?<![\d.])(?:25[0-5]|2[0-4]\d|1?\d?\d)(?:\.(?:25[0-5]|2[0-4]\d|1?\d?\d)){3}(?![\d.])")
# 序號：括號裡 6–20 碼英數（至少一個數字，擋掉「(VMware)」）；或「-」後 7 碼（IBM 機型序號寫法）
_RE_SERIAL = re.compile(r"\((?=[0-9A-Z-]*\d)([0-9A-Z][0-9A-Z-]{5,19})\)|-([0-9A-Z]{7})(?=/|$|\s|\))")
STATE_MARKERS = (("已拔電", "已拔電"), ("已關機", "已關機"), ("關機", "已關機"), ("閒置", "閒置"),
                 ("待下線", "待下線"), ("可下架", "可下架"))
PLANNED_MARKERS = ("vSan預留", "暫定", "預留")
SKIP_WORDS = ("理線槽", "panel", "電源排插", "限制使用", "層板", "孔位", "空白")


def is_skip(text: str) -> bool:
    t = (text or "").strip()
    if not t or t.upper() in ("X", "-"):
        return True
    low = t.lower()
    return any(w.lower() in low for w in SKIP_WORDS)


def parse_device(text: str) -> dict[str, Any]:
    t = (text or "").strip()
    side = None
    m = re.match(r"^\s*(前|後)\s*[:：]\s*", t)
    if m:
        side = m.group(1)
        t = t[m.end():]
    # 0.0.0.0／255.x 不是設備位址（CIA 沒 IP 時有人填 0.0.0.0），不拿來比、也不算自相矛盾
    ips = [ip for ip in dict.fromkeys(_RE_IP.findall(t)) if ip != "0.0.0.0" and not ip.startswith("255.")]
    serials = [a or b for a, b in _RE_SERIAL.findall(t.upper())]
    state = next((lab for k, lab in STATE_MARKERS if k in t), None)
    planned = any(k.lower() in t.lower() for k in PLANNED_MARKERS)
    unreadable = bool(re.search(r"未辨識|看不清|[?？]{1,}", t))
    return {"text": t, "ips": ips, "main_ip": ips[0] if ips else None, "serials": serials,
            "state": state, "planned": planned, "side": side, "confidence": "未辨識" if unreadable else None}


# ===================================================================== 匯入：解析 Excel
def _int(v) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)) and float(v).is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


# 櫃頭：C1、M01(期自營)、D-1 櫃、1-1、B10(VMware)；後面只能接結尾／空白／括號／「櫃」（擋掉「3PAR 8000」這種設備名）
_RE_RACK_HEAD = re.compile(r"^\s*(\d{1,2}[A-Za-z]{1,2}\d{1,2}|[A-Za-z]{0,3}\s*-?\s*\d{1,2}(?:\s*-\s*\d{1,2})?)(?=$|[\s(（櫃])")


def parse_workbook(data: bytes, racks_out: list | None = None) -> list[dict]:
    """每一張分頁找「櫃頭」：一格字串（像 C1、M01(期自營)、D-1 櫃、1-1），下一列左右兩側有 U 數字（頂端 ≥ 30）。
    U 欄往下讀到沒有數字為止；設備＝U 欄之間那一欄（合併格跨幾 U 就是幾 U）。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    out: list[dict] = []
    for ws in wb.worksheets:
        # 表格式分頁（例如系統匯出的「需比對清單」：第一列是「項次／機櫃／U…」表頭）不是機櫃圖，不讀
        first = {str(c.value).strip() for c in next(ws.iter_rows(min_row=1, max_row=1), []) if c.value is not None}
        if ws.title == "需比對清單" or {"項次", "機櫃"} <= first:
            continue
        fm = re.search(r"(\d{1,2})\s*F", ws.title, re.I)
        floor = f"{fm.group(1)}F" if fm else None
        merged_at: dict[tuple, Any] = {}
        for rg in ws.merged_cells.ranges:
            for r in range(rg.min_row, rg.max_row + 1):
                for c in range(rg.min_col, rg.max_col + 1):
                    merged_at[(r, c)] = rg
        seen_heads: set = set()
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str) or not v.strip():
                    continue
                m = _RE_RACK_HEAD.match(v)
                if not m or not re.search(r"\d", m.group(1)):
                    continue
                r0 = cell.row
                rg = merged_at.get((r0, cell.column))
                c1, c2 = (rg.min_col, rg.max_col) if rg else (cell.column, cell.column)
                if (r0, c1) in seen_heads:
                    continue
                ucols = [c for c in range(max(1, c1 - 1), c2 + 2) if (_int(ws.cell(r0 + 1, c).value) or 0) >= 30]
                # 櫃頭那一列自己旁邊不會有 U 數字（有的話是某一 U 的設備格，不是櫃頭）
                if not ucols or any(_int(ws.cell(r0, c).value) is not None for c in ucols):
                    continue
                left = min(ucols)
                right = max(ucols) if len(ucols) > 1 else None
                dcols = [c for c in range(left + 1, (right or left + 2)) if c not in ucols] or [left + 1]
                dcol = dcols[0]
                seen_heads.add((r0, c1))
                rack = re.sub(r"\s+", "", m.group(1))
                if racks_out is not None:
                    racks_out.append({"sheet": ws.title, "rack": rack, "rack_norm": norm_rack(rack)})
                urow: dict[int, int] = {}
                r = r0 + 1
                while r <= ws.max_row:
                    u = _int(ws.cell(r, left).value)
                    if u is None and right:
                        u = _int(ws.cell(r, right).value)
                    if u is None:
                        break
                    urow[r] = u
                    r += 1
                # 櫃下方另外列的設備（沒有 U 位、例如「位置不明」或「已關機可下架」那幾格）：
                # 只收有 IP／序號／狀態標記的字，圖例、說明文字不算
                below = []
                blank = 0
                for rb in range(r, min(ws.max_row, r + 25) + 1):
                    vb = ws.cell(rb, dcol).value
                    if vb is None or not str(vb).strip():
                        blank += 1
                        if blank >= 4:
                            break
                        continue
                    blank = 0
                    tb = str(vb).strip().lstrip("⚠ ").strip()
                    if is_skip(tb):
                        continue
                    d = parse_device(tb)
                    if d["ips"] or d["serials"] or d["state"]:
                        below.append(d)
                for d in below:
                    out.append({"sheet": ws.title, "floor": floor, "rack": rack, "rack_norm": norm_rack(rack),
                                "u_start": None, "u_end": None, **d})
                done: set = set()
                for rr in sorted(urow):
                    rg2 = merged_at.get((rr, dcol))
                    top = (rg2.min_row, rg2.min_col) if rg2 else (rr, dcol)
                    if top in done:
                        continue
                    done.add(top)
                    val = ws.cell(*top).value
                    if val is None:
                        continue
                    us = [urow[x] for x in range(rg2.min_row, rg2.max_row + 1) if x in urow] if rg2 else [urow[rr]]
                    if not us:
                        continue
                    text = str(val).strip()
                    if is_skip(text):
                        continue
                    lines = [x.strip() for x in text.split("\n") if x.strip()]
                    # 一格寫好幾台（每行各有 IP）→ 拆開；一台寫好幾行（型號／用途／序號）→ 一筆
                    parts = lines if sum(1 for x in lines if _RE_IP.search(x)) >= 2 else [text.replace("\n", " ")]
                    for p in parts:
                        if is_skip(p):
                            continue
                        d = parse_device(p.lstrip("⚠ ").strip())
                        out.append({"sheet": ws.title, "floor": floor, "rack": rack, "rack_norm": norm_rack(rack),
                                    "u_start": min(us), "u_end": max(us), **d})
    return out


# 清單格式（照片判讀結果常做成這樣）：一列一台。表頭用「包含」比對，順序不拘；機櫃、內容必填
FLAT_HEADERS = {"機房": "room", "樓層": "floor", "機櫃": "rack", "U起": "u_start", "U迄": "u_end", "內容": "text",
                "IP": "ip", "序號": "serial", "來源": "source", "照片": "source", "拍攝日期": "shot_date", "日期": "shot_date",
                "把握度": "confidence", "信心": "confidence", "註記": "note", "備註": "note", "分頁": "page"}
# 只認「整格就是這個字」的表頭（「U」單獨一欄寫 12-15；「系統U」那種不算）
FLAT_EXACT = {"U": "u_range", "U位": "u_range"}
# 樓層 → 機房（使用者 10-23 的截圖清單：3F／4F＝內湖、21F＝敦南、寫板橋的＝板橋）
FLOOR_SITE = {"3F": "內湖", "4F": "內湖", "21F": "敦南"}


def _flat_map(header: list) -> dict[str, int] | None:
    m: dict[str, int] = {}
    for i, h in enumerate(header):
        t = str(h or "").strip()
        if t in FLAT_EXACT and FLAT_EXACT[t] not in m:
            m[FLAT_EXACT[t]] = i
            continue
        for k, f in sorted(FLAT_HEADERS.items(), key=lambda x: -len(x[0])):
            if k.lower() in t.lower() and f not in m:
                m[f] = i
                break
    return m if "rack" in m and "text" in m else None


def row_site(room, floor) -> str | None:
    """這一列是哪個機房：機房欄寫了就用（含內湖／敦南／板橋）；沒有就看樓層（3F／4F＝內湖、21F＝敦南、含板橋＝板橋）。"""
    for v in (room, floor):
        t = str(v or "").strip()
        for s_ in SITES:
            if s_ in t:
                return s_
    return FLOOR_SITE.get(str(floor or "").strip().upper())


def _u_range(v) -> tuple[int | None, int | None]:
    """「12-15」「12~15」「12」「12U」都收。"""
    t = str(v or "").strip().upper().replace("U", "")
    m = re.fullmatch(r"\s*(\d{1,2})\s*(?:[-~～－]\s*(\d{1,2}))?\s*", t)
    if not m:
        return None, None
    a, b = int(m.group(1)), int(m.group(2) or m.group(1))
    return min(a, b), max(a, b)


def parse_flat(table: list[list], sheet: str = "清單", site: str | None = None,
               defaults: dict | None = None, other_site: list | None = None) -> tuple[list[dict], list[dict]] | None:
    """清單格式：找到含「機櫃」「內容」的表頭列才算；一列一台，可帶來源照片、拍攝日期、判讀把握度。
    site 給了：機房欄／樓層看得出是別的機房的列不收（數量記在 other_site）；看不出來的照收。
    defaults：檔案沒有來源／拍攝日期／把握度欄時用的預設值（匯入畫面填的）。"""
    dft = defaults or {}
    for hi, header in enumerate(table[:10]):
        mp = _flat_map(header)
        if mp:
            break
    else:
        return None
    rows, racks = [], {}
    for line in table[hi + 1:]:
        g = lambda f: (line[mp[f]] if f in mp and mp[f] < len(line) else None)
        rack, text = str(g("rack") or "").strip(), str(g("text") or "").strip()
        if not rack or not text or is_skip(text):
            continue
        rs = row_site(g("room"), g("floor"))
        if site and rs and rs != site:
            if other_site is not None:
                other_site.append(rs)
            continue
        note = str(g("note") or "").strip()
        d = parse_device(text + (f" {note}" if note else ""))
        d["text"] = text + (f"（{note}）" if note else "")
        extra_ip = [ip for ip in _RE_IP.findall(str(g("ip") or "")) if ip != "0.0.0.0"]
        d["ips"] = list(dict.fromkeys(extra_ip + d["ips"]))
        d["main_ip"] = d["ips"][0] if d["ips"] else None
        sn = str(g("serial") or "").strip().upper()
        if sn:
            d["serials"] = list(dict.fromkeys([sn] + d["serials"]))
        conf = norm_confidence(g("confidence")) if "confidence" in mp else norm_confidence(dft.get("confidence"))
        d["confidence"] = "未辨識" if d["confidence"] == "未辨識" else conf
        u1, u2 = _int(g("u_start")), _int(g("u_end"))
        if u1 is None and u2 is None and "u_range" in mp:
            u1, u2 = _u_range(g("u_range"))
        shot = g("shot_date") if "shot_date" in mp else dft.get("shot_date")
        if hasattr(shot, "strftime"):
            shot = shot.strftime("%Y-%m-%d")
        else:
            shot = str(shot).strip() if shot is not None and str(shot).strip() else None
        rn = norm_rack(rack)
        racks.setdefault(rn, {"sheet": sheet, "rack": rack, "rack_norm": rn})
        rows.append({"sheet": sheet, "floor": str(g("floor") or "").strip() or None, "rack": rack, "rack_norm": rn,
                     "u_start": u1 if u1 or not u2 else u2, "u_end": u2 or u1, **d,
                     "source": str(g("source") or "").strip() or dft.get("source") or None, "shot_date": shot})
    return rows, list(racks.values())


def _read_table(file_name: str, data: bytes) -> list[tuple[str, list[list]]]:
    if file_name.lower().endswith(".csv"):
        import csv as _csv
        txt = data.decode("utf-8-sig", errors="replace")
        return [("清單", [list(r) for r in _csv.reader(io.StringIO(txt))])]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]


def import_diagram(conn, site: str, file_name: str, data: bytes, by: str | None, defaults: dict | None = None) -> dict:
    """現場格式（櫃並排）或清單格式（一列一台，可帶來源照片／拍攝日期／把握度）都收。照片本身不收（敏感），只記檔名。"""
    ensure_tables(conn)
    if site not in SITES:
        raise ValueError(f"機房要是 {'／'.join(SITES)}")
    low = file_name.lower()
    if not (low.endswith(".xlsx") or low.endswith(".csv")):
        raise ValueError("只收 .xlsx 或 .csv（舊版 .xls 請先另存成 .xlsx）")
    racks: list[dict] = []
    rows: list[dict] = []
    try:
        tables = _read_table(file_name, data)
        skipped: list[str] = []
        other_site: list[str] = []
        dft = {k: v for k, v in (defaults or {}).items() if v}
        dft.setdefault("source", file_name)
        for sheet, table in tables:          # 先看是不是清單格式（有「機櫃」「內容」表頭）
            flat = parse_flat(table, sheet, site, dft, other_site)
            if flat:
                rows += flat[0]
                racks += flat[1]
            else:
                skipped.append(sheet)
        notes = []
        if rows and skipped:                 # 清單格式的檔裡夾了別種分頁：講出來，不靜靜略過
            notes.append(f"這些分頁不是清單格式、沒讀：{'、'.join(skipped)}")
        if other_site:
            c = Counter(other_site)
            notes.append("別的機房的列沒收（請另外選那個機房再匯一次）：" + "、".join(f"{k} {v} 列" for k, v in c.items()))
        skipped_note = "；".join(notes) or None
        if not rows and low.endswith(".xlsx"):
            rows = parse_workbook(data, racks)
            for r in rows:                   # 現場格式：來源＝分頁名稱
                r.setdefault("source", r["sheet"])
    except Exception as exc:  # noqa: BLE001 - 檔案壞掉／不是 Excel：回給使用者看，不吞
        raise ValueError(f"讀不了這個檔：{exc}") from exc
    if not racks:
        raise ValueError("這個檔找不到任何機櫃（現場格式：櫃頭下一列要有 U 數字；清單格式：表頭要有「機櫃」「內容」）")
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cur = conn.execute("INSERT INTO rack_diagram_import (site,file_name,sha256,sheets,rows,imported_by,imported_at,racks_json) "
                       "VALUES (?,?,?,?,?,?,?,?)", (site, file_name, hashlib.sha256(data).hexdigest(),
                                                    len({r["sheet"] for r in racks}), len(rows), by, now,
                                                    json.dumps(racks, ensure_ascii=False)))
    iid = cur.lastrowid
    for r in rows:
        conn.execute("INSERT INTO rack_diagram_row (import_id,sheet,floor,rack,rack_norm,u_start,u_end,text,ips,main_ip,"
                     "serials,state,planned,side,source,shot_date,confidence) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                     (iid, r["sheet"], r["floor"], r["rack"], r["rack_norm"], r["u_start"], r["u_end"], r["text"],
                      json.dumps(r["ips"]), r["main_ip"], json.dumps(r["serials"]), r["state"], int(r["planned"]), r["side"],
                      r.get("source"), r.get("shot_date"), r.get("confidence")))
    conn.commit()
    return {"import_id": iid, "rows": len(rows), "sheets": len({r["sheet"] for r in racks}), "racks": len(racks),
            "skipped": skipped_note}


def list_imports(conn) -> list[dict]:
    ensure_tables(conn)
    return [dict(r) for r in conn.execute("SELECT * FROM rack_diagram_import ORDER BY id DESC")]


def import_rows(conn, iid: int) -> list[dict]:
    ensure_tables(conn)
    out = []
    for r in conn.execute("SELECT * FROM rack_diagram_row WHERE import_id = ? ORDER BY id", (iid,)):
        d = dict(r)
        d["ips"] = json.loads(d["ips"] or "[]")
        d["serials"] = json.loads(d["serials"] or "[]")
        d["planned"] = bool(d["planned"])
        out.append(d)
    return out


def diagram_racks(imp, rows: list[dict]) -> set[str]:
    """這份現場圖有哪些櫃：櫃頭都算（空櫃、只有位置不明清單的櫃也算）；舊匯入沒記櫃頭就退回用有設備的櫃。"""
    try:
        heads = json.loads(imp["racks_json"] or "null")
    except (KeyError, IndexError, ValueError):
        heads = None          # 舊匯入沒有這欄
    return {h["rack_norm"] for h in heads} if heads else {r["rack_norm"] for r in rows}


def last_import(conn, site: str) -> dict | None:
    ensure_tables(conn)
    r = conn.execute("SELECT * FROM rack_diagram_import WHERE site = ? ORDER BY id DESC LIMIT 1", (site,)).fetchone()
    return dict(r) if r else None


# ===================================================================== 對帳
# 配對順序（2026-10-24 使用者核可，照片／現場圖為準）：資產序號 → 主機名稱貼紙 → IP（有寫才用）→ 機型＋位置。
# 機型＋位置：同一櫃有兩台以上同型號就不靠機型配（兩台同型號的 IBM 會配錯）。判讀把握度「未辨識」的格子不配。
RESULTS = ("一致", "同櫃但U位不同", "機櫃不同", "機房不同", "機型不同", "系統沒填機房", "系統沒有機櫃資料", "系統登記成VM",
           "照片有系統沒有", "多筆候選（只能人工）", "沒有可比對的編號（只能人工）", "未辨識", "規劃中", "現場圖自相矛盾")
POSITION_RESULTS = ("同櫃但U位不同", "機櫃不同", "機房不同")
# CIA 待更正只收「位置／機型不同」而且「有人按過確認（照片對）」的列
FIX_RESULTS = POSITION_RESULTS + ("機型不同", "系統沒有機櫃資料", "系統沒填機房")
CONFIDENCE = ("高", "中", "低", "未辨識")
# 判讀把握度欄位寫法 → 標準值
_CONF_ALIAS = {"高": "高", "high": "高", "h": "高", "中": "中", "mid": "中", "medium": "中", "m": "中",
               "低": "低", "low": "低", "l": "低", "未辨識": "未辨識", "看不清": "未辨識", "?": "未辨識", "？": "未辨識"}


def norm_confidence(v) -> str | None:
    """空白＝沒填（None）；認不得的寫法（「很高」、打錯字）一律當「未辨識」，不悄悄放行。"""
    t = str(v or "").strip()
    if not t:
        return None
    return _CONF_ALIAS.get(t.lower(), _CONF_ALIAS.get(t, "未辨識"))


def _model_tokens(s: str | None) -> set[str]:
    """機型裡有數字的那幾段（R750、S922、X3650、9200L…），拿來比機型；純英文字（DELL、IBM、Server）不算。"""
    return {t for t in re.findall(r"[A-Z0-9]+", (s or "").upper()) if re.search(r"\d", t) and re.search(r"[A-Z]", t) and len(t) >= 3}


class _Sys:
    def __init__(self, conn, ms=None):
        import blast_radius
        import rack_layout

        self.conn = conn
        self.ms = ms if ms is not None else rack_layout.machines(conn)
        self.by_key = {m["key"]: m for m in self.ms}
        self.by_serial = {s: m for m in self.ms for s in m["serials"]}
        self.ip_serials: dict[str, set] = {}
        for r in conn.execute("SELECT asset_serial, ip FROM hardware WHERE ip IS NOT NULL"):
            for ip in _RE_IP.findall(r["ip"]) or [r["ip"].strip()]:      # CIA 有時一欄寫好幾個 IP（逗號隔開）
                self.ip_serials.setdefault(ip, set()).add(r["asset_serial"])
        for r in conn.execute("SELECT asset_serial, ip, ips FROM asset_nic"):
            for ip in [r["ip"]] + (r["ips"] or "").split(","):
                if ip and ip.strip():
                    self.ip_serials.setdefault(ip.strip(), set()).add(r["asset_serial"])
        self.hw_ids: list[tuple] = [(r["asset_serial"], " ".join(x or "" for x in (r["hw_serial"], r["asset_serial"], r["hardware_no"])).upper())
                                    for r in conn.execute("SELECT asset_serial, hw_serial, hardware_no FROM hardware")]
        self.host_key = blast_radius.host_match_key
        self.by_host: dict[str, list] = {}
        for m in self.ms:
            for h in {m["name"], *m.get("asset_names", [])}:
                k = self.host_key(h)
                if k:
                    self.by_host.setdefault(k, []).append(m)
        self.model_tokens_all = set().union(*(_model_tokens(m.get("model")) for m in self.ms)) if self.ms else set()
        # 系統主機名的開頭英文字（取前 3 碼，例：SVRB21 → SVR）：判斷照片上的字像不像「我們的主機名」
        self.host_prefixes = set()
        for m in self.ms:
            for h in {m["name"], *m.get("asset_names", [])}:
                mm = re.match(r"([A-Za-z]{3,})\d", (h or "").strip())
                if mm:
                    self.host_prefixes.add(mm.group(1)[:3].upper())

    def by_sn(self, sn: str) -> set[str]:
        """序號只做「整段完全相等」（不做子字串，避免誤中）；回所有命中的資產序號，由呼叫端判斷是否唯一。"""
        sn = sn.upper()
        return {a for a, txt in self.hw_ids if sn in txt.split()}

    def machines_of(self, serials) -> list[dict]:
        cands = {self.by_serial[s]["key"]: self.by_serial[s] for s in serials if s in self.by_serial}
        return list(({k: m for k, m in cands.items() if m["kind"] not in ("vm", "vm_model")} or cands).values())


def _hostname_tokens(text: str, models: set[str] | None = None) -> list[str]:
    """設備文字裡像主機名稱貼紙的字（英文字開頭、含數字，至少 4 碼）；IP 先拿掉，已知機型字（S922、R750）不算。"""
    t = _RE_IP.sub(" ", text or "")
    return [w for w in dict.fromkeys(re.findall(r"[A-Za-z][A-Za-z0-9_.-]*\d[A-Za-z0-9_-]*", t))
            if len(w) >= 4 and w.upper() not in (models or set())]


# 開頭是這些字的是描述（設備種類／廠牌），不是主機名貼紙：「Server-M3650」「DELL-R750」
GENERIC_WORDS = {"SERVER", "STORAGE", "SWITCH", "ROUTER", "FIREWALL", "NODE", "DISK", "TAPE", "CONSOLE", "KVM", "PDU", "UPS",
                 "DELL", "IBM", "HP", "HPE", "LENOVO", "CISCO", "EMC", "ORACLE", "FUJITSU", "JUNIPER", "ARUBA", "FORTINET",
                 "NETAPP", "HUAWEI", "INSPUR", "SUPERMICRO", "BROCADE"}


def is_label_key(w: str, sysd: _Sys) -> bool:
    """照片上的字算不算「具體的主機名稱貼紙」（拿來決定配不到時是「照片有系統沒有」還是「沒有可比對的編號」）：
    英文＋數字，不是已知機型字，而且：有 - 或 _ 分段（NH-R3-03-ACI-L207、APP_SW_01），
    或開頭 3 個英文字跟系統主機名的開頭一樣（SVRW21 ↔ 系統有 SVRB21）。
    純描述（Cisco switch、KVM、M3650、DD6900、Node1）不算。"""
    u = w.upper()
    if u in sysd.model_tokens_all or not re.search(r"[A-Z]", u) or not re.search(r"\d", u):
        return False
    segs = [x for x in re.split(r"[-_]", u) if x]
    if segs and segs[0] in GENERIC_WORDS:
        return False
    digit_segs = [x for x in segs if re.search(r"[0-9]", x)]
    # 「Server-M3650」：有數字的段全是已知機型字 → 是型號描述，不是主機名
    if len(segs) >= 2 and re.match(r"[A-Z]", segs[0]) and any(x not in sysd.model_tokens_all for x in digit_segs):
        return True
    mm = re.match(r"([A-Z]{3,})\d", u)
    return bool(mm and mm.group(1)[:3] in sysd.host_prefixes)


def label_keys(text: str, sysd: _Sys) -> list[str]:
    """照片文字裡「具體的主機名稱貼紙」：is_label_key 成立，而且前面不是緊貼中文字（「會館DB3_Node1」是描述，不是貼紙）。"""
    t = _RE_IP.sub(" ", text or "")
    out = []
    for m in re.finditer(r"[A-Za-z][A-Za-z0-9_.-]*\d[A-Za-z0-9_-]*", t):
        w = m.group(0)
        prev = t[m.start() - 1] if m.start() else " "
        if len(w) >= 4 and not ("\u4e00" <= prev <= "\u9fff") and is_label_key(w, sysd) and w not in out:
            out.append(w)
    return out


def _match(sysd: _Sys, r: dict, site: str, rules: dict) -> tuple[list[dict], str | None]:
    """回（候選機器, 配對方式）。順序：資產序號 → 主機名稱 → IP → 機型＋位置。
    候選有好幾台實體機時不挑（呼叫端標「多筆候選（只能人工）」），不會拿第一台去配。"""
    for s in r["serials"]:
        a = sysd.by_sn(s)
        if a:
            ms = sysd.machines_of(a)
            if ms:
                return ms, f"序號 {s}"
    for w in _hostname_tokens(r["text"], sysd.model_tokens_all):
        k = sysd.host_key(w)
        hits = list({m["key"]: m for m in sysd.by_host.get(k or "", [])}.values())
        if hits:
            hits = [m for m in hits if m["kind"] not in ("vm", "vm_model")] or hits
            return hits, f"主機名 {w}"
    for ip in r["ips"]:
        if sysd.ip_serials.get(ip):
            ms = sysd.machines_of(sysd.ip_serials[ip])
            if ms:
                return ms, f"IP {ip}"
    toks = _model_tokens(r["text"]) & sysd.model_tokens_all
    if toks and r["u_start"]:
        same_rack = [m for m in sysd.ms if m["kind"] == "rack" and m["rack"] and site_of(m["room"]) == site
                     and alias(site, m["rack"], rules) == r["rack_norm"] and toks & _model_tokens(m.get("model"))]
        if len(same_rack) == 1:          # 同櫃兩台以上同型號：不靠機型配
            m = same_rack[0]
            if m.get("u_start") and max(r["u_start"], m["u_start"]) <= min(r["u_end"], m["u_end"]):
                return [m], f"機型＋位置 {'／'.join(sorted(toks))}"
    return [], None


def reconcile(conn, iid: int, sysd: _Sys | None = None) -> dict[str, Any]:
    ensure_tables(conn)
    imp = conn.execute("SELECT * FROM rack_diagram_import WHERE id = ?", (iid,)).fetchone()
    if not imp:
        raise LookupError(f"找不到匯入 {iid}")
    site = imp["site"]
    rules = alias_rules(conn)
    sysd = sysd or _Sys(conn)
    rows = import_rows(conn, iid)
    reviews = latest_reviews(conn, iid)
    ip_places: dict[str, set] = defaultdict(set)
    for r in rows:
        for ip in r["ips"]:
            ip_places[ip].add((r["rack_norm"], r["u_start"], r["u_end"]))
    out = []
    for n, r in enumerate(rows, 1):
        res: dict[str, Any] = {"no": n, "row_id": r["id"], "sheet": r["sheet"], "floor": r["floor"], "rack": r["rack"],
                               "rack_norm": r["rack_norm"],
                               "u": ("" if not r["u_start"] else f"{r['u_start']}" if r["u_start"] == r["u_end"] else f"{r['u_start']}-{r['u_end']}"),
                               "u_start": r["u_start"], "u_end": r["u_end"], "text": r["text"], "ip": r["main_ip"],
                               "state": r["state"], "source": r.get("source"), "shot_date": r.get("shot_date"),
                               "confidence": r.get("confidence"), "match_by": None, "sys_key": None, "sys_serial": None,
                               "review": reviews.get(r["id"])}
        if r.get("confidence") == "未辨識":
            res["result"] = "未辨識"
        elif r["planned"]:
            res["result"] = "規劃中"
        elif r["main_ip"] and len(ip_places.get(r["main_ip"], ())) > 1:
            res["result"] = "現場圖自相矛盾"
            res["detail"] = f"同一個 IP 在現場圖出現 {len(ip_places[r['main_ip']])} 個位置"
        else:
            cands, how = _match(sysd, r, site, rules)
            phys = [m for m in cands if m["kind"] not in ("vm", "vm_model")]
            if len(phys) > 1:
                res["result"] = "多筆候選（只能人工）"
                res["detail"] = f"{how} 對到 {len(phys)} 台：" + "、".join(m["name"] for m in phys[:5])
                res["match_by"] = how
                cands = []
            if res.get("result"):
                pass
            elif not cands:
                # 有寫序號／IP／像主機名的字（不是已知機型）才算「照片有系統沒有」；只寫機型的配不到＝只能人工
                has_id = r["ips"] or r["serials"] or label_keys(r["text"], sysd)
                res["result"] = "照片有系統沒有" if has_id else "沒有可比對的編號（只能人工）"
                if has_id:
                    keys = r["ips"] + r["serials"] + label_keys(r["text"], sysd)
                    res["detail"] = "找過但對不到系統：" + "、".join(keys[:4])
            else:
                m = sorted(cands, key=lambda m: m["kind"] != "rack")[0]
                res.update({"sys_key": m["key"], "sys_serial": m["asset_serial"], "sys_name": m["name"],
                            "sys_room": m["room"], "sys_rack": m["rack"] or (m["rack_raw"] or ""),
                            "sys_u": (f"{m['u_start']}" if m["u_start"] == m["u_end"] else f"{m['u_start']}-{m['u_end']}") if m["u_start"] else "",
                            "sys_model": m.get("model"), "sys_state": m.get("state_label") or m.get("asset_status") or "",
                            "match_by": how, "same_ip_count": len(cands)})
                if m["kind"] in ("vm", "vm_model"):
                    res["result"] = "系統登記成VM"
                elif (m["room"] or "未填") == "未填":
                    res["result"] = "系統沒填機房"
                elif site not in (m["room"] or ""):
                    res["result"] = "機房不同"
                elif not m["rack"]:
                    res["result"] = "系統沒有機櫃資料"
                elif alias(site, m["rack"], rules) != r["rack_norm"]:
                    res["result"] = "機櫃不同"
                elif m["u_start"] and r["u_start"] and not max(r["u_start"], m["u_start"]) <= min(r["u_end"], m["u_end"]):
                    res["result"] = "同櫃但U位不同"
                else:
                    seen = _model_tokens(r["text"]) & sysd.model_tokens_all
                    mine = _model_tokens(m.get("model"))
                    if seen and mine and not (seen & mine):
                        res["result"] = "機型不同"
                        res["detail"] = f"現場寫 {'／'.join(sorted(seen))}，系統登記 {m.get('model')}"
                    else:
                        res["result"] = "一致"
                        if not m["u_start"] or not r["u_start"]:
                            res["detail"] = "同櫃（缺 U 位）"
        out.append(res)
    # 現場圖／照片有涵蓋的櫃 vs 系統：沒拍到的櫃（不等於空櫃）、拍到的櫃裡系統有但照片沒有
    diag_racks = diagram_racks(imp, rows)
    matched = {x["sys_key"] for x in out if x.get("sys_key")}
    sys_racks: dict[str, dict] = {}
    sys_only = []
    for m in sysd.ms:
        if m["kind"] == "rack" and m["rack"] and site_of(m["room"]) == site:
            k = alias(site, m["rack"], rules)
            x = sys_racks.setdefault(k, {"rack": m["rack"], "rack_norm": k, "room": m["room"], "keys": [], "names": []})
            x["keys"].append(m["key"])
            x["names"].append(m["name"])
            if k in diag_racks and m["key"] not in matched and not _is_oob(m) and m.get("state") == "active":
                sys_only.append({"key": m["key"], "serial": m["asset_serial"], "name": m["name"], "rack": m["rack"],
                                 "rack_norm": k, "u": f"{m['u_start']}-{m['u_end']}" if m.get("u_start") else "",
                                 "model": m.get("model"), "ip": m.get("ip"),
                                 "note": "不一定是錯：可能裝在背面、被擋住，或已下架"})
    missing = [{**v, "count": len(v["keys"])} for k, v in sorted(sys_racks.items()) if k not in diag_racks]
    # 一致率只算「兩邊都放得到」的：配到系統機器、而且兩邊都有機櫃（機房不同也算進分母）
    PLACED = ("一致", "同櫃但U位不同", "機型不同", "機櫃不同", "機房不同")
    SAME_RACK = ("一致", "同櫃但U位不同", "機型不同")
    for x in out:
        x["placed"] = x["result"] in PLACED
        x["same_rack"] = x["result"] in SAME_RACK
        x["u_both"] = x["same_rack"] and bool(x["u_start"]) and bool(x.get("sys_u"))
        x["u_ok"] = x["u_both"] and x["result"] != "同櫃但U位不同"

    def rates(xs: list[dict]) -> dict:
        placed = [x for x in xs if x["placed"]]
        same = [x for x in placed if x["same_rack"]]
        ub = [x for x in same if x["u_both"]]
        return {"placed": len(placed), "same_rack": len(same), "u_both": len(ub), "u_ok": sum(1 for x in ub if x["u_ok"]),
                "same_rack_rate": round(len(same) * 100 / len(placed)) if placed else None,
                "u_rate": round(sum(1 for x in ub if x["u_ok"]) * 100 / len(ub)) if ub else None}
    per: dict[tuple, list] = defaultdict(list)
    for x in out:
        per[(x["sheet"], x["rack"])].append(x)
    rate = []
    for (sh, rk), xs in sorted(per.items()):
        c = Counter(x["result"] for x in xs)
        rt = rates(xs)
        rate.append({"sheet": sh, "rack": rk, "total": len(xs), "comparable": rt["placed"], "ok": rt["same_rack"],
                     "rate": rt["same_rack_rate"], "u_rate": rt["u_rate"], "counts": dict(c)})
    counts = Counter(x["result"] for x in out)
    groups = {"位置不同": sum(counts.get(k, 0) for k in POSITION_RESULTS), "照片有系統沒有": counts.get("照片有系統沒有", 0),
              "系統有照片沒有": len(sys_only), "機型不同": counts.get("機型不同", 0), "沒拍到的櫃": len(missing)}
    return {"import": dict(imp), "rows": out, "counts": {k: counts.get(k, 0) for k in RESULTS}, "groups": groups,
            "rates": rates(out),
            "missing_racks": missing, "system_only": sys_only, "rate": rate, "matched_keys": sorted(matched)}


# ===================================================================== 逐筆確認（只增不改，稽核用）
REVIEW_ACTIONS = {"confirm": "確認（照片對）", "misread": "判讀錯了"}


REVIEWABLE = FIX_RESULTS + ("照片有系統沒有",)


def review_row(conn, row_id: int, action: str, note: str | None, by: str | None) -> dict:
    """使用者對一筆差異的判定。只新增一筆紀錄，不改不刪；最新一筆為準。不寫 CIA。
    同時記下「當時這列配到哪台、結果是什麼」：之後系統資料變了、改配到別台，這個確認就不算數（匯出不收）。"""
    ensure_tables(conn)
    if not by:
        raise PermissionError("要登入才能確認（確認人是稽核欄位，不能空白）")
    if action not in REVIEW_ACTIONS:
        raise ValueError("動作要是「確認（照片對）」或「判讀錯了」")
    r = conn.execute("SELECT id, import_id, source, shot_date, confidence FROM rack_diagram_row WHERE id = ?", (row_id,)).fetchone()
    if not r:
        raise LookupError(f"找不到這一列 {row_id}")
    cur = next((x for x in reconcile(conn, r["import_id"])["rows"] if x["row_id"] == row_id), None)
    if not cur or cur["result"] not in REVIEWABLE:
        raise ValueError(f"這一列現在是「{(cur or {}).get('result', '?')}」，不是要確認的差異")
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn.execute("INSERT INTO rack_diagram_review (row_id, import_id, action, note, source, shot_date, confidence, reviewed_by, reviewed_at, "
                 "sys_key, sys_serial, result, sys_rack, sys_u, sys_model) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                 (row_id, r["import_id"], action, (note or "").strip() or None, r["source"], r["shot_date"], r["confidence"],
                  by, now, cur.get("sys_key"), cur.get("sys_serial"), cur["result"], cur.get("sys_rack"), cur.get("sys_u"),
                  cur.get("sys_model")))
    conn.commit()
    return latest_reviews(conn, r["import_id"]).get(row_id)


def latest_reviews(conn, iid: int) -> dict[int, dict]:
    ensure_tables(conn)
    out: dict[int, dict] = {}
    for r in conn.execute("SELECT * FROM rack_diagram_review WHERE import_id = ? ORDER BY id", (iid,)):
        d = dict(r)
        d["action_label"] = REVIEW_ACTIONS.get(d["action"], d["action"])
        out[d["row_id"]] = d
    return out


def review_history(conn, iid: int) -> list[dict]:
    ensure_tables(conn)
    if not conn.execute("SELECT 1 FROM rack_diagram_import WHERE id = ?", (iid,)).fetchone():
        raise LookupError(f"找不到匯入 {iid}")
    return [{**dict(r), "action_label": REVIEW_ACTIONS.get(r["action"], r["action"])} for r in conn.execute(
        "SELECT * FROM rack_diagram_review WHERE import_id = ? ORDER BY id DESC", (iid,))]


def diff_previous(conn, iid: int) -> dict[str, Any]:
    """跟同一個機房上一次匯入比：多了／少了／搬了（同 IP 換櫃或換 U）。"""
    imp = conn.execute("SELECT * FROM rack_diagram_import WHERE id = ?", (iid,)).fetchone()
    if not imp:
        raise LookupError(f"找不到匯入 {iid}")
    prev = conn.execute("SELECT id, imported_at, file_name FROM rack_diagram_import WHERE site = ? AND id < ? ORDER BY id DESC LIMIT 1",
                        (imp["site"], iid)).fetchone()
    if not prev:
        return {"previous": None, "added": [], "removed": [], "moved": []}
    cur, old = import_rows(conn, iid), import_rows(conn, prev["id"])
    key = lambda r: r["main_ip"] or f"{r['rack_norm']}|{r['u_start']}|{r['text']}"
    ci, oi = {key(r): r for r in cur}, {key(r): r for r in old}
    added = [ci[k] for k in ci if k not in oi]
    removed = [oi[k] for k in oi if k not in ci]
    moved = [{"ip": k, "from": f"{oi[k]['rack']} {oi[k]['u_start']}-{oi[k]['u_end']}", "to": f"{ci[k]['rack']} {ci[k]['u_start']}-{ci[k]['u_end']}",
              "text": ci[k]["text"]} for k in ci if k in oi and ci[k]["main_ip"]
             and (ci[k]["rack_norm"], ci[k]["u_start"], ci[k]["u_end"]) != (oi[k]["rack_norm"], oi[k]["u_start"], oi[k]["u_end"])]
    return {"previous": dict(prev), "added": added, "removed": removed, "moved": moved}


def fix_rows(conn, iid: int) -> list[dict]:
    """CIA 待更正的列，全部條件都要成立才收：
      · 現在的結果是位置／機型不同（FIX_RESULTS），而且配到一台系統機器
      · 最新一筆判定是「確認（照片對）」
      · 判定當時配到的就是這台、結果也一樣（系統資料變了就失效，要重新確認）
      · 同一台機器沒有被兩列以上確認（衝突的都不收）
    沒確認、判讀錯了、未辨識、照片有系統沒有、多筆候選的一律不收。"""
    rc = reconcile(conn, iid)
    ok = []
    for x in rc["rows"]:
        rv = x.get("review") or {}
        if (x["result"] in FIX_RESULTS and x.get("sys_key") and rv.get("action") == "confirm"
                and rv.get("sys_key") == x["sys_key"] and rv.get("result") == x["result"]
                and all((rv.get(k) or "") == (x.get(k) or "") for k in ("sys_rack", "sys_u", "sys_model"))):
            ok.append(x)
    # 同一台被兩列確認、或一列確認一列判讀錯了：有衝突，都不收
    misread = {x["sys_key"] for x in rc["rows"] if x.get("sys_key") and (x.get("review") or {}).get("action") == "misread"}
    dup = {k for k, n in Counter(x["sys_key"] for x in ok).items() if n > 1} | misread
    return [x for x in ok if x["sys_key"] not in dup]


def fix_xlsx(conn, iid: int) -> tuple[io.BytesIO, int]:
    """CIA 待更正（給 CIA 窗口人工更正用；系統不寫 CIA）：只收按過「確認（照片對）」的差異，帶來源、拍攝日期、把握度、確認人。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    imp = conn.execute("SELECT * FROM rack_diagram_import WHERE id = ?", (iid,)).fetchone()
    if not imp:
        raise LookupError(f"找不到匯入 {iid}")
    rows = fix_rows(conn, iid)
    wb = Workbook()
    ws = wb.active
    ws.title = "CIA待更正"
    ws.append(["項次", "結果", "資產序號", "主機", "IP", "系統機房", "系統機櫃", "系統U", "系統機型",
               "現場機房", "現場機櫃", "現場U", "現場文字", "來源", "拍攝日期", "判讀把握度", "確認人", "確認時間", "備註", "配對方式"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0C6B37")
    for n, x in enumerate(rows, 1):
        rv = x["review"]
        ws.append([n, x["result"], x.get("sys_serial"), x.get("sys_name"), x["ip"], x.get("sys_room"), x.get("sys_rack"),
                   x.get("sys_u"), x.get("sys_model"), imp["site"], x["rack"], x["u"], x["text"], x.get("source") or x["sheet"],
                   x.get("shot_date"), x.get("confidence"), rv.get("reviewed_by"), rv.get("reviewed_at"), rv.get("note"),
                   x.get("match_by")])
    if not rows:
        ws.append(["", "還沒有按過「確認（照片對）」的差異，所以這份是空的"])
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, len(rows)


# ===================================================================== 匯出現場格式
FILL = {"net": "92D050", "esxi": "FFC000", "db": "BDD7EE", "off": "FF7C80", "storage": "FFE699", "test": "CCC0DA", "other": "FFFFFF"}
LEGEND = [("net", "網路相關設備（含 SAN 交換器）"), ("esxi", "虛擬環境（ESXi）"), ("db", "DB 相關"), ("storage", "儲存"),
          ("test", "測試環境"), ("off", "關機未下架／待確認"), ("other", "其他")]
FIRST = ("機櫃不同", "U位不同", "機房不同", "現場圖沒有這一櫃", "關機未下架", "待確認")


def category(m: dict) -> str:
    if m.get("state") in ("off", "pending"):
        return "off"
    t = m.get("machine_type") or ""
    if t in ("network", "san"):
        return "net"
    if t == "esxi":
        return "esxi"
    if t == "storage":
        return "storage"
    names = " ".join([m.get("name") or "", *m.get("asset_names", [])]).upper()
    if re.search(r"DB|SQL|ORA|MONGO", names):
        return "db"
    if any(e in ("測試", "UAT") for e in m.get("envs", [])):
        return "test"
    return "other"


def _is_oob(m: dict) -> bool:
    """iDRAC／OOB 管理口的登記（主機名含 OOB、或 OS 寫 iDRAC）：現場圖不會畫，不算「找不到」。"""
    return "OOB" in (m.get("name") or "").upper() or "IDRAC" in (m.get("os") or "").upper()


def _dev_text(m: dict) -> str:
    """設備格文字：型號 主機名 IP；沒有 IP 就附硬體序號（括號），匯回來才對得到。"""
    sn = None if m.get("ip") else next(iter(m.get("hw_serials") or []), None)
    return " ".join(x for x in (m.get("model"), m.get("name"), m.get("ip"), f"({sn})" if sn else None) if x)


def _natkey(s: str):
    return [(0, int(t), "") if t.isdigit() else (1, 0, t) for t in re.findall(r"\d+|\D+", s)]


def export_xlsx(conn, room: str | None = None, ms: list[dict] | None = None) -> tuple[io.BytesIO, int]:
    """系統資料 → 跟現場機櫃圖同格式的 Excel。room 給了就只出那個機房。旗標依該機房最後一次匯入的對帳結果。"""
    import rack_layout
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ensure_tables(conn)
    rules = alias_rules(conn)
    ms = ms if ms is not None else rack_layout.machines(conn)
    ms = [m for m in ms if m["kind"] == "rack" and m["rack"] and site_of(m["room"]) and (not room or m["room"] == room)]
    # 各機房最後一次匯入的對帳：系統機器 key → (結果, 現場圖位置)；現場圖有哪些櫃
    cmp_by_key: dict[str, tuple] = {}
    diag_racks: dict[str, set] = {}
    for site in {site_of(m["room"]) for m in ms}:
        imp = last_import(conn, site)
        if not imp:
            continue
        rc = reconcile(conn, imp["id"])
        diag_racks[site] = diagram_racks(imp, rc["rows"])
        for x in rc["rows"]:
            if x.get("sys_key"):
                cmp_by_key.setdefault(x["sys_key"], (x["result"], f"{x['rack']} {x['u']}U"))

    def reasons(m: dict, site: str, rack_in_diag: bool | None) -> list[str]:
        out = []
        if rack_in_diag is False:
            out.append("現場圖沒有這一櫃")
        elif rack_in_diag:
            res = cmp_by_key.get(m["key"])
            if res is None and not _is_oob(m):
                out.append("現場圖找不到這台（IP／序號都對不到）")
            elif res and res[0] in ("機櫃不同", "同櫃但U位不同", "機房不同"):
                out.append(f"{res[0]}：現場圖寫 {res[1]}")
        if m.get("state") in ("off", "pending"):
            out.append(m.get("state_label") or "狀態待確認")
        return out

    thin = Side(style="thin", color="808080")
    BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    RED = Side(style="thick", color="FF0000")
    REDBOX = Border(left=RED, right=RED, top=RED, bottom=RED)
    UFILL = PatternFill("solid", fgColor="808080")
    HFILL = PatternFill("solid", fgColor="00B0F0")
    by_sheet: dict[str, dict[str, list]] = defaultdict(lambda: defaultdict(list))
    for m in ms:
        site = site_of(m["room"])
        rk = label_rack(site, m["rack"], rules)
        prefix = re.match(r"\d*[A-Za-z]*", rk).group(0) or rk.split("-")[0]
        by_sheet[f"{m['room']}_{prefix}排"][rk].append(m)       # 一個機房＋排一張分頁
    wb = Workbook()
    summary = wb.active
    summary.title = "需比對清單"
    todo: list[list] = []
    for sheet_name in sorted(by_sheet, key=_natkey):
        ws = wb.create_sheet(sheet_name[:31])
        col = 1
        for rk in sorted(by_sheet[sheet_name], key=_natkey):
            devs = by_sheet[sheet_name][rk]
            site = site_of(devs[0]["room"])
            in_diag = (alias(site, devs[0]["rack"], rules) in diag_racks[site]) if site in diag_racks else None
            h = max([45] + [d["u_end"] or 0 for d in devs])
            cl, cc, cr = col, col + 1, col + 2
            for c, w in ((cl, 4), (cc, 36), (cr, 4), (cr + 1, 2)):
                ws.column_dimensions[get_column_letter(c)].width = w
            ws.merge_cells(start_row=1, start_column=cl, end_row=1, end_column=cr)
            hc = ws.cell(1, cl, rk + ("（現場圖沒有這一櫃）" if in_diag is False else ""))
            hc.fill, hc.font, hc.alignment = HFILL, Font(bold=True, color="FF0000" if in_diag is False else "0000FF", size=12), Alignment(horizontal="center")
            for u in range(h, 0, -1):
                row = 2 + (h - u)
                for c in (cl, cr):
                    x = ws.cell(row, c, u)
                    x.fill, x.font, x.alignment, x.border = UFILL, Font(color="FFFFFF", size=9), Alignment(horizontal="center"), BOX
                ws.cell(row, cc).border = BOX
            groups: dict[tuple, list] = defaultdict(list)
            nou = []
            for d in devs:
                if d["u_start"]:
                    groups[(d["u_start"], d["u_end"] or d["u_start"])].append(d)
                else:
                    nou.append(d)
            blocks: list[list] = []   # 重疊的 U 範圍併成一格（同一格列出全部設備）
            for (a, b), ds in sorted(groups.items()):
                if blocks and a <= blocks[-1][1]:
                    blocks[-1][1] = max(blocks[-1][1], b)
                    blocks[-1][2].extend(ds)
                else:
                    blocks.append([a, b, list(ds)])
            for a, b, ds in sorted(blocks, key=lambda x: -x[1]):
                top, bot = 2 + (h - b), 2 + (h - a)
                if bot > top:
                    ws.merge_cells(start_row=top, start_column=cc, end_row=bot, end_column=cc)
                txt = "\n".join(_dev_text(d) for d in ds)
                why = []
                for d in ds:
                    rs = reasons(d, site, in_diag)
                    if rs:
                        why.append(f'{d["name"]}：' + "；".join(rs))
                        todo.append([site, rk, f"{a}-{b}" if a != b else str(a), d["name"], d["ip"] or "", d["model"] or "", "；".join(rs)])
                x = ws.cell(top, cc, ("⚠ " if why else "") + txt)
                x.fill = PatternFill("solid", fgColor=FILL[category(ds[0])])
                x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                x.font = Font(size=9, bold=bool(why), color="C00000" if why else "000000")
                if why:
                    x.comment = Comment("\n".join(why), "系統")
                    for rr in range(top, bot + 1):
                        ws.cell(rr, cc).border = REDBOX
            r0 = h + 3
            if nou:
                ws.cell(r0, cc, "位置不明（有機櫃、沒 U 位）").font = Font(bold=True, color="C00000")
                for i, d in enumerate(nou, 1):
                    ws.cell(r0 + i, cc, f'⚠ {_dev_text(d)}').font = Font(size=9, color="C00000")
                    todo.append([site, rk, "", d["name"], d["ip"] or "", d["model"] or "", "系統沒有 U 位，請現場確認在第幾 U"])
            col += 4
        lr = max(ws.max_row, 50) + 2
        for i, (k, lab) in enumerate(LEGEND):
            ws.cell(lr + i, 2, lab).fill = PatternFill("solid", fgColor=FILL[k])
        ws.cell(lr + len(LEGEND), 2, "紅框＋⚠＝需要現場比對（滑到格子上看原因）").font = Font(bold=True, color="C00000")
        ws.freeze_panes = "A2"
    summary.append(["項次", "優先", "機房", "機櫃", "U", "主機", "IP", "機型", "為什麼要比對"])
    for c in summary[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0C6B37")

    def prio(t):
        return "先看" if any(k in t[6] for k in FIRST) else "其次"
    todo.sort(key=lambda t: (prio(t) != "先看", t[0], _natkey(t[1]), -(int(t[2].split("-")[-1]) if t[2] else 0)))
    for i, t in enumerate(todo, 1):
        summary.append([i, prio(t), *t])
        if prio(t) == "先看":
            summary.cell(summary.max_row, 2).font = Font(bold=True, color="C00000")
    if not todo:
        summary.append(["", "", "沒有需要現場比對的項目" if diag_racks else "這個機房還沒匯入過現場圖，所以沒有比對旗標"])
    summary.auto_filter.ref = summary.dimensions
    summary.freeze_panes = "A2"
    for i, w in enumerate((6, 6, 6, 10, 8, 30, 16, 26, 60), 1):
        summary.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, len(todo)
