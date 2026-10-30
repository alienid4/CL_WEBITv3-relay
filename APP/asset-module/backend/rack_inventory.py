"""機櫃實體盤點流程（2026-10-27，v1.503.0）。

使用方式（使用者原話整理）：
  1. 「開始盤點」：選一個機房，或幾個機櫃。系統把今天的帳凍結成「基準」：每櫃一箱一列（U、機型、資產編號、
     原廠序號、主機名、APID；併進主機的管理口、同箱的 LPAR／元件寫在那一箱裡面）。
  2. 「下載盤點檔」：一本 xlsx、一櫃一張工作表、從上到下排（跟 A4 盤點表同一份列）；結果欄預設「在」，
     現場只改例外列；每張表下面 5 列空白給「櫃上有但沒登記」；隱藏工作表記盤點編號＋櫃＋每列的 key，
     上傳時用它驗；key 欄鎖住。機房裡 NB 沒網路，所以整個流程要能離線。
  3. 回到座位「上傳」：驗檔（是不是這次盤點的檔、key 有沒有被動），錯誤講到「工作表 B2，第 5 列：…」。
     結果跟「基準」比，不是跟今天的資料比。可以只交一部分櫃；同一櫃重交＝取代，舊的留歷史。
     網頁也可以直接填一櫃（同一張表，預設在，只點例外）。
  4. 「差異清單」自動產生，每一列都有下一步（標記已不在櫃上／位置更正／CIA 待更正／補登 CIA 期限），
     狀態 待處理／已處理／不處理（要原因），留誰、何時。盤點期間系統裡被搬過的標「盤點期間異動」。
  5. 主管頁：進度、差異未處理、序號待補；「匯出盤點報告」。
  6. 每一箱有「上次盤點日期」；4-1 機房那一行顯示「超過一年未盤」。

口徑：一列＝一個實體箱子（rack_layout.rack_boxes 的 box／orphan／suspect，含在用、關機未下架、待確認；
已不在櫃上的不列）。所以「盤點列數＝基準箱子數」，verify_numbers 會核。
"""
from __future__ import annotations

import io
import json
import re
import secrets
import sqlite3
from datetime import date, datetime, timedelta
from typing import Any

RESULTS = ("在", "不在", "不符")
DIFF_STATUS = ("待處理", "已處理", "不處理")
EXTRA_ROWS = 5
STALE_DAYS = 365
DEFAULT_EXTRA_DAYS = 7
META_SHEET = "_meta"
META_MAGIC = "webit3-rack-audit"
HEAD_ROW = 4                       # 表頭在第 4 列，資料從第 5 列開始
COLS = ["項次", "U", "機型", "資產編號", "原廠序號", "主機名", "結果", "正確 U", "正確型號", "正確序號", "備註"]
EDITABLE = {"結果", "正確 U", "正確型號", "正確序號", "備註"}
EXTRA_EDITABLE = {"U", "機型", "原廠序號", "備註"}
DIFF_LABEL = {"missing": "不在", "u": "U 位不符", "model": "型號不符", "serial": "序號不符", "serial_new": "序號補登",
              "extra": "櫃上有但沒登記", "changed": "盤點期間異動"}
NEXT_STEP = {"missing": "確認是否已下架 → 標記已不在櫃上", "u": "位置更正（進 CIA 待更正）",
             "model": "已列入 CIA 待更正（型號）", "serial": "已列入 CIA 待更正（序號）",
             "serial_new": "已列入 CIA 待更正（序號補登）", "extra": "補登 CIA",
             "changed": "確認系統裡的異動是否正確"}
ACTIONS = {"missing": "mark_removed", "u": "fix_position"}


# ===== 資料表 =====

def ensure_tables(conn: sqlite3.Connection) -> None:
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_audit ("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT, scope_json TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'open',"
        " token TEXT NOT NULL, baseline_at TEXT NOT NULL, created_by TEXT, closed_by TEXT, closed_at TEXT)")
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_audit_rack ("
        " audit_id INTEGER NOT NULL, room TEXT NOT NULL, rack TEXT NOT NULL, snapshot_json TEXT NOT NULL,"
        " book_count INTEGER NOT NULL, PRIMARY KEY (audit_id, room, rack))")
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_audit_submit ("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, audit_id INTEGER NOT NULL, room TEXT NOT NULL, rack TEXT NOT NULL,"
        " version INTEGER NOT NULL, source TEXT NOT NULL, counted_by TEXT, counted_on TEXT, rows_json TEXT NOT NULL,"
        " extras_json TEXT NOT NULL, filename TEXT, submitted_by TEXT,"
        " submitted_at TEXT NOT NULL DEFAULT (datetime('now','localtime')), superseded INTEGER NOT NULL DEFAULT 0)")
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_audit_diff_state ("
        " audit_id INTEGER NOT NULL, diff_key TEXT NOT NULL, status TEXT NOT NULL, reason TEXT, updated_by TEXT,"
        " updated_at TEXT NOT NULL DEFAULT (datetime('now','localtime')), PRIMARY KEY (audit_id, diff_key))")
    conn.execute(
        "CREATE TABLE IF NOT EXISTS rack_audit_log ("
        " id INTEGER PRIMARY KEY AUTOINCREMENT, audit_id INTEGER NOT NULL, action TEXT NOT NULL, detail TEXT,"
        " by_user TEXT, at TEXT NOT NULL DEFAULT (datetime('now','localtime')))")


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _log(conn, audit_id: int, action: str, detail: Any, by: str | None) -> None:
    conn.execute("INSERT INTO rack_audit_log (audit_id, action, detail, by_user) VALUES (?,?,?,?)",
                 (audit_id, action, json.dumps(detail, ensure_ascii=False), by))


def extra_days(conn: sqlite3.Connection) -> int:
    import rack_layout
    try:
        return max(1, int(rack_layout._setting_json(conn, "rack_audit_extra_days") or DEFAULT_EXTRA_DAYS))
    except (TypeError, ValueError):
        return DEFAULT_EXTRA_DAYS


# ===== 基準（凍結的帳）=====

def _snapshot_rack(devs: list[dict], plan: dict) -> list[dict]:
    """一櫃的盤點列：一箱一列，從上到下（跟 A4 盤點表同一個排序）。"""
    import rack_layout as rl

    bx = rl.rack_boxes(devs, plan)
    by = {d["key"]: d for d in devs}
    members_of = {b["head"]: b["keys"][1:] for b in bx["boxes"]}
    rows = []
    for m in sorted(devs, key=lambda m: (m["u_start"] is None, -(m["u_end"] or 0), m["name"].lower())):
        c = bx["category"][m["key"]]
        if c not in rl.COUNTED:
            continue
        mem = [by[k] for k in members_of.get(m["key"], [])]
        ports = [by[k] for k in bx["ports_of"].get(m["key"], [])]
        rows.append({
            "key": m["key"], "category": c, "state": m["state"],
            "u_start": m["u_start"], "u_end": m["u_end"], "u": rl._u_text(m), "extra_ranges": m.get("extra_ranges") or [],
            "model": m["model"], "asset_serial": m["asset_serial"], "hw_serial": ", ".join(m["hw_serials"]) or "",
            "name": m["name"], "api_ids": list(m["api_ids"]),
            "members": [{"key": x["key"], "name": x["name"], "asset_serial": x["asset_serial"],
                         "role": rl.ROLE_LABEL.get(bx["roles"].get(x["key"], ""), "")} for x in mem],
            "ports": [{"key": x["key"], "name": x["name"], "ip": x["ip"], "kind": bx["mgmt"][x["key"]]["kind"]} for x in ports],
            "all_keys": [m["key"]] + [x["key"] for x in mem] + [x["key"] for x in ports],
        })
    for i, r in enumerate(rows, 1):
        r["no"] = i
    return rows


def _row_name_text(r: dict) -> str:
    """主機名格：本體＋同箱的其他登記＋管理口（離線檔裡看得到這一箱裝了什麼）。"""
    t = r["name"] or ""
    if r["members"]:
        t += "／" + "、".join(f"{x['name']}（{x['role']}）" for x in r["members"])
    if r["ports"]:
        t += "；" + "、".join(f"管理 IP：{x['ip'] or '無 IP'}（{x['kind']}）" for x in r["ports"])
    return t


def start_audit(conn: sqlite3.Connection, room: str | None, racks: list[dict] | None, by: str,
                title: str | None = None) -> dict:
    """開始盤點：凍結選定範圍的帳。room 給了、racks 沒給＝那個機房每一櫃；racks＝[{room, rack}]。"""
    import overlap_queue
    import rack_layout as rl

    ensure_tables(conn)
    ms = rl.machines(conn)
    have = {(m["room"], m["rack"]) for m in ms if m["kind"] == "rack" and m.get("rack")}
    if racks:
        scope = []
        for x in racks:
            k = (str(x.get("room") or room or "").strip(), rl._norm_rack(str(x.get("rack") or "").strip()))
            if k not in have:
                raise ValueError(f"沒有這一櫃：{k[0]} {k[1]}")
            if k not in scope:
                scope.append(k)
    elif room:
        scope = sorted({k for k in have if k[0] == room}, key=lambda k: rl._natural(k[1]))
        if not scope:
            raise ValueError(f"{room} 沒有機櫃資料")
    else:
        raise ValueError("要選一個機房，或至少一櫃")
    acts = overlap_queue.active_decisions(conn)
    baseline = _now()
    token = secrets.token_hex(12)
    rooms = sorted({k[0] for k in scope})
    title = (title or "").strip() or (f"{'、'.join(rooms)} 盤點 {baseline[:10]}")
    cur = conn.execute("INSERT INTO rack_audit (title, scope_json, token, baseline_at, created_by) VALUES (?,?,?,?,?)",
                       (title, json.dumps([list(k) for k in scope], ensure_ascii=False), token, baseline, by))
    aid = cur.lastrowid
    total = 0
    for rm, rk in scope:
        devs = [m for m in ms if m["kind"] == "rack" and m["room"] == rm and m["rack"] == rk]
        rows = _snapshot_rack(devs, overlap_queue.merge_plan(devs, rm, rk, acts))
        total += len(rows)
        conn.execute("INSERT INTO rack_audit_rack (audit_id, room, rack, snapshot_json, book_count) VALUES (?,?,?,?,?)",
                     (aid, rm, rk, json.dumps(rows, ensure_ascii=False), len(rows)))
    _log(conn, aid, "start", {"racks": len(scope), "boxes": total}, by)
    conn.commit()
    return {"id": aid, "title": title, "racks": len(scope), "boxes": total, "baseline_at": baseline}


def _audit(conn, aid: int) -> sqlite3.Row:
    ensure_tables(conn)
    a = conn.execute("SELECT * FROM rack_audit WHERE id = ?", (aid,)).fetchone()
    if not a:
        raise LookupError("查無這次盤點")
    return a


def _racks(conn, aid: int) -> list[dict]:
    out = []
    for r in conn.execute("SELECT * FROM rack_audit_rack WHERE audit_id = ?", (aid,)):
        out.append({"room": r["room"], "rack": r["rack"], "rows": json.loads(r["snapshot_json"]), "book_count": r["book_count"]})
    import rack_layout as rl
    out.sort(key=lambda x: (x["room"], rl._natural(x["rack"])))
    return out


def _active_submits(conn, aid: int) -> dict[tuple, dict]:
    out = {}
    for s in conn.execute("SELECT * FROM rack_audit_submit WHERE audit_id = ? AND superseded = 0", (aid,)):
        out[(s["room"], s["rack"])] = {**dict(s), "rows": json.loads(s["rows_json"]), "extras": json.loads(s["extras_json"])}
    return out


# ===== 清單、單次盤點、單櫃 =====

def list_audits(conn: sqlite3.Connection) -> list[dict]:
    ensure_tables(conn)
    out = []
    for a in conn.execute("SELECT * FROM rack_audit ORDER BY id DESC"):
        racks = conn.execute("SELECT COUNT(*), COALESCE(SUM(book_count), 0) FROM rack_audit_rack WHERE audit_id = ?",
                             (a["id"],)).fetchone()
        done = conn.execute("SELECT COUNT(DISTINCT room || '|' || rack) FROM rack_audit_submit WHERE audit_id = ? AND superseded = 0",
                            (a["id"],)).fetchone()[0]
        d = diffs(conn, a["id"])
        out.append({"id": a["id"], "title": a["title"], "status": a["status"], "baseline_at": a["baseline_at"],
                    "created_by": a["created_by"], "rack_total": racks[0], "book_count": racks[1], "rack_done": done,
                    "diff_count": len(d), "diff_open": sum(1 for x in d if x["status"] == "待處理"),
                    "rooms": sorted({r[0] for r in json.loads(a["scope_json"])})})
    return out


def audit_detail(conn: sqlite3.Connection, aid: int) -> dict:
    a = _audit(conn, aid)
    subs = _active_submits(conn, aid)
    d = diffs(conn, aid)
    racks = []
    for r in _racks(conn, aid):
        s = subs.get((r["room"], r["rack"]))
        dd = [x for x in d if (x["room"], x["rack"]) == (r["room"], r["rack"])]
        res = {k: 0 for k in RESULTS}
        if s:
            for row in r["rows"]:
                res[s["rows"].get(row["key"], {}).get("result", "在")] += 1
        racks.append({"room": r["room"], "rack": r["rack"], "book_count": r["book_count"], "done": bool(s),
                      "counted": sum(res.values()) if s else 0, "result_counts": res if s else None,
                      "extra_count": len(s["extras"]) if s else 0,
                      "counted_by": s["counted_by"] if s else None, "counted_on": s["counted_on"] if s else None,
                      "source": s["source"] if s else None, "version": s["version"] if s else 0,
                      "submitted_at": s["submitted_at"] if s else None,
                      "diff_count": len(dd), "diff_open": sum(1 for x in dd if x["status"] == "待處理"),
                      "row_keys": [x["key"] for x in r["rows"]]})
    return {"id": aid, "title": a["title"], "status": a["status"], "baseline_at": a["baseline_at"],
            "created_by": a["created_by"], "racks": racks, "extra_days": extra_days(conn)}


def rack_sheet(conn: sqlite3.Connection, aid: int, room: str, rack: str) -> dict:
    """網頁填寫用：基準列＋目前的結果（沒交過＝全部預設在）＋歷史版本。"""
    _audit(conn, aid)
    r = next((x for x in _racks(conn, aid) if (x["room"], x["rack"]) == (room, rack)), None)
    if r is None:
        raise LookupError(f"這次盤點沒有這一櫃：{room} {rack}")
    s = _active_submits(conn, aid).get((room, rack))
    rows = []
    for x in r["rows"]:
        got = (s["rows"].get(x["key"]) if s else None) or {}
        rows.append({**x, "name_text": _row_name_text(x), "result": got.get("result", "在"),
                     "correct_u": got.get("correct_u") or "", "correct_model": got.get("correct_model") or "",
                     "correct_serial": got.get("correct_serial") or "", "note": got.get("note") or ""})
    hist = [dict(h) for h in conn.execute(
        "SELECT id, version, source, counted_by, counted_on, submitted_by, submitted_at, superseded, filename "
        "FROM rack_audit_submit WHERE audit_id = ? AND room = ? AND rack = ? ORDER BY version DESC", (aid, room, rack))]
    return {"room": room, "rack": rack, "rows": rows, "extras": s["extras"] if s else [], "done": bool(s),
            "counted_by": s["counted_by"] if s else "", "counted_on": s["counted_on"] if s else "", "history": hist}


# ===== 交結果（上傳與網頁共用）=====

_RE_U = re.compile(r"^\s*(\d{1,2})\s*(?:[-~～]\s*(\d{1,2}))?\s*U?\s*$", re.I)


def parse_u(v) -> tuple[int, int] | None:
    """「12」「12U」「12-15」「12-15U」→ (12, 15)；空白回 None；格式錯丟 ValueError。"""
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (int, float)) and float(v).is_integer():
        return int(v), int(v)
    m = _RE_U.match(str(v))
    if not m:
        raise ValueError("U 要寫成 12 或 12-15")
    a, b = int(m.group(1)), int(m.group(2) or m.group(1))
    if a < 1 or b < 1:
        raise ValueError("U 至少是 1")
    return min(a, b), max(a, b)


def _norm(v) -> str:
    return " ".join(str(v or "").split())


def _clean_rows(snapshot: list[dict], rows: dict, where) -> tuple[dict, list[str], list[str]]:
    """檢查一櫃的結果。where(row_key, 欄) → 錯誤訊息前綴。回 (乾淨的結果, 錯誤, 提醒)。"""
    errs, warns, out = [], [], {}
    for x in snapshot:
        got = rows.get(x["key"]) or {}
        res = _norm(got.get("result")) or "在"
        if res not in RESULTS:
            errs.append(f"{where(x['key'], '結果')}結果欄只能是 在/不在/不符")
            continue
        cu, cm, cs = got.get("correct_u"), _norm(got.get("correct_model")), _norm(got.get("correct_serial"))
        try:
            uu = parse_u(cu)
        except ValueError as exc:
            errs.append(f"{where(x['key'], '正確 U')}正確 U {exc}")
            continue
        changed = (uu and uu != (x["u_start"], x["u_end"])) or (cm and cm.upper() != _norm(x["model"]).upper()) \
            or (cs and cs.upper() not in {s.strip().upper() for s in (x["hw_serial"] or "").split(",") if s.strip()})
        if res == "不符" and not (uu or cm or cs):
            errs.append(f"{where(x['key'], '結果')}結果「不符」要填正確 U、正確型號或正確序號其中一個")
            continue
        if res == "在" and changed:
            res = "不符"
            warns.append(f"{where(x['key'], '結果')}有填正確值，視為「不符」")
        if res == "不在" and (uu or cm or cs):
            warns.append(f"{where(x['key'], '結果')}結果是「不在」，正確值不採用")
            uu, cm, cs = None, "", ""
        out[x["key"]] = {"result": res, "correct_u": f"{uu[0]}-{uu[1]}" if uu and uu[0] != uu[1] else (str(uu[0]) if uu else ""),
                         "correct_model": cm, "correct_serial": cs, "note": _norm(got.get("note"))}
    return out, errs, warns


def _clean_extras(extras: list[dict], where) -> tuple[list[dict], list[str]]:
    errs, out = [], []
    for i, e in enumerate(extras or []):
        model, serial, note = _norm(e.get("model")), _norm(e.get("serial")), _norm(e.get("note"))
        try:
            uu = parse_u(e.get("u"))
        except ValueError as exc:
            errs.append(f"{where(i, 'U')}U {exc}")
            continue
        if not (model or serial or uu or note):
            continue
        if not (model or serial):
            errs.append(f"{where(i, '機型')}櫃上有但沒登記：機型或序號至少填一個")
            continue
        out.append({"model": model, "serial": serial, "u": (f"{uu[0]}-{uu[1]}" if uu[0] != uu[1] else str(uu[0])) if uu else "",
                    "note": note})
    return out, errs


def submit_rack(conn: sqlite3.Connection, aid: int, room: str, rack: str, rows: dict, extras: list[dict],
                counted_by: str, counted_on: str | None, by: str, source: str = "web", filename: str | None = None,
                commit: bool = True) -> dict:
    """存一櫃的結果。同一櫃再交＝新版本取代舊的（舊的留著、標 superseded）。"""
    a = _audit(conn, aid)
    if a["status"] != "open":
        raise ValueError("這次盤點已結案，不能再交結果")
    r = next((x for x in _racks(conn, aid) if (x["room"], x["rack"]) == (room, rack)), None)
    if r is None:
        raise LookupError(f"這次盤點沒有這一櫃：{room} {rack}")
    if not _norm(counted_by):
        raise ValueError("要填盤點人")
    no = {x["key"]: x["no"] for x in r["rows"]}
    clean, errs, warns = _clean_rows(r["rows"], rows or {}, lambda k, c: f"{rack} 第 {no[k]} 項：")
    ex, errs2 = _clean_extras(extras, lambda i, c: f"{rack} 新增第 {i + 1} 列：")
    errs += errs2
    if errs:
        raise ValueError("；".join(errs[:20]))
    on = _norm(counted_on) or date.today().isoformat()
    return _save(conn, aid, room, rack, clean, ex, _norm(counted_by), on, by, source, filename, warns, commit)


def _save(conn, aid, room, rack, clean, ex, counted_by, on, by, source, filename, warns, commit=True) -> dict:
    v = conn.execute("SELECT COALESCE(MAX(version), 0) FROM rack_audit_submit WHERE audit_id = ? AND room = ? AND rack = ?",
                     (aid, room, rack)).fetchone()[0]
    conn.execute("UPDATE rack_audit_submit SET superseded = 1 WHERE audit_id = ? AND room = ? AND rack = ?", (aid, room, rack))
    conn.execute("INSERT INTO rack_audit_submit (audit_id, room, rack, version, source, counted_by, counted_on, rows_json, "
                 "extras_json, filename, submitted_by) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                 (aid, room, rack, v + 1, source, counted_by, on, json.dumps(clean, ensure_ascii=False),
                  json.dumps(ex, ensure_ascii=False), filename, by))
    _log(conn, aid, "submit", {"room": room, "rack": rack, "version": v + 1, "source": source,
                               "not_ok": sum(1 for x in clean.values() if x["result"] != "在"), "extras": len(ex)}, by)
    if commit:
        conn.commit()
    return {"room": room, "rack": rack, "version": v + 1, "replaced": v > 0, "warnings": warns,
            "counted": len(clean), "extras": len(ex)}


# ===== 離線 Excel =====

def _sheet_name(room: str, rack: str, multi_room: bool, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", f"{room}-{rack}" if multi_room else rack)[:28] or "rack"
    n, i = base, 2
    while n in used or n == META_SHEET:
        n = f"{base[:26]}~{i}"
        i += 1
    used.add(n)
    return n


def build_xlsx(conn: sqlite3.Connection, aid: int) -> tuple[io.BytesIO, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    a = _audit(conn, aid)
    racks = _racks(conn, aid)
    multi = len({r["room"] for r in racks}) > 1
    wb = Workbook()
    wb.remove(wb.active)
    meta = []
    used: set = set()
    head_fill = PatternFill("solid", fgColor="00806A")
    edit_fill = PatternFill("solid", fgColor="FFF8E1")
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    unlocked = Protection(locked=False)
    widths = [6, 8, 22, 16, 16, 34, 8, 9, 18, 16, 24]
    for r in racks:
        name = _sheet_name(r["room"], r["rack"], multi, used)
        ws = wb.create_sheet(name)
        ws["A1"] = f"機櫃盤點　{r['room']}　{r['rack']}　（盤點 #{aid}，基準 {a['baseline_at'][:16]}，帳上 {r['book_count']} 台）"
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"], ws["D2"] = "盤點人", "盤點日期"
        for c in ("B2", "E2"):
            ws[c].protection = unlocked
            ws[c].fill = edit_fill
            ws[c].border = border
        ws["G2"] = "只改例外列；盤點人沒填＝這櫃還沒盤（上傳時略過）"
        ws["G2"].font = Font(size=9, color="666666")
        for i, h in enumerate(COLS, 1):
            c = ws.cell(HEAD_ROW, i, h)
            c.fill, c.font, c.border = head_fill, Font(color="FFFFFF", bold=True), border
            c.alignment = Alignment(horizontal="center")
            ws.column_dimensions[c.column_letter].width = widths[i - 1]
        dv = DataValidation(type="list", formula1='"在,不在,不符"', allow_blank=True, showErrorMessage=True,
                            errorTitle="結果", error="只能是 在／不在／不符")
        ws.add_data_validation(dv)
        rown = HEAD_ROW + 1
        row_meta = []
        for x in r["rows"]:
            vals = [x["no"], x["u"] or "位置不明", x["model"] or "", x["asset_serial"] or "", x["hw_serial"],
                    _row_name_text(x), "在", "", "", "", ""]
            for i, v in enumerate(vals, 1):
                c = ws.cell(rown, i, v)
                c.border = border
                if COLS[i - 1] in EDITABLE:
                    c.protection = unlocked
                    c.fill = edit_fill
            dv.add(ws.cell(rown, 7))
            row_meta.append([rown, x["key"], x["no"], x["asset_serial"] or ""])
            rown += 1
        ws.cell(rown, 1, "櫃上有但沒登記（有才填：機型／序號／U／備註）").font = Font(bold=True, color="00806A")
        rown += 1
        extra_first = rown
        for j in range(EXTRA_ROWS):
            ws.cell(rown, 1, f"新增{j + 1}").border = border
            for i, h in enumerate(COLS, 1):
                c = ws.cell(rown, i)
                c.border = border
                if h in EXTRA_EDITABLE:
                    c.protection = unlocked
                    c.fill = edit_fill
            rown += 1
        ws.freeze_panes = ws.cell(HEAD_ROW + 1, 1)
        ws.protection.sheet = True                     # key 欄（項次、U、機型、資產編號、序號、主機名）鎖住
        ws.protection.formatColumns = False
        meta.append({"sheet": name, "room": r["room"], "rack": r["rack"], "rows": row_meta, "extra_first": extra_first})
    ms = wb.create_sheet(META_SHEET)
    ms["A1"], ms["B1"], ms["C1"] = META_MAGIC, aid, a["token"]
    ms["A2"] = json.dumps(meta, ensure_ascii=False)
    ms.sheet_state = "veryHidden"
    ms.protection.sheet = True
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, f"機櫃盤點_{aid}_{a['baseline_at'][:10]}.xlsx"


def parse_upload(conn: sqlite3.Connection, aid: int, data: bytes, by: str, filename: str | None = None) -> dict:
    """上傳離線檔。任何一格有錯就整份不存（講出每一格錯在哪）；沒填盤點人的櫃略過（部分上傳）。"""
    from openpyxl import load_workbook

    a = _audit(conn, aid)
    if a["status"] != "open":
        raise ValueError("這次盤點已結案，不能再交結果")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:                         # noqa: BLE001 - 檔案壞掉要講給使用者，不吞
        raise ValueError(f"讀不了這個檔（不是 xlsx，或檔案壞了）：{exc}") from None
    if META_SHEET not in wb.sheetnames:
        raise ValueError("這不是系統下載的盤點檔（找不到隱藏的盤點資訊）；請用「下載盤點檔」那一份")
    mws = wb[META_SHEET]
    if mws["A1"].value != META_MAGIC:
        raise ValueError("這不是系統下載的盤點檔")
    if str(mws["B1"].value) != str(aid) or mws["C1"].value != a["token"]:
        raise ValueError(f"這個檔是盤點 #{mws['B1'].value} 的，不是這次（#{aid}）")
    try:
        meta = json.loads(mws["A2"].value or "[]")
    except (TypeError, ValueError):
        raise ValueError("盤點檔的隱藏資訊被改壞了，請重新下載") from None
    racks = {(r["room"], r["rack"]): r for r in _racks(conn, aid)}
    errs, warns, todo, skipped = [], [], [], []
    for sm in meta:
        sheet, key = sm["sheet"], (sm["room"], sm["rack"])
        if key not in racks:
            errs.append(f"工作表 {sheet}：這次盤點沒有這一櫃")
            continue
        if sheet not in wb.sheetnames:
            skipped.append({"room": key[0], "rack": key[1], "why": "工作表被刪掉"})
            continue
        ws = wb[sheet]
        counted_by = _norm(ws["B2"].value)
        if not counted_by:
            skipped.append({"room": key[0], "rack": key[1], "why": "盤點人沒填，視為還沒盤"})
            continue
        on = ws["E2"].value
        on = on.strftime("%Y-%m-%d") if isinstance(on, (datetime, date)) else (_norm(on) or date.today().isoformat())
        if [rm[1] for rm in sm["rows"]] != [x["key"] for x in racks[key]["rows"]]:
            errs.append(f"工作表 {sheet}：盤點資訊跟系統的基準對不上，請重新下載")
            continue
        rows, rowat, bad = {}, {}, False
        for rown, k, no, aser in sm["rows"]:
            rowat[k] = rown
            if str(ws.cell(rown, 1).value) != str(no) or _norm(ws.cell(rown, 4).value) != _norm(aser):
                errs.append(f"工作表 {sheet}，第 {rown} 列：項次或資產編號被改過（列被移動、刪除或插入），請用原檔重填")
                bad = True
                continue
            rows[k] = {"result": ws.cell(rown, 7).value, "correct_u": ws.cell(rown, 8).value,
                       "correct_model": ws.cell(rown, 9).value, "correct_serial": ws.cell(rown, 10).value,
                       "note": ws.cell(rown, 11).value}
        if bad:
            continue
        clean, e1, w1 = _clean_rows(racks[key]["rows"], rows,
                                    lambda k, c: f"工作表 {sheet}，第 {rowat[k]} 列：")
        extras = []
        for j in range(EXTRA_ROWS):
            rn = sm["extra_first"] + j
            extras.append({"u": ws.cell(rn, 2).value, "model": ws.cell(rn, 3).value, "serial": ws.cell(rn, 5).value,
                           "note": ws.cell(rn, 11).value})
        ex, e2 = _clean_extras(extras, lambda i, c: f"工作表 {sheet}，第 {sm['extra_first'] + i} 列：")
        errs += e1 + e2
        warns += w1
        if not (e1 or e2):
            todo.append((key, clean, ex, counted_by, on))
    if errs:
        return {"ok": False, "errors": errs, "warnings": warns, "saved": [], "skipped": skipped}
    saved = []
    for (rm, rk), clean, ex, cb, on in todo:
        saved.append(_save(conn, aid, rm, rk, clean, ex, cb, on, by, "xlsx", filename, [], commit=False))
    conn.commit()
    return {"ok": True, "errors": [], "warnings": warns, "saved": saved, "skipped": skipped}


# ===== 差異清單 =====

def _changes_after(conn, aid: int, baseline: str, keys: list[str]) -> list[str]:
    """基準之後系統裡對這些設備的異動（不含這次盤點自己按的一鍵處理）。"""
    import rack_layout as rl

    rl.ensure_position_table(conn)
    rl.ensure_state_table(conn)
    mine = f"盤點#{aid}"
    out = []
    q = ",".join("?" for _ in keys)
    for r in conn.execute(f"SELECT key, action, note, updated_by, updated_at FROM rack_position_audit "
                          f"WHERE key IN ({q}) AND updated_at >= ?", (*keys, baseline)):
        if not (r["note"] or "").startswith(mine):
            out.append(f"位置{'更正' if r['action'] == 'set' else r['action']}（{r['updated_by'] or '—'} {r['updated_at'][:16]}）")
    for r in conn.execute(f"SELECT key, new_state, note, updated_by, updated_at FROM rack_state_audit "
                          f"WHERE key IN ({q}) AND updated_at >= ?", (*keys, baseline)):
        if not (r["note"] or "").startswith(mine):
            out.append(f"狀態改成 {r['new_state'] or '清除'}（{r['updated_by'] or '—'} {r['updated_at'][:16]}）")
    return out


def diffs(conn: sqlite3.Connection, aid: int) -> list[dict]:
    """自動產生的差異清單（只看已交結果的櫃；跟基準比）。每一列：種類、下一步、狀態。"""
    a = _audit(conn, aid)
    subs = _active_submits(conn, aid)
    states = {r["diff_key"]: dict(r) for r in conn.execute("SELECT * FROM rack_audit_diff_state WHERE audit_id = ?", (aid,))}
    days = extra_days(conn)
    today = date.today()
    out = []

    def add(rk, row, kind, book, found, s, extra_i=None):
        dk = f"{rk['room']}|{rk['rack']}|{row['key'] if row else f'extra{extra_i}'}|{kind}"
        st = states.get(dk) or {}
        d = {"diff_key": dk, "room": rk["room"], "rack": rk["rack"], "kind": kind, "label": DIFF_LABEL[kind],
             "next_step": NEXT_STEP[kind], "action": ACTIONS.get(kind), "book": book, "found": found,
             "key": row["key"] if row else None, "all_keys": row["all_keys"] if row else [],
             "no": row["no"] if row else None, "name": row["name"] if row else (found or ""),
             "asset_serial": row["asset_serial"] if row else "", "model": row["model"] if row else "",
             "u": row["u"] if row else "", "counted_by": s["counted_by"], "counted_on": s["counted_on"],
             "note": (s["rows"].get(row["key"], {}).get("note") if row else s["extras"][extra_i].get("note")) or "",
             "status": st.get("status") or "待處理", "reason": st.get("reason"), "status_by": st.get("updated_by"),
             "status_at": st.get("updated_at"), "during_change": [], "due": None, "aging_days": None, "overdue": False,
             "extra": s["extras"][extra_i] if extra_i is not None else None}
        if kind == "extra":
            try:
                d0 = date.fromisoformat(s["counted_on"][:10])
            except (TypeError, ValueError):
                d0 = today
            d["due"] = (d0 + timedelta(days=days)).isoformat()
            d["aging_days"] = (today - d0).days
            d["overdue"] = d["status"] == "待處理" and today > d0 + timedelta(days=days)
        out.append(d)
        return d

    for rk in _racks(conn, aid):
        s = subs.get((rk["room"], rk["rack"]))
        if not s:
            continue
        for row in rk["rows"]:
            got = s["rows"].get(row["key"], {"result": "在"})
            made = []
            if got["result"] == "不在":
                made.append(add(rk, row, "missing", f"{row['u']} {row['model'] or ''}".strip(), "現場沒有", s))
            elif got["result"] == "不符":
                uu = parse_u(got.get("correct_u")) if got.get("correct_u") else None
                if uu and uu != (row["u_start"], row["u_end"]):
                    made.append(add(rk, row, "u", row["u"] or "位置不明", f"{uu[0]}U" if uu[0] == uu[1] else f"{uu[0]}-{uu[1]}U", s))
                if got.get("correct_model") and got["correct_model"].upper() != _norm(row["model"]).upper():
                    made.append(add(rk, row, "model", row["model"] or "（空白）", got["correct_model"], s))
                book_sn = {x.strip().upper() for x in (row["hw_serial"] or "").split(",") if x.strip()}
                if got.get("correct_serial") and got["correct_serial"].upper() not in book_sn:
                    made.append(add(rk, row, "serial" if book_sn else "serial_new", row["hw_serial"] or "（帳上沒有）",
                                    got["correct_serial"], s))
            ch = _changes_after(conn, aid, a["baseline_at"], row["all_keys"])
            if ch:
                if not made:
                    made.append(add(rk, row, "changed", f"{row['u']}（基準）", "；".join(ch[:3]), s))
                for d in made:
                    d["during_change"] = ch
        for i, e in enumerate(s["extras"]):
            add(rk, None, "extra", "（沒登記）", " ".join(x for x in (e["model"], e["serial"], e["u"] and f"{e['u']}U") if x), s, i)
    for i, d in enumerate(out, 1):
        d["id"] = i
    return out


def set_diff_status(conn: sqlite3.Connection, aid: int, diff_key: str, status: str, reason: str | None, by: str) -> dict:
    _audit(conn, aid)
    if status not in DIFF_STATUS:
        raise ValueError("狀態只能是 待處理／已處理／不處理")
    if status == "不處理" and not _norm(reason):
        raise ValueError("選「不處理」要寫原因")
    if not any(d["diff_key"] == diff_key for d in diffs(conn, aid)):
        raise LookupError("查無這筆差異（可能那一櫃重交後已經沒有這個差異）")
    conn.execute("INSERT INTO rack_audit_diff_state (audit_id, diff_key, status, reason, updated_by, updated_at) "
                 "VALUES (?,?,?,?,?,datetime('now','localtime')) ON CONFLICT(audit_id, diff_key) DO UPDATE SET "
                 "status = excluded.status, reason = excluded.reason, updated_by = excluded.updated_by, "
                 "updated_at = excluded.updated_at", (aid, diff_key, status, _norm(reason) or None, by))
    _log(conn, aid, "diff_status", {"diff_key": diff_key, "status": status, "reason": reason}, by)
    conn.commit()
    return {"diff_key": diff_key, "status": status}


def act(conn: sqlite3.Connection, aid: int, diff_key: str, action: str, by: str) -> dict:
    """一鍵處理：不在 → 整箱標記已不在櫃上；U 不符 → 整箱位置更正（走拖拉搬移，會進 CIA 待更正、可以復原）。"""
    import rack_layout as rl

    d = next((x for x in diffs(conn, aid) if x["diff_key"] == diff_key), None)
    if d is None:
        raise LookupError("查無這筆差異")
    if d["status"] != "待處理":
        raise ValueError(f"這筆已經是「{d['status']}」")
    note = f"盤點#{aid} {d['counted_on']} {d['counted_by']}：{d['label']}"
    if action == "mark_removed" and d["kind"] == "missing":
        for k in d["all_keys"]:
            rl.set_state_override(conn, k, "removed", note, by, commit=False)
        out = {"done": "已標記已不在櫃上", "keys": d["all_keys"]}
    elif action == "fix_position" and d["kind"] == "u":
        uu = parse_u(d["found"])
        r = rl.move_box(conn, d["key"], d["room"], d["rack"], uu[0], uu[1] - uu[0] + 1, note, by)
        out = {"done": f"已更正位置到 {d['found']}", "move_id": r["move_id"], "moved": r["moved"]}
    else:
        raise ValueError("這種差異沒有一鍵處理")
    conn.execute("INSERT INTO rack_audit_diff_state (audit_id, diff_key, status, reason, updated_by, updated_at) "
                 "VALUES (?,?,'已處理',?,?,datetime('now','localtime')) ON CONFLICT(audit_id, diff_key) DO UPDATE SET "
                 "status='已處理', reason=excluded.reason, updated_by=excluded.updated_by, updated_at=excluded.updated_at",
                 (aid, diff_key, out["done"], by))
    _log(conn, aid, "act", {"diff_key": diff_key, "action": action}, by)
    conn.commit()
    return out


# ===== 主管頁、報告 =====

def summary(conn: sqlite3.Connection, aid: int) -> dict:
    det = audit_detail(conn, aid)
    d = diffs(conn, aid)
    subs = _active_submits(conn, aid)
    no_sn, sn_found = [], []
    for rk in _racks(conn, aid):
        s = subs.get((rk["room"], rk["rack"]))
        for row in rk["rows"]:
            if not row["hw_serial"]:
                got = (s["rows"].get(row["key"], {}) if s else {}).get("correct_serial")
                (sn_found if got else no_sn).append({"room": rk["room"], "rack": rk["rack"], "key": row["key"],
                                                     "no": row["no"], "name": row["name"], "asset_serial": row["asset_serial"],
                                                     "model": row["model"], "u": row["u"], "serial_found": got or ""})
    racks = det["racks"]
    by_kind = {}
    for x in d:
        by_kind.setdefault(x["label"], {"total": 0, "待處理": 0, "已處理": 0, "不處理": 0})
        by_kind[x["label"]]["total"] += 1
        by_kind[x["label"]][x["status"]] += 1
    return {"id": aid, "title": det["title"], "baseline_at": det["baseline_at"], "status": det["status"],
            "rack_total": len(racks), "rack_done": sum(1 for r in racks if r["done"]),
            "book_count": sum(r["book_count"] for r in racks), "counted": sum(r["counted"] for r in racks),
            "people": sorted({r["counted_by"] for r in racks if r["counted_by"]}),
            "diff_total": len(d), "diff_open": sum(1 for x in d if x["status"] == "待處理"),
            "diff_overdue": sum(1 for x in d if x["overdue"]), "by_kind": by_kind,
            "serial_missing": no_sn, "serial_found": sn_found}


def report_xlsx(conn: sqlite3.Connection, aid: int) -> tuple[io.BytesIO, str]:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    s = summary(conn, aid)
    det = audit_detail(conn, aid)
    d = diffs(conn, aid)
    wb = Workbook()
    ws = wb.active
    ws.title = "摘要"
    rows = [("盤點", f"#{aid} {s['title']}"), ("範圍", "、".join(f"{r['room']} {r['rack']}" for r in det["racks"])),
            ("基準日", s["baseline_at"]), ("盤點人", "、".join(s["people"]) or "—"),
            ("進度", f"{s['rack_done']}／{s['rack_total']} 櫃"), ("帳上台數（範圍內）", s["book_count"]),
            ("已盤台數", s["counted"]), ("差異", s["diff_total"]), ("差異未處理", s["diff_open"]),
            ("逾期未補登", s["diff_overdue"]), ("序號待補", len(s["serial_missing"])), ("現場抄到序號", len(s["serial_found"]))]
    for k, v in rows:
        ws.append([k, v])
    ws.append([])
    ws.append(["差異種類", "總數", "待處理", "已處理", "不處理"])
    for k, v in s["by_kind"].items():
        ws.append([k, v["total"], v["待處理"], v["已處理"], v["不處理"]])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 20, 60
    hdr = PatternFill("solid", fgColor="00806A")

    def table(title, cols, data):
        w = wb.create_sheet(title)
        w.append(cols)
        for c in w[1]:
            c.fill, c.font = hdr, Font(color="FFFFFF", bold=True)
        for r in data:
            w.append(r)
        w.freeze_panes = "A2"
    table("各櫃", ["項次", "機房", "機櫃", "帳上", "已盤", "在", "不在", "不符", "沒登記", "盤點人", "盤點日期", "差異", "未處理"],
          [[i, r["room"], r["rack"], r["book_count"], r["counted"],
            *([r["result_counts"][k] for k in RESULTS] if r["result_counts"] else ["", "", ""]),
            r["extra_count"], r["counted_by"] or "", r["counted_on"] or "", r["diff_count"], r["diff_open"]]
           for i, r in enumerate(det["racks"], 1)])
    table("差異", ["項次", "機房", "機櫃", "種類", "設備", "資產編號", "帳上", "現場", "下一步", "狀態", "原因／處理", "處理人", "處理時間",
                 "盤點期間異動", "補登期限"],
          [[x["id"], x["room"], x["rack"], x["label"], x["name"], x["asset_serial"], x["book"], x["found"], x["next_step"],
            x["status"], x["reason"] or "", x["status_by"] or "", x["status_at"] or "", "；".join(x["during_change"]), x["due"] or ""]
           for x in d])
    table("序號待補", ["項次", "機房", "機櫃", "盤點項次", "設備", "資產編號", "機型", "U", "現場抄到的序號"],
          [[i, x["room"], x["rack"], x["no"], x["name"], x["asset_serial"], x["model"], x["u"], x["serial_found"]]
           for i, x in enumerate(s["serial_missing"] + s["serial_found"], 1)])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, f"盤點報告_{aid}_{date.today().strftime('%Y%m%d')}.xlsx"


def close_audit(conn: sqlite3.Connection, aid: int, by: str) -> dict:
    a = _audit(conn, aid)
    if a["status"] != "open":
        raise ValueError("已經結案")
    conn.execute("UPDATE rack_audit SET status = 'closed', closed_by = ?, closed_at = datetime('now','localtime') WHERE id = ?",
                 (by, aid))
    _log(conn, aid, "close", {}, by)
    conn.commit()
    return {"id": aid, "status": "closed"}


# ===== 上次盤點日期、CIA 待更正 =====

def last_audit_map(conn: sqlite3.Connection) -> dict[str, str]:
    """每一筆登記 key → 最近一次有結果的盤點日期（整箱一起算：本體、同箱、管理口）。"""
    ensure_tables(conn)
    snaps = {(r["audit_id"], r["room"], r["rack"]): json.loads(r["snapshot_json"])
             for r in conn.execute("SELECT * FROM rack_audit_rack WHERE (audit_id, room, rack) IN "
                                   "(SELECT audit_id, room, rack FROM rack_audit_submit WHERE superseded = 0)")}
    out: dict[str, str] = {}
    for s in conn.execute("SELECT audit_id, room, rack, counted_on FROM rack_audit_submit WHERE superseded = 0"):
        on = (s["counted_on"] or "")[:10]
        for row in snaps.get((s["audit_id"], s["room"], s["rack"]), []):
            for k in row["all_keys"]:
                if on > out.get(k, ""):
                    out[k] = on
    return out


def stale_keys(conn: sqlite3.Connection, keys: list[str], last: dict[str, str] | None = None) -> list[str]:
    """超過一年未盤（含從來沒盤過）。"""
    last = last if last is not None else last_audit_map(conn)
    cut = (date.today() - timedelta(days=STALE_DAYS)).isoformat()
    return [k for k in keys if last.get(k, "") < cut]


def cia_rows(conn: sqlite3.Connection) -> list[dict]:
    """盤點出來、要 CIA 改的（型號不符、序號不符／補登、櫃上有但沒登記）；不處理的不列。欄名同 CIA 待更正。"""
    ensure_tables(conn)
    out = []
    for a in conn.execute("SELECT id FROM rack_audit ORDER BY id"):
        for d in diffs(conn, a["id"]):
            if d["kind"] not in ("model", "serial", "serial_new", "extra") or d["status"] == "不處理":
                continue
            out.append({"更正類型": f"盤點：{d['label']}", "資產序號": d["asset_serial"], "主機名稱": d["name"] if d["key"] else "",
                        "原機房": d["room"], "原Rack": d["rack"], "原U位": d["u"],
                        "新機房": d["room"] if d["kind"] == "extra" else "", "新Rack": d["rack"] if d["kind"] == "extra" else "",
                        "新U位": "", "設備U數": "", "盤點人": d["counted_by"], "現場盤點日期": d["counted_on"],
                        "備註": f"盤點 #{a['id']}（{d['status']}）" + (f"：{d['note']}" if d["note"] else "")
                                + (f"；補登期限 {d['due']}" if d["due"] else ""),
                        "CIA原值（機櫃欄）": "", "正確型號": d["found"] if d["kind"] == "model" else "",
                        "正確序號": d["found"] if d["kind"] in ("serial", "serial_new") else ""})
            if d["kind"] == "extra":
                e = d["extra"]
                out[-1].update({"正確型號": e["model"], "正確序號": e["serial"], "新U位": f"{e['u']}U" if e["u"] else "",
                                "原機房": "", "原Rack": ""})
    return out
