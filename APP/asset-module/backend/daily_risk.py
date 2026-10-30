"""設備風險清單＋出事聯絡資訊（v1.499.0，2026-10-23）。

使用者怎麼用：每天打開一頁，看今天哪些機器要注意；出事時立刻看到要打給誰。

四個分頁，全部以「台」（system_stats.machine_key，全站唯一的同一台規則）為單位，退役不算：

* 長時間未重開機：host_spec.spec_json.uptime（收集當下的開機時間）→ 推回開機時刻 → 到今天幾天。
* 硬體異常：只用已經收到的資料（SMART、軟體 RAID、硬體 RAID、多路徑、健康基準的磁碟使用率），
  每一筆標來源；沒收到的來源講覆蓋台數，不猜。
* 單點（無備援）：每個 APID 的「正式」環境現役機器，1 台＝單機；全部在同一機房＝同機房。
  機房未填無法判定的另列，不算單點。**不寫 ci_node.is_spof**（那是人工維護欄位，4-2 等頁在讀）。
* 保固／合約到期：聯絡資訊（手動三層＋合約匯入）解析出來的合約迄日分級。0 筆資料也要能顯示。

聯絡資訊（使用者 2026-10-23：沒有合約清冊，「聯絡資訊先設計，以後慢慢補」）：
主路徑是手動逐筆補，三層——型號（套到該型號全部機器）／系統（APID）／本機；越具體的贏。
合約匯入是次要路徑，匯入結果夾在對應層級之後（見 _resolve_fields 的順序）。每次異動都留稽核。
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import json
import re
import sqlite3

import manage_state
import system_report
import system_stats

# ---------------------------------------------------------------- 設定
SET_UPTIME_DAYS = "risk_uptime_days"
SET_DISK_PCT = "risk_disk_pct"
SET_CONTRACT_MAPPING = "contract_import_mapping"
DEFAULT_UPTIME_DAYS = 365
DEFAULT_DISK_PCT = 90

CONTACT_FIELDS = ("vendor", "phone", "contact", "contract_no", "contract_end", "note")
#: 「有聯絡或合約資料」看這三欄任一有值（窗口、合約編號單獨存在打不了電話，也看不出到期）
KEY_FIELDS = ("vendor", "phone", "contract_end")
LEVELS = ("asset", "system", "model")
LEVEL_LABEL = {"asset": "本機", "system": "系統", "model": "型號"}

NO_APID = "（未填 APID）"
NO_MODEL = "（型號未填）"
BAND_EXPIRED, BAND_90, BAND_180, BAND_OK, BAND_NONE = "expired", "d90", "d180", "ok", "none"
BAND_LABEL = {BAND_EXPIRED: "已過期", BAND_90: "90 天內", BAND_180: "180 天內",
              BAND_OK: "180 天以上", BAND_NONE: "無合約資料"}


def _setting(conn, key: str):
    try:
        r = conn.execute("SELECT value FROM app_settings WHERE key = ?", (key,)).fetchone()
    except sqlite3.Error:
        return None
    return r[0] if r else None


def _int_setting(conn, key: str, default: int, lo: int, hi: int) -> int:
    try:
        v = int(str(_setting(conn, key)).strip())
    except (TypeError, ValueError):
        return default
    return v if lo <= v <= hi else default


def settings(conn) -> dict:
    return {"uptime_days": _int_setting(conn, SET_UPTIME_DAYS, DEFAULT_UPTIME_DAYS, 1, 3650),
            "disk_pct": _int_setting(conn, SET_DISK_PCT, DEFAULT_DISK_PCT, 50, 100)}


def save_settings(conn, uptime_days: int | None, disk_pct: int | None) -> dict:
    if uptime_days is not None:
        if not 1 <= int(uptime_days) <= 3650:
            raise ValueError("未重開機天數要在 1～3650 之間")
        conn.execute("INSERT OR REPLACE INTO app_settings(key, value) VALUES (?,?)",
                     (SET_UPTIME_DAYS, str(int(uptime_days))))
    if disk_pct is not None:
        if not 50 <= int(disk_pct) <= 100:
            raise ValueError("磁碟使用率門檻要在 50～100 之間")
        conn.execute("INSERT OR REPLACE INTO app_settings(key, value) VALUES (?,?)",
                     (SET_DISK_PCT, str(int(disk_pct))))
    conn.commit()
    return settings(conn)


# ---------------------------------------------------------------- 機器清單（台）
def norm_model(s) -> str:
    return " ".join(str(s or "").split()).upper()


def norm_serial(s) -> str:
    return re.sub(r"\s+", "", str(s or "")).upper()


def norm_host(s) -> str:
    return system_stats.strip_domain(str(s or "").strip()).lower()


def _first(rows, field):
    for r in rows:
        v = r[field]
        if v is not None and str(v).strip():
            return str(v).strip()
    return None


def machines(conn) -> list[dict]:
    """現役（非退役）登記依 machine_key 收成「台」。代表列優先取帳內（CIA）那筆。"""
    rows = conn.execute(
        "SELECT asset_serial, hostname, ip, api_id, environment, physical_location, rack_no, os, "
        "device_model, hw_serial, asset_status, is_vm FROM hardware").fetchall()
    rows = [r for r in rows if not system_stats.is_retired_row(r)]
    spec_model = {}
    try:
        for r in conn.execute("SELECT asset_serial, model FROM host_spec WHERE TRIM(COALESCE(model,'')) <> ''"):
            spec_model[r[0]] = r[1]
    except sqlite3.Error:
        pass
    out = []
    for key, g in system_stats.collapse_machines(rows).items():
        inb = [r for r in g if not system_report.is_off_book(r["asset_serial"])]
        ordered = inb + [r for r in g if r not in inb]
        rep = ordered[0]
        models = []
        for r in ordered:
            for m in (r["device_model"], spec_model.get(r["asset_serial"])):
                if m and str(m).strip() and norm_model(m) not in [norm_model(x) for x in models]:
                    models.append(str(m).strip())
        out.append({
            "key": key,
            "serial": rep["asset_serial"],
            "serials": [r["asset_serial"] for r in ordered],
            "in_book": bool(inb),
            "is_vm": any(manage_state.is_vm_value(r["is_vm"], r["device_model"]) for r in g),
            "hostname": _first(ordered, "hostname"),
            "ip": _first(ordered, "ip"),
            "api_ids": sorted({(r["api_id"] or "").strip() for r in g if (r["api_id"] or "").strip()}),
            "env": system_stats.normalize_env(rep["environment"]),
            "room": system_stats.normalize_location(_first(ordered, "physical_location")),
            "rack": _first(ordered, "rack_no"),
            "os": _first(ordered, "os"),
            "model": models[0] if models else None,
            "models": models,
            "hw_serials": sorted({norm_serial(r["hw_serial"]) for r in g if (r["hw_serial"] or "").strip()}),
        })
    out.sort(key=lambda m: m["serial"])
    return out


def _system_names(conn) -> dict[str, str]:
    try:
        return {(r[0] or "").strip(): r[1] for r in conn.execute("SELECT api_id, name FROM business_system")}
    except sqlite3.Error:  # 系統名稱只是顯示用：表不在就只顯示 APID，不影響任何數字
        return {}


def _brief(m: dict, names: dict) -> dict:
    """表格每一列都帶的機器欄位（下鑽與匯出共用）。"""
    return {"key": m["key"], "serial": m["serial"], "hostname": m["hostname"], "ip": m["ip"],
            "os": m["os"], "room": m["room"], "rack": m["rack"], "model": m["model"],
            "env": m["env"], "api_ids": m["api_ids"],
            "systems": "、".join(f"{a} {names.get(a) or ''}".strip() for a in m["api_ids"]) or None}


# ---------------------------------------------------------------- 1. 長時間未重開機
_UP_UNITS = {"year": 365 * 86400, "week": 7 * 86400, "day": 86400, "hour": 3600, "minute": 60, "min": 60}


def _parse_dt(s) -> _dt.datetime | None:
    if not s:
        return None
    s = str(s).strip().replace("T", " ")[:19]
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return _dt.datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def parse_uptime(text, collected_at) -> dict | None:
    """收集當下的開機時間文字 → {seconds, boot}。認得：

    * `uptime -p`：`up 1 year, 23 weeks, 6 days, 10 hours, 16 minutes`
    * 傳統 `uptime`：`10:30 up 120 days,  3:45,  2 users, load average…`
    * `/proc/uptime`：`12345.67 2345.6`
    * Windows 收集器：`up 12 days, 3 hours（開機 2026-10-01 08:00:00）` → 直接用開機時刻

    認不出來回 None（畫面歸「有收集、讀不出開機時間」，不猜）。
    """
    t = str(text or "").strip()
    col = _parse_dt(collected_at)
    if not t:
        return None
    m = re.search(r"開機\s*(\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(?::\d{2})?)", t)
    if m:
        boot = _parse_dt(m.group(1))
        if boot:
            secs = int((col - boot).total_seconds()) if col else None
            return {"seconds": secs, "boot": boot}
    secs = None
    if re.match(r"^\d+(\.\d+)?(\s|$)", t):
        secs = int(float(t.split()[0]))
    elif "up" in t.lower():
        tail = t.lower().split("up", 1)[1]
        tail = re.split(r"\d+\s+users?|load average", tail)[0]
        total, hit = 0, False
        for n, unit in re.findall(r"(\d+)\s*(year|week|day|hour|minute|min)s?\b", tail):
            total += int(n) * _UP_UNITS[unit]
            hit = True
        hm = re.search(r"(?<![\d:])(\d{1,2}):(\d{2})(?![\d:])", tail)
        if hm:
            total += int(hm.group(1)) * 3600 + int(hm.group(2)) * 60
            hit = True
        secs = total if hit else None
    if secs is None or col is None:
        return None
    return {"seconds": secs, "boot": col - _dt.timedelta(seconds=secs)}


def _ip_keys(conn, serial_key: dict) -> dict[str, set]:
    """IP → 現役的台（一個 IP 可能對到不只一台，呼叫端要自己判斷唯一才用）。"""
    out: dict[str, set] = {}
    for r in conn.execute("SELECT asset_serial, ip FROM hardware WHERE TRIM(COALESCE(ip,'')) <> ''"):
        k = serial_key.get(r[0])
        if k:
            out.setdefault(r[1].strip(), set()).add(k)
    return out


def _latest_spec_by_machine(conn, serial_key: dict) -> dict[str, dict]:
    """每台最新一份 host_spec（同一台多筆登記各收過時取最新）。

    收集當時掛的序號後來不在資產庫（掃描空殼併回 CIA 後，DYN- 那筆已刪），就用收集時的 IP
    找現役的台——IP 只對到一台才算，對到多台不猜。"""
    out: dict[str, dict] = {}
    try:
        rows = conn.execute("SELECT asset_serial, ip, platform, spec_json, collected_at FROM host_spec").fetchall()
    except sqlite3.Error:
        return out
    ip_keys = None
    for r in rows:
        k = serial_key.get(r["asset_serial"])
        if not k and r["ip"]:
            if ip_keys is None:
                ip_keys = _ip_keys(conn, serial_key)
            cand = ip_keys.get(r["ip"].strip()) or set()
            k = next(iter(cand)) if len(cand) == 1 else None
        if not k:
            continue
        if k in out and (out[k]["collected_at"] or "") >= (r["collected_at"] or ""):
            continue
        try:
            spec = json.loads(r["spec_json"] or "{}") or {}
        except (ValueError, TypeError):
            spec = {}
        out[k] = {"serial": r["asset_serial"], "platform": r["platform"], "spec": spec,
                  "collected_at": r["collected_at"]}
    return out


def uptime_tab(conn, ms: list[dict], names: dict, now: _dt.datetime, threshold: int) -> dict:
    by_key = {m["key"]: m for m in ms}
    serial_key = {s: m["key"] for m in ms for s in m["serials"]}
    specs = _latest_spec_by_machine(conn, serial_key)
    rows, unreadable = [], []
    with_uptime = 0
    for k, sp in specs.items():
        up = parse_uptime(sp["spec"].get("uptime"), sp["collected_at"])
        if not up:
            unreadable.append({**_brief(by_key[k], names), "collected_at": sp["collected_at"],
                               "raw": sp["spec"].get("uptime")})
            continue
        with_uptime += 1
        days = (now - up["boot"]).days
        if days >= threshold:
            rows.append({**_brief(by_key[k], names), "days": days,
                         "last_boot": up["boot"].strftime("%Y-%m-%d %H:%M"),
                         "collected_at": sp["collected_at"], "platform": sp["platform"]})
    rows.sort(key=lambda r: (-r["days"], r["serial"]))
    unreadable.sort(key=lambda r: r["serial"])
    return {"threshold": threshold, "rows": rows, "count": len(rows),
            "collected": len(specs), "with_uptime": with_uptime,
            "unreadable": unreadable, "unreadable_count": len(unreadable)}


# ---------------------------------------------------------------- 2. 硬體異常
_SMART_OK = {"PASSED", "OK", "HEALTHY"}
_VD_OK = {"OPTL", "OK", "OPTIMAL"}
_PD_OK = {"ONLN", "OK", "UGOOD", "GHS", "DHS", "JBOD", "ONLINE"}


def hw_findings(spec: dict) -> list[dict]:
    """一份 spec_json 裡看得出的硬體異常。只判讀已收到的欄位，沒收到的不算異常。"""
    out = []
    for d in spec.get("disks") or []:
        h = (d.get("health") or "").strip()
        if h and h.upper() not in _SMART_OK:
            out.append({"source": "SMART", "item": "磁碟", "detail": f"{d.get('disk') or '?'} {d.get('model') or ''}：{h}".strip()})
    for a in spec.get("raid_sw_detail") or []:
        bad = [x["dev"] for x in (a.get("members") or []) if (x.get("flag") or "") == "F"]
        if a.get("degraded") or bad:
            out.append({"source": "軟體 RAID", "item": "RAID",
                        "detail": f"{a.get('name')} {a.get('level') or ''} {a.get('status') or ''}"
                                  + (f" 故障成員 {','.join(bad)}" if bad else "")})
    hw = spec.get("raid_hw_detail") or {}
    if isinstance(hw, dict) and hw.get("parsed"):
        for vd in hw.get("virtual_drives") or []:
            st = (vd.get("state") or "").strip()
            if st and st.upper() not in _VD_OK:
                out.append({"source": "硬體 RAID", "item": "RAID",
                            "detail": f"VD {vd.get('dg', '')}/{vd.get('vd', '')} {vd.get('raid_level') or ''}：{st}"})
        for pd in hw.get("drives") or []:
            st = (pd.get("state") or "").strip()
            if st and st.upper() not in _PD_OK:
                out.append({"source": "硬體 RAID", "item": "磁碟",
                            "detail": f"實體碟 {pd.get('eid_slot') or pd.get('bay') or ''}：{st}"})
    for mp in spec.get("multipath_detail") or []:
        bad = [p for p in (mp.get("paths") or [])
               if re.search(r"fail|faulty|offline", p.get("state") or "", re.I)]
        if bad:
            out.append({"source": "多路徑", "item": "SAN 路徑",
                        "detail": f"{mp.get('name')}：{len(bad)}/{mp.get('path_count') or len(mp.get('paths') or [])} 條路徑異常"})
    return out


def hardware_tab(conn, ms: list[dict], names: dict, disk_pct: int) -> dict:
    by_key = {m["key"]: m for m in ms}
    serial_key = {s: m["key"] for m in ms for s in m["serials"]}
    specs = _latest_spec_by_machine(conn, serial_key)
    found: dict[str, list] = {}
    cover = {"SMART": 0, "軟體 RAID": 0, "硬體 RAID": 0, "多路徑": 0, "健康基準": 0, "PSU": 0}
    for k, sp in specs.items():
        s = sp["spec"]
        cover["SMART"] += bool(s.get("disks"))
        cover["軟體 RAID"] += bool(s.get("raid_sw_detail"))
        cover["硬體 RAID"] += bool(isinstance(s.get("raid_hw_detail"), dict) and s["raid_hw_detail"].get("parsed"))
        cover["多路徑"] += bool(s.get("multipath_detail"))
        for f in hw_findings(s):
            found.setdefault(k, []).append({**f, "at": sp["collected_at"]})
    try:
        cover["PSU"] = conn.execute("SELECT COUNT(*) FROM host_spec WHERE psu_json IS NOT NULL").fetchone()[0]
    except sqlite3.Error:
        pass
    # 健康基準：每個 IP 最新一份的磁碟使用率
    ip_keys = _ip_keys(conn, serial_key)
    try:
        snaps = conn.execute(
            "SELECT s.ip, s.taken_at, s.disks_json FROM health_snapshot s JOIN ("
            " SELECT ip, MAX(taken_at) AS t FROM health_snapshot GROUP BY ip) x "
            "ON x.ip = s.ip AND x.t = s.taken_at").fetchall()
    except sqlite3.Error:
        snaps = []
    covered_keys = set()
    for sn in snaps:
        keys = ip_keys.get((sn[0] or "").strip()) or set()
        covered_keys |= keys
        try:
            disks = json.loads(sn[2] or "{}") or {}
        except (ValueError, TypeError):
            disks = {}
        import optical_media
        hot = sorted(((mnt, p) for mnt, p in disks.items()
                      if isinstance(p, (int, float)) and p >= disk_pct and not optical_media.is_optical(mnt)),
                     key=lambda x: -x[1])
        if not hot:
            continue
        for k in sorted(keys):
            found.setdefault(k, []).append({
                "source": "健康基準", "item": "磁碟空間", "at": sn[1],
                "detail": "、".join(f"{mnt} {round(p)}%" for mnt, p in hot)})
    cover["健康基準"] = len(covered_keys)
    rows = []
    for k, fs in found.items():
        rows.append({**_brief(by_key[k], names), "findings": fs, "finding_count": len(fs),
                     "sources": "、".join(sorted({f["source"] for f in fs})),
                     "items": "、".join(sorted({f["item"] for f in fs})),
                     "detail": "；".join(f["detail"] for f in fs)})
    rows.sort(key=lambda r: (-r["finding_count"], r["serial"]))
    return {"rows": rows, "count": len(rows), "disk_pct": disk_pct, "collected": len(specs),
            "coverage": [{"source": s, "hosts": n} for s, n in cover.items()]}


# ---------------------------------------------------------------- 3. 單點（無備援）
SPOF_SINGLE = "單機"
SPOF_SAME_ROOM = "同機房"


def spof_rule(rooms: list[str]) -> tuple[str | None, bool]:
    """一個系統正式環境每台機器的機房（一台一個值，未填＝'未填'）→ (單點類型或 None, 能不能判定)。

    * 1 台 → 單機（不管機房填沒填）
    * ≥2 台、機房都有填、只有一間 → 同機房
    * ≥2 台、有台機房未填、有填的不超過一間 → 無法判定（不算單點，另列）
    * 其餘 → 不是單點
    """
    n = len(rooms)
    if n == 0:
        return None, True
    if n == 1:
        return SPOF_SINGLE, True
    known = {r for r in rooms if r != "未填"}
    if len(known) >= 2:
        return None, True
    if len(known) == 1 and "未填" not in rooms:
        return SPOF_SAME_ROOM, True
    return None, False


def spof_tab(conn, names: dict) -> dict:
    import blast_radius

    rows = conn.execute(
        "SELECT asset_serial, hostname, ip, api_id, environment, physical_location, rack_no, os, asset_status "
        "FROM hardware WHERE TRIM(COALESCE(api_id,'')) <> ''").fetchall()
    rows = [r for r in rows if not system_stats.is_retired_row(r)]
    per: dict[str, list] = {}
    for r in rows:
        per.setdefault(r["api_id"].strip(), []).append(r)
    flagged, undetermined, ok = [], [], []
    for api in sorted(per):
        prod, dr_rooms = [], []
        for key, g in system_stats.collapse_machines(per[api]).items():
            env = blast_radius.map_env_group(g[0]["environment"])
            room = blast_radius.map_room(_first(g, "physical_location"))
            item = {"key": key, "serial": g[0]["asset_serial"], "hostname": _first(g, "hostname"),
                    "ip": _first(g, "ip"), "room": room, "rack": _first(g, "rack_no"), "os": _first(g, "os"),
                    "env": env}
            if env == "正式":
                prod.append(item)
            elif env == "備援":
                dr_rooms.append(item)
        if not prod:
            continue
        kind, decidable = spof_rule([p["room"] for p in prod])
        prod.sort(key=lambda x: x["serial"])
        prod_rooms = sorted({p["room"] for p in prod})
        rec = {"api_id": api, "name": names.get(api), "kind": system_stats.system_kind(api),
               "prod_count": len(prod), "rooms": "、".join(prod_rooms), "room_count": len(prod_rooms),
               "dr_other_room": sum(1 for d in dr_rooms if d["room"] not in prod_rooms and d["room"] != "未填"),
               "reason": kind, "machines": prod}
        if kind:
            flagged.append(rec)
        elif not decidable:
            rec["reason"] = "機房未填，無法判定"
            undetermined.append(rec)
        else:
            ok.append(rec)
    flagged.sort(key=lambda r: (r["reason"] != SPOF_SINGLE, r["api_id"]))
    return {"rows": flagged, "count": len(flagged), "undetermined": undetermined,
            "undetermined_count": len(undetermined), "ok_count": len(ok),
            "systems_with_prod": len(flagged) + len(undetermined) + len(ok)}


# ---------------------------------------------------------------- 聯絡資訊（三層＋合約匯入）
def parse_date(v) -> str | None:
    """各種日期寫法 → YYYY-MM-DD。民國年（<1911）自動 +1911。認不出回 None。"""
    if v is None:
        return None
    if isinstance(v, _dt.datetime):
        return v.date().isoformat()
    if isinstance(v, _dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 <= v <= 80000:      # Excel 序號日期
            return (_dt.date(1899, 12, 30) + _dt.timedelta(days=int(v))).isoformat()
        v = str(int(v))
    s = str(v).strip()
    if not s:
        return None
    m = re.fullmatch(r"(\d{4})(\d{2})(\d{2})", s)
    if not m:
        m = re.match(r"^(\d{2,4})\s*[/\-.年]\s*(\d{1,2})\s*[/\-.月]\s*(\d{1,2})", s)
    if not m:
        return None
    y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 1911:
        y += 1911
    try:
        return _dt.date(y, mo, d).isoformat()
    except ValueError:
        return None


def _entries(conn) -> dict[tuple[str, str], dict]:
    out = {}
    try:
        for r in conn.execute("SELECT * FROM contact_info"):
            d = dict(r)
            k = d["key"] if d["level"] != "model" else norm_model(d["key"])
            out[(d["level"], k)] = d
    except sqlite3.Error:
        pass
    return out


def _current_contract_rows(conn) -> list[dict]:
    try:
        imp = conn.execute("SELECT id FROM contract_import WHERE is_current = 1 ORDER BY id DESC LIMIT 1").fetchone()
        if not imp:
            return []
        return [dict(r) for r in conn.execute(
            "SELECT * FROM contract_row WHERE import_id = ? ORDER BY row_no", (imp[0],))]
    except sqlite3.Error:  # 舊庫還沒建合約表（db.py init 前）＝沒有匯入過；畫面顯示「尚無合約資料」
        return []


def _resolve_fields(m: dict, ent: dict, contract_by_key: dict) -> dict:
    """一台機器每個欄位的值與來源。順序（越前面越優先）：
    本機（手動）→ 合約匯入（序號／主機名對到這台）→ 系統（手動）→ 型號（手動）→ 合約匯入（型號）。
    手動在同一層贏過匯入：人補的比檔案新。"""
    layers = []
    for s in m["serials"]:
        e = ent.get(("asset", s))
        if e:
            layers.append(("本機（手動）", e))
            break
    for c in contract_by_key.get(m["key"], {}).get("direct", []):
        layers.append((f"合約匯入（{'序號' if c['match_by'] == 'serial' else '主機名'}）", c))
    for a in m["api_ids"]:
        e = ent.get(("system", a))
        if e:
            layers.append((f"系統 {a}（手動）", e))
            break
    for mdl in m["models"]:
        e = ent.get(("model", norm_model(mdl)))
        if e:
            layers.append(("型號（手動）", e))
            break
    for c in contract_by_key.get(m["key"], {}).get("model", []):
        layers.append(("合約匯入（型號）", c))
    out = {}
    for f in CONTACT_FIELDS:
        for label, e in layers:
            v = e.get("end_date") if (f == "contract_end" and "end_date" in e) else e.get(f)
            if v is not None and str(v).strip():
                out[f] = {"value": str(v).strip(), "source": label}
                break
        else:
            out[f] = {"value": None, "source": None}
    return out


def resolve_all(conn, ms: list[dict] | None = None) -> dict[str, dict]:
    """每台機器的聯絡／合約欄位（含來源）。"""
    ms = ms if ms is not None else machines(conn)
    serial_key = {s: m["key"] for m in ms for s in m["serials"]}
    contract_by_key: dict[str, dict] = {}
    for c in _current_contract_rows(conn):
        if c["match_status"] != "matched":
            continue
        try:
            serials = json.loads(c["matched_serials"] or "[]")
        except (ValueError, TypeError):
            serials = []
        bucket = "model" if c["match_by"] == "model" else "direct"
        for s in serials:
            k = serial_key.get(s)
            if k:
                contract_by_key.setdefault(k, {"direct": [], "model": []})[bucket].append(c)
    ent = _entries(conn)
    return {m["key"]: _resolve_fields(m, ent, contract_by_key) for m in ms}


def has_info(res: dict) -> bool:
    return any(res[f]["value"] for f in KEY_FIELDS)


def band_of(end: str | None, today: _dt.date) -> str:
    d = parse_date(end) if end else None
    if not d:
        return BAND_NONE
    left = (_dt.date.fromisoformat(d) - today).days
    if left < 0:
        return BAND_EXPIRED
    if left <= 90:
        return BAND_90
    if left <= 180:
        return BAND_180
    return BAND_OK


def contract_scope(m: dict) -> bool:
    """保固／合約看的是帳內（CIA）現役**實體機**——VM 沒有自己的硬體保固（看它跑的那台 ESXi）。"""
    return m["in_book"] and not m["is_vm"]


def contract_tab(conn, ms: list[dict], names: dict, today: _dt.date, resolved: dict) -> dict:
    book = [m for m in ms if contract_scope(m)]
    bands = {b: [] for b in BAND_LABEL}
    for m in book:
        r = resolved[m["key"]]
        b = band_of(r["contract_end"]["value"], today)
        end = parse_date(r["contract_end"]["value"]) if r["contract_end"]["value"] else None
        bands[b].append({**_brief(m, names), "band": b, "band_label": BAND_LABEL[b],
                         "contract_end": end, "days_left": (_dt.date.fromisoformat(end) - today).days if end else None,
                         "end_source": r["contract_end"]["source"],
                         "vendor": r["vendor"]["value"], "phone": r["phone"]["value"],
                         "contract_no": r["contract_no"]["value"]})
    for b in bands.values():
        b.sort(key=lambda x: (x["days_left"] if x["days_left"] is not None else 99999, x["serial"]))
    counts = {b: len(v) for b, v in bands.items()}
    return {"total": len(book), "counts": counts, "labels": BAND_LABEL,
            "has_data": any(counts[b] for b in (BAND_EXPIRED, BAND_90, BAND_180, BAND_OK)),
            "rows": bands[BAND_EXPIRED] + bands[BAND_90] + bands[BAND_180],
            "none_rows": bands[BAND_NONE], "ok_rows": bands[BAND_OK]}


def _api_contact_text(conn, api_ids: list[str]) -> list[str]:
    """業務系統層的窗口：AP 負責人（部門）＋值班電話。查不到就沒有這一段，不猜。"""
    out = []
    for a in api_ids:
        b = conn.execute("SELECT ap_owner, ap_department FROM business_system WHERE api_id = ?", (a,)).fetchone()
        who = " ".join(x for x in ((b["ap_owner"] or "").strip(), (b["ap_department"] or "").strip()) if x) if b else ""
        try:
            o = conn.execute("SELECT phone FROM system_oncall WHERE api_id = ?", (a,)).fetchone()
        except sqlite3.OperationalError:   # 舊庫沒有值班表：只少這一段
            o = None
        ph = (o["phone"] or "").strip() if o else ""
        if who or ph:
            out.append(f"{a} AP：{who or '—'}" + (f"（值班 {ph}）" if ph else ""))
    return out


def notify_map(conn, ms: list[dict]) -> dict[str, str]:
    """每台機器「出事通知誰」一行字：保管者／使用者（hardware）＋所屬業務系統的 AP 窗口與值班電話。
    1.3 設備風險每一列都帶（2.007.17 使用者：「1.3 沒有通知誰的資訊」）；兩邊都沒登記的機器不在表裡（畫面寫「未登記」）。"""
    people: dict[str, tuple[set, set]] = {}
    for r in conn.execute("SELECT asset_serial, custodian, user_name FROM hardware"):
        people[r["asset_serial"]] = ((r["custodian"] or "").strip(), (r["user_name"] or "").strip())
    cache: dict[str, list[str]] = {}
    out: dict[str, str] = {}
    for m in ms:
        cust = sorted({people[x][0] for x in m["serials"] if x in people and people[x][0]})
        usr = sorted({people[x][1] for x in m["serials"] if x in people and people[x][1]})
        parts = []
        if cust:
            parts.append("保管 " + "、".join(cust))
        if usr:
            parts.append("使用 " + "、".join(usr))
        for a in m["api_ids"]:
            if a not in cache:
                cache[a] = _api_contact_text(conn, [a])
            parts += cache[a]
        if parts:
            out[m["key"]] = "；".join(parts)
    return out


def risk_overview(conn, now: _dt.datetime | None = None) -> dict:
    now = now or _dt.datetime.now()
    st = settings(conn)
    ms = machines(conn)
    names = _system_names(conn)
    resolved = resolve_all(conn, ms)
    out = {"generated_at": now.strftime("%Y-%m-%d %H:%M:%S"), "settings": st,
           "uptime": uptime_tab(conn, ms, names, now, st["uptime_days"]),
           "hardware": hardware_tab(conn, ms, names, st["disk_pct"]),
           "spof": spof_tab(conn, names),
           "contract": contract_tab(conn, ms, names, now.date(), resolved)}
    nm = notify_map(conn, ms)
    for tab in ("uptime", "hardware", "contract"):
        for r in out[tab]["rows"]:
            r["notify"] = nm.get(r.get("key")) or None
    for r in out["spof"]["rows"]:   # 單點風險的單位是業務系統：通知該系統的窗口
        r["notify"] = "；".join(_api_contact_text(conn, [r["api_id"]])) or None
    return out


# ---------------------------------------------------------------- 待補清單
def todo(conn) -> dict:
    """還沒有任何聯絡／合約資料（廠商、維護電話、合約迄日都沒有）的帳內現役機器，受影響台數多的在前。

    * 依系統：帳內現役機器（含 VM——系統窗口跟虛實無關）。一台掛多個 APID 時每個系統各算一次。
    * 依型號：只算實體機（VM 的型號就是「VM」，補了也沒意義）。一台只歸一組，各組加總＝缺資料實體機台數。"""
    ms = machines(conn)
    names = _system_names(conn)
    res = resolve_all(conn, ms)
    missing = [m for m in ms if m["in_book"] and not has_info(res[m["key"]])]
    by_sys: dict[str, list] = {}
    by_model: dict[str, dict] = {}
    for m in missing:
        for a in (m["api_ids"] or [NO_APID]):
            by_sys.setdefault(a, []).append(m)
        if m["is_vm"]:
            continue
        nm = norm_model(m["model"]) or NO_MODEL
        g = by_model.setdefault(nm, {"model": m["model"] or NO_MODEL, "machines": []})
        g["machines"].append(m)
    systems = [{"api_id": a, "name": names.get(a), "count": len(v),
                "machines": [_brief(x, names) for x in v]} for a, v in by_sys.items()]
    models = [{"model": g["model"], "model_key": k, "count": len(g["machines"]),
               "machines": [_brief(x, names) for x in g["machines"]]} for k, g in by_model.items()]
    systems.sort(key=lambda r: (-r["count"], r["api_id"]))
    models.sort(key=lambda r: (-r["count"], r["model"]))
    book = [m for m in ms if m["in_book"]]
    phys = [m for m in book if not m["is_vm"]]
    return {"in_book_total": len(book), "missing_count": len(missing), "filled_count": len(book) - len(missing),
            "physical_total": len(phys),
            "physical_missing": sum(1 for m in missing if not m["is_vm"]),
            "systems": systems, "models": models}


# ---------------------------------------------------------------- 手動維護（admin）
def _audit(conn, kind: str, level: str, key: str, before, after, user: str | None) -> None:
    conn.execute(
        "INSERT INTO contact_audit (kind, level, key, before_json, after_json, changed_by, changed_at) "
        "VALUES (?,?,?,?,?,?,datetime('now','localtime'))",
        (kind, level, key, json.dumps(before, ensure_ascii=False) if before else None,
         json.dumps(after, ensure_ascii=False) if after else None, user))


def save_entry(conn, level: str, key: str, fields: dict, user: str | None) -> dict | None:
    """新增／修改／（全部欄位清空＝刪除）一筆手動聯絡資訊。回存好的那筆（刪除回 None）。"""
    level = (level or "").strip()
    key = (key or "").strip()
    if level not in LEVELS:
        raise ValueError("層級只能是 asset／system／model")
    if not key:
        raise ValueError("缺少對象（序號／APID／型號）")
    if level == "asset" and not conn.execute(
            "SELECT 1 FROM hardware WHERE asset_serial = ?", (key,)).fetchone():
        raise ValueError(f"找不到資產序號 {key}")
    vals = {f: (str(fields.get(f)).strip() if fields.get(f) is not None else "") or None for f in CONTACT_FIELDS}
    if vals["contract_end"]:
        d = parse_date(vals["contract_end"])
        if not d:
            raise ValueError(f"合約迄日看不懂：{vals['contract_end']}")
        vals["contract_end"] = d
    lookup_key = key if level != "model" else norm_model(key)
    before = None
    for r in conn.execute("SELECT * FROM contact_info WHERE level = ?", (level,)):
        if (r["key"] if level != "model" else norm_model(r["key"])) == lookup_key:
            before = dict(r)
            break
    if not any(vals.values()):
        if before:
            conn.execute("DELETE FROM contact_info WHERE id = ?", (before["id"],))
            _audit(conn, "contact", level, key, before, None, user)
            conn.commit()
        return None
    if before:
        conn.execute(
            "UPDATE contact_info SET vendor=?, phone=?, contact=?, contract_no=?, contract_end=?, note=?, "
            "updated_by=?, updated_at=datetime('now','localtime') WHERE id=?",
            (*[vals[f] for f in CONTACT_FIELDS], user, before["id"]))
        rid = before["id"]
    else:
        cur = conn.execute(
            "INSERT INTO contact_info (level, key, vendor, phone, contact, contract_no, contract_end, note, "
            "updated_by, updated_at) VALUES (?,?,?,?,?,?,?,?,?,datetime('now','localtime'))",
            (level, key, *[vals[f] for f in CONTACT_FIELDS], user))
        rid = cur.lastrowid
    after = dict(conn.execute("SELECT * FROM contact_info WHERE id = ?", (rid,)).fetchone())
    _audit(conn, "contact", level, key, before, after, user)
    conn.commit()
    return after


def list_entries(conn) -> list[dict]:
    """手動登錄的全部筆數，含每筆套到幾台（下鑽用 keys）。"""
    ms = machines(conn)
    names = _system_names(conn)
    out = []
    for r in conn.execute("SELECT * FROM contact_info ORDER BY level, key"):
        d = dict(r)
        if d["level"] == "asset":
            hit = [m for m in ms if d["key"] in m["serials"]]
        elif d["level"] == "system":
            hit = [m for m in ms if d["key"] in m["api_ids"]]
        else:
            nk = norm_model(d["key"])
            hit = [m for m in ms if any(norm_model(x) == nk for x in m["models"])]
        d["level_label"] = LEVEL_LABEL[d["level"]]
        d["target_name"] = names.get(d["key"]) if d["level"] == "system" else None
        d["machine_count"] = len(hit)
        d["machines"] = [_brief(m, names) for m in hit]
        out.append(d)
    return out


def _roles_of(row) -> dict:
    """system_oncall.roles_json → {角色: {name, phone, note}}；壞掉或沒有＝空。"""
    try:
        raw = row["roles_json"] if row else None
    except (KeyError, IndexError):
        raw = None
    if not raw:
        return {}
    try:
        d = json.loads(raw)
    except ValueError:      # roles_json 壞掉只是看不到備援／負責人，不影響值班電話本身（phone／note 在別的欄）
        return {}
    return d if isinstance(d, dict) else {}


def save_oncall(conn, api_id: str, phone: str | None, note: str | None, user: str | None) -> dict | None:
    api_id = (api_id or "").strip()
    if not api_id:
        raise ValueError("缺少 APID")
    phone = (phone or "").strip() or None
    note = (note or "").strip() or None
    before = conn.execute("SELECT * FROM system_oncall WHERE api_id = ?", (api_id,)).fetchone()
    before = dict(before) if before else None
    if not phone and not note:
        if before:
            roles = _roles_of(before)
            if before["phone"]:      # 電話被清掉＝電話實際改變，舊姓名不再對應
                roles.pop("值班", None)
            if roles:       # 還有備援／負責人（名冊匯入的）：只清掉值班電話，不連他們一起刪
                conn.execute("UPDATE system_oncall SET phone=NULL, note=NULL, roles_json=?, updated_by=?, "
                             "updated_at=datetime('now','localtime') WHERE api_id=?",
                             (json.dumps(roles, ensure_ascii=False), user, api_id))
                after = dict(conn.execute("SELECT * FROM system_oncall WHERE api_id = ?", (api_id,)).fetchone())
                _audit(conn, "oncall", "system", api_id, before, after, user)
            else:
                conn.execute("DELETE FROM system_oncall WHERE api_id = ?", (api_id,))
                _audit(conn, "oncall", "system", api_id, before, None, user)
            conn.commit()
        return None
    roles = _roles_of(before) if before else {}
    if before and (before["phone"] or None) != phone:
        roles.pop("值班", None)  # 電話實際改變才清值班姓名（舊姓名在稽核 before 裡還在）；只改備註不清
    conn.execute(
        "INSERT INTO system_oncall (api_id, phone, note, roles_json, updated_by, updated_at) "
        "VALUES (?,?,?,?,?,datetime('now','localtime')) ON CONFLICT(api_id) DO UPDATE SET "
        "phone=excluded.phone, note=excluded.note, roles_json=excluded.roles_json, "
        "updated_by=excluded.updated_by, updated_at=excluded.updated_at",
        (api_id, phone, note, json.dumps(roles, ensure_ascii=False) if roles else None, user))
    after = dict(conn.execute("SELECT * FROM system_oncall WHERE api_id = ?", (api_id,)).fetchone())
    _audit(conn, "oncall", "system", api_id, before, after, user)
    conn.commit()
    return after


def list_oncall(conn) -> list[dict]:
    names = _system_names(conn)
    return [{**dict(r), "name": names.get(r["api_id"]), "roles": _roles_of(r)}
            for r in conn.execute("SELECT * FROM system_oncall ORDER BY api_id")]


def audit_log(conn, limit: int = 500) -> list[dict]:
    return [dict(r) for r in conn.execute(
        "SELECT * FROM contact_audit ORDER BY id DESC LIMIT ?", (int(limit),))]


# ---------------------------------------------------------------- 出事聯絡卡
def _people_asset(conn, serials: list[str]) -> dict:
    ph = ",".join("?" for _ in serials)
    rows = conn.execute(f"SELECT asset_serial, api_id, custodian, user_name, usage_unit FROM hardware "
                        f"WHERE asset_serial IN ({ph})", serials).fetchall()
    apis = sorted({(r["api_id"] or "").strip() for r in rows if (r["api_id"] or "").strip()})
    owners = []
    for a in apis:
        b = conn.execute("SELECT ap_owner, ap_department FROM business_system WHERE api_id = ?", (a,)).fetchone()
        if b and (b["ap_owner"] or b["ap_department"]):
            owners.append({"api_id": a, "name": b["ap_owner"], "department": b["ap_department"]})
    reg = []
    for p in conn.execute(f"SELECT DISTINCT person_name, phone, belong_department FROM personnel "
                          f"WHERE asset_serial IN ({ph}) AND TRIM(COALESCE(person_name,'')) <> ''", serials):
        reg.append({"name": p["person_name"], "phone": p["phone"], "department": p["belong_department"]})
    uniq = lambda f: sorted({(r[f] or "").strip() for r in rows if (r[f] or "").strip()})  # noqa: E731
    return {"owners": owners, "custodians": uniq("custodian"), "ap_users": uniq("user_name"),
            "registered": reg, "api_ids": apis}


def contact_card(conn, asset_serial: str | None = None, api_id: str | None = None) -> dict:
    names = _system_names(conn)
    if asset_serial:
        hw = conn.execute("SELECT asset_serial FROM hardware WHERE asset_serial = ?", (asset_serial,)).fetchone()
        if not hw:
            raise LookupError(f"找不到資產 {asset_serial}")
        ms = machines(conn)
        m = next((x for x in ms if asset_serial in x["serials"]), None)
        if m is None:      # 退役：照樣給聯絡卡，只看這一筆
            r = conn.execute("SELECT * FROM hardware WHERE asset_serial = ?", (asset_serial,)).fetchone()
            mdl = (r["device_model"] or "").strip()
            m = {"key": "sn:" + asset_serial, "serial": asset_serial, "serials": [asset_serial],
                 "api_ids": [(r["api_id"] or "").strip()] if (r["api_id"] or "").strip() else [],
                 "models": [mdl] if mdl else [], "model": mdl or None, "in_book": True}
            ms = [m]
        res = resolve_all(conn, ms)[m["key"]]
        people = _people_asset(conn, m["serials"])
        apis = m["api_ids"]
        scope = {"type": "asset", "serial": m["serial"], "serials": m["serials"], "model": m["model"]}
    elif api_id:
        api_id = api_id.strip()
        ms = machines(conn)
        sys_ms = [x for x in ms if api_id in x["api_ids"]]
        ent = _entries(conn).get(("system", api_id))
        res = {f: {"value": (ent or {}).get(f), "source": f"系統 {api_id}（手動）" if (ent or {}).get(f) else None}
               for f in CONTACT_FIELDS}
        all_res = resolve_all(conn, sys_ms) if sys_ms else {}
        serials = [s for x in sys_ms for s in x["serials"]]
        people = _people_asset(conn, serials) if serials else {"owners": [], "custodians": [], "ap_users": [],
                                                               "registered": [], "api_ids": []}
        b = conn.execute("SELECT ap_owner, ap_department FROM business_system WHERE api_id = ?", (api_id,)).fetchone()
        people["owners"] = [{"api_id": api_id, "name": b["ap_owner"], "department": b["ap_department"]}] \
            if b and (b["ap_owner"] or b["ap_department"]) else []
        apis = [api_id]
        scope = {"type": "system", "api_id": api_id, "name": names.get(api_id),
                 "machines": len(sys_ms), "machines_with_info": sum(1 for v in all_res.values() if has_info(v))}
    else:
        raise ValueError("要給 asset_serial 或 api_id")
    oncall = []
    for a in apis:
        r = conn.execute("SELECT phone, note, updated_by, updated_at FROM system_oncall WHERE api_id = ?",
                         (a,)).fetchone()
        oncall.append({"api_id": a, "name": names.get(a), **(dict(r) if r else
                       {"phone": None, "note": None, "updated_by": None, "updated_at": None})})
    end = res["contract_end"]["value"]
    return {"scope": scope, "people": people, "maint": res, "oncall": oncall,
            "contract_band": BAND_LABEL[band_of(end, _dt.date.today())] if end else None}


# ---------------------------------------------------------------- 合約匯入（次要路徑）
#: (欄位, 中文名, 表頭同義字)。同義字越長越先比（「資產序號」要先於「序號」）。
IMPORT_FIELDS = [
    ("serial", "序號", ["硬體序號", "機器序號", "序號", "serial number", "serial", "s/n", "sn"]),
    ("model", "型號", ["設備機型", "型號", "機型", "model"]),
    ("asset_code", "資產編號", ["資產編號", "資產序號", "財產編號", "asset no", "asset"]),
    ("hostname", "主機名", ["主機名稱", "主機名", "hostname", "host name", "host"]),
    ("vendor", "廠商", ["維護廠商", "原廠", "廠商", "供應商", "vendor"]),
    ("contract_no", "合約編號", ["合約編號", "合約號碼", "合約號", "contract no", "contract"]),
    ("start_date", "起日", ["合約起日", "起日", "開始日", "生效日", "start date", "start"]),
    ("end_date", "迄日", ["合約迄日", "保固到期", "到期日", "迄日", "結束日", "end date", "expire", "expiry", "end"]),
    ("phone", "維護電話", ["維護電話", "服務電話", "電話", "專線", "phone", "tel"]),
    ("contact", "窗口", ["聯絡窗口", "窗口", "聯絡人", "contact"]),
    ("note", "備註", ["備註", "說明", "remark", "note"]),
]
IMPORT_FIELD_KEYS = [f[0] for f in IMPORT_FIELDS]
MATCH_FIELDS = ("serial", "asset_code", "hostname", "model")


def _cell(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def read_table(filename: str, data: bytes, max_rows: int | None = None) -> tuple[list[str], list[list]]:
    """xlsx／csv → (表頭, 資料列)。表頭列＝前 10 列裡非空字串最多的那列（前面可以有標題列）。"""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        grid = []
        for r in ws.iter_rows(values_only=True):
            grid.append([_cell(c) for c in r])
            if max_rows is not None and len(grid) > max_rows + 20:     # 解析途中就中止，不把整張表讀完
                wb.close()
                raise ValueError(f"列數太多（上限 {max_rows} 列）")
        wb.close()
    elif name.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "cp950", "big5", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        dialect = csv.excel
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            pass
        grid = []
        for r in csv.reader(io.StringIO(text), dialect):
            grid.append([_cell(c) for c in r])
            if max_rows is not None and len(grid) > max_rows + 20:
                raise ValueError(f"列數太多（上限 {max_rows} 列）")
    else:
        raise ValueError("只收 .xlsx 或 .csv")
    grid = [r for r in grid if any(c is not None for c in r)]
    if not grid:
        raise ValueError("檔案是空的")
    best = max(range(min(10, len(grid))),
               key=lambda i: (sum(1 for c in grid[i] if isinstance(c, str)), -i))
    width = max(len(r) for r in grid)
    headers = []
    for i in range(width):
        h = grid[best][i] if i < len(grid[best]) else None
        h = str(h).strip() if h is not None else ""
        headers.append(h or f"欄{i + 1}")
    seen: dict[str, int] = {}
    for i, h in enumerate(headers):      # 表頭重複時加編號，對應才分得開
        if h in seen:
            seen[h] += 1
            headers[i] = f"{h}({seen[h]})"
        else:
            seen[h] = 1
    rows = [list(r) + [None] * (width - len(r)) for r in grid[best + 1:]]
    return headers, rows


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    cands = []
    for f, _lbl, syns in IMPORT_FIELDS:
        for s in syns:
            for h in headers:
                hl = h.lower()
                if s == hl or (s in hl and (not s.isascii() or len(s) >= 4)):
                    cands.append((-len(s), s != hl, f, h))
    cands.sort()
    used_f, used_h, out = set(), set(), {}
    for _l, _e, f, h in cands:
        if f in used_f or h in used_h:
            continue
        out[f] = h
        used_f.add(f)
        used_h.add(h)
    return out


def apply_mapping(headers: list[str], rows: list[list], mapping: dict) -> list[dict]:
    idx = {f: headers.index(h) for f, h in (mapping or {}).items() if f in IMPORT_FIELD_KEYS and h in headers}
    out = []
    for n, r in enumerate(rows, start=1):
        rec = {f: r[i] for f, i in idx.items()}
        if not any(v is not None and str(v).strip() for v in rec.values()):
            continue
        clean = {}
        for f in IMPORT_FIELD_KEYS:
            v = rec.get(f)
            if f in ("start_date", "end_date"):
                clean[f + "_raw"] = None if v is None else str(v)
                clean[f] = parse_date(v)
            else:
                clean[f] = None if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)).strip() or None
        clean["row_no"] = n
        out.append(clean)
    return out


def match_index(ms: list[dict]) -> dict:
    idx = {"serial": {}, "host": {}, "model": {}}
    for m in ms:
        if not m["in_book"]:
            continue
        for s in list(m["hw_serials"]) + [norm_serial(x) for x in m["serials"]]:
            if s:
                idx["serial"].setdefault(s, set()).add(m["key"])
        h = norm_host(m["hostname"])
        if h:
            idx["host"].setdefault(h, set()).add(m["key"])
        if not m["is_vm"]:
            for mdl in m["models"]:
                idx["model"].setdefault(norm_model(mdl), set()).add(m["key"])
    return idx


def match_row(rec: dict, idx: dict) -> tuple[str, str | None, set]:
    """一列合約 → (matched/ambiguous/unmatched, 對到的依據, 台的 key)。順序：序號 → 主機名 → 型號。
    序號、主機名對到不只一台＝ambiguous（不猜）；型號一次套全部同型號＝model 層級。"""
    keys: set = set()
    for f in ("serial", "asset_code"):
        if rec.get(f):
            keys |= idx["serial"].get(norm_serial(rec[f]), set())
    if keys:
        return ("matched" if len(keys) == 1 else "ambiguous"), "serial", keys
    if rec.get("hostname"):
        keys = idx["host"].get(norm_host(rec["hostname"]), set())
        if keys:
            return ("matched" if len(keys) == 1 else "ambiguous"), "hostname", set(keys)
    if rec.get("model"):
        keys = idx["model"].get(norm_model(rec["model"]), set())
        if keys:
            return "matched", "model", set(keys)
    return "unmatched", None, set()


def saved_mapping(conn) -> dict:
    try:
        return json.loads(_setting(conn, SET_CONTRACT_MAPPING) or "{}") or {}
    except (ValueError, TypeError):  # 記住的對應壞掉就當沒記，改用自動建議（使用者看得到「對應來源」）
        return {}


def preview(conn, filename: str, data: bytes, mapping: dict | None = None) -> dict:
    headers, raw = read_table(filename, data)
    sm = saved_mapping(conn)
    if not mapping:
        mapping = sm if sm and all(h in headers for h in sm.values()) else suggest_mapping(headers)
        mapping_from = "saved" if mapping is sm else "suggested"
    else:
        mapping_from = "user"
    mapping = {f: h for f, h in mapping.items() if f in IMPORT_FIELD_KEYS and h in headers}
    recs = apply_mapping(headers, raw, mapping)
    ms = machines(conn)
    names = _system_names(conn)
    by_key = {m["key"]: m for m in ms}
    idx = match_index(ms)
    counts = {"matched": 0, "unmatched": 0, "ambiguous": 0}
    for r in recs:
        st, by, keys = match_row(r, idx)
        counts[st] += 1
        r["match_status"], r["match_by"] = st, by
        r["machines"] = [_brief(by_key[k], names) for k in sorted(keys)]
        r["matched_serials"] = sorted(by_key[k]["serial"] for k in keys)
        r["bad_end"] = bool(r.get("end_date_raw") and not r.get("end_date"))
    problems = []
    if not any(f in mapping for f in MATCH_FIELDS):
        problems.append("至少要對應 序號／資產編號／主機名／型號 其中一欄，才對得到機器")
    if not any(f in mapping for f in ("end_date", "vendor", "phone")):
        problems.append("至少要對應 迄日／廠商／維護電話 其中一欄")
    return {"filename": filename, "headers": headers, "mapping": mapping, "mapping_from": mapping_from,
            "fields": [{"key": f, "label": lbl} for f, lbl, _s in IMPORT_FIELDS],
            "sample": raw[:5], "rows": recs, "row_count": len(recs), "counts": counts,
            "by_model": sum(1 for r in recs if r["match_by"] == "model"),
            "bad_end": sum(1 for r in recs if r["bad_end"]), "problems": problems}


def save_import(conn, filename: str, data: bytes, mapping: dict, user: str | None) -> dict:
    p = preview(conn, filename, data, mapping)
    if p["problems"]:
        raise ValueError("；".join(p["problems"]))
    if not p["row_count"]:
        raise ValueError("沒有資料列")
    conn.execute("UPDATE contract_import SET is_current = 0 WHERE is_current = 1")
    cur = conn.execute(
        "INSERT INTO contract_import (filename, imported_by, imported_at, row_count, matched, unmatched, ambiguous, "
        "mapping_json, is_current) VALUES (?,?,datetime('now','localtime'),?,?,?,?,?,1)",
        (filename, user, p["row_count"], p["counts"]["matched"], p["counts"]["unmatched"],
         p["counts"]["ambiguous"], json.dumps(p["mapping"], ensure_ascii=False)))
    iid = cur.lastrowid
    for r in p["rows"]:
        conn.execute(
            "INSERT INTO contract_row (import_id, row_no, serial, model, asset_code, hostname, vendor, contract_no, "
            "start_date, end_date, end_date_raw, phone, contact, note, match_status, match_by, matched_serials) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (iid, r["row_no"], r["serial"], r["model"], r["asset_code"], r["hostname"], r["vendor"],
             r["contract_no"], r["start_date"], r["end_date"], r.get("end_date_raw"), r["phone"], r["contact"],
             r["note"], r["match_status"], r["match_by"], json.dumps(r["matched_serials"], ensure_ascii=False)))
    conn.execute("INSERT OR REPLACE INTO app_settings(key, value) VALUES (?,?)",
                 (SET_CONTRACT_MAPPING, json.dumps(p["mapping"], ensure_ascii=False)))
    _audit(conn, "contract_import", "import", str(iid), None,
           {"filename": filename, "rows": p["row_count"], **p["counts"]}, user)
    conn.commit()
    return {"import_id": iid, "row_count": p["row_count"], "counts": p["counts"]}


def contract_current(conn) -> dict:
    try:
        imp = conn.execute("SELECT * FROM contract_import WHERE is_current = 1 ORDER BY id DESC LIMIT 1").fetchone()
        history = [dict(r) for r in conn.execute("SELECT * FROM contract_import ORDER BY id DESC")]
    except sqlite3.Error:
        imp, history = None, []
    rows = []
    if imp:
        names = _system_names(conn)
        ms = machines(conn)
        serial_m = {s: m for m in ms for s in m["serials"]}
        for r in _current_contract_rows(conn):
            d = dict(r)
            try:
                ss = json.loads(d.pop("matched_serials") or "[]")
            except (ValueError, TypeError):
                ss = []
            d["machines"] = [_brief(serial_m[s], names) for s in ss if s in serial_m]
            rows.append(d)
    return {"current": dict(imp) if imp else None, "rows": rows, "history": history,
            "saved_mapping": saved_mapping(conn)}


# ---------------------------------------------------------------- 值班電話名冊匯入（W221 #11）
# 沿用 system_oncall（每系統一列）：值班電話＝原本的 phone／note；姓名與備援、負責人放 roles_json。
# 流程跟合約匯入同一套：先 dry-run 預覽（不寫庫）→ 確認才寫；寫庫與稽核同一個交易。
ROSTER_ROLES = ("值班", "備援", "負責人")
ROSTER_COLS = [("api_id", "系統代號 api_id", ("api_id", "apid", "系統代號", "系統代碼", "系統代號api_id")),
               ("name", "姓名", ("姓名", "name", "人員", "聯絡人")),
               ("phone", "電話", ("電話", "phone", "手機", "號碼")),
               ("role", "角色", ("角色", "role")),
               ("note", "備註", ("備註", "note", "說明"))]
ROSTER_MAX_BYTES = 2 * 1024 * 1024
ROSTER_MAX_ROWS = 5000
_PHONE_OK = re.compile(r"^[0-9+\-()#*、 ]+$")


class RosterChanged(Exception):
    """預覽之後檔案或資料庫變了：確認匯入作廢，要重新預覽。"""


ROSTER_ZIP_MAX_TOTAL = 30 * 1024 * 1024
ROSTER_ZIP_MAX_FILES = 100
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def read_limited(fobj, limit: int = ROSTER_MAX_BYTES, exc=ValueError) -> bytes:
    """有上限的讀取：最多讀 limit+1 位元組，超過就拒，不把整份大檔載進記憶體。"""
    data = fobj.read(limit + 1)
    if len(data) > limit:
        raise exc(f"檔案太大（上限 {limit // (1024 * 1024)} MB）")
    return data


def clean_filename(name) -> str:
    n = str(name or "").replace("\\", "/").split("/")[-1]
    n = "".join(ch for ch in n if ch >= " " and ch != "\x7f")
    return n.strip()[:100]


def _check_xlsx_zip(data: bytes, max_total: int = ROSTER_ZIP_MAX_TOTAL,
                    max_files: int = ROSTER_ZIP_MAX_FILES, exc=ValueError) -> None:
    """xlsx 是 zip：先看解壓後總大小與檔案數，過關才交給 openpyxl（防 zip 炸彈）。"""
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > max_files or sum(i.file_size for i in infos) > max_total:
                raise exc(f"xlsx 內容過大或結構異常（解壓後上限 {max_total // (1024 * 1024)} MB、{max_files} 個檔案）")
    except zipfile.BadZipFile:
        raise ValueError("xlsx 檔案毀損或不是 xlsx")


def roster_template_rows() -> list[list[str]]:
    """範本：表頭＋三列說明用範例（代號 TEST-*，真實環境不存在這個系統，直接匯入會被擋，不會誤寫）。"""
    return [[c[1] for c in ROSTER_COLS],
            ["TEST-001", "王小明", "0900-000-001", "值班", "平日 18:00 後"],
            ["TEST-001", "李小華", "0900-000-002", "備援", ""],
            ["TEST-001", "陳小美", "0900-000-003", "負責人", ""]]


def norm_phone(v) -> str:
    """只做去空白與統一分隔，不改號碼本身。回傳空字串＝沒填。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).strip()
    t = re.sub(r"\s*[,，;；/／、]\s*", "、", t)
    t = re.sub(r"\s+", " ", t)
    return t.strip("、 ")


def _known_systems(conn) -> set[str]:
    out = set(_system_names(conn))
    try:
        out |= {(r[0] or "").strip() for r in conn.execute(
            "SELECT DISTINCT api_id FROM hardware WHERE TRIM(COALESCE(api_id,'')) <> ''")}
    except sqlite3.Error:
        pass
    out.discard("")
    return out


def _roster_current(row, role: str) -> dict:
    """某系統某角色現在登記的 {name, phone, note}（沒有＝全空字串）。"""
    roles = _roles_of(row)
    if role == "值班":
        return {"name": (roles.get("值班") or {}).get("name") or "",
                "phone": (row["phone"] if row else "") or "", "note": (row["note"] if row else "") or ""}
    d = roles.get(role) or {}
    return {"name": d.get("name") or "", "phone": d.get("phone") or "", "note": d.get("note") or ""}


def roster_plan(conn, filename: str, data: bytes) -> dict:
    """解析名冊檔，逐列判定 新增／更新／不變／衝突／格式錯誤；不寫資料庫。"""
    if len(data) > ROSTER_MAX_BYTES:
        raise ValueError("檔案太大（上限 2 MB）")
    filename = clean_filename(filename)
    if filename.lower().endswith((".xlsx", ".xlsm")):
        _check_xlsx_zip(data)
    headers, rows = read_table(filename, data, ROSTER_MAX_ROWS)
    idx: dict[str, int] = {}
    for key, _lbl, syns in ROSTER_COLS:
        for i, h in enumerate(headers):
            hl = h.strip().lower().replace(" ", "")
            if i not in idx.values() and hl in syns:
                idx[key] = i
                break
    missing = [lbl for key, lbl, _s in ROSTER_COLS if key in ("api_id", "phone", "role") and key not in idx]
    if missing:
        raise ValueError("缺少必要欄位：" + "、".join(missing) + "（請用範本）")
    if len(rows) > ROSTER_MAX_ROWS:
        raise ValueError(f"列數太多（上限 {ROSTER_MAX_ROWS} 列）")
    known = _known_systems(conn)

    def cell(r, k):
        i = idx.get(k)
        v = r[i] if i is not None and i < len(r) else None
        if k == "phone":
            return norm_phone(v)
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return ("" if v is None else str(v)).strip()

    out: list[dict] = []
    for n, r in enumerate(rows, start=1):
        if not any(c is not None and str(c).strip() for c in r):
            continue
        rec = {"row_no": n, **{k: cell(r, k) for k, _l, _s in ROSTER_COLS}, "status": "", "reason": "",
               "before": None, "after": None}
        err = None
        if any(str(rec[k]).startswith(_FORMULA_START) for k in ("api_id", "name", "note")):
            err = "格式錯誤：不可以公式字元（= + - @）開頭"
        elif not rec["api_id"]:
            err = "缺系統代號"
        elif rec["api_id"] not in known:
            err = f"找不到系統代號 {rec['api_id']}"
        elif rec["role"] not in ROSTER_ROLES:
            err = "角色只能是 " + "／".join(ROSTER_ROLES)
        elif not rec["phone"]:
            err = "缺電話"
        elif not _PHONE_OK.match(rec["phone"]) or len(re.findall(r"\d", rec["phone"])) < 3:
            err = f"電話格式看不懂：{rec['phone']}"
        if err:
            rec["status"], rec["reason"] = "error", err
        out.append(rec)
    # 同系統同角色多筆：內容不同＝衝突（全部不寫）；內容相同＝重複，只留第一筆
    groups: dict[tuple, list[dict]] = {}
    for rec in out:
        if rec["status"] != "error":
            groups.setdefault((rec["api_id"], rec["role"]), []).append(rec)
    for g in groups.values():
        if len({(x["name"], x["phone"], x["note"]) for x in g}) > 1:
            for x in g:
                x["status"] = "conflict"
                x["reason"] = "同系統同角色在檔案裡有多筆不同內容（第 " + "、".join(str(y["row_no"]) for y in g) + " 列），不自動擇一"
            continue
        for x in g[1:]:
            x["status"], x["reason"] = "unchanged", f"與第 {g[0]['row_no']} 列重複，略過"
    cur = {r["api_id"]: r for r in conn.execute("SELECT * FROM system_oncall")}
    for rec in out:
        if rec["status"]:
            continue
        before = _roster_current(cur.get(rec["api_id"]), rec["role"])
        after = {"name": rec["name"], "phone": rec["phone"], "note": rec["note"]}
        rec["before"], rec["after"] = before, after
        empty = not any(before.values())
        rec["status"] = "new" if empty else ("unchanged" if before == after else "update")
    counts = {k: sum(1 for x in out if x["status"] == k) for k in ("new", "update", "unchanged", "conflict", "error")}
    import hashlib
    digest = hashlib.sha256(json.dumps(
        [[x["api_id"], x["role"], x["status"], x["before"], x["after"]] for x in out if x["status"] in ("new", "update")],
        ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
    return {"filename": filename, "rows": out, "counts": counts, "row_count": len(out),
            "file_sha256": hashlib.sha256(data).hexdigest(), "plan_digest": digest}


def roster_apply(conn, filename: str, data: bytes, user: str | None,
                 file_sha256: str | None = None, plan_digest: str | None = None) -> dict:
    """確認匯入：重新解析同一份檔，只寫 新增／更新 的列（衝突、格式錯誤不寫）。稽核與資料同一個交易。"""
    filename = clean_filename(filename)      # 寫稽核前先清洗檔名
    # 「重新規劃＋比對 plan_digest＋寫入」包在同一個寫入鎖裡（BEGIN IMMEDIATE）：
    # 比對完到寫入之間，別的連線不可能先改了名冊，確認的就是預覽的那一份（TOCTOU）。
    if not conn.in_transaction:
        conn.execute("BEGIN IMMEDIATE")
    try:
        plan = roster_plan(conn, filename, data)
        if file_sha256 is not None or plan_digest is not None:     # 內部直接呼叫（測試）可不帶；端點一定帶
            if plan["file_sha256"] != file_sha256 or plan["plan_digest"] != plan_digest:
                raise RosterChanged("資料或檔案已變動，請重新預覽")
        todo = [x for x in plan["rows"] if x["status"] in ("new", "update")]
        by_api: dict[str, list[dict]] = {}
        for x in todo:
            by_api.setdefault(x["api_id"], []).append(x)
        for api_id, items in by_api.items():
            row = conn.execute("SELECT * FROM system_oncall WHERE api_id = ?", (api_id,)).fetchone()
            before = dict(row) if row else None
            roles = _roles_of(row)
            phone, note = (row["phone"], row["note"]) if row else (None, None)
            for x in items:
                a = x["after"]
                if x["role"] == "值班":
                    phone, note = a["phone"] or None, a["note"] or None
                    roles["值班"] = {"name": a["name"]}
                else:
                    roles[x["role"]] = dict(a)
            conn.execute(
                "INSERT INTO system_oncall (api_id, phone, note, roles_json, updated_by, updated_at) "
                "VALUES (?,?,?,?,?,datetime('now','localtime')) ON CONFLICT(api_id) DO UPDATE SET "
                "phone=excluded.phone, note=excluded.note, roles_json=excluded.roles_json, "
                "updated_by=excluded.updated_by, updated_at=excluded.updated_at",
                (api_id, phone, note, json.dumps(roles, ensure_ascii=False), user))
            after = dict(conn.execute("SELECT * FROM system_oncall WHERE api_id = ?", (api_id,)).fetchone())
            _audit(conn, "oncall", "system", api_id, before, after, user)
        # 摘要只記檔名與各類筆數，不記檔案內容
        _audit(conn, "oncall_roster_import", "import", filename, None,
               {"filename": filename, "rows": plan["row_count"], **plan["counts"], "systems_written": len(by_api)}, user)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return {"counts": plan["counts"], "written_rows": len(todo), "systems_written": len(by_api),
            "skipped": plan["counts"]["conflict"] + plan["counts"]["error"]}


# ---------------------------------------------------------------- 夜間速查設定（W221-N1）
SET_NIGHT_AUTO_LIVE = "night_auto_live"


def night_settings(conn) -> dict:
    """auto_live：預設開（維持原行為）；只有明確存成 '0' 才算關。"""
    return {"auto_live": _setting(conn, SET_NIGHT_AUTO_LIVE) != "0"}


def save_night_auto_live(conn, on: bool, user: str | None) -> dict:
    before = night_settings(conn)
    conn.execute("INSERT OR REPLACE INTO app_settings(key, value) VALUES (?,?)",
                 (SET_NIGHT_AUTO_LIVE, "1" if on else "0"))
    _audit(conn, "night_setting", "setting", SET_NIGHT_AUTO_LIVE, before, {"auto_live": bool(on)}, user)
    conn.commit()
    return night_settings(conn)
