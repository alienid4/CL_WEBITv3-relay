"""選單：套用新分類／還原（備份最近 5 版）＋ 我的最愛／最近用過（每人一份）。2026-10-23。

全站選單順序存 app_settings.nav_layout（管理員拖拉存的自訂版面）。新分類上線時**不覆蓋**自訂版面：
管理員在「選單編輯」按「套用新分類」才換，換之前把當下版面（可能是 null＝內建預設）備份進
app_settings.nav_layout_backups（最近 5 版、帶時間與操作人），「還原上一版」取回最新那份。

我的最愛：nav_favorite 表，以登入帳號為 key；登入停用時前端改存瀏覽器（畫面上會講），
這裡仍接受讀寫（key 用 _noauth），API 不會壞。
"""
from __future__ import annotations

import io
import json
import sqlite3
from datetime import datetime
from typing import Any

LAYOUT_KEY = "nav_layout"
BACKUP_KEY = "nav_layout_backups"
MAX_BACKUPS = 5
DEFAULT_FAVORITES = ["/changes", "/oncall", "/assets", "/datacenter", "/topology"]
MAX_ITEMS = 50
NOAUTH_USER = "_noauth"


def _get_json(conn, key: str):
    r = conn.execute("SELECT value FROM app_settings WHERE key = ?", (key,)).fetchone()
    if not r or not r[0]:
        return None
    try:
        return json.loads(r[0])
    except (ValueError, TypeError):  # 壞掉的設定值：當沒有
        return None


def _set_json(conn, key: str, value) -> None:
    if value is None:
        conn.execute("DELETE FROM app_settings WHERE key = ?", (key,))
    else:
        conn.execute("INSERT INTO app_settings (key, value) VALUES (?, ?) "
                     "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                     (key, json.dumps(value, ensure_ascii=False)))


def backups(conn) -> list[dict]:
    return _get_json(conn, BACKUP_KEY) or []


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def apply_default(conn, by: str) -> dict:
    """先備份當下版面（null 也備份＝「當時是內建預設」），再清掉自訂版面＝改用內建新分類。"""
    cur = _get_json(conn, LAYOUT_KEY)
    bk = [{"at": _now(), "by": by, "layout": cur, "reason": "套用新分類前"}] + backups(conn)
    _set_json(conn, BACKUP_KEY, bk[:MAX_BACKUPS])
    _set_json(conn, LAYOUT_KEY, None)
    conn.commit()
    return {"ok": True, "backed_up": cur is not None, "backups": len(bk[:MAX_BACKUPS])}


def restore_latest(conn, by: str) -> dict:
    """取回最新一份備份（它會從備份清單移掉）；當下版面也先存一份，免得還原錯了回不來。"""
    bk = backups(conn)
    if not bk:
        raise ValueError("沒有可還原的備份")
    latest, rest = bk[0], bk[1:]
    cur = _get_json(conn, LAYOUT_KEY)
    # 當下版面放最前面：再按一次「還原上一版」就回到還原前（等於復原這次還原）
    rest = [{"at": _now(), "by": by, "layout": cur, "reason": "還原前的版面"}] + rest
    _set_json(conn, BACKUP_KEY, rest[:MAX_BACKUPS])
    _set_json(conn, LAYOUT_KEY, latest["layout"])
    conn.commit()
    return {"ok": True, "restored_at": latest["at"], "restored_by": latest.get("by"),
            "restored_default": latest["layout"] is None}


# ===== 我的最愛 =====

def ensure_table(conn) -> None:
    conn.execute("CREATE TABLE IF NOT EXISTS nav_favorite (username TEXT PRIMARY KEY, items TEXT, recent TEXT, "
                 "updated_at TEXT)")


def _clean(lst, cap: int) -> list[str]:
    out: list[str] = []
    for x in lst or []:
        if isinstance(x, str) and x.startswith("/") and x not in out and len(x) < 200:
            out.append(x)
    return out[:cap]


def get_favorites(conn, username: str | None) -> dict[str, Any]:
    ensure_table(conn)
    u = username or NOAUTH_USER
    r = conn.execute("SELECT items, recent FROM nav_favorite WHERE username = ?", (u,)).fetchone()
    if not r:
        return {"username": u, "items": list(DEFAULT_FAVORITES), "recent": [], "is_default": True}
    return {"username": u, "items": json.loads(r["items"] or "[]"), "recent": json.loads(r["recent"] or "[]"),
            "is_default": False}


def set_favorites(conn, username: str | None, items, recent=None) -> dict[str, Any]:
    ensure_table(conn)
    u = username or NOAUTH_USER
    cur = get_favorites(conn, u)
    items = _clean(items, MAX_ITEMS)
    recent = _clean(recent if recent is not None else cur["recent"], 5)
    conn.execute("INSERT INTO nav_favorite (username, items, recent, updated_at) VALUES (?,?,?,?) "
                 "ON CONFLICT(username) DO UPDATE SET items = excluded.items, recent = excluded.recent, "
                 "updated_at = excluded.updated_at",
                 (u, json.dumps(items, ensure_ascii=False), json.dumps(recent, ensure_ascii=False), _now()))
    conn.commit()
    return get_favorites(conn, u)


# ===== 新舊選單編號對照（xlsx）=====

MAPPING_COLS = ["項次", "功能", "路徑", "舊編號（舊預設）", "目前畫面編號", "新編號（新分類）", "新群組"]


def mapping_xlsx(rows: list[dict]) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "新舊選單編號對照"
    for c, name in enumerate(MAPPING_COLS, 1):
        cell = ws.cell(1, c, name)
        cell.fill = PatternFill("solid", fgColor="00806A")
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = max(10, min(28, len(name) + 6))
    keys = ["no", "label", "to", "old_no", "cur_no", "new_no", "new_group"]
    for i, r in enumerate(rows, 2):
        for c, k in enumerate(keys, 1):
            ws.cell(i, c, r.get(k, ""))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
