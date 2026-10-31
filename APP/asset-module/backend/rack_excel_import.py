"""機櫃圖 Excel 對照與回寫（2.007.20）。

使用者 2026-10-28：機房提供的 Excel 機櫃圖跟系統差太多（嚴格相同只有約 26%），要以 Excel 為準、
目標 9 成相同。做成「可重複」的流程，不是一次性腳本：
  上傳 Excel → 逐格比對系統（IP／硬體序號）→ 列出每一格的狀態與建議 → 使用者勾選 → 套用。

狀態（status）：
  OK         同櫃同 U、對得上同一台
  POSONLY    同櫃同 U 系統有設備，但名稱／IP 對不上（不確定是不是同一台；不當成一致，也不動）
  NOPOS      系統有這台，但沒填櫃位              → 可套用（預設勾）
  OTHERROOM  系統有這台，但登記在別的機房        → 可套用（預設勾）
  MOVED      系統有這台，但放在別的櫃／別的 U    → 可套用，但兩邊誰對要人判斷（預設不勾）
  MISSING    系統沒有這台                        → 只列出，之後另外新增

絕不自動寫的情況（hard blocker，套用時伺服器端再驗一次）：
  mgmt   對到的是管理卡（iDRAC／OOB…）不是主機本體
  vm     對到的是 VM（不是實體機）
  ambig  一格對到好幾台
  vendor Excel 寫的廠牌跟系統型號廠牌不同
  dup    同一台被好幾格指到
軟性（可勾但預設不勾）：overlap（寫入後與別台重疊）、new_rack（系統沒有這個櫃）、corrected（這台的位置已被人工更正過）、moved。

寫入走既有的 rack_layout.set_position（留稽核紀錄、CIA 匯入不會蓋掉、可還原）；套用前先做資料庫快照。
"""
from __future__ import annotations

import io
import json
import re
import sqlite3
from pathlib import Path
from typing import Any

import rack_layout

SKIP_SHEET = re.compile(r"關機|xxx|調整|備份|copy", re.I)
_HEADER_JUNK = re.compile(r"^(U\s*數|panel|牆壁)$", re.I)
JUNK = re.compile(r"預留|電源排插|^panel|Cable Management|^孔$|^位$|牆壁|Patch Panel|^U 數$|螢幕$|^Desktop|KVM$|理線槽|限制使用|上方|擴充櫃")
HARD = {"mgmt", "vm", "ambig", "vendor", "dup", "urange"}
_IP = re.compile(r"(?<![\d.])(?:\d{1,3}\.){3}\d{1,3}(?![\d.])")

BLOCKER_TEXT = {
    "mgmt": "對到的是管理卡（iDRAC／OOB），不是主機本體",
    "vm": "對到的是 VM，不是實體機",
    "ambig": "這一格對到好幾台，無法自動判斷",
    "vendor": "Excel 寫的廠牌跟系統型號不同",
    "dup": "同一台被好幾格指到",
    "urange": "U 編號超出 1-60，Excel 可能讀錯了",
    "overlap": "寫入後會跟別台重疊",
    "new_rack": "系統沒有這個櫃（會新增一個櫃）",
    "corrected": "這台的位置已被人工更正過",
    "moved": "兩邊位置不同，要人判斷誰對",
}


# ---------------------------------------------------------------- 解析 Excel
def _isu(v: Any) -> int | None:
    if isinstance(v, (int, float)) and float(v).is_integer():
        return int(v)
    if isinstance(v, str) and re.fullmatch(r"\d{1,3}", v.strip()):
        return int(v.strip())
    return None


def _parse_sheet(ws) -> list[dict]:
    """一張工作表 → 機櫃區塊。機櫃的畫法：左右各一欄 U 編號（兩欄相隔 2），中間一欄放設備，
    設備佔多個 U 時中間那欄是合併儲存格；U 編號由大到小至少連續 10 列才算一個櫃。"""
    merged = {(m.min_row, m.min_col): (m.max_row, m.max_col) for m in ws.merged_cells.ranges}
    covered = set()
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                if (r, c) != (m.min_row, m.min_col):
                    covered.add((r, c))
    maxr, maxc = ws.max_row, ws.max_column
    found: list[tuple[int, list[tuple[int, int]], int]] = []
    used: set[tuple[int, int]] = set()
    for c in range(1, maxc + 1):
        r = 1
        while r <= maxr:
            u = _isu(ws.cell(r, c).value)
            if u is None:
                r += 1
                continue
            run = [(r, u)]
            rr = r + 1
            while rr <= maxr:
                if (rr, c) in covered:
                    rr += 1
                    continue
                u2 = _isu(ws.cell(rr, c).value)
                if u2 is not None and u2 == run[-1][1] - 1:
                    run.append((rr, u2))
                    rr += 1
                else:
                    break
            if len(run) >= 10 and (c, run[0][0]) not in used:
                twin = c + 2
                if sum(1 for rw, uu in run if _isu(ws.cell(rw, twin).value) == uu) >= len(run) * 0.6:
                    found.append((c, run, rr - 1))
                    used.add((twin, run[0][0]))
            r = rr if rr > r else r + 1
    out = []
    for c, run, last in found:
        r0 = run[0][0]
        name = None
        for cols in ((c + 1,), (c, c + 2)):              # 櫃名優先看設備那欄（合併儲存格的左上角），再看兩側 U 欄
            for rr in range(r0 - 1, max(0, r0 - 4), -1):
                for cc in cols:
                    v = ws.cell(rr, cc).value
                    if isinstance(v, str) and v.strip() and _isu(v) is None and not _HEADER_JUNK.match(v.strip()):
                        name = v.strip()
                        break
                if name:
                    break
            if name:
                break
        urow: dict[int, int | None] = {}
        cur = None
        starts = dict(run)
        for rw in range(r0, last + 1):
            cur = starts.get(rw, cur)
            urow[rw] = cur
        devs = []
        for rw, _uu in run:
            if (rw, c + 1) in covered:
                continue
            v = ws.cell(rw, c + 1).value
            if v in (None, ""):
                continue
            er, _ec = merged.get((rw, c + 1), (rw, c + 1))
            us = [urow[x] for x in range(rw, er + 1) if urow.get(x)]
            if not us:
                continue
            devs.append({"u_end": max(us), "u_start": min(us), "label": re.sub(r"\s+", " ", str(v).replace("\n", " / ")).strip()})
        out.append({"rack_raw": name, "devices": devs})
    return out


def rack_norm(name: str | None) -> str:
    """Excel 的櫃名 → 系統的櫃名寫法：去掉括號說明、「櫃」、前導 Colo-，B08＝B8。"""
    n = re.sub(r"[（(].*?[）)]", "", name or "").strip()
    n = re.sub(r"^Colo-", "", n, flags=re.I).strip()
    return rack_layout._norm_rack(n) if n else ""


def parse_workbook(data: bytes) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    racks: list[dict] = []
    skipped, merged, conflicts = [], [], []
    by_rack: dict[str, dict] = {}
    for ws in wb.worksheets:
        if SKIP_SHEET.search(ws.title):
            skipped.append(ws.title)
            continue
        for blk in _parse_sheet(ws):
            rk = rack_norm(blk["rack_raw"])
            if not rk:
                continue
            if rk not in by_rack:
                by_rack[rk] = {"sheet": ws.title, "rack": rk, "rack_raw": blk["rack_raw"], "devices": list(blk["devices"])}
                racks.append(by_rack[rk])
                continue
            # 內湖：同一個櫃的網路設備畫在 A_NETWORK 分頁、主機畫在別的分頁 → 設備合併；U 重疊的以先出現的為準並記下來
            merged.append(f"{ws.title}/{blk['rack_raw']}")
            have = by_rack[rk]["devices"]
            for d in blk["devices"]:
                if any(d["u_start"] <= e["u_end"] and e["u_start"] <= d["u_end"] for e in have):
                    conflicts.append(f"{rk} {d['u_start']}-{d['u_end']}U：{ws.title} 的「{d['label'][:20]}」跟前面分頁重疊，已略過")
                else:
                    have.append(d)
    dup_racks = merged
    return {"racks": racks, "skipped_sheets": skipped, "duplicate_racks": dup_racks, "conflicts": conflicts}


# ---------------------------------------------------------------- 比對
def _rack_alias(ms: list[dict], room: str) -> dict[str, str]:
    """系統同一個實體櫃有兩種寫法（內湖「1-1」＝「A1」）：由 rack_raw 的括號代碼建對照 {A1: '1-1'}。"""
    out: dict[str, str] = {}
    for m in ms:
        if m["room"] != room or not m.get("rack"):
            continue
        hit = re.search(r"\(([A-Za-z]\d{1,3})\)", m.get("rack_raw") or "")
        if hit and rack_layout._norm_rack(hit.group(1)) != m["rack"]:
            out[rack_layout._norm_rack(hit.group(1))] = m["rack"]
    return out


def _index(ms: list[dict]):
    by_ip: dict[str, list[dict]] = {}
    by_sn: dict[str, list[dict]] = {}
    for m in ms:
        if m["kind"] == "removed":
            continue
        for ip in _IP.findall(str(m.get("ip") or "")):
            by_ip.setdefault(ip, []).append(m)
        for sn in list(m.get("hw_serials") or []):
            if len(sn) >= 6:
                by_sn.setdefault(sn.upper(), []).append(m)
    return by_ip, by_sn


def _tokens(label: str) -> tuple[list[str], list[str]]:
    ips = _IP.findall(label)
    sns = []
    for t in re.findall(r"[0-9A-Za-z]{6,12}", label):
        u = t.upper()
        if re.search(r"[A-Z]", u) and re.search(r"\d", u):
            sns.append(u)
    return ips, sns


def _is_vm(m: dict) -> bool:
    return (m.get("model") or "").strip().upper() in ("(VM)", "VM", "（VM）") or m.get("kind") in ("vm", "vm_model")


def _vendor_clash(label: str, m: dict) -> bool:
    a, b = rack_layout.vendor_of({"model": label, "name": ""}), rack_layout.vendor_of(m)
    if {a, b} == {"ibm", "lenovo"}:          # System x 2014 年起是 Lenovo：Excel 寫 IBM X3550M5、系統寫 Lenovo x3550 M5 是同一台
        return False
    return not rack_layout.vendor_compatible({"model": label, "name": ""}, m)


def _quality(h: dict) -> int:
    """同一個 IP 在系統裡有好幾筆時挑哪一筆：有機櫃位置的 > 有型號的 > 自動建立的空殼（DYN-…／沒型號沒櫃位）。"""
    if h["kind"] == "rack" and h.get("rack"):
        return 3
    return 2 if (h.get("model") or "").strip() else 1


def _same_rack(a: str | None, b: str | None, alias: dict[str, str]) -> bool:
    if not a or not b:
        return False
    return a == b or alias.get(a) == b or alias.get(b) == a


def _hit_machines(label: str, idx, ms) -> tuple[list[dict], list[str]]:
    by_ip, by_sn = idx
    ips, sns = _tokens(label)
    found: dict[str, dict] = {}
    for ip in ips:
        for m in by_ip.get(ip, []):
            found[m["key"]] = m
    for sn in sns:
        for m in by_sn.get(sn, []):
            found[m["key"]] = m
    return list(found.values()), ips


def build_preview(conn: sqlite3.Connection, room: str, parsed: dict, ms: list[dict] | None = None) -> dict:
    ms = ms if ms is not None else rack_layout.machines(conn)
    alias = _rack_alias(ms, room)
    idx = _index(ms)
    sys_racks = {m["rack"] for m in ms if m["room"] == room and m.get("rack")}
    occ: dict[tuple[str, int], list[dict]] = {}
    for m in ms:
        if m["room"] == room and m["kind"] == "rack" and m.get("rack") and m.get("u_start"):
            for u in range(m["u_start"], (m["u_end"] or m["u_start"]) + 1):
                occ.setdefault((m["rack"], u), []).append(m)
    rows: list[dict] = []
    for blk in parsed["racks"]:
        rk = blk["rack"]
        to_rack = rk if rk in sys_racks else alias.get(rk, rk)           # 寫進哪個櫃（別名只在系統沒有同名櫃時才用）
        same_names = {rk, to_rack} | {k for k, v in alias.items() if v == rk or k == rk}   # 同一個實體櫃的所有寫法
        for d in blk["devices"]:
            label = d["label"]
            if JUNK.search(label):
                continue
            us = set(range(d["u_start"], d["u_end"] + 1))
            hits, _ips = _hit_machines(label, idx, ms)
            blockers: list[str] = []
            target = None
            if hits:
                real = [h for h in hits if not rack_layout.is_mgmt_record(h) and not _is_vm(h)]
                if not real:
                    blockers.append("vm" if any(_is_vm(h) for h in hits) else "mgmt")
                    target = hits[0]
                else:
                    ok = [h for h in real if not _vendor_clash(label, h)]
                    if not ok:
                        blockers.append("vendor")
                        target = real[0]
                    else:
                        top = max(_quality(h) for h in ok)
                        best = [h for h in ok if _quality(h) == top]
                        target = best[0]
                        if len(best) > 1:
                            blockers.append("ambig")
            row = {"rack": rk, "to_rack": to_rack, "u_start": d["u_start"], "u_end": d["u_end"], "label": label,
                   "sheet": blk["sheet"], "key": None, "sys_name": None, "sys_ip": None, "sys_now": None}
            if d["u_start"] < 1 or d["u_end"] > 60:
                blockers.append("urange")
            if target is not None:
                row.update(key=target["key"], sys_name=target["name"], sys_ip=target.get("ip"),
                           sys_now=_where(target))
            if target is not None and not (set(blockers) & HARD):
                cur_rack = target.get("rack") if target["kind"] == "rack" else None
                has_pos = bool(cur_rack and target.get("u_start"))
                if target["room"] == room and has_pos and _same_rack(cur_rack, rk, alias) \
                        and us & set(range(target["u_start"], (target["u_end"] or target["u_start"]) + 1)):
                    status = "OK"
                elif not has_pos:
                    status = "NOPOS"
                elif target["room"] != room:
                    status = "OTHERROOM"
                else:
                    status = "MOVED"
                    blockers.append("moved")
                if status != "OK":
                    if target.get("position_corrected"):
                        blockers.append("corrected")
                    if rk not in sys_racks and alias.get(rk) is None:
                        blockers.append("new_rack")
                    others = [o for rn in same_names for u in us for o in occ.get((rn, u), []) if o["key"] != target["key"]
                              and not rack_layout.is_mgmt_record(o)]
                    if others:
                        blockers.append("overlap")
                row["status"] = status
            elif target is not None:
                row["status"] = "BLOCKED"
            else:
                same = [o for rr in same_names for u in us for o in occ.get((rr, u), [])]
                if same:
                    row["status"] = "POSONLY"
                    row["sys_name"] = same[0]["name"]
                    row["sys_ip"] = same[0].get("ip")
                    row["sys_now"] = _where(same[0])
                    row["occupant_keys"] = sorted({o["key"] for o in same})       # 佔著這格的系統設備（面板「確認」欄用）
                else:
                    row["status"] = "MISSING"
            row["blockers"] = blockers
            row["action"] = "set" if row["status"] in ("NOPOS", "OTHERROOM", "MOVED") and not (set(blockers) & HARD) else None
            rows.append(row)
    # 同一台被好幾格指到 → 全部降成 dup
    cnt: dict[str, int] = {}
    for r in rows:
        if r["action"] or r["status"] == "OK":
            cnt[r["key"]] = cnt.get(r["key"], 0) + 1
    for r in rows:
        if r["action"] and cnt[r["key"]] > 1:
            r["blockers"].append("dup")
            r["action"] = None
            r["status"] = "BLOCKED"
    for i, r in enumerate(rows, 1):
        r["id"] = i
        r["auto"] = bool(r["action"] and r["status"] in ("NOPOS", "OTHERROOM") and not r["blockers"])
        r["blocker_text"] = "；".join(BLOCKER_TEXT[b] for b in r["blockers"])
    return {"rows": rows, "summary": _summary(rows, parsed), "alias": alias}


def _where(m: dict) -> str:
    if m["kind"] == "rack" and m.get("rack"):
        lo, hi = m.get("u_start"), m.get("u_end") or m.get("u_start")
        u = ("沒填U" if not lo else f"{lo}U" if lo == hi else f"{lo}-{hi}U")
        return f"{m['room']}/{m['rack']}/{u}"
    return f"{m['room']}/沒有櫃位"


def _summary(rows: list[dict], parsed: dict) -> dict:
    c: dict[str, int] = {}
    for r in rows:
        c[r["status"]] = c.get(r["status"], 0) + 1
    n = len(rows)
    return {"total": n, "by_status": c, "strict_ok_pct": round(c.get("OK", 0) / n * 100) if n else 0,
            "auto_count": sum(1 for r in rows if r["auto"]),
            "racks": len(parsed["racks"]), "skipped_sheets": parsed["skipped_sheets"],
            "duplicate_racks": parsed["duplicate_racks"]}


# ---------------------------------------------------------------- 儲存與套用
def ensure(conn: sqlite3.Connection) -> None:
    conn.execute("CREATE TABLE IF NOT EXISTS rack_excel_run ("
                 " id INTEGER PRIMARY KEY AUTOINCREMENT, filename TEXT, room TEXT NOT NULL, created_by TEXT,"
                 " created_at TEXT NOT NULL DEFAULT (datetime('now','localtime')), summary_json TEXT, rows_json TEXT,"
                 " applied_json TEXT)")


def preview(conn: sqlite3.Connection, room: str, data: bytes, filename: str, by: str) -> dict:
    ensure(conn)
    room = (room or "").strip()
    if not room:
        raise ValueError("要選機房")
    try:
        parsed = parse_workbook(data)
    except Exception as exc:  # noqa: BLE001 - openpyxl 對壞檔丟各種例外，統一轉成給使用者看的訊息
        raise ValueError(f"讀不了這個 Excel：{exc}") from exc
    if not parsed["racks"]:
        raise ValueError("這個檔案裡找不到機櫃圖（要有由大到小的 U 編號欄，左右各一欄、中間放設備）")
    res = build_preview(conn, room, parsed)
    cur = conn.execute("INSERT INTO rack_excel_run (filename, room, created_by, summary_json, rows_json) VALUES (?,?,?,?,?)",
                       (filename, room, by, json.dumps(res["summary"], ensure_ascii=False),
                        json.dumps(res["rows"], ensure_ascii=False)))
    conn.commit()
    return {"run_id": cur.lastrowid, "room": room, "filename": filename, **res}


def get_run(conn: sqlite3.Connection, run_id: int) -> dict:
    ensure(conn)
    r = conn.execute("SELECT * FROM rack_excel_run WHERE id = ?", (run_id,)).fetchone()
    if r is None:
        raise ValueError("查無這次對照")
    return {"run_id": r["id"], "room": r["room"], "filename": r["filename"], "created_by": r["created_by"],
            "created_at": r["created_at"], "summary": json.loads(r["summary_json"] or "{}"),
            "rows": json.loads(r["rows_json"] or "[]"), "applied": json.loads(r["applied_json"] or "[]")}


def list_runs(conn: sqlite3.Connection, limit: int = 30) -> list[dict]:
    ensure(conn)
    return [{"run_id": r["id"], "room": r["room"], "filename": r["filename"], "created_by": r["created_by"],
             "created_at": r["created_at"], "summary": json.loads(r["summary_json"] or "{}"),
             "applied_count": len(json.loads(r["applied_json"] or "[]"))}
            for r in conn.execute("SELECT id, room, filename, created_by, created_at, summary_json, applied_json "
                                  "FROM rack_excel_run ORDER BY id DESC LIMIT ?", (limit,))]


def apply(conn: sqlite3.Connection, run_id: int, row_ids: list[int], by: str, backup_to: Path | None = None) -> dict:
    """套用使用者勾選的格。伺服器端依「現在」的資料重算一次，硬性不安全的一律跳過（不信任畫面傳來的勾選）。"""
    run = get_run(conn, run_id)
    want = set(int(x) for x in row_ids)
    if not want:
        raise ValueError("沒有勾選任何一格")
    already = {a["id"] for a in run["applied"]}
    # 以「現在」的狀態重算整份（資料可能在預覽後被改過）
    fresh = {r["id"]: r for r in _recompute(conn, run)}
    todo, skipped = [], []
    for rid in sorted(want):
        old = next((r for r in run["rows"] if r["id"] == rid), None)
        r = fresh.get(rid)
        if old is None or r is None:
            skipped.append({"id": rid, "why": "查無這一格"})
        elif rid in already:
            skipped.append({"id": rid, "why": "這一格已經套用過"})
        elif r["action"] != "set" or (set(r["blockers"]) & HARD):
            skipped.append({"id": rid, "why": "現在不能自動寫：" + (r["blocker_text"] or r["status"])})
        elif r["key"] != old["key"]:
            skipped.append({"id": rid, "why": "預覽後對到的設備變了，請重新上傳"})
        else:
            todo.append(r)
    # 同一批裡不能互相重疊
    taken: dict[tuple[str, int], int] = {}
    final = []
    for r in todo:
        clash = next((taken[(r["to_rack"], u)] for u in range(r["u_start"], r["u_end"] + 1) if (r["to_rack"], u) in taken), None)
        if clash:
            skipped.append({"id": r["id"], "why": f"跟本批第 {clash} 格的位置重疊"})
            continue
        for u in range(r["u_start"], r["u_end"] + 1):
            taken[(r["to_rack"], u)] = r["id"]
        final.append(r)
    if not final:
        return {"applied": 0, "skipped": skipped, "backup": None}
    backup_file = None
    if backup_to is not None:
        import backup
        from db import get_db_path
        backup_to.mkdir(parents=True, exist_ok=True)
        backup_file = backup_to / f"pre_rack_excel_run{run_id}_{__import__('datetime').datetime.now():%Y%m%d_%H%M%S_%f}.db"
        backup.snapshot(get_db_path(), backup_file)
    ms = rack_layout.machines(conn)
    note = f"Excel 機櫃圖匯入：{run['filename']}（第 {run['run_id']} 次對照）"
    done = []
    for r in final:
        rack_layout.set_position(conn, r["key"], run["room"], r["to_rack"], r["u_start"], r["u_end"] - r["u_start"] + 1,
                                 note, by, ms=ms, commit=False)
        done.append({"id": r["id"], "key": r["key"], "name": r["sys_name"], "to": f"{run['room']}/{r['to_rack']}/{r['u_start']}-{r['u_end']}U",
                     "from": r["sys_now"]})
    conn.execute("UPDATE rack_excel_run SET applied_json = ? WHERE id = ?",
                 (json.dumps(run["applied"] + done, ensure_ascii=False), run_id))
    conn.commit()
    return {"applied": len(done), "skipped": skipped, "backup": str(backup_file) if backup_file else None, "done": done}


def _recompute(conn: sqlite3.Connection, run: dict, ms: list[dict] | None = None) -> list[dict]:
    """用預覽時存下的 Excel 格（櫃、U、文字）再比一次現在的系統。id 依原順序，與預覽一致。"""
    racks: dict[str, dict] = {}
    for r in run["rows"]:
        b = racks.setdefault(r["rack"], {"sheet": r["sheet"], "rack": r["rack"], "rack_raw": r["rack"], "devices": []})
        b["devices"].append({"u_start": r["u_start"], "u_end": r["u_end"], "label": r["label"]})
    res = build_preview(conn, run["room"], {"racks": list(racks.values()), "skipped_sheets": [], "duplicate_racks": []}, ms)
    # build_preview 依櫃分組後重新編號：用 (櫃, U, 文字) 對回原本的 id
    key = {(r["rack"], r["u_start"], r["u_end"], r["label"]): r["id"] for r in run["rows"]}
    out = []
    for r in res["rows"]:
        rid = key.get((r["rack"], r["u_start"], r["u_end"], r["label"]))
        if rid is not None:
            r["id"] = rid
            out.append(r)
    return out


# ---------------------------------------------------------------- 機櫃正面圖的「Excel 層」（2026-10-28：以 Excel 為主，先全部畫出來）
# Excel 是人工盤點的結果，系統（CIA）沒登或沒填位置不能因此不畫。這層只是「畫出來給人看」：
# 不新增資產、不算台數、不算用電；使用者確認後才另外轉成正式資產（口徑 A）。
SUGGEST = {
    "NOPOS": "系統有這台、但沒登位置：到「Excel 對照」勾選回寫位置",
    "OTHERROOM": "系統把這台登在別的機房：確認後到「Excel 對照」回寫位置",
    "MOVED": "系統登的位置跟 Excel 不同：請人到現場確認誰對",
    "MISSING": "系統沒有這台：請補登 CIA，或確認後新增成資產",
    "BLOCKED": "系統有對到，但不能自動寫入（見原因）：請人工確認",
}
_DRAWN = ("OK", "POSONLY")        # 這兩種系統面板已經在同一格，不用再畫一個


def _serial_of(m: dict | None) -> str:
    """面板上的序號＝硬體序號（service tag）；資產編號是另一欄，不拿來頂替。"""
    return ((m or {}).get("hw_serials") or [""])[0]


def _asset_of(m: dict | None) -> str:
    """資產編號（HW-…）；DYN-／VC-／AUTO- 是系統自己產的，不顯示。"""
    a = (m or {}).get("asset_serial") or ""
    return "" if re.match(r"(?i)^(DYN|VC|AUTO)-", a) else a


def excel_only_racks(conn: sqlite3.Connection, room: str, ms: list[dict]) -> list[dict]:
    """該機房最近一次 Excel 對照裡有、但系統（CIA）完全沒有這個櫃的設備登記的櫃（例：Excel 的 B02(VMware)）。
    讓這些櫃也能在機房平面圖點開（只畫 Excel 層，不新增資產、不算台數）。內湖那種同一實體櫃兩種寫法的，先對回系統的寫法。"""
    ensure(conn)
    r = conn.execute("SELECT id FROM rack_excel_run WHERE room = ? ORDER BY id DESC LIMIT 1", (room,)).fetchone()
    if r is None:
        return []
    run = get_run(conn, r["id"])
    alias = _rack_alias(ms, room)
    have = {m["rack"] for m in ms if m["room"] == room and m.get("rack") and m["kind"] in ("rack", "removed")}
    out: dict[str, dict] = {}
    for x in run["rows"]:
        rk = alias.get(x["rack"], x["rack"])
        if not rk or rk in have:
            continue
        e = out.setdefault(rk, {"rack": rk, "cells": 0})
        e["cells"] += 1
    return sorted(out.values(), key=lambda e: rack_layout._natural(e["rack"]))


def rack_cells(conn: sqlite3.Connection, room: str, rack: str, ms: list[dict]) -> dict:
    """這個櫃在「該機房最近一次 Excel 對照」裡、系統面板還沒畫到的格（依現在的系統資料重算，回寫過的就會消失）。"""
    ensure(conn)
    r = conn.execute("SELECT id FROM rack_excel_run WHERE room = ? ORDER BY id DESC LIMIT 1", (room,)).fetchone()
    if r is None:
        return {"run": None, "cells": [], "total": 0}
    run = get_run(conn, r["id"])
    here = [x for x in run["rows"] if x["rack"] == rack or x.get("to_rack") == rack]
    if not here:
        return {"run": {"run_id": run["run_id"], "filename": run["filename"], "created_at": run["created_at"]},
                "cells": [], "total": 0}
    fresh = _recompute(conn, {**run, "rows": here}, ms)
    by_key = {m["key"]: m for m in ms}
    apis = sorted({a for x in fresh if x.get("key") in by_key for a in by_key[x["key"]]["api_ids"]})
    sysnames = rack_layout._system_names(conn, apis)
    guess = rack_layout._system_guess(conn, apis)
    cells = [{"id": x["id"], "key": x.get("key"), "u_start": x["u_start"], "u_end": x["u_end"], "label": x["label"], "status": x["status"],
              "sys_name": x["sys_name"], "sys_ip": x["sys_ip"], "sys_now": x["sys_now"], "blocker_text": x["blocker_text"],
              "sys_serial": _serial_of(by_key.get(x.get("key"))),
              "sys_asset": _asset_of(by_key.get(x.get("key"))),
              "sys_systems": [(a if a in guess else sysnames[a]) for a in (by_key[x["key"]]["api_ids"] if x.get("key") in by_key else []) if a in sysnames],
              "suggest": SUGGEST.get(x["status"], "")}
             for x in fresh if x["status"] not in _DRAWN]
    cells.sort(key=lambda c: -c["u_end"])
    # 每台系統設備跟 Excel 對的結果（面板燈號與「確認」欄用）：一台被好幾格指到時取最差的那個
    order = ["MISSING", "BLOCKED", "MOVED", "OTHERROOM", "NOPOS", "POSONLY", "OK"]
    per_key: dict[str, str] = {}
    for x in fresh:
        for k in ([x["key"]] if x.get("key") else x.get("occupant_keys", [])):
            if k not in per_key or order.index(x["status"]) < order.index(per_key[k]):
                per_key[k] = x["status"]
    return {"run": {"run_id": run["run_id"], "filename": run["filename"], "created_at": run["created_at"]},
            "cells": cells, "total": len(fresh), "by_key": per_key}
