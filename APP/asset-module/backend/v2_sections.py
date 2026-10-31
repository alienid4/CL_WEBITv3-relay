"""A2 病歷表的「全部資料」（2026-10-27 使用者：要完整病歷表、祖宗十八代）。

舊資產頁 /assets/[serial] 的每一個頁簽，在這裡都有一段（OLD_TAB_MAP 釘住，測試會擋）；外加 2.x 帶來的
連線、異動、所有登記、盤點／位置更正紀錄、存在證據。每一段只呼叫既有的計算（舊頁同一支），這裡只做「整理成表」。

每一段回同一個形狀：{key, group, label, count, unknown, kind, cols, rows, note}
  count＝rows 的筆數（verify_numbers 核：頁簽上的數字＝點進去的列數）；
  查不到的段 count=None、unknown＝原因（畫面寫「無資料（原因）」，不留白）。
"""
from __future__ import annotations

import io
import json
import sqlite3
from typing import Any, Callable

GROUPS = [
    # 主機履歷的橫排頁籤（藍圖定稿，2026-10-27）：概要（前端）｜影響範圍｜硬體與位置｜軟體與服務｜帳號與安全｜
    # 關聯設備｜異動紀錄｜原始資料｜系統重建資料
    ("impact", "影響範圍"), ("hw", "硬體與位置"), ("sw", "軟體與服務"), ("sec", "帳號與安全"),
    ("near", "關聯設備"), ("chg", "異動紀錄"), ("raw", "原始資料"), ("rebuild", "系統重建資料"),
]
#: 舊資產頁的頁簽（＋IIS 樹、OCP 同叢集區塊）→ A2 的哪一段。拿掉任何一個，test_v2_sections 會紅。
OLD_TAB_MAP = {
    "advanced": ["cia", "personnel"],    # 重點欄位＝上方 6 格＋CIA 原始欄位（聯絡卡在第 3 格）
    "hwinfo": ["hw_spec", "disks", "nics"], "personnel": ["personnel"], "software": ["software"],
    "services": ["services"], "accounts": ["accounts"], "vulns": ["vulns"], "audit": ["audit"], "cia": ["cia"],
    "health": ["health"], "san": ["san"], "rebuild": ["rebuild"], "schedule": ["schedule"], "history": ["changes"],
    "iis": ["services"], "ocp_siblings": ["ocp"],
}


class Ctx:
    """一台機器的共用資料（同一個請求裡只算一次）。"""

    def __init__(self, conn: sqlite3.Connection, key: str):
        import manage_state as ms
        import oncall

        self.conn = conn
        groups, _ = oncall._groups(conn)
        if key not in groups:
            how, keys = oncall.resolve(conn, key, groups)
            if len(keys) != 1:
                raise LookupError("查無這台（或對到不只一台）")
            key = keys[0]
        self.key = key
        self.groups = groups
        self.rows = [dict(r) for r in groups[key]]
        self.serials = [r["asset_serial"] for r in self.rows]
        self.serial = self.serials[0]
        self.ips = []
        for r in self.rows:
            for ip in ms.ips_of(r["ip"]):
                if ip not in self.ips:
                    self.ips.append(ip)
        self._detail = None

    @property
    def mkey(self) -> str:
        return self.key[4:] if self.key.startswith("ret:") else self.key

    def detail(self) -> dict:
        """舊資產頁的 /api/assets/{serial}（同一支函式）。"""
        if self._detail is None:
            from routers import assets as ra
            self._detail = ra.asset_detail(self.serial, None, self.conn)
        return self._detail

    def impact(self) -> dict | None:
        """4-2／3-2 影響範圍同一支計算（blast_radius.resolve＋impact，mode=probe 順便拿「相依元件」）。
        CI 圖譜裡沒有這台回 None。這裡不另寫走訪邏輯。"""
        if not hasattr(self, "_imp"):
            import blast_radius
            self._imp = None
            try:
                res = blast_radius.resolve(self.conn, self.serial)
            except sqlite3.OperationalError as exc:
                if "no such table" not in str(exc):
                    raise
                res = {}                       # 還沒建過 CI 圖譜
            if res.get("status") == "resolved":
                self._imp = blast_radius.impact(self.conn, res["node_id"], mode="probe")
        return self._imp

    def node_label(self, nid: str) -> dict:
        """CI 節點 → 畫面用的名字（主機用 name_of，跟 A2 其他地方同一條）。"""
        import v2_views

        n = self.conn.execute("SELECT node_type, label, asset_serial FROM ci_node WHERE node_id = ?", (nid,)).fetchone()
        serial = n["asset_serial"] if n else None
        row = self.conn.execute("SELECT * FROM hardware WHERE asset_serial = ?", (serial,)).fetchone() if serial else None
        name = v2_views.name_of(row) if row else ((n["label"] if n else None) or nid)
        return {"node_type": n["node_type"] if n else nid.split(":")[0], "name": name, "serial": serial,
                "ip": row["ip"] if row else None, "key": serial or name,
                "badge": v2_views.type_badge(row) if row else None}

    def onboarded(self) -> bool:
        """已納管＝existence 的 onboarded（收集帳號 7 天內連得上），跟 6 格第 5 格同一個判斷。"""
        if not hasattr(self, "_onb"):
            import existence
            m = existence.for_serial(self.conn, self.serial)
            self._onb = bool(m and m.get("onboarded"))
        return self._onb


def flat_reason(t: str | None) -> str | None:
    """原因字串會被畫面包成「無資料（原因）」，原因裡的全形括號改成冒號，避免「（…（…））」兩層。"""
    if not t:
        return t
    import re
    prev = None
    while prev != t:
        prev = t
        t = re.sub(r"（([^（）]*)）$", r"：\1", t)          # 結尾的括號 → 冒號
        t = re.sub(r"（([^（）]*)）", r"，\1，", t)          # 中間的括號 → 逗號
    return t.replace("，，", "，").rstrip("，")


def _sec(key, label, cols, rows, unknown=None, kind="table", note=None) -> dict:
    unknown = flat_reason(unknown)
    return {"key": key, "label": label, "kind": kind, "cols": cols, "rows": rows,
            "count": None if unknown else len(rows), "unknown": unknown, "note": note}


def _c(k, label, mono=False):
    return {"k": k, "label": label, "mono": mono}


def _not_collected(ctx: Ctx, what: str) -> str:
    return f"未納管，尚未收集{what}" if not ctx.onboarded() else f"已納管，尚未收集到{what}"


def _flat(d: dict, prefix: str = "") -> list[tuple[str, str]]:
    out = []
    for k, v in (d or {}).items():
        if v is None or v == "" or v == [] or v == {}:
            continue
        if isinstance(v, dict):
            out += _flat(v, f"{prefix}{k}.")
        elif isinstance(v, list):
            out.append((f"{prefix}{k}", "、".join(_render(x) for x in v)))
        else:
            out.append((f"{prefix}{k}", str(v)))
    return out


def _render(v: Any) -> str:
    if isinstance(v, dict):
        return " ".join(f"{k}={_render(x)}" for k, x in v.items() if x not in (None, "", [], {}))
    if isinstance(v, list):
        return "\n".join(_render(x) for x in v)
    return "" if v is None else str(v)


# ===== 1 硬體與位置 =====

def s_hw_spec(ctx: Ctx) -> dict:
    d = ctx.detail()
    hs = d.get("host_spec")
    rows = []
    for k in ("device_model", "hw_serial", "asset_purpose", "is_vm", "os", "environment"):
        v = d["hardware"].get(k)
        if v not in (None, ""):
            rows.append({"source": "CIA 登記", "field": k, "value": str(v)})
    if hs:
        for k, v in _flat({x: hs.get(x) for x in ("platform", "cpu", "cores", "mem_mb", "storage", "model", "serial",
                                                  "vendor", "collected_at")}):
            rows.append({"source": "系統讀到", "field": k, "value": v})
        for k, v in _flat(hs.get("spec_json") or {}):
            rows.append({"source": "系統讀到", "field": k, "value": v})
        for k, v in _flat(hs.get("psu_json") or {}):
            rows.append({"source": "系統收集（電源）", "field": k, "value": v})
    note = None if hs else _not_collected(ctx, "硬體規格")
    return _sec("hw_spec", "硬體規格", [_c("source", "來源"), _c("field", "欄位", True), _c("value", "值")], rows, note=note)


def s_disks(ctx: Ctx) -> dict:
    hs = ctx.detail().get("host_spec")
    st = (hs or {}).get("storage_json") or {}
    rows = [{"kind": st.get("by") or "", "name": x.get("name"), "size_gb": x.get("size_gb"), "free_gb": x.get("free_gb")}
            for x in st.get("items") or []]
    rows += [{"kind": "HBA", "name": h.get("hba_name") or h.get("name") or h.get("wwpn"), "size_gb": None,
              "free_gb": None, "detail": _render({k: v for k, v in h.items() if k not in ("id", "asset_serial")})}
             for h in ctx.detail().get("hbas") or []]
    if not rows:
        return _sec("disks", "磁碟／儲存", [], [], unknown=_not_collected(ctx, "磁碟"))
    return _sec("disks", "磁碟／儲存", [_c("kind", "種類"), _c("name", "名稱", True), _c("size_gb", "大小 GB", True),
                                     _c("free_gb", "剩餘 GB", True), _c("detail", "細節")], rows)


def s_nics(ctx: Ctx) -> dict:
    rows = [{"nic": n.get("nic_name"), "ip": n.get("ip"), "mac": n.get("mac"), "subnet": n.get("subnet"),
             "gateway": n.get("gateway"), "vlan": n.get("vlan"), "bond": n.get("bond"), "speed": n.get("speed_cur"),
             "state": n.get("state"), "source": "系統讀到"} for n in ctx.detail().get("nics") or []]
    seen = {r["ip"] for r in rows}
    rows += [{"nic": None, "ip": ip, "mac": None, "subnet": None, "gateway": None, "vlan": None, "bond": None,
              "speed": None, "state": None, "source": "CIA 登記"} for ip in ctx.ips if ip not in seen]
    if not rows:
        return _sec("nics", "網路卡／IP", [], [], unknown="CIA 未填 IP，也尚未收集網路卡")
    return _sec("nics", "網路卡／IP", [_c("nic", "網卡", True), _c("ip", "IP", True), _c("mac", "MAC", True),
                                     _c("subnet", "遮罩", True), _c("gateway", "閘道", True), _c("vlan", "VLAN"),
                                     _c("bond", "Bond"), _c("speed", "速度 Mb", True), _c("state", "狀態"),
                                     _c("source", "來源")], rows)


def s_san(ctx: Ctx) -> dict:
    from fastapi import HTTPException
    from routers import assets as ra

    rows = []
    for ip in ctx.ips:
        try:
            d = ra.san_switch_detail_endpoint(ip, None, ctx.conn)
        except HTTPException:
            continue
        data = d.get("data") or {}
        for k, v in _flat({x: y for x, y in d.items() if x != "data"}):
            rows.append({"ip": ip, "field": k, "value": v})
        for k, v in _flat(data):
            rows.append({"ip": ip, "field": k, "value": v[:2000]})
    if not rows:
        return {**_sec("san", "SAN", [], [], unknown="這台不是收過資料的 SAN 交換器"), "na": True}   # 跟這台無關：畫面不顯示（找碴 C18）
    return _sec("san", "SAN", [_c("ip", "IP", True), _c("field", "欄位", True), _c("value", "值")], rows)


def s_location(ctx: Ctx) -> dict:
    """位置與盤點紀錄：位置更正、拖拉搬移、機櫃狀態標記、機櫃實體盤點的結果。"""
    import rack_inventory
    import rack_layout as rl

    conn = ctx.conn
    rl.ensure_position_table(conn)
    rl.ensure_state_table(conn)
    rl.ensure_move_log(conn)
    rack_inventory.ensure_tables(conn)
    keys = {ctx.key, ctx.mkey}
    rows = []
    for r in conn.execute("SELECT * FROM rack_position_audit WHERE key IN (?, ?) ORDER BY id", (ctx.key, ctx.mkey)):
        after = json.loads(r["after_json"]) if r["after_json"] else {}
        rows.append({"at": r["updated_at"], "kind": f"位置更正（{r['action']}）", "by": r["updated_by"],
                     "what": _render(after) or "清除更正", "note": r["note"]})
    for r in conn.execute("SELECT * FROM rack_state_audit WHERE key IN (?, ?) ORDER BY id", (ctx.key, ctx.mkey)):
        rows.append({"at": r["updated_at"], "kind": "機櫃狀態標記", "by": r["updated_by"],
                     "what": f"{r['old_state'] or '—'} → {r['new_state'] or '清除'}", "note": r["note"]})
    for r in conn.execute("SELECT * FROM rack_move_log ORDER BY id"):
        if keys & set(json.loads(r["after_json"])):
            frm, to = json.loads(r["from_json"]), json.loads(r["to_json"])
            rows.append({"at": r["moved_at"], "kind": "拖拉搬移" + ("（已復原）" if r["undone_at"] else ""), "by": r["moved_by"],
                         "what": f"{frm['rack']} {rl._u_txt(frm['u_start'], frm['u_end'])} → {to['rack']} "
                                 f"{rl._u_txt(to['u_start'], to['u_end'])}", "note": r["note"]})
    for s in conn.execute("SELECT * FROM rack_audit_submit WHERE superseded = 0 ORDER BY id"):
        snap = conn.execute("SELECT snapshot_json FROM rack_audit_rack WHERE audit_id = ? AND room = ? AND rack = ?",
                            (s["audit_id"], s["room"], s["rack"])).fetchone()
        for row in json.loads(snap[0]) if snap else []:
            if keys & set(row["all_keys"]):
                got = json.loads(s["rows_json"]).get(row["key"], {"result": "在"})
                rows.append({"at": s["counted_on"], "kind": f"機櫃盤點 #{s['audit_id']}", "by": s["counted_by"],
                             "what": f"{s['room']} {s['rack']}：{got['result']}"
                                     + (f"（正確 U {got.get('correct_u')}）" if got.get("correct_u") else ""),
                             "note": got.get("note")})
    rows.sort(key=lambda x: x["at"] or "", reverse=True)
    m = next((x for x in rl.machines(conn) if x["key"] == ctx.mkey), None)
    note = None
    if m:
        note = (f"現在：{m['room']} {m.get('rack') or ''} {rl._u_text(m) if m.get('u_start') else ''}".strip()
                + ("（現場更正過）" if m.get("position_corrected") else ""))
    return _sec("location", "位置與盤點紀錄", [_c("at", "時間", True), _c("kind", "種類"), _c("what", "內容"),
                                          _c("by", "誰"), _c("note", "備註")], rows, note=note)


# ===== 2 軟體與服務 =====

def s_software(ctx: Ctx) -> dict:
    d = ctx.detail()
    rows = [{"name": p.get("name"), "version": p.get("version"), "arch": p.get("arch"), "vendor": p.get("vendor"),
             "source": p.get("source"), "last_seen": p.get("last_seen")} for p in d.get("packages") or []]
    rows += [{"name": s.get("software_name") or s.get("name") or s.get("db_name"), "version": s.get("version"),
              "arch": None, "vendor": s.get("vendor"), "source": "CIA 登記", "last_seen": None}
             for s in d.get("software") or []]
    if not rows:
        return _sec("software", "軟體", [], [], unknown=_not_collected(ctx, "軟體清單"))
    return _sec("software", "軟體", [_c("name", "軟體"), _c("version", "版本", True), _c("arch", "架構"),
                                  _c("vendor", "廠商"), _c("source", "來源"), _c("last_seen", "最後看到", True)], rows)


def s_services(ctx: Ctx) -> dict:
    import iis_inventory
    import service_inventory

    rows = []
    for s in ctx.serials[:1]:          # 範圍本來就是「整台」（service_inventory 依同一台展開），只查代表那筆
        for r in service_inventory.list_services(ctx.conn, asset_serial=s, include_gone=False, include_infra=True):
            rows.append({"type": "埠", "name": f"{r.get('proto') or ''}/{r.get('port')}", "detail": r.get("service_guess"),
                         "process": r.get("process"), "exposure": r.get("exposure"), "last_seen": r.get("last_seen")})
        for r in service_inventory.list_winservices(ctx.conn, asset_serial=s):
            rows.append({"type": "Windows 服務", "name": r.get("name") or r.get("service_name"),
                         "detail": r.get("display_name"), "process": r.get("path") or r.get("binary_path"),
                         "exposure": r.get("start_mode") or r.get("state"), "last_seen": r.get("last_seen")})
        iis = iis_inventory.for_asset(ctx.conn, s)
        for st in iis.get("sites") or []:
            rows.append({"type": "IIS 網站", "name": st.get("name"), "detail": _render(st.get("bindings")) or st.get("physical_path"),
                         "process": st.get("app_pool") or st.get("pool"), "exposure": st.get("state"), "last_seen": st.get("last_seen")})
        for p in iis.get("pools") or []:
            rows.append({"type": "IIS 應用程式集區", "name": p.get("name"), "detail": p.get("identity") or p.get("runtime"),
                         "process": None, "exposure": p.get("state"), "last_seen": p.get("last_seen")})
    if not rows:
        return _sec("services", "服務／埠（含 IIS）", [], [], unknown=_not_collected(ctx, "服務"))
    return _sec("services", "服務／埠（含 IIS）", [_c("type", "種類"), _c("name", "埠／名稱", True), _c("detail", "服務"),
                                             _c("process", "程式／集區"), _c("exposure", "對外／狀態"),
                                             _c("last_seen", "最後看到", True)], rows)


def s_schedule(ctx: Ctx) -> dict:
    import schedule_inventory

    d = schedule_inventory.detail_for(ctx.conn, ctx.serial)
    rows = [{"source": r.get("source"), "account": r.get("account"), "name": r.get("name"), "command": r.get("command"),
             "schedule": r.get("schedule") or r.get("raw"), "next_run": r.get("next_run"), "last_run": r.get("last_run")}
            for r in d.get("rows") or []]
    if not rows:
        return _sec("schedule", "排程", [], [], unknown=d.get("reason") or _not_collected(ctx, "排程"))
    return _sec("schedule", "排程", [_c("source", "來源"), _c("account", "帳號"), _c("name", "名稱", True),
                                  _c("command", "指令", True), _c("schedule", "多久跑一次"), _c("next_run", "下次", True),
                                  _c("last_run", "上次", True)], rows)


# ===== 3 帳號與安全 =====

def s_accounts(ctx: Ctx) -> dict:
    import account_inventory

    rows, seen = [], set()
    for s in ctx.serials:
        for r in account_inventory.list_accounts(ctx.conn, asset_serial=s):
            if (r.get("ip"), r.get("username"), r.get("uid")) in seen:
                continue                 # 同一台的多筆登記查到同一批帳號，只列一次
            seen.add((r.get("ip"), r.get("username"), r.get("uid")))
            rows.append({"username": r.get("username"), "uid": r.get("uid"), "kind": r.get("kind"),
                         "admin": "是" if r.get("is_sudoer") else "否", "login": "可" if r.get("can_login") else "不可",
                         "last_login": r.get("last_login"), "pw_status": r.get("pw_status"), "shell": r.get("shell")})
    if not rows:
        return _sec("accounts", "帳號", [], [], unknown=_not_collected(ctx, "帳號"))
    return _sec("accounts", "帳號", [_c("username", "帳號", True), _c("uid", "UID", True), _c("kind", "種類"),
                                  _c("admin", "管理員權限"), _c("login", "能登入"), _c("last_login", "最後登入", True),
                                  _c("pw_status", "密碼狀態"), _c("shell", "Shell", True)], rows)


def s_vulns(ctx: Ctx) -> dict:
    import vuln_import

    d = vuln_import.by_asset(ctx.conn, ctx.serial)
    if not d.get("has_report"):
        return _sec("vulns", "弱點", [], [], unknown="還沒匯入過弱點報告")
    rows = [{"status": "未修", **{k: v.get(k) for k in ("plugin_name", "name", "severity", "risk", "port", "cve")}}
            for v in d.get("open") or []]
    rows += [{"status": "修過", **{k: v.get(k) for k in ("plugin_name", "name", "severity", "risk", "port", "cve")}}
             for v in d.get("closed") or []]
    for r in rows:
        r["title"] = r.pop("plugin_name") or r.pop("name", None)
    note = None if d.get("in_report") else "最近一次弱點報告裡沒有這台"
    return _sec("vulns", "弱點", [_c("status", "狀態"), _c("title", "弱點"), _c("severity", "嚴重度"), _c("risk", "風險"),
                               _c("port", "埠", True), _c("cve", "CVE", True)], rows, note=note)


def s_audit(ctx: Ctx) -> dict:
    import config_audit

    d = config_audit.by_asset(ctx.conn, ctx.serial)
    if not d.get("run"):
        return _sec("audit", "組態檢核", [], [], unknown=d.get("state") or "沒有檢核結果")
    rows = [{"rule": i.get("rule_id") or i.get("item_id") or i.get("code"), "title": i.get("title") or i.get("name"),
             "result": i.get("result_label") or i.get("result") or i.get("status"), "detail": i.get("detail") or i.get("actual")}
            for i in d.get("items") or []]
    run = d["run"]
    return _sec("audit", "組態檢核", [_c("rule", "條目", True), _c("title", "項目"), _c("result", "結果"), _c("detail", "實際值")],
                rows, note=f"{run.get('file_name')}・{run.get('imported_at')}・{run.get('baseline') or ''}")


def s_health(ctx: Ctx) -> dict:
    d = ctx.detail()
    h = d.get("health") or {}
    rows = [{"kind": i.get("label"), "level": i.get("level"), "detail": i.get("detail"), "basis": i.get("basis"),
             "action": i.get("action")} for i in h.get("issues") or []]
    hs = conn_health(ctx)
    if hs:
        rows.append({"kind": "最近一次健檢", "level": hs.get("overall"), "detail": _render(hs), "basis": hs.get("source"),
                     "action": None})
    if not rows and not h:
        return _sec("health", "健康檢查", [], [], unknown="沒有健康檢查資料")
    return _sec("health", "健康檢查", [_c("kind", "項目"), _c("level", "燈號"), _c("detail", "內容"), _c("basis", "依據"),
                                    _c("action", "要做什麼")], rows,
                note=None if rows else "沒有發現問題")


def conn_health(ctx: Ctx) -> dict | None:
    out = None
    try:
        for ip in ctx.ips:
            h = ctx.conn.execute("SELECT taken_at, overall, reachable, mem_pct, load_per, source FROM health_snapshot "
                                 "WHERE ip = ? ORDER BY taken_at DESC LIMIT 1", (ip,)).fetchone()
            if h and (out is None or h["taken_at"] > out["taken_at"]):
                out = dict(h)
    except sqlite3.OperationalError as exc:
        if "no such table" not in str(exc):
            raise
    return out


# ===== 關聯設備（不是影響範圍）＋連線原始列 =====

def s_conn(ctx: Ctx) -> dict:
    import v2_views

    d = v2_views.connections(ctx.conn, ctx.key)
    if d.get("reason") and not d["rows"]:
        return _sec("conn", "連線", [], [], unknown=d["reason"])
    return _sec("conn", "連線", [_c("direction", "方向"), _c("local_port", "本機埠", True), _c("peer_ip", "對方 IP", True),
                               _c("peer_name", "對方是誰"), _c("peer_port", "對方埠", True), _c("count", "次數", True),
                               _c("last_seen", "最後看到", True)], d["rows"])


def s_same_system(ctx: Ctx) -> dict:
    import v2_views

    names = v2_views._sysnames(ctx.conn)
    apis = sorted({(r["api_id"] or "").strip() for r in ctx.rows if (r["api_id"] or "").strip()})
    if not apis:
        return _sec("same_system", "同業務系統主機", [], [], unknown="CIA 沒填系統")
    rows = []
    for a in apis:
        for m in v2_views._same_system(ctx.conn, ctx.groups, a, ctx.key, names)["machines"]:
            rows.append({"key": m["key"], "hostname": m["hostname"], "type": f"{m['badge'].get('version_label') or m['badge']['label']}{' VM' if m['badge']['vm'] else ''}",
                         "ip": m["ip"], "where": m["where"], "system": f"{a} {names.get(a) or ''}"})
    return _sec("same_system", "同業務系統主機", [_c("hostname", "主機名"), _c("type", "類型"), _c("ip", "IP", True),
                                           _c("where", "在哪裡"), _c("system", "系統")], rows)


def s_same_rack(ctx: Ctx) -> dict:
    import v2_views

    ri = v2_views._rack_info(ctx.conn, ctx.mkey, ctx.serials)
    if not ri or ri.get("kind") != "rack":
        return _sec("same_rack", "同機櫃設備", [], [], unknown="不在機櫃上（VM 或 CIA 沒填機櫃）")
    rows = [{"key": k, "hostname": n} for k, n in zip(ri["other_keys"], ri["other_names"])]
    return _sec("same_rack", "同機櫃設備", [_c("hostname", "主機名")], rows, note=f"{ri['room']} {ri['rack']}")


def s_ocp(ctx: Ctx) -> dict:
    import arch

    d = arch.ocp_siblings(ctx.conn, ctx.serial)
    if not d.get("is_ocp"):
        return {**_sec("ocp", "OCP 同叢集", [], [], unknown="不是 OpenShift 節點"), "na": True}
    rows = [{"key": s.get("asset_serial"), "hostname": s.get("hostname"), "role": s.get("role"), "ip": s.get("ip"),
             "state": s.get("state") or s.get("status")} for s in d.get("siblings") or []]
    return _sec("ocp", "OCP 同叢集", [_c("hostname", "主機名"), _c("role", "角色"), _c("ip", "IP", True), _c("state", "狀態")],
                rows, note=f"叢集 {d.get('cluster')}（{d.get('basis')}）" + ("；" + "；".join(d.get("risks") or []) if d.get("risks") else ""))


# ===== 影響範圍（2026-10-27 使用者：「關係是指這台死掉了會影響哪一些主機」）=====
# 全部從 blast_radius.impact（3-2 影響範圍同一支）拿；這裡只依邊的種類分成四個小頁簽，不另寫走訪邏輯。
_NODE_KIND = {"hw": "主機", "esxi": "ESXi", "ds": "Datastore", "cluster": "叢集", "seg": "網段", "rack": "機櫃",
              "bizsys": "業務系統", "sys": "系統"}


def _no_graph(ctx: Ctx) -> str:
    return "CI 圖譜沒有此主機：尚未重建圖譜或不在圖譜範圍"


def s_imp_on(ctx: Ctx) -> dict:
    """承載的主機：這台 ESXi 上的 VM（runs_on）、同一個箱子的 LPAR／元件（4-5 箱子）。"""
    import rack_layout as rl
    import v2_views

    rows = []
    imp = ctx.impact()
    for d in (imp or {}).get("dependents") or []:
        if d["edge_type"] == "runs_on" and d["depth"] == 1:
            n = ctx.node_label(d["node_id"])
            rows.append({"key": n["key"], "hostname": n["name"], "type": "VM（跑在這台 ESXi 上）", "ip": n["ip"],
                         "how": "CI 圖譜 runs_on（RVTools／vCenter）"})
    ri = v2_views._rack_info(ctx.conn, ctx.mkey, ctx.serials)
    if ri and ri.get("kind") == "rack":
        det = rl.rack_detail(ctx.conn, ri["room"], ri["rack"])
        me = next((x for x in det["devices"] if x["key"] == ctx.mkey), None)
        for mm in (me or {}).get("merged") or []:
            rows.append({"key": mm["key"], "hostname": mm["name"], "type": mm.get("role") or "同一個箱子",
                         "ip": mm.get("ip"), "how": "4-5 機櫃：系統猜是同一個箱子"})
    if not rows:
        why = "無承載的主機：此主機不是 ESXi，也不是 Power 機身" if imp else _no_graph(ctx) + "；機櫃上也沒有同箱的登記"
        return _sec("imp_on", "承載的主機", [], [], unknown=why)
    return _sec("imp_on", "承載的主機", [_c("hostname", "主機名"), _c("type", "是什麼"), _c("ip", "IP", True),
                                       _c("how", "依據")], rows)


def s_imp_storage(ctx: Ctx) -> dict:
    """使用其儲存：這台 ESXi 上的 VM 放在哪些 datastore（stores_on）。儲存陣列 → datastore 還沒有對應資料。"""
    imp = ctx.impact()
    if not imp:
        return _sec("imp_storage", "使用其儲存", [], [], unknown=_no_graph(ctx))
    if imp.get("node_type") != "esxi":
        return _sec("imp_storage", "使用其儲存", [], [],
                    unknown="只有 ESXi 算得出來：儲存陣列跟 datastore 的對應還沒有資料")
    vms = [d["node_id"] for d in imp["dependents"] if d["edge_type"] == "runs_on" and d["depth"] == 1]
    rows = []
    if vms:
        q = ",".join("?" for _ in vms)
        for r in ctx.conn.execute(
                f"SELECT dst_node_id, COUNT(DISTINCT src_node_id) n FROM ci_edge WHERE edge_type = 'stores_on' "
                f"AND gone_at IS NULL AND src_node_id IN ({q}) GROUP BY dst_node_id ORDER BY n DESC", vms):
            rows.append({"datastore": ctx.node_label(r[0])["name"], "vms": r[1], "how": "RVTools vDisk（stores_on）"})
    if not rows:
        return _sec("imp_storage", "使用其儲存", [], [], unknown="這台 ESXi 上的 VM 沒有 datastore 資料")
    return _sec("imp_storage", "使用其儲存", [_c("datastore", "Datastore", True), _c("vms", "使用的 VM 數", True),
                                            _c("how", "依據")], rows)


def s_imp_callers(ctx: Ctx) -> dict:
    """連入主機：觀察到的連線（連進來）＋連線推論邊。"""
    import v2_views

    rows = []
    for d in (ctx.impact() or {}).get("dependents") or []:
        if d["edge_type"] == "observed_conn":
            n = ctx.node_label(d["node_id"])
            rows.append({"key": n["key"], "hostname": n["name"], "ip": n["ip"], "port": None, "count": None,
                         "how": f"連線推論（第 {d['depth']} 層）"})
    c = v2_views.connections(ctx.conn, ctx.key)
    for r in c.get("rows") or []:
        if r["direction"] == "inbound":
            rows.append({"key": r["peer_ip"], "hostname": r.get("peer_name") or r["peer_ip"], "ip": r["peer_ip"],
                         "port": r["local_port"], "count": r["count"], "how": "觀察到的連線（連進來）"})
    if not rows:
        return _sec("imp_callers", "連入主機", [], [],
                    unknown="沒有觀察到連入的連線：需已納管並收集服務資料" if not c.get("rows")
                    else "觀察到的連線都是它連出去的，沒有連進來的")
    return _sec("imp_callers", "連入主機", [_c("hostname", "主機名"), _c("ip", "IP", True), _c("port", "連入埠", True),
                                            _c("count", "次數", True), _c("how", "依據")], rows)


def s_imp_depends(ctx: Ctx) -> dict:
    """相依元件：它跑在哪台 ESXi、放在哪個 datastore、哪個叢集／機櫃／網段，加上它連出去的。"""
    import v2_views

    rows = []
    imp = ctx.impact()
    for d in (imp or {}).get("dependencies") or []:
        n = ctx.node_label(d["node_id"])
        kind = _NODE_KIND.get(d["node_id"].split(":")[0], n["node_type"])
        rows.append({"key": n["key"] if n["serial"] else None, "name": n["name"], "kind": kind, "port": None,
                     "how": f"CI 圖譜 {d['edge_type']}（第 {d['depth']} 層）"})
    c = v2_views.connections(ctx.conn, ctx.key)
    for r in c.get("rows") or []:
        if r["direction"] == "outbound":
            rows.append({"key": r["peer_ip"], "name": r.get("peer_name") or r["peer_ip"], "kind": "它連出去的主機",
                         "port": r["peer_port"], "how": "觀察到的連線（連出去）"})
    if not rows:
        return _sec("imp_depends", "相依元件", [], [], unknown=_no_graph(ctx) if not imp else "CI 圖譜沒有相依元件")
    return _sec("imp_depends", "相依元件", [_c("name", "名稱"), _c("kind", "是什麼"), _c("port", "埠", True),
                                        _c("how", "依據")], rows)


def _own_systems(ctx: Ctx) -> list[dict]:
    """這台自己登記的業務系統（含系統等級、重要度＝與 3-2 同一條：可用性 ≥ 3 為重大）。"""
    own = sorted({(r["api_id"] or "").strip() for r in ctx.rows if (r["api_id"] or "").strip()})
    out = []
    for a in own:
        b = ctx.conn.execute("SELECT name, sys_class, sys_class_override, ap_owner, ap_department FROM business_system "
                             "WHERE api_id = ?", (a,)).fetchone()
        av = ctx.conn.execute("SELECT MAX(availability) FROM hardware WHERE TRIM(api_id) = ?", (a,)).fetchone()[0]
        out.append({"api_id": a, "name": b["name"] if b else None,
                    "sys_class": ((b["sys_class_override"] or b["sys_class"]) if b else None),
                    "severity": "重大" if (av or 0) >= 3 else "一般",
                    "ap_owner": b["ap_owner"] if b else None, "ap_department": b["ap_department"] if b else None})
    return out


def s_imp_systems(ctx: Ctx) -> dict:
    """受影響業務系統（2.006.5，找碴 C9）：本系統＋上下游（3-2 影響範圍同一份 by_biz_system），
    每套寫重要度、系統等級、AP 部門／主管——原本要點到舊版 /blast 才看得到。"""
    imp = ctx.impact()
    rows = [{**s, "relation": "本系統", "hosts": None, "via": "主機登記"} for s in _own_systems(ctx)]
    seen = {r["api_id"] for r in rows}
    for b in (imp or {}).get("summary", {}).get("by_biz_system") or []:
        if b["api_id"] in seen:
            continue
        seen.add(b["api_id"])
        bs = ctx.conn.execute("SELECT sys_class, sys_class_override, ap_owner, ap_department FROM business_system "
                              "WHERE api_id = ?", (b["api_id"],)).fetchone()
        rows.append({"api_id": b["api_id"], "name": b.get("name"), "relation": "受影響（上下游）",
                     "sys_class": ((bs["sys_class_override"] or bs["sys_class"]) if bs else None),
                     "severity": b.get("severity"), "hosts": len(b.get("assets") or []),
                     "via": "連線推論" if b.get("via") == "observed" else "登記的關聯",
                     "ap_owner": bs["ap_owner"] if bs else None, "ap_department": bs["ap_department"] if bs else None})
    if not rows:
        return _sec("imp_systems", "受影響業務系統", [], [],
                    unknown=(_no_graph(ctx) + "；主機也沒有登記業務系統") if not imp else "主機沒有登記業務系統，CI 圖譜也沒有走到任何系統")
    note = None if (imp or {}).get("deps_available") else "業務系統之間的上下游資料不足，實際影響可能更多"
    return _sec("imp_systems", "受影響業務系統", [_c("api_id", "代號", True), _c("name", "業務系統"), _c("relation", "關係"),
                                            _c("severity", "重要度"), _c("sys_class", "系統等級"), _c("hosts", "受影響主機（台）", True),
                                            _c("ap_department", "AP 部門"), _c("ap_owner", "AP 主管"), _c("via", "依據")], rows, note=note)


def s_imp_notify(ctx: Ctx) -> dict:
    """通知對象：3-2 影響範圍同一份 notify（受影響主機的負責人／AP）。"""
    imp = ctx.impact()
    if not imp:
        return _sec("imp_notify", "通知對象", [], [], unknown=_no_graph(ctx))
    rows = [{"name": n.get("name"), "role": n.get("role"), "department": n.get("department"), "phone": n.get("phone")}
            for n in imp["summary"].get("notify") or []]
    if not rows:
        return _sec("imp_notify", "通知對象", [], [], unknown="受影響的主機都沒有登錄負責人")
    return _sec("imp_notify", "通知對象", [_c("name", "姓名"), _c("role", "角色"), _c("department", "部門"),
                                        _c("phone", "電話", True)], rows)


def impact_total(ctx: Ctx) -> dict:
    """合計一行：本系統（重要度）、上下游幾套、承載幾台主機（3-2 影響範圍同一份 summary）。
    2.006.5：不再連到舊版 3-2，完整清單在下方「受影響業務系統」「通知對象」小頁籤。"""
    imp = ctx.impact()
    own = _own_systems(ctx)
    own_txt = "、".join(f"{s['api_id']} {s['name'] or ''}".strip() + f"（{s['severity']}）" for s in own)
    link = None
    if not imp:
        return {"known": False, "hosts": None, "systems": len(own), "link": link,
                "text": f"無法計算影響範圍：{_no_graph(ctx)}" + (f"；至少影響本系統 {own_txt}" if own else "")}
    others = {b["api_id"] for b in imp["summary"]["by_biz_system"]} - {s["api_id"] for s in own}
    hosts = int((imp["summary"].get("counts") or {}).get("host") or 0)
    parts = [f"本系統 {own_txt}" if own else "未登記業務系統",
             f"上下游 {len(others)} 套" if imp.get("deps_available") or others else "上下游無資料",
             f"承載／相依主機 {hosts} 台（不含此主機）"]
    return {"known": True, "hosts": hosts, "systems": len(own) + len(others), "link": link,
            "deps_available": imp.get("deps_available"), "text": "影響：" + "；".join(parts)}


# ===== 5 紀錄與原始資料 =====

def s_changes(ctx: Ctx) -> dict:
    import daily_digest

    d = ctx.detail()
    rows = [{"at": x["date"], "source": "每日快照比對", "what": f"{x['label']}"
             + (f"：{x.get('before') if x.get('before') is not None else '（空）'} → {x.get('after') if x.get('after') is not None else '（空）'}"
                if x.get("before") is not None or x.get("after") is not None else "")
             + (f"（{x['detail']}）" if x.get("detail") else "")} for x in daily_digest.machine_history(ctx.conn, ctx.mkey, 365)]
    rows += [{"at": h.get("detected_at"), "source": "異常紀錄", "what": f"{h.get('issue_type')}：{h.get('hostname') or ''} {h.get('ip') or ''}"
              + ("（已處理）" if h.get("handled_at") else "")} for h in d.get("history") or []]
    rows += [{"at": h.get("created_at"), "source": "納管紀錄", "what": f"{'成功' if h.get('ok') else '失敗'}・{h.get('stage') or ''}・"
              f"{h.get('message') or ''}（{h.get('triggered_by') or ''}）"} for h in d.get("onboard_history") or []]
    rows.sort(key=lambda x: x["at"] or "", reverse=True)
    return _sec("changes", "異動紀錄", [_c("at", "時間", True), _c("source", "來源"), _c("what", "變了什麼")], rows)


_NOT_CIA = ("DYN-", "VC-", "AUTO-")


def s_cia(ctx: Ctx) -> dict:
    rows = []
    for r in ctx.rows:
        if (r["asset_serial"] or "").upper().startswith(_NOT_CIA):
            continue
        for k, v in r.items():
            if v not in (None, ""):
                rows.append({"serial": r["asset_serial"], "field": k, "value": str(v)})
    if not rows:
        return _sec("cia", "CIA 原始欄位", [], [], unknown="CIA 沒有登記這台（只有其他來源看到）")
    return _sec("cia", "CIA 原始欄位", [_c("serial", "資產編號", True), _c("field", "欄位", True), _c("value", "值")], rows)


def s_sources(ctx: Ctx) -> dict:
    """所有登記來源。純讀（Ctx 只 SELECT hardware；本段不呼叫任何 ensure_*／不寫庫，W221-009 Codex 53 已驗）。
    位置三套口徑各司其職，這裡全列出讓人核對（權威入口見 docs/codex-review/2026-10-30/位置與來源時間_權威入口.md）：
    raw_location＝physical_location 原值；location_group＝manage_state.group_location（含「（未歸類：原值）」「未填」）；
    rack_room＝blast_radius.rack_room（機櫃／容量／預留的機房鍵，分點保留）。"""
    import blast_radius
    import manage_state as ms

    rows = []
    for r in ctx.rows:
        s = r["asset_serial"] or ""
        src = "DY 動態資產" if s.startswith("DYN-") else "vCenter／RVTools" if s.startswith("VC-") else \
            "自動建立" if s.startswith("AUTO-") else "CIA"
        rows.append({"serial": s, "source": src, "hostname": r.get("hostname"), "ip": r.get("ip"), "os": r.get("os"),
                     "status": r.get("asset_status"), "api_id": r.get("api_id"), "rack": r.get("rack_no"),
                     "raw_location": r.get("physical_location"), "location_group": ms.group_location(r.get("physical_location")),
                     "rack_room": blast_radius.rack_room(r.get("physical_location")),
                     "source_updated_at": r.get("updated_at"), "manual_updated_at": r.get("manual_updated_at")})
    return _sec("sources", "所有登記來源", [_c("serial", "登記", True), _c("source", "來源"), _c("hostname", "主機名"),
                                    _c("ip", "IP", True), _c("os", "作業系統"), _c("status", "狀態"), _c("api_id", "系統", True),
                                    _c("rack", "機櫃欄"), _c("raw_location", "原始位置（原值）"), _c("location_group", "位置分組"),
                                    _c("rack_room", "機櫃機房鍵"), _c("source_updated_at", "資料時間", True),
                                    _c("manual_updated_at", "人工修改時間", True)], rows, note=f"同一台的 {len(rows)} 筆登記（系統猜是同一台）")


def s_personnel(ctx: Ctx) -> dict:
    rows = []
    for s in ctx.serials:
        for p in ctx.conn.execute("SELECT * FROM personnel WHERE asset_serial = ?", (s,)):
            p = dict(p)
            rows.append({"serial": s, "name": p.get("person_name"), "phone": p.get("phone"), "proxy": p.get("proxy1"),
                         "proxy_phone": p.get("proxy1_phone"), "dept": " ".join(x for x in (p.get("belong_division"), p.get("belong_department")) if x),
                         "job": p.get("job_desc")})
    if not rows:
        return _sec("personnel", "人員（CIA）", [], [], unknown="CIA 沒有登錄聯絡人")
    return _sec("personnel", "人員（CIA）", [_c("name", "姓名"), _c("phone", "電話", True), _c("proxy", "代理人"),
                                         _c("proxy_phone", "代理人電話", True), _c("dept", "單位"), _c("job", "職掌"),
                                         _c("serial", "登記", True)], rows)


def s_existence(ctx: Ctx) -> dict:
    import existence

    m = existence.for_serial(ctx.conn, ctx.serial)
    if not m:
        return _sec("existence", "偵測證據", [], [], unknown="存在可信度算不出來（這台不在計算範圍）")
    label = existence._EV_LABEL if hasattr(existence, "_EV_LABEL") else {}
    rows = [{"source": label.get(k, k), "last_seen": v} for k, v in sorted((m.get("last_seen") or {}).items(),
                                                                       key=lambda kv: kv[1] or "", reverse=True) if v]
    note = f"偵測狀態：{m.get('level_label')}（{m.get('score')} 分）" + ("・計入台數" if m.get("counted") else "・沒計入台數")
    return _sec("existence", "偵測證據", [_c("source", "哪個來源看到"), _c("last_seen", "最後看到", True)], rows, note=note)


# ===== 系統重建資料 =====

def s_rebuild(ctx: Ctx) -> dict:
    from routers import rebuild as rr

    d = rr.detail_for(ctx.conn, ctx.serial)
    if not d:
        return _sec("rebuild", "系統重建資料", [], [], kind="rebuild", unknown=_not_collected(ctx, "系統重建資料"))
    rows = []
    meta = {"state", "id", "group", "title", "necessity", "label", "note", "truncated"}
    for g in (d.get("data") or {}).get("groups") or []:
        for it in g.get("items") or []:
            body = []
            for k, v in it.items():
                if k in meta:
                    continue
                if k == "files" and isinstance(v, list):
                    for f in v:
                        body.append(f"### {f.get('path')}\n{f.get('content') or ''}")
                else:
                    body.append(f"[{k}]\n{_render(v)}")
            code = it.get("state") or ""
            st = "ok" if code == "collected" else ("none" if code == "empty" else "bad")
            rows.append({"group": g.get("group"), "id": it.get("id"), "title": it.get("title"),
                         "necessity": "必要" if it.get("necessity") == "A" else "參考", "state": it.get("label") or code,
                         "status": st, "problem": st == "bad" and it.get("necessity") == "A",
                         "note": it.get("note"), "content": "\n".join(body).strip()})
    s = (d.get("data") or {}).get("summary") or {}
    return _sec("rebuild", "系統重建資料", [_c("group", "分組"), _c("title", "項目"), _c("necessity", "必要"),
                                       _c("state", "狀態"), _c("content", "內容")], rows, kind="rebuild",
                note=f"{d.get('platform') or ''}・最後收集 {d.get('collected_at') or '—'}・必要項已收 {s.get('a_collected', '—')}/{s.get('total_a', '—')}")


SECTIONS: list[tuple[str, str, Callable[[Ctx], dict]]] = [
    ("hw", "hw_spec", s_hw_spec), ("hw", "disks", s_disks), ("hw", "nics", s_nics), ("hw", "san", s_san),
    ("hw", "location", s_location),
    ("sw", "software", s_software), ("sw", "services", s_services), ("sw", "schedule", s_schedule),
    ("sec", "accounts", s_accounts), ("sec", "vulns", s_vulns), ("sec", "audit", s_audit), ("sec", "health", s_health),
    ("impact", "imp_systems", s_imp_systems), ("impact", "imp_notify", s_imp_notify),
    ("impact", "imp_on", s_imp_on), ("impact", "imp_storage", s_imp_storage), ("impact", "imp_callers", s_imp_callers),
    ("impact", "imp_depends", s_imp_depends),
    ("near", "same_system", s_same_system), ("near", "same_rack", s_same_rack), ("near", "ocp", s_ocp), ("near", "conn", s_conn),
    ("chg", "changes", s_changes), ("raw", "cia", s_cia), ("raw", "sources", s_sources), ("raw", "personnel", s_personnel),
    ("raw", "existence", s_existence),
    ("rebuild", "rebuild", s_rebuild),
]
SECTION_KEYS = [k for _, k, _ in SECTIONS]


def _run(ctx: Ctx, group: str, key: str, fn) -> dict:
    try:
        out = fn(ctx)
    except Exception as exc:  # noqa: BLE001 - 一段讀不到不能讓整張病歷表掛掉，但要講出原因
        out = _sec(key, key, [], [], unknown=f"讀取失敗：{type(exc).__name__}: {exc}")
    out["group"] = group
    return out


def section(conn: sqlite3.Connection, key: str, name: str) -> dict:
    fn = next(((g, f) for g, k, f in SECTIONS if k == name), None)
    if fn is None:
        raise ValueError(f"沒有這一段：{name}")
    return _run(Ctx(conn, key), fn[0], name, fn[1])


def all_sections(conn: sqlite3.Connection, key: str) -> list[dict]:
    ctx = Ctx(conn, key)
    return [_run(ctx, g, k, f) for g, k, f in SECTIONS]


def summary(conn: sqlite3.Connection, key: str) -> dict:
    """每一段的筆數或「無資料（原因）」（頁簽上的數字；跟點進去的 rows 同一次計算）。"""
    ctx = Ctx(conn, key)
    secs = [_run(ctx, g, k, f) for g, k, f in SECTIONS]
    return {"groups": [{"key": g, "label": lab} for g, lab in GROUPS], "impact_total": impact_total(ctx),
            "sections": [{**{k: s[k] for k in ("key", "group", "label", "count", "unknown", "note")}, "na": bool(s.get("na"))}
                         for s in secs]}


def export_xlsx(conn: sqlite3.Connection, key: str, only: str | None = None) -> tuple[io.BytesIO, str]:
    """匯出完整病歷（一段一張工作表）；only='rebuild' 只匯系統重建資料。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    ctx = Ctx(conn, key)
    secs = [_run(ctx, g, k, f) for g, k, f in SECTIONS if not only or k == only]
    name = ctx.rows[0]["hostname"] or ctx.serial
    wb = Workbook()
    ws = wb.active
    ws.title = "摘要"
    ws.append(["機器", name])
    ws.append(["資產編號", "、".join(ctx.serials)])
    ws.append(["IP", "、".join(ctx.ips)])
    ws.append([])
    ws.append(["段", "筆數", "說明"])
    for s in secs:
        ws.append([s["label"], s["count"] if s["count"] is not None else "無資料", s["unknown"] or s["note"] or ""])
    ws.column_dimensions["A"].width, ws.column_dimensions["C"].width = 22, 80
    hdr = PatternFill("solid", fgColor="00806A")
    used = {"摘要"}
    for s in secs:
        t = s["label"].replace("／", "-").replace("（", "(").replace("）", ")")[:28]
        while t in used:
            t = t[:26] + "_2"
        used.add(t)
        w = wb.create_sheet(t)
        if s["unknown"]:
            w.append([f"無資料（{s['unknown']}）"])
            continue
        if s["note"]:
            w.append([s["note"]])
        w.append(["項次"] + [c["label"] for c in s["cols"]])
        for c in w[w.max_row]:
            c.fill, c.font = hdr, Font(color="FFFFFF", bold=True)
        for i, r in enumerate(s["rows"], 1):
            vals = [i] + [r.get(c["k"]) for c in s["cols"]]
            w.append([v if not isinstance(v, str) else v[:32000] for v in vals])
        if s["kind"] == "rebuild":
            w.column_dimensions["E"].width = 120
            for row in w.iter_rows(min_row=2):
                row[-1].alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    import export_safe
    export_safe.neutralize_workbook(wb)     # 防公式注入（審查 B1）
    wb.save(buf)
    buf.seek(0)
    from datetime import date
    return buf, f"{'系統重建資料' if only == 'rebuild' else '主機履歷'}_{name}_{date.today().strftime('%Y%m%d')}.xlsx"
