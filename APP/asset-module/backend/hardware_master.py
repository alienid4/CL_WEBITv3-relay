"""資產主檔匯出：全欄位＋同台判定，外加一張欄位說明（2026-09-18）。

## 為什麼要有這支

2026-09-18 早上另一個 session 用一次性腳本產了「WEBIT3_資產主檔_hardware_YYYYMMDD.xlsx」
給使用者（兩張分頁：資產主檔、欄位說明），使用者問「這表格系統能直接產嗎」。
一次性腳本不進版控、交接也沒寫，下一個人根本不知道它存在——所以做成系統功能。

## 比那份腳本多做的兩件事

1. **有效值另計**：匯入來源大量用「NA」填空。原本的「已填筆數」把 NA 算成已填，
   填充率會虛胖（221 實測 request_no 大半是 NA）。欄位說明多一欄「有效值」，
   兩個數字並列，由人判斷——不替使用者決定哪個才算數。
2. **同台筆數不跨退役**：使用者 2026-09-18 拍板「報廢跟使用中不能算同一台」
   （IP／主機名會被回收給新機器，退役那筆是上一台）。所以同台筆數以
   `machine_key ＋ 是否退役` 分組，跟資產詳細頁的別名（v1.213.1）同一套規則。

## 欄位意義從哪來

直接讀 `schema.sql` 的 `--` 註解，不另外維護一份——兩份遲早會不一致。
schema.sql 沒寫到的欄位（後來用 ALTER 加的）標「schema 無註解」，不猜。
"""
from __future__ import annotations

import io
import re
from pathlib import Path

import manage_state as ms
import system_stats

SCHEMA = Path(__file__).parent / "schema.sql"

#: 這些字串是「沒有值」的填空字，不是值
#: 「無」刻意**不**列入：對 big_ip_vip 這類欄位，「無」是有意義的回答（＝沒有 VIP），
#: 不是漏填（使用者 2026-09-18 拍板拿掉）。
NOT_A_VALUE = {"", "na", "n/a", "none", "null", "-", "--"}

#: 衍生欄位（不在資料表裡、由系統算出來的），欄位說明要一併講清楚怎麼算
DERIVED = (
    ("來源", "序號前綴判定：DYN-＝DY、VC-＝vCenter、AUTO-＝系統自動補，其餘＝CIA 正式登記"),
    ("machine_key", "同一台的判定鍵：主機名（小寫）|IP；缺任一或 IP 為 0.0.0.0 時退回 sn:序號（各算一台）"),
    ("同台筆數", "同一個 machine_key、且同為退役或同為非退役的登記筆數。報廢與使用中不算同一台（IP／主機名會被回收）"),
)


def _source_of(serial) -> str:
    s = str(serial or "")
    if s.startswith("DYN-"):
        return "DY"
    if s.startswith("VC-"):
        return "vCenter"
    if s.startswith("AUTO-"):
        return "AUTO"
    return "CIA"


def schema_comments() -> dict[str, str]:
    """讀 schema.sql 的 hardware 區塊：欄位 → 註解。

    行尾有 `--` 註解就用它；沒有的話，用緊接在上面那段區塊註解的第一行
    （像 vi_sdk_server 那種整段說明寫在上方的）。
    """
    out: dict[str, str] = {}
    try:
        lines = SCHEMA.read_text(encoding="utf-8").splitlines()
    except OSError:
        return out
    inside, pending = False, []
    for ln in lines:
        s = ln.strip()
        if not inside:
            if re.match(r"CREATE TABLE (IF NOT EXISTS )?hardware\b", s):
                inside = True
            continue
        if s.startswith(");"):
            break
        if s.startswith("--"):
            txt = s.lstrip("-").strip()
            if txt:
                pending.append(txt)
            continue
        m = re.match(r"([a-z_][a-z0-9_]*)\s+[A-Z]", s)
        if not m:
            continue
        col = m.group(1)
        note = s.split("--", 1)[1].strip() if "--" in s else ""
        if not note and pending:
            note = pending[0]
        out[col] = note
        pending = []
    return out


def _columns(conn) -> list[str]:
    return [r[1] for r in conn.execute("PRAGMA table_info(hardware)")]


def _is_value(v) -> bool:
    if v is None:
        return False
    return str(v).strip().lower() not in NOT_A_VALUE


def field_dictionary(conn) -> list[dict]:
    cols = _columns(conn)
    notes = schema_comments()
    total = conn.execute("SELECT COUNT(*) FROM hardware").fetchone()[0] or 0
    out = []
    for i, c in enumerate(cols, 1):
        filled = valid = 0
        for (v,) in conn.execute(f'SELECT "{c}" FROM hardware'):
            if v is not None and str(v).strip() != "":
                filled += 1
                if _is_value(v):
                    valid += 1
        out.append({
            "no": i,
            "column": c,
            "meaning": notes.get(c) or "（schema 無註解）",
            "filled": filled,
            "fill_rate": (filled / total) if total else None,
            "valid": valid,
            "valid_rate": (valid / total) if total else None,
        })
    return out


def master_rows(conn) -> tuple[list[str], list[list]]:
    """資產主檔：hardware 全欄位（原樣）＋來源＋machine_key＋同台筆數。"""
    cols = _columns(conn)
    rows = conn.execute("SELECT " + ", ".join(f'"{c}"' for c in cols) + " FROM hardware ORDER BY id").fetchall()
    idx = {c: i for i, c in enumerate(cols)}

    def key_of(r):
        mk = system_stats.machine_key(r[idx["hostname"]], r[idx["ip"]], r[idx["asset_serial"]])
        retired = (str(r[idx["asset_status"]] or "").strip() in ms.RETIRED_STATUS)
        return mk, retired

    counts: dict[tuple, int] = {}
    for r in rows:
        k = key_of(r)
        counts[k] = counts.get(k, 0) + 1

    header = cols + [d[0] for d in DERIVED]
    body = []
    for r in rows:
        mk, retired = key_of(r)
        body.append(list(r) + [_source_of(r[idx["asset_serial"]]), mk, counts[(mk, retired)]])
    return header, body


def export_xlsx(conn) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0B7A5B")

    wb = Workbook()
    ws = wb.active
    ws.title = "資產主檔"
    header, body = master_rows(conn)
    ws.append(header)
    for r in body:
        ws.append(r)
    for cell in ws[1]:
        cell.font, cell.fill = head_font, head_fill
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("欄位說明")
    ws2.append(["#", "欄位", "意義（取自 schema 註解）", "已填筆數", "填充率",
                "有效值筆數（排除 NA 等填空字）", "有效率"])
    for d in field_dictionary(conn):
        ws2.append([d["no"], d["column"], d["meaning"], d["filled"], d["fill_rate"],
                    d["valid"], d["valid_rate"]])
    n = ws2.max_row
    for row in ws2.iter_rows(min_row=2, max_row=n):
        row[4].number_format = "0%"
        row[6].number_format = "0%"
    ws2.append([])
    ws2.append(["", "衍生欄位", "（系統計算，不在資料表裡）"])
    for name, how in DERIVED:
        ws2.append(["", name, how])
    ws2.append([])
    ws2.append(["", "填空字", "以下值不算有效值：" + "、".join(sorted(x for x in NOT_A_VALUE if x) ) + "（不分大小寫）與空白"])
    for cell in ws2[1]:
        cell.font, cell.fill = head_font, head_fill
    ws2.column_dimensions["B"].width = 24
    ws2.column_dimensions["C"].width = 60
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
